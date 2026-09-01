# -*- coding: utf-8 -*-
"""Структурный разбор DOCX: части документа вместе с их МЕСТОМ в разметке.

Зачем. Разбор заявлений сегодня работает по плоской строке: DOCX схлопывается
через join по переводу строки, а раскладка (таблица? какая ячейка? какой абзац?)
потом кропотливо восстанавливается регулярками. Отсюда семь паттернов на одно
поле — по одному на каждую вёрстку.

Между тем сама раскладка в файле ЕСТЬ, её просто выбрасывают. Доказательство —
в этом же проекте: `_apply_table_amounts` берёт суммы из настоящих ячеек и стоит
последним, самым доверенным слоем финансового каскада. Он применён к пяти
меткам; этот модуль обобщает приём.

Контракт совместимости (проверяется `tests/test_doc_structure.py`): плоский
текст, собранный из этих частей, байт-в-байт равен тому, что отдаёт
`_docx_text_parts`. Пока это держится, структурный слой НЕ МОЖЕТ сдвинуть
эталон — плоский путь получает ровно тот же текст, что и раньше.

На PDF модуль не работает: там разметки нет в принципе (см. `analyze_from_text`).
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from functools import lru_cache
from typing import Any, Dict, List, Optional, Pattern, Tuple

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class DocPart:
    """Одна часть документа с провенансом.

    `origin` повторяет старый плоский контракт (body / table / colophon / raw),
    остальные поля — то, что раньше терялось:

    * `para_index` — порядковый номер абзаца в теле, для запроса «значение =
      следующий непустой абзац после абзаца-метки»;
    * `table_id` / `row` / `col` — координаты ячейки, для запроса «значение =
      ячейка справа от ячейки с меткой»;
    * `col_span` — сколько колонок перекрывает объединённая ячейка. Обход
      `row.cells` в python-docx отдаёт такую ячейку по разу на каждую
      перекрытую колонку; в плоском тексте это лишние повторы (замерено:
      5 документов корпуса, 85 повторов), а здесь — честная ширина;
    * `style` — имя стиля абзаца; заголовки разделов отличаются от текста
      именно им, а не написанием.
    """

    text: str
    origin: str
    para_index: Optional[int] = None
    table_id: Optional[int] = None
    row: Optional[int] = None
    col: Optional[int] = None
    col_span: int = 1
    style: Optional[str] = None


def _paragraph_style(paragraph: Any) -> Optional[str]:
    """Имя стиля абзаца; None, если стиль недоступен.

    Битые и урезанные DOCX встречаются регулярно (конвертер, чужие генераторы),
    а стиль здесь — подсказка, а не обязательное поле: падать из-за него нельзя.
    """
    try:
        style = paragraph.style
        return style.name if style is not None else None
    except Exception:
        return None


def _cells_with_span(row: Any) -> List[Tuple[Any, int, int]]:
    """Ячейки строки как (cell, col, col_span) — БЕЗ повторов объединённых.

    `row.cells` отдаёт объединённую ячейку по разу на каждую перекрытую колонку.
    Отличаем повтор от соседней ячейки с тем же текстом по тождеству XML-узла
    (`cell._tc`), а не по тексту: в таблицах расчёта задолженности одинаковые
    значения в соседних колонках — норма.

    Сравнение идёт по `is`, а не по `id()`: прокси lxml создаются по требованию,
    и CPython переиспользует освободившийся `id` (тот же дефект — §S.3.5).
    Список `seen` держит узлы живыми на время обхода.
    """
    out: List[Tuple[Any, int, int]] = []
    seen: List[Any] = []
    for col, cell in enumerate(row.cells):
        try:
            tc = cell._tc
        except Exception:
            tc = None

        if tc is not None and any(tc is prev for prev in seen):
            # Повтор объединённой ячейки: расширяем span уже добавленной.
            for i in range(len(out) - 1, -1, -1):
                prev_cell, prev_col, prev_span = out[i]
                try:
                    if prev_cell._tc is tc:
                        out[i] = (prev_cell, prev_col, prev_span + 1)
                        break
                except Exception:
                    continue
            continue

        if tc is not None:
            seen.append(tc)
        out.append((cell, col, 1))
    return out


def docx_parts(
    file_path: str,
    doc: Any = None,
    raw_chunks: Optional[List[str]] = None,
) -> List[DocPart]:
    """Части DOCX с провенансом, В ТОМ ЖЕ ПОРЯДКЕ, что и плоский экстрактор.

    Порядок обязан совпадать: `document_analyzer._docx_text_parts` строит из
    этого списка пары (text, origin), а из них — плоский текст анализа.
    Перестановка частей сдвинула бы эталон.

    `raw_chunks` — куски, недоступные объектной модели python-docx (надписи,
    сноски). Их собирает вызывающий (`_extract_docx_raw_xml_text`): разбор
    сырого XML остаётся в анализаторе, чтобы не заводить второй обход zip.
    """
    from docx import Document

    if doc is None:
        doc = Document(file_path)

    parts: List[DocPart] = []

    # 1. Тело: сначала все абзацы, затем все таблицы — как в плоском обходе.
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            parts.append(DocPart(
                text, "body", para_index=index, style=_paragraph_style(paragraph),
            ))

    for table_id, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            for cell, col, span in _cells_with_span(row):
                text = cell.text.strip()
                if text:
                    parts.append(DocPart(
                        text, "table",
                        table_id=table_id, row=row_index, col=col, col_span=span,
                    ))

    # 2. Колонтитулы всех секций, с дедупликацией по содержимому: секции часто
    # ссылаются на один и тот же колонтитул.
    seen_headers = set()
    for section in doc.sections:
        for container in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if container is None:
                continue
            try:
                chunk_parts: List[str] = []
                for paragraph in container.paragraphs:
                    if paragraph.text.strip():
                        chunk_parts.append(paragraph.text.strip())
                for table in container.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                chunk_parts.append(cell.text.strip())
                chunk = "\n".join(chunk_parts)
                if chunk and chunk not in seen_headers:
                    seen_headers.add(chunk)
                    parts.append(DocPart(chunk, "colophon"))
            except Exception as exc:
                logger.warning(f"Не удалось обработать колонтитул: {exc}")

    # 3. Надписи и сноски — как есть, без координат: разметки у них нет.
    for chunk in (raw_chunks or []):
        parts.append(DocPart(chunk, "raw"))

    return parts


def flat_pairs(parts: List[DocPart]) -> List[Tuple[str, str]]:
    """Старый контракт (text, origin) из структурных частей.

    ВАЖНО: объединённые ячейки здесь РАЗВОРАЧИВАЮТСЯ обратно в повторы. Плоский
    текст — вход эталона, и убрать из него дубликаты значит сдвинуть golden;
    это отдельная осознанная правка (§S.3.2), а не побочный эффект провенанса.
    """
    pairs: List[Tuple[str, str]] = []
    for part in parts:
        pairs.append((part.text, part.origin))
        for _ in range(part.col_span - 1):
            pairs.append((part.text, part.origin))
    return pairs


# --- Структурные запросы -----------------------------------------------------

def cells_of(parts: List[DocPart]) -> List[DocPart]:
    """Только ячейки таблиц."""
    return [p for p in parts if p.origin == "table"]


def index_cells(parts: List[DocPart]) -> Dict[Tuple[int, int, int], DocPart]:
    """Ячейки по координате (table_id, row, col) — для соседских запросов."""
    return {
        (p.table_id, p.row, p.col): p
        for p in parts
        if p.origin == "table" and p.table_id is not None
    }


# Ячейка-метка — это ПОДПИСЬ, а не значение: «ИНН», «ИНН:», «ИНН должника».
# Верхняя граница длины отсекает абзац-простыню, случайно попавший в таблицу.
_LABEL_CELL_MAX = 60

# Хвост после метки внутри одной ячейки: «ИНН: 7707083893». Двоеточие/тире
# обязательны — иначе «ИНН организации» дало бы значение «организации».
_INLINE_SEP = r"[:\-–—]"


@lru_cache(maxsize=1)
def _next_label_re() -> Pattern[str]:
    """Начало ПОДПИСИ следующего поля — граница значения.

    Одна ячейка часто несёт несколько полей подряд: «545 418,01 рублей
    Государственная пошлина: 16 136 рублей». Без обрезки значение уезжает вместе
    с чужой подписью — теневой отчёт показал это на 18 документах по
    `claimAmount`.

    Метки берутся из реестра, длинные раньше коротких (`all_labels` это уже
    гарантирует), иначе «Адрес» съел бы «Адрес регистрации».
    """
    from label_synonyms import all_labels, labels_alternation

    return re.compile(
        r"[\s,;]*(?:" + labels_alternation(all_labels()) + r")\s*" + _INLINE_SEP,
        re.IGNORECASE,
    )


def _trim_at_next_label(value: str) -> str:
    """Обрезает значение на подписи следующего поля."""
    return _next_label_re().split(value, maxsplit=1)[0].strip().rstrip(",;")


@lru_cache(maxsize=256)
def _label_re(label: str) -> Pattern[str]:
    """Регулярка метки: начало части, целым словом, регистронезависимо."""
    return re.compile(r"^\s*" + re.escape(label) + r"\b", re.IGNORECASE)


@lru_cache(maxsize=256)
def _bare_label_re(label: str) -> Pattern[str]:
    """Часть, которая ЦЕЛИКОМ является меткой: «Адрес», «Адрес:», «Адрес —».

    Разделяющий признак для запросов «сосед» и «следующий абзац». Первый заход
    требовал лишь начала с метки — и теневой отчёт сразу показал цену: абзац
    «Адрес регистрации: 413105, …» считался подписью, а значением объявлялся
    СЛЕДУЮЩИЙ абзац («Дата государственной регистрации: …»). Из 69 расхождений
    отчёта почти все были этим.
    """
    return re.compile(
        r"^\s*" + re.escape(label) + r"\s*" + _INLINE_SEP + r"?\s*$", re.IGNORECASE,
    )


@lru_cache(maxsize=256)
def _inline_re(label: str) -> Pattern[str]:
    """Метка + обязательный разделитель + значение в той же части."""
    return re.compile(
        r"^\s*" + re.escape(label) + r"\b\s*" + _INLINE_SEP + r"\s*(.+)$",
        re.IGNORECASE | re.DOTALL,
    )


def _is_label_cell(part: DocPart, label: str) -> bool:
    """Часть — подпись поля, а не текст: коротка и состоит ТОЛЬКО из метки."""
    return len(part.text) <= _LABEL_CELL_MAX and bool(_bare_label_re(label).match(part.text))


def value_right_of_label(
    parts: List[DocPart], label: str, max_gap: int = 3,
) -> Optional[str]:
    """Значение из ячейки СПРАВА от ячейки-метки.

    Основная форма таблицы реквизитов: подпись слева, значение справа. Пустые
    ячейки между ними пропускаем — вёрстка часто оставляет разделительную
    колонку, — но не дальше `max_gap`, иначе поедем в соседний блок.

    Объединённую метку перешагиваем по её ширине: «ИНН» на три колонки значит,
    что значение стоит в четвёртой, а не во второй.
    """
    idx = index_cells(parts)
    for part in cells_of(parts):
        if not _is_label_cell(part, label):
            continue
        start = (part.col or 0) + part.col_span
        for step in range(max_gap):
            neighbour = idx.get((part.table_id, part.row, start + step))
            if neighbour is not None and neighbour.text.strip():
                return _trim_at_next_label(neighbour.text)
    return None


def value_below_label(parts: List[DocPart], label: str) -> Optional[str]:
    """Значение из ячейки ПОД ячейкой-меткой — форма «шапка таблицы сверху»."""
    idx = index_cells(parts)
    for part in cells_of(parts):
        if not _is_label_cell(part, label):
            continue
        below = idx.get((part.table_id, (part.row or 0) + 1, part.col))
        if below is not None and below.text.strip():
            return _trim_at_next_label(below.text)
    return None


def value_inline_in_cell(parts: List[DocPart], label: str) -> Optional[str]:
    """Значение в ТОЙ ЖЕ части, после метки и разделителя: «ИНН: 7707083893».

    Смотрит и ячейки, и абзацы: самая частая форма в заявлениях — «метка:
    значение» одной строкой, и разметка тут ничем не помогает. Разделитель
    обязателен — без него подпись «ИНН организации» превратилась бы в значение
    «организации».
    """
    pattern = _inline_re(label)
    for part in parts:
        if part.origin not in ("table", "body"):
            continue
        m = pattern.match(part.text)
        if m and m.group(1).strip():
            return _trim_at_next_label(m.group(1))
    return None


def paragraph_after_label(parts: List[DocPart], label: str) -> Optional[str]:
    """Следующий непустой АБЗАЦ после абзаца, который ЦЕЛИКОМ является меткой.

    Узкая шапка заявления печатает подпись отдельной строкой, а значение —
    следующей. По плоскому тексту это неотличимо от переноса внутри значения;
    по номеру абзаца — отличимо точно.

    Требование «абзац целиком метка» тут не придирка: без него «Адрес: 413105…»
    считался бы подписью, а значением стал бы соседний абзац (см. `_bare_label_re`).
    """
    body = [p for p in parts if p.origin == "body"]
    for i, part in enumerate(body):
        if _is_label_cell(part, label):
            for nxt in body[i + 1:]:
                if nxt.text.strip():
                    return _trim_at_next_label(nxt.text)
            return None
    return None


def structural_value(parts: List[DocPart], field: str) -> Optional[Tuple[str, str]]:
    """Значение поля по разметке: (значение, каким запросом взято) или None.

    Метки берём из реестра `label_synonyms.FIELD_LABELS` — здесь их быть не
    должно, иначе реестр снова начнёт расходиться с извлечением (§S.2.D).

    Порядок запросов — по убыванию однозначности разметки: своя ячейка справа,
    ячейка под меткой, хвост в той же ячейке, следующий абзац. Первый ответ
    и выигрывает; уровень возвращается наружу, чтобы спорные случаи были видны
    в теневом отчёте, а не молча проглатывались.
    """
    from label_synonyms import FIELD_LABELS

    labels = FIELD_LABELS.get(field) or []
    for query_name, query in (
        ("cell_right", value_right_of_label),
        ("cell_below", value_below_label),
        ("cell_inline", value_inline_in_cell),
        ("para_after", paragraph_after_label),
    ):
        for label in labels:
            value = query(parts, label)
            if value:
                return value, query_name
    return None


# --- Шапка «метки стопкой» ----------------------------------------------------
#
# Часть заявлений печатает шапку в две колонки, и при конвертации колонки
# разъезжаются: сначала идут ПОДРЯД все метки, затем ПОДРЯД все значения в том
# же порядке.
#
#     Заинтересованное лицо (должник):        <- метка 1
#     Адрес регистрации:                      <- метка 2
#     Петросян Генрих Суренович (ИНН …)       <- значение 1
#     344049, г. Ростов-на-Дону, ул. Еляна…   <- значение 2
#
# В плоском тексте это неразличимо: метка и чужое значение оказываются на одной
# строке, и построчный разбор выдаёт «Адрес регистрации: Петросян Генрих
# Суренович». В разметке DOCX границы абзацев целы — отсюда и решение.
#
# В проекте этот приём уже есть, но захардкожен под ОДНУ сигнатуру ВТБ
# (`_STACKED_SIG_RE`). Здесь то же правило, но метки берутся из реестра.

_LABEL_TAIL_RE = re.compile(r":\s*$")


@lru_cache(maxsize=1)
def _any_bare_label_re() -> Pattern[str]:
    """Часть целиком является меткой реестра (возможно с уточнением в скобках)."""
    from label_synonyms import all_labels, labels_alternation

    return re.compile(
        r"^\s*(?:" + labels_alternation(all_labels()) + r")"
        r"[^\n:]{0,30}\s*:\s*$",
        re.IGNORECASE | re.DOTALL,
    )


def _merge_split_labels(texts: List[str]) -> List[str]:
    """Склеивает метку, разорванную переносом колонки.

    «ФИНАНСОВЫЙ» + «УПРАВЛЯЮЩИЙ:» — одна метка, разложенная на две части.
    Без склейки хвост метки принимается за ЗНАЧЕНИЕ, и все пары съезжают.
    """
    out: List[str] = []
    i = 0
    while i < len(texts):
        cur = texts[i]
        if (i + 1 < len(texts) and not _LABEL_TAIL_RE.search(cur)
                and _any_bare_label_re().match(f"{cur} {texts[i + 1]}")):
            out.append(f"{cur} {texts[i + 1]}")
            i += 2
            continue
        out.append(cur)
        i += 1
    return out


def stacked_label_pairs(parts: List[DocPart]) -> List[Tuple[str, str]]:
    """Пары «метка → значение» из шапки, где метки идут стопкой.

    Возвращает пары в порядке следования. Стопкой считается серия из ДВУХ и
    более меток подряд: одна метка со значением ниже — это обычная раскладка,
    и трогать её нельзя.

    Чистая функция: один вход — один результат, побочных эффектов нет.
    """
    texts = _merge_split_labels(
        [" ".join(p.text.split()) for p in parts if p.text and p.text.strip()]
    )
    bare = _any_bare_label_re()
    pairs: List[Tuple[str, str]] = []
    i = 0
    while i < len(texts):
        run = []
        while i < len(texts) and bare.match(texts[i]):
            run.append(texts[i])
            i += 1
        if len(run) < 2:
            # Одна метка (или ни одной) — обычная раскладка. Сдвиг обязателен:
            # без него цикл не двигается и виснет.
            i += 1
            continue
        # Значения — столько же частей подряд, и ни одна из них не метка:
        # иначе стопка не кончилась и пары съедут.
        values = []
        for text in texts[i:i + len(run)]:
            if bare.match(text):
                break
            values.append(text)
        # strict=False намеренно: значений может оказаться МЕНЬШЕ, чем меток
        # (стопка оборвалась следующей меткой), и лишние метки остаются без
        # пары — додумывать за документ нельзя.
        pairs.extend(zip(run, values, strict=False))
        i += len(values)
    return pairs

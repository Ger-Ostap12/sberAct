import re
import sys
import spacy
from docx import Document
import field_contract
from requisites_validation import is_valid_inn
from fio_detector import (
    extract_debtor_name,
    is_person_name,
    _normalize_fio,
    extract_debtor_details,
    extract_third_party_details,
    extract_debtors,
    extract_third_parties,
    extract_heirs,
)
from org_normalizer import (
    base_org_name,
    norm_org_key,
    looks_like_law_ref,
    date_in_law_context,
)
from morph_utils import detect_gender, inflect_surname
from label_synonyms import all_labels, labels_alternation
from semantic_classifier import classify_procedure_family, classify_debtor_name, debtor_names_match
from classify_mixin import (
    ClassifyMixin,
    _procedure_family_from_document_type,
    _entity_type_from_document_type,
    _collateral_expected_from_document_type,
)
from parties_mixin import PartiesMixin
from amounts_mixin import AmountsMixin
from ip_mixin import IpExtractionMixin
from obligations_mixin import ObligationsMixin
from inflection_mixin import InflectionMixin
from prior_collection_mixin import PriorCollectionMixin
from patterns import build_patterns
from creditor_registry import _match_creditor_registry
import nlp_natasha as _nlp
import fns_registry as _fns_reg
from typing import Dict, Any, List, Tuple, Optional, Union
import logging
import os
from pathlib import Path
import json

# Настройка логирования (должно быть до использования logger)
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


try:
    from pymorphy3 import MorphAnalyzer
except ImportError:  # pragma: no cover - pymorphy3 может отсутствовать
    MorphAnalyzer = None

# Реестр известных банков/кредиторов: при совпадении названия подставляем ИНН, ОГРН, адрес из кода


class DocumentAnalyzer(ClassifyMixin, PartiesMixin, AmountsMixin, IpExtractionMixin, ObligationsMixin, InflectionMixin, PriorCollectionMixin):
    def __init__(self):
        """
        Инициализация анализатора документов
        """
        self.nlp = None
        self.morph = None
        self.load_nlp_model()
        self.patterns = self.load_patterns()

    def load_nlp_model(self):
        """
        Загружает модель spaCy для русского языка.
        При запуске из exe (PyInstaller) модель ищется в _MEIPASS или рядом с exe (_internal).
        """
        try:
            if getattr(sys, "frozen", False):
                # Во frozen данные модели лежат в версионной подпапке
                # (ru_core_news_sm/ru_core_news_sm-X.Y.Z), spacy.load по пути пакета
                # её не находит. Грузим через сам пакет — он знает свой data-путь.
                import ru_core_news_sm
                self.nlp = ru_core_news_sm.load()
                logger.info("Модель spaCy загружена из bundle (exe)")
            else:
                self.nlp = spacy.load("ru_core_news_sm")
                logger.info("Модель spaCy загружена успешно")
        except OSError as e:
            logger.warning("Русская модель spaCy не найдена, используем базовую модель: %s", e)
            try:
                self.nlp = spacy.load("en_core_web_sm")
            except OSError:
                logger.error("Не удалось загрузить модель spaCy")
                self.nlp = None

    def load_patterns(self) -> Dict[str, List[Dict[str, Any]]]:
        """
        Загружает паттерны для извлечения данных
        """
        return build_patterns()

    def analyze(self, file_path: str) -> Dict[str, Any]:
        """
        Анализирует документ и извлекает данные

        Args:
            file_path: Путь к файлу документа

        Returns:
            Словарь с результатами анализа
        """
        try:
            logger.info(f"Начинаем анализ документа: {file_path}")

            text = self.extract_text(file_path)
            result = self.analyze_from_text(text, page_count=self.get_page_count(file_path))
            # Таблица расчёта задолженности разбирается точнее, чем плоский текст:
            # см. _apply_table_amounts.
            if str(file_path).lower().endswith(".docx"):
                self._apply_table_amounts(file_path, result)
            return result
        except Exception as e:
            logger.error(f"Ошибка при анализе документа: {str(e)}")
            raise

    # Метки таблицы расчёта задолженности (формат Сбербанка) → поля.
    _DEBT_TABLE_LABELS = (
        ("итого задолженность", "totalDebt"),
        ("ссудная задолженность", "loanDebt"),
        ("проценты за кредит", "interest"),
        ("задолженность по неустойке", "forfeit"),
        # Госпошлина ВНУТРИ таблицы расчёта — ссудная ([17]), а не банкротная
        # ([16]): она входит в ИТОГО требований (что и проверяет сверка сумм
        # ниже). Банкротная платится отдельно за подачу заявления и в состав
        # требований не входит, её берёт разбор просительной части.
        ("госпошлина", "loanStateDuty17"),
    )

    # Денежная ячейка целиком: «2 438 262,70» (в т.ч. с неразрывным пробелом,
    # которым конвертер разделяет разряды).
    _DEBT_AMOUNT_RE = re.compile(r"[\d\s ]+[,.]\d{2}")
    # Та же сумма в конце строки — таблица, схлопнутая в «метка значение».
    # Формат жёстче, чем у отдельной ячейки: слева стоит текст метки, и без
    # границы «не цифра и не разделитель» в сумму затекала дата — из
    # «…на 01.01.2026 126 000,00» получалось «2026 126 000,00».
    _DEBT_TRAILING_AMOUNT_RE = re.compile(
        r"(?<![\d.,])(\d{1,3}(?:[\s\u00a0]\d{3})+[,.]\d{2}|\d+[,.]\d{2})\s*$"
    )

    def _debt_amounts_from_rows(self, rows: List[Tuple[str, str]]) -> Dict[str, str]:
        """Поля сумм из пар «метка → значение» таблицы расчёта задолженности.

        Строки «в т.ч. …» — подпункты, их пропускаем: именно из-за них в акты
        уезжала 7 559,11 («в т.ч. на просроченные проценты» — часть неустойки,
        а не проценты по кредиту).

        Возвращает {}, если таблица не наша или доверять ей нельзя: арифметика
        обязана сходиться (ссудная + проценты + неустойка + госпошлина = ИТОГО).
        Это страхует и от таблиц, где конвертер сдвинул ячейки, и от плоского
        текста с перемешанным порядком строк (текстовый слой PDF через pypdf).
        """
        found: Dict[str, str] = {}
        for label_raw, value_raw in rows:
            label = re.sub(r"\s+", " ", label_raw).strip().lower()
            value = value_raw.strip()
            if label.startswith("в т.ч") or not value:
                continue
            if not self._DEBT_AMOUNT_RE.fullmatch(value):
                continue
            for needle, field in self._DEBT_TABLE_LABELS:
                if needle in label and field not in found:
                    found[field] = re.sub(r"\s+", " ", value).replace(".", ",")
                    break

        if "totalDebt" not in found or len(found) < 3:
            return {}  # не наша таблица

        def _num(key: str) -> float:
            try:
                return float(found.get(key, "0").replace(" ", "").replace(" ", "").replace(",", "."))
            except ValueError:
                return 0.0

        total = _num("totalDebt")
        parts = _num("loanDebt") + _num("interest") + _num("forfeit") + _num("loanStateDuty17")
        if total <= 0 or abs(total - parts) > 0.05:
            logger.info(
                "Таблица расчёта: ИТОГО %.2f не сходится с суммой строк %.2f — суммы из "
                "таблицы не применяем", total, parts,
            )
            return {}

        # Основной долг живёт в двух полях ([13] читает оба).
        found["principalDebt"] = found.get("loanDebt", "")
        # Общая сумма требований = ИТОГО из таблицы: именно её банк просит
        # включить в реестр (п. 6 просительной части). Пересчитывать её на сумму
        # частей не надо — расхождение как раз на госпошлину, которая в ИТОГО есть.
        found["debtAmount"] = found["totalDebt"]
        found["requirementsSum"] = found["totalDebt"]
        return found

    def _write_debt_amounts(self, fields: Dict[str, Any], found: Dict[str, str], source: str) -> None:
        """Переносит выверенные суммы в поля, вытесняя добытое регулярками."""
        for key, value in found.items():
            if value:
                fields[key] = value

        # Ту же строку «Госпошлина …» общий слой регулярок уже записал в
        # stateDuty ([16], банкротная). Раз таблица опознала её как ссудную,
        # это один и тот же рубль в двух маркерах — снимаем дубль. Иное
        # значение в stateDuty не трогаем: там настоящая банкротная пошлина
        # из просительной части.
        loan_duty = found.get("loanStateDuty17")
        if loan_duty:
            for key in ("stateDuty", "stateDuty16"):
                if fields.get(key) == loan_duty:
                    fields.pop(key, None)

        logger.info("Суммы взяты из %s: %s", source, found)

    def _debt_rows_from_text(self, text: str) -> List[Tuple[str, str]]:
        """Пары «метка → значение» таблицы расчёта, восстановленные из ПЛОСКОГО текста.

        Экстрактор разворачивает таблицу по ячейкам подряд, поэтому метка и её
        сумма становятся соседними строками:

            Ссудная задолженность
            2 438 262,70

        Реже (сжатая вёрстка, текстовый слой PDF) обе части лежат в одной строке:
        «Проценты за кредит 0,00» — берём и такой вид. Пары отдаём как есть,
        доверие к ним проверяет _debt_amounts_from_rows.
        """
        rows: List[Tuple[str, str]] = []
        pending_label: Optional[str] = None
        for raw in text.splitlines():
            line = raw.strip()
            if not line:
                continue
            if self._DEBT_AMOUNT_RE.fullmatch(line):
                if pending_label:
                    rows.append((pending_label, line))
                    pending_label = None  # одна метка — одно значение
                continue
            match = self._DEBT_TRAILING_AMOUNT_RE.search(line)
            if match and line[:match.start()].strip():
                rows.append((line[:match.start()].strip(), match.group(1).strip()))
                pending_label = None
                continue
            pending_label = line
        return rows

    def _apply_text_debt_table(self, fields: Dict[str, Any], text: str) -> None:
        """Суммы таблицы расчёта, когда на входе только текст (файла нет).

        Конвертерный путь (PDF → DOCX → правка текста в предпросмотре →
        /analyze-text) файла-источника не имеет, и `_apply_table_amounts` там не
        отрабатывает — а таблица к этому моменту уже развёрнута в строки.
        Регулярки по такому тексту цепляли первую подходящую подпись:
        «в т.ч. на просроченные проценты\\n7 559,11» читалось как проценты по
        кредиту, и 7 559,11 уезжала во ВСЕ акты вместе с общей суммой долга.
        """
        found = self._debt_amounts_from_rows(self._debt_rows_from_text(text))
        if found:
            self._write_debt_amounts(fields, found, "таблицы расчёта в тексте")

    def _apply_table_amounts(self, file_path: str, result: Dict[str, Any]) -> None:
        """Суммы из НАСТОЯЩЕЙ таблицы DOCX — источник надёжнее плоского текста.

        Пары «метка → значение» здесь заданы разметкой, а не соседством строк,
        поэтому при наличии файла берём их отсюда, поверх текстового разбора.
        """
        try:
            from docx import Document as _Docx
            doc = _Docx(file_path)
        except Exception as exc:
            logger.debug(f"Таблицы DOCX недоступны: {exc}")
            return

        rows: List[Tuple[str, str]] = []
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) < 2:
                    continue
                rows.append((cells[0], cells[-1]))

        found = self._debt_amounts_from_rows(rows)
        if not found:
            return

        fields = result.get("fields")
        if not isinstance(fields, dict):
            return
        self._write_debt_amounts(fields, found, "таблицы расчёта задолженности")

    def analyze_from_text(self, text: str, page_count: Optional[int] = None) -> Dict[str, Any]:
        """
        Анализирует уже извлечённый текст заявления (без обращения к файлу).

        Отдельная точка входа для текста, отредактированного пользователем в
        предпросмотре после OCR-конвертации PDF (эндпоинт /analyze-text):
        файла-источника на этом пути нет, только текст.

        Args:
            text: Плоский текст документа (формат — как у extract_text)
            page_count: Число страниц исходника; None — неизвестно
                (используется 1, как и при ошибке подсчёта в get_page_count)

        Returns:
            Словарь с результатами анализа (идентичен analyze)
        """
        try:
            if not text:
                raise ValueError("Не удалось извлечь текст из документа")

            logger.info(f"Извлеченный текст (первые 500 символов): {text[:500]}")


            raw_text = text
            text = self._mask_attachment_list_amounts(text)
            text = self._normalize_whitespace_for_matching(text)
            text = self._normalize_label_wrap_for_matching(text)

            text_lower = text.lower()

            is_kfh_detected, kfh_head_name = self._detect_kfh_head(text)

            document_type = self.classify_document(text)

            if document_type in ("ip_collection", "ip_collection_collateral", "ip_collection_collateral_auto"):
                extracted_fields = self._analyze_ip_collection(text, document_type)
            elif document_type in ("legal_collection", "legal_collection_collateral", "legal_collection_collateral_auto"):
                extracted_fields = self._analyze_legal_collection(text, document_type)
            elif document_type in ["ip_enforcement_statement", "ip_enforcement_statement_collateral",
                                 "ip_enforcement_realization", "ip_enforcement_realization_collateral",
                                 "ip_enforcement_restructuring", "ip_enforcement_restructuring_collateral"]:
                extracted_fields = self._analyze_ip_enforcement(text, document_type)
            elif document_type in ["physical_realization_collateral", "physical_restructuring_collateral", "observation_collateral", "competition_collateral"]:
                extracted_fields = self._analyze_physical_collateral(text, document_type)
            else:
                # Универсальный fallback: если для типа документа нет собственных паттернов,
                # используем rtk_application как базовый шаблон, чтобы всё равно извлекать поля.
                effective_type = document_type
                if document_type not in self.patterns:
                    logger.info(
                        f"Тип документа '{document_type}' не имеет собственных паттернов, "
                        f"используем 'rtk_application' как базовый тип для извлечения полей"
                    )
                    effective_type = "rtk_application"

                extracted_fields = self.extract_fields(text, effective_type)
                collateral_detected = False

            self._contract_checkpoint(extracted_fields, "извлечение по паттернам")

            document_type = self._maybe_upgrade_to_observation_collateral(extracted_fields, text, document_type, text_lower)

            if is_kfh_detected:
                extracted_fields["isKfh"] = True
                if kfh_head_name:
                    extracted_fields["kfhHeadName"] = kfh_head_name
                logger.info(f" Установлен флаг isKfh=True, глава КФХ: {kfh_head_name}")

            if document_type:
                extracted_fields["sourceDocumentType"] = document_type

            self._fill_creditor_requisites(extracted_fields, text)

            confidence = self.calculate_confidence(document_type, extracted_fields, text)

            self._postprocess_mortgage_and_case_number(extracted_fields, text, document_type)

            recommended_acts = self._get_recommended_acts(document_type, extracted_fields, text)

            collateral_description = extracted_fields.get("mortgageCollateralDescription1221")
            if collateral_description and re.match(r"^Кому\s+выдана\s+", (collateral_description or "").strip(), re.IGNORECASE):
                collateral_description = None
                if "mortgageCollateralDescription1221" in extracted_fields:
                    del extracted_fields["mortgageCollateralDescription1221"]
            # Все предметы залога по документу (недвижимость/авто, каждый карточкой);
            # дубли одного объекта схлопываем, оставляя самый заполненный.
            collateral_descs = self._extract_all_collateral_items(text)
            if not collateral_descs and collateral_description:
                collateral_descs = self._split_collateral_items(collateral_description)
            collaterals_list = self._dedupe_collaterals(
                [self._make_collateral_obj(idx, d) for idx, d in enumerate(collateral_descs)]
            )
            # Отбрасываем нереальные «иное» (шаблонные «Предметом залога является …»).
            collaterals_list = [c for c in collaterals_list if self._collateral_has_substance(c)]
            if collaterals_list:
                logger.info(f" Создано {len(collaterals_list)} предметов залога: {[c['collateralType'] for c in collaterals_list]}")


            self._fix_debtor_name(text, extracted_fields)

            details = self._finalize_debtor_person_requisites(extracted_fields, text)

            self._cleanup_extracted_fields(extracted_fields, text)

            # Своп сторон: applicantName ошибочно = кредитор (раскладки «Заявитель:» «Должник:»)
            self._reconcile_applicant_is_debtor(extracted_fields, text)

            # Косметика артефактов сторон: роль-суффикс «(заёмщик)», хвост метки в courtName, мусорный managerName
            self._cleanup_party_artifacts(extracted_fields, text)


            self._apply_stacked_party_layout(extracted_fields, text)


            if (not extracted_fields.get("creditorInn")
                    and not extracted_fields.get("creditorOgrn")):
                self._fill_creditor_requisites(extracted_fields, text)
            # Реквизиты из документа уже есть, но адрес кредитора пуст — дозаполняем
            # (адрес из блока кредитора или из реестра известных банков).
            elif not extracted_fields.get("creditorAddress"):
                self._fill_creditor_requisites(extracted_fields, text)


            self._apply_fns_authority(extracted_fields, text)


            self._fill_creditor_nlp(extracted_fields, text)


            is_self_bk = self._detect_self_bankruptcy(text)
            sb_third_parties = None
            if is_self_bk:
                sb_third_parties = self._apply_self_bankruptcy_layout(extracted_fields, text)
                # Тип лица мог поменяться (legal individual) — рекомендации заново.
                recommended_acts = self._get_recommended_acts(document_type, extracted_fields, text)

            if is_self_bk:
                # Должник один и уже выверен layout-парсером — генерик-извлечение
                # extract_debtors по такой шапке тащит кредиторов, минуем его.
                debtors_result = [self._single_debtor_from_fields(extracted_fields)]
                third_parties_result = sb_third_parties or []
            else:
                debtors_result, third_parties_result = self._resolve_debtors_and_third_parties(extracted_fields, text, details)


            _raw_cols = collaterals_list if collaterals_list else extracted_fields.get('collaterals', [])
            collaterals_final = [c for c in _raw_cols if self._collateral_has_substance(c)]
            if not collaterals_final:
                extracted_fields.pop("mortgageCollateralDescription1221", None)
                extracted_fields.pop("collaterals", None)


            self._normalize_financial_block(extracted_fields, text)


            self._apply_prayer_finances(extracted_fields, text)


            self._apply_stacked_finances(extracted_fields, text)


            self._apply_table_breakdown_finances(extracted_fields, text)

            # Таблица расчёта задолженности, развёрнутая в строки текста («метка»,
            # следом «значение»). Последней в каскаде: пары из таблицы точнее
            # любой регулярки по плоскому тексту, а гейт на сходимость арифметики
            # не даёт ей сработать на чужом документе.
            self._apply_text_debt_table(extracted_fields, text)

            # Самобанкротство: свой источник итога (сводная фраза «общий объём
            # задолженности составляет …»). Последним — ни один слой выше такой
            # раскладки не знает, а его итог перекрывает сумму первого кредитора.
            if is_self_bk:
                self._apply_self_bankruptcy_total(extracted_fields, text)

            self._contract_checkpoint(extracted_fields, "финансовый каскад")


            if not is_self_bk:
                self._apply_fns_queue_finances(extracted_fields, text)

            # Банкротная госпошлина двумя слагаемыми под одной меткой («… 1 490 913
            # руб.+ 100 000 руб.») — суммируем в банкротную, убираем ложную ссудную.
            if not is_self_bk:
                self._sum_bankruptcy_duty(extracted_fields, text)


            self._fix_truncated_creditor_name(extracted_fields, text)


            self._clean_creditor_artifacts(extracted_fields, text)

            # Адрес управляющего — фолбэк для многострочного «Адрес регистрации:»,
            # когда ФИО и адрес на разных строках (общие паттерны не справляются).
            self._fill_manager_address(extracted_fields, text)

            # Саморегулируемая организация: сопоставляем упоминание СРО из текста с
            # реестром sro_data и подставляем каноничное полное имя (NAIM_FULL).
            self._resolve_manager_sro(extracted_fields, text)

            # Вспомогательный NLP-слой (Natasha, второстепенно): когда regex не
            # извлёк адрес совсем или обрезал его — достраиваем по окну роли.
            self._refine_addresses_nlp(extracted_fields, text)

            # Адрес не может содержать реквизиты (ИНН/ОГРН/ОГРНИП/СНИЛС/КПП) —
            # обрезаем хвост, а полностью мусорные адреса (без букв) убираем.
            self._sanitize_address_fields(extracted_fields)

            # Адреса в записях должников/третьих лиц строятся отдельным путём (не через
            # fields) — чистим их той же обрезкой склейки (хвост «В лице ликвидатора: …
            # Сообщение №… о намерении обратиться в суд» при однострочной PDF docx склейке).
            for _entry_list in (debtors_result, third_parties_result):
                for _entry in (_entry_list or []):
                    _ea = _entry.get("address")
                    if isinstance(_ea, str) and _ea.strip():
                        _clean = re.split(r"\b(?:ИНН|ОГРНИП|ОГРН|СНИЛС|КПП)\b", _ea,
                                          flags=re.IGNORECASE)[0].strip(" ,;-")
                        _entry["address"] = self._truncate_glued_address(_clean)

            # Банкротная госпошлина не должна совпадать с итогом/осн.долгом —
            # это мусор (в документе отдельной банкротной госпошлины нет). Чистим.
            self._clear_garbage_bankruptcy_duty(extracted_fields)

            # Ипотека: госпошлина из шапки, если финансовый каскад её не заполнил.
            if document_type == "mortgage_claim":
                self._fill_mortgage_state_duty(extracted_fields, text)

            # Ранее вынесенное решение другого суда (взыскание до банкротства).
            prior_decision = self._extract_prior_court_decision(text)
            if prior_decision:
                for k, v in prior_decision.items():
                    if v:
                        extracted_fields[k] = v
                # Если основной номер дела совпал с «ранее вынесенным» — это его
                # контекст («…по делу №… взыскана…»), а не дело текущего заявления.
                pc = prior_decision.get("priorCaseNumber")
                if pc and extracted_fields.get("caseNumber") == pc:
                    extracted_fields.pop("caseNumber", None)


            if not any(k.startswith("fnsQ") for k in extracted_fields):
                _brk = extracted_fields.get("financeBreakdown") or {}
                for _fk in ("principalDebt", "interest", "forfeit", "penalties", "loanStateDuty17"):
                    _v = extracted_fields.get(_fk)
                    if _fk == "principalDebt" and not _v:
                        _v = extracted_fields.get("loanDebt")  # поле фронта: principalDebt OR loanDebt
                    if _fk not in _brk and _v and self._fin_amount(_v) > 0:
                        _brk[_fk] = [_v]
                if _brk:
                    extracted_fields["financeBreakdown"] = _brk


            finance_breakdown = extracted_fields.pop("financeBreakdown", None)


            application_kind = "self_bankruptcy" if is_self_bk else None

            document_type_warning = None
            if not is_self_bk:
                expected_family = _procedure_family_from_document_type(document_type)
                if expected_family is not None:
                    semantic_family, _clause_details = classify_procedure_family(
                        raw_text, extracted_fields.get("debtorName", "")
                    )
                    if semantic_family is not None and semantic_family != expected_family:
                        document_type_warning = {
                            "documentType": document_type,
                            "regexFamily": expected_family,
                            "semanticFamily": semantic_family,
                            "message": (
                                "Автоматическое определение типа документа не подтверждено "
                                "вторым способом проверки — рекомендуем перепроверить тип "
                                "заявления вручную."
                            ),
                        }

            # Сверка имени должника вторым способом (NER-харвест + ролевой якорь) —
            # shadow-баннер по образцу document_type_warning: только предупреждение,
            # `debtorName` не трогаем (источник истины — regex). Golden-безопасно
            # (поле не в `_TOP_FIELDS`). Не зависит от эмбеддинг-модели.
            debtor_name_warning = None
            _regex_debtor = extracted_fields.get("debtorName", "") or ""
            _sem_debtor, _sem_meta = classify_debtor_name(raw_text)
            if _sem_debtor and not debtor_names_match(_regex_debtor, _sem_debtor):
                debtor_name_warning = {
                    "regexName": _regex_debtor,
                    "semanticName": _sem_debtor,
                    "message": (
                        "Имя должника, найденное вторым способом проверки, отличается "
                        "от определённого автоматически — рекомендуем перепроверить "
                        "ФИО/наименование должника вручную."
                    ),
                }

            _mgr_cur = (extracted_fields.get("managerName") or "").split("(")[0].strip()
            if not is_person_name(_mgr_cur):
                _mgr_cand = self._extract_manager_candidate(text)
                if _mgr_cand:
                    extracted_fields["managerName"] = _mgr_cand


            debtor_status_hint = None
            heirs_result = []
            if not is_self_bk and extracted_fields.get("entityType") == "legal":
                if self._detect_absent_debtor(text):
                    debtor_status_hint = "absent"
                elif self._detect_liquidation(text):
                    debtor_status_hint = "liquidation"
                    _liq = self._extract_liquidator(text)
                    if _liq:
                        extracted_fields["liquidatorName"] = _liq
            elif not is_self_bk and extracted_fields.get("entityType") == "individual":
                if self._detect_deceased(text):
                    debtor_status_hint = "deceased"
                    # Сведения о смерти извлекаем ТОЛЬКО внутри этой ветки: якорь
                    # «нотариус»/«умер» вне контекста смерти должника — чужие факты.
                    extracted_fields.update(self._extract_death_details(text))
                    heirs_result = extract_heirs(text)
            provenance = extracted_fields.pop(self._PROVENANCE_KEY, {}) or {}
            contract_issues = field_contract.apply_contract(extracted_fields)

            for _entries, _label in ((debtors_result, "debtors"),
                                     (third_parties_result, "thirdParties"),
                                     (heirs_result, "heirs")):
                contract_issues += field_contract.check_entries(_entries, _label)
            field_issues = [i.as_dict() for i in contract_issues]

            entity_type_warning = None
            _expected_entity = _entity_type_from_document_type(document_type)
            _actual_entity = extracted_fields.get("entityType")
            if (_expected_entity is not None and _actual_entity
                    and _actual_entity != "kfh" and _actual_entity != _expected_entity):
                entity_type_warning = {
                    "documentType": document_type,
                    "expectedEntityType": _expected_entity,
                    "actualEntityType": _actual_entity,
                    "message": (
                        "Тип документа предполагает тип лица должника, отличный от "
                        "определённого по реквизитам — рекомендуем перепроверить "
                        "вручную."
                    ),
                }

            collateral_warning = None
            _expected_collateral = _collateral_expected_from_document_type(document_type)
            _actual_collateral = bool(collaterals_final)
            if _expected_collateral is not None and _expected_collateral != _actual_collateral:
                collateral_warning = {
                    "documentType": document_type,
                    "expectedCollateral": _expected_collateral,
                    "actualCollateral": _actual_collateral,
                    "message": (
                        "Тип документа предполагает "
                        + ("наличие" if _expected_collateral else "отсутствие")
                        + " залога, но по извлечённым данным "
                        + ("залог не найден" if _expected_collateral else "залог обнаружен")
                        + " — рекомендуем перепроверить вручную."
                    ),
                }

            recommended_acts = self._get_recommended_acts(document_type, extracted_fields, text)

            # Ипотека: структурированные предметы залога для формы (стоимость/НПЦ/
            # отчёт/ЕГРН реконсилируются из абзаца оценки и записей ЕГРН). Top-level,
            # вне fields и collaterals — golden-снимок его не фиксирует.
            # Вид ипотеки (военная/ДДУ/гражданская) — top-level, фронт инициализирует
            # переключатель (пользователь может сменить). Считаем ДО предмета: разбор
            # предмета для ДДУ иной (права требования по договору долевого участия).
            mortgage_kind = (
                self._detect_mortgage_kind(text)
                if document_type == "mortgage_claim" else None
            )
            mortgage_properties = (
                self._build_mortgage_properties(text, collaterals_final, mortgage_kind)
                if document_type == "mortgage_claim" else []
            )

            # [1221] описание предмета — чистое (из структурного mortgageProperties),
            # а НЕ жадный блоб со всей правовой «водой» секции залога (ст.334 ГК, чужие
            # ФИО из шаблонного текста заявления). Блоб уже отработал на разбиении
            # collaterals выше, дальше в акт должно идти короткое описание объекта.
            # Адреса предметов ипотеки контракт до сих пор не смотрел вовсе:
            # check_entries звался только для debtors/thirdParties/heirs, а
            # mortgageProperties строятся ПОЗЖЕ этого цикла. В итоге в адрес
            # объекта могла доехать любая строка. Правило то же самое, что и
            # для остальных адресов, — отдельного изобретать не нужно.
            if mortgage_properties:
                field_issues += [
                    i.as_dict() for i in field_contract.check_entries(
                        mortgage_properties, "mortgageProperties")
                ]

            if document_type == "mortgage_claim" and mortgage_properties:
                _descs = [str(p.get("description") or "").strip() for p in mortgage_properties]
                _descs = [d for d in _descs if d]
                if _descs:
                    extracted_fields["mortgageCollateralDescription1221"] = "; ".join(_descs)

            # Формируем результат
            result = {
                "documentType": document_type,
                "confidence": confidence,
                "fields": extracted_fields,
                "obligations": extracted_fields.get('obligations', []),
                "collaterals": collaterals_final,
                "mortgageProperties": mortgage_properties,
                "mortgageKind": mortgage_kind,
                "financeBreakdown": finance_breakdown,
                "applicationKind": application_kind,
                "debtorStatusHint": debtor_status_hint,

                "documentTypeWarning": document_type_warning,
                "debtorNameWarning": debtor_name_warning,
                "entityTypeWarning": entity_type_warning,
                "collateralWarning": collateral_warning,
                "fieldIssues": field_issues,
                "fieldQuality": field_contract.assess_quality(
                    extracted_fields, contract_issues, provenance
                ),
                # ИСХОДНЫЙ текст, не маскированный: на нём держатся посекционный
                # предпросмотр и построчная вставка правок в docx (§A).
                "rawText": raw_text,
                "metadata": {
                    "pageCount": page_count if page_count is not None else 1,
                    "wordCount": len(raw_text.split()),
                    "language": "ru"
                },
                "recommendedActs": recommended_acts,
                "debtors": debtors_result,
                "thirdParties": third_parties_result,
                # Наследники умершего должника — массив, как thirdParties: их может
                # быть несколько, и у каждого свои реквизиты.
                "heirs": heirs_result
            }

            entity_type = extracted_fields.get("entityType")
            if entity_type:
                result["entityType"] = entity_type

            logger.info(f"Анализ завершен. Тип: {document_type}, уверенность: {confidence:.2f}")
            logger.info(f"Рекомендуемые акты: {recommended_acts}")
            return result

        except Exception as e:
            logger.error(f"Ошибка при анализе документа: {str(e)}")
            raise

    # Префиксы перед ФИО, которые убираем перед склонением.
    _NAME_PREFIX_RE = r'^(ИП\s+|ГЛАВА\s+КФХ\s+ИП\s+|ГЛАВА\s+КФХ\s+|КФХ\s+ИП\s+)'

    # Типы судов РФ (основы прилагательных перед словом «суд»).
    _COURT_KIND_RE = (
        r"(?:городск|районн|межрайонн|областн|краев|окружн|верховн|гарнизонн|военн|"
        r"арбитражн|конституционн|уставн|апелляционн|кассационн|третейск|мирск)\w*"
    )

    # Общие слова в названии суда, которые пишутся со строчной буквы.
    _COURT_LOWER_WORDS = {
        "суд", "суда", "суде", "суду", "судом", "арбитражный", "районный", "городской",
        "межрайонный", "областной", "краевой", "окружной", "верховный", "военный",
        "гарнизонный", "апелляционный", "кассационный", "конституционный", "уставный",
        "мировой", "мировому", "судье", "судьи", "судебного", "судебный", "участка",
        "участок", "области", "область", "край", "края", "краю", "республики",
        "республика", "республике", "округа", "округ", "округе", "города", "город",
        "автономного", "автономной", "общей", "юрисдикции", "района", "район", "и", "по",
    }


    def _clean_court_line(self, s: str) -> Optional[str]:
        """Очищает строку с названием суда: убирает приставку «В …» и хвост-адрес."""
        s = re.sub(r"\s+", " ", s).strip()
        s = re.sub(r"^(?:В|Во)\s+(?=[А-ЯЁ])", "", s)          # адресная приставка «В …»
        s = re.sub(r"^(?:от\s+истца|истец|заявитель)[\s:,-]*", "", s, flags=re.IGNORECASE)
        s = re.split(r",?\s*\d{5,6}\b", s)[0]                  # обрезаем по почтовому индексу
        s = re.split(r"\s+(?:от\s+истца|от\s+заявител)", s, flags=re.IGNORECASE)[0]
        # обрезаем адресный хвост без индекса («…, ул. Ленина», «…, г. Москва»)
        s = re.split(r",\s*(?:ул\.|улиц|г\.|город|пр-?кт|проспект|пер\.|переул|пл\.|площад|наб\.|бул|шоссе|д\.\s*\d)",
                     s, flags=re.IGNORECASE)[0]
        s = s.strip(" ,.;")
        if 6 <= len(s) <= 140 and re.search(r"(?:суд|участ|судь)", s, re.IGNORECASE):
            return s
        return None






    def extract_text(self, file_path: str) -> str:
        """
        Извлекает текст из документа. Поддерживаются форматы: .docx (Word), .pdf.
        """
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix == ".pdf":
            return self._extract_text_from_pdf(file_path)
        if suffix in (".docx", ".doc"):
            return self._extract_text_from_docx(file_path)
        raise ValueError(f"Неподдерживаемый формат файла: {suffix}. Используйте .docx или .pdf")

    def _extract_text_from_pdf(self, file_path: str) -> str:
        """Извлекает текст из PDF."""
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            text_parts = []
            for page in reader.pages:
                part = page.extract_text()
                if part and part.strip():
                    text_parts.append(part.strip())
            if not text_parts:
                raise ValueError("PDF не содержит извлекаемого текста (возможно, скан или пустой файл)")
            return "\n".join(text_parts)
        except ValueError:
            raise
        except Exception as e:
            logger.error(f"Ошибка при извлечении текста из PDF: {str(e)}")
            raise ValueError("Не удалось извлечь текст из PDF (возможно, файл поврежден или скан без OCR)") from e

    def _extract_text_from_docx(self, file_path: str) -> str:
        """Извлекает ВЕСЬ текст из Word документа (.docx).

        Тонкая обёртка над :meth:`_docx_text_parts` — тот же обход, только
        соединяет части в плоский текст. Порядок и дедуп — единый источник
        истины для анализа, предпросмотра и baseline «Скачать с правками».
        """
        try:
            parts = self._docx_text_parts(file_path)
            if not parts:
                raise ValueError("Документ пуст или не содержит извлекаемого текста")
            return "\n".join(text for text, _origin in parts)
        except ValueError:
            raise
        except Exception as e:
            logger.error(f"Ошибка при извлечении текста из Word: {str(e)}")
            raise ValueError("Не удалось извлечь текст из документа (возможно, файл поврежден или пустой)") from e

    def _docx_text_parts(self, file_path: str) -> List[Tuple[str, str]]:
        """Части текста DOCX с указанием происхождения (провенанс) каждой части.

        Возвращает список ``(text, origin)`` В ТОМ ЖЕ ПОРЯДКЕ, что и старый
        плоский экстрактор — так `"\\n".join(text ...)` байт-в-байт совпадает с
        `extract_text` (гарант инварианта для секций и golden). ``origin`` ∈
        ``{"body", "table", "colophon", "raw"}``:
          1. тело: параграфы (``body``), затем таблицы (``table``);
          2. колонтитулы всех секций (``colophon``) — с дедупликацией;
          3. надписи (w:txbxContent) и сноски/концевые сноски (``raw``).
        Провенанс нужен посекционному предпросмотру (`extract_sections`), плоский
        текст его игнорирует.
        """
        doc = Document(file_path)
        parts: List[Tuple[str, str]] = []

        # 1. Тело документа: параграфы, затем таблицы
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append((paragraph.text.strip(), "body"))

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append((cell.text.strip(), "table"))

        # 2. Колонтитулы всех секций: header/footer + first_page + even_page.
        # Секции часто ссылаются на один и тот же колонтитул — дедуплицируем.
        seen_headers = set()
        for section in doc.sections:
            for container in (
                section.header, section.footer,
                section.first_page_header, section.first_page_footer,
                section.even_page_header, section.even_page_footer,
            ):
                if container is None:
                    continue
                try:
                    chunk_parts = []
                    for paragraph in container.paragraphs:
                        if paragraph.text.strip():
                            chunk_parts.append(paragraph.text.strip())
                    for table in container.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                if cell.text.strip():
                                    chunk_parts.append(cell.text.strip())
                    chunk = "\n".join(chunk_parts)
                    if chunk and chunk not in seen_headers:
                        seen_headers.add(chunk)
                        parts.append((chunk, "colophon"))
                except Exception as exc:
                    logger.warning(f"Не удалось обработать колонтитул: {exc}")

        # 3 + 4. Надписи (w:txbxContent) и сноски/концевые сноски —
        # части DOCX, недоступные через объектную модель python-docx.
        for chunk in self._extract_docx_raw_xml_text(file_path):
            parts.append((chunk, "raw"))

        return parts

    def _extract_docx_raw_xml_text(self, file_path: str) -> List[str]:
        """Собирает текст из частей DOCX, не покрытых моделью python-docx.

        Открывает .docx как zip и через lxml извлекает:
          - надписи / текстовые поля (w:txbxContent) из document.xml и всех
            header*/footer* — модель python-docx их не отдаёт;
          - сноски и концевые сноски (footnotes.xml / endnotes.xml).
        Полностью оффлайн (zipfile + lxml, lxml уже идёт зависимостью python-docx).
        """
        from zipfile import ZipFile
        from lxml import etree

        W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        ns = {"w": W}
        results: List[str] = []
        seen = set()

        def paragraph_lines(node) -> List[str]:
            """Текст узла по абзацам: каждый w:p одна строка (склейка w:t)."""
            lines = []
            for p in node.findall(".//w:p", ns):
                line = "".join(t.text or "" for t in p.findall(".//w:t", ns)).strip()
                if line:
                    lines.append(line)
            return lines

        def add(chunk: str) -> None:
            if chunk and chunk not in seen:
                seen.add(chunk)
                results.append(chunk)

        try:
            with ZipFile(file_path) as zf:
                names = set(zf.namelist())

                # 3. Надписи: document.xml + все header*/footer*.
                # Берём ТОЛЬКО w:txbxContent, иначе продублируем тело документа.
                txbx_parts = [
                    n for n in names
                    if n == "word/document.xml"
                    or (n.startswith(("word/header", "word/footer")) and n.endswith(".xml"))
                ]
                for part in txbx_parts:
                    root = etree.fromstring(zf.read(part))
                    for tb in root.findall(".//w:txbxContent", ns):
                        add("\n".join(paragraph_lines(tb)))

                # 4. Сноски и концевые сноски — целиком.
                for part in ("word/footnotes.xml", "word/endnotes.xml"):
                    if part not in names:
                        continue
                    root = etree.fromstring(zf.read(part))
                    for line in paragraph_lines(root):
                        add(line)
        except Exception as exc:
            logger.warning(f"Не удалось извлечь сырой XML-текст из DOCX: {exc}")

        return results


    _TITLE_START_RE = re.compile(
        r"^\s*(?:исковое\s+)?(?:заявлени[ея]|ходатайство)\b",
        re.IGNORECASE,
    )
    _TITLE_CAPS_RE = re.compile(r"\b(?:ЗАЯВЛЕНИЕ|ХОДАТАЙСТВО)\b")

    def _is_title_line(self, line: str) -> bool:
        return bool(self._TITLE_START_RE.match(line) or self._TITLE_CAPS_RE.search(line))
    _PRAYER_VERB = r"(?:прошу|просим|просит|просят|ходатайству(?:ю|ем|ет))(?![а-яёА-ЯЁ])"

    _PRAYER_COLON_RE = re.compile(_PRAYER_VERB + r"(?:\s+суд)?\s*:", re.IGNORECASE)
    _PRAYER_START_RE = re.compile(r"^\s*(?:\d+[.)]\s*)?" + _PRAYER_VERB, re.IGNORECASE)
    # Якорь блока приложений — строка, начинающаяся с «Приложение(я/й)».
    _ATTACH_ANCHOR_RE = re.compile(r"^\s*приложени[еяй]", re.IGNORECASE)

    _ATTACH_ITEM_RE = re.compile(r"^\s*\d+[.)]\s+\S")


    _HSPACE_CHARS = "          ﻿"
    _HSPACE_RUN_RE = re.compile(r"[ \t" + _HSPACE_CHARS + r"]{2,}")
    _HSPACE_ONE_RE = re.compile(r"[" + _HSPACE_CHARS + r"]")


    # Строка — ЦЕЛИКОМ метка («Должник:»), значит её значение на следующей строке.
    # Метки берём из РЕЕСТРА (`label_synonyms.all_labels`) — он и заведён для того,
    # чтобы поддержка нового банка была правкой данных, а не логики (CLAUDE.md).
    #
    # Границы правила подобраны ЗАМЕРОМ по корпусу; оба ослабления дают регрессию
    # (оба проверены, дифф эталона был на 2 файлах с мусором вместо имён):
    #   1. «строка ЗАКАНЧИВАЕТСЯ меткой» (чтобы ловить «…, адрес регистрации:» в
    #      прозе) — в прозе двоеточие после слова-метки не значит «дальше значение»:
    #      склеивались куски шапки, applicantName становился «рбитражный суд
    #      Ростовско», кредитор терялся;
    #   2. IGNORECASE — начинала матчиться строка «ФИНАНСОВЫЙ УПРАВЛЯЮЩИЙ:» и
    #      съедала следующую строку, разваливая разбор двух заявлений ВТБ.
    # Поэтому: только строка-метка и только в том регистре, в котором метка
    # записана в реестре.
    _LABEL_ONLY_LINE_RE = re.compile(
        r"^[ \t]*(?:" + labels_alternation(all_labels()) + r")[ \t]*:[ \t]*$"
    )
    # Начало строки — метка (с двоеточием). Нужно как ГАРД: значение не может
    # начинаться с чужой метки.
    _LABEL_STARTS_LINE_RE = re.compile(
        r"^[ \t]*(?:" + labels_alternation(all_labels()) + r")[ \t]*:"
    )

    def _normalize_label_wrap_for_matching(self, text: str) -> str:
        """Поднять значение на строку его метки: «Должник:\\nИванов» «Должник: Иванов».

        Один и тот же блок шапки банки печатают и в строку, и с переносом; при
        PDF docx перенос появляется ещё и сам, когда шапка двухколоночная. Замер
        (`measure_format_robustness.py`): вставка переноса после метки меняла
        результат у 8 документов из 74 — ехали `applicantName` с падежами,
        `debtorName`, адреса.

         Направление канонизации выбрано ЗАМЕРОМ, а не рассуждением. Обратный
        вариант («всегда перенос после метки») давал ту же устойчивость, но менял
        эталон у 6 файлов: часть паттернов требует значение на строке метки.

         ГАРДЫ обязательны: без них «СНИЛС:» (метка БЕЗ значения) склеивалась со
        следующей строкой — заголовком «ЗАЯВЛЕНИЕ» — и метка получала выдуманное
        значение. Значением не может быть титул документа и не может быть чужая
        метка; пустая строка означает, что значения нет вовсе.
        """
        lines = text.split("\n")
        result: List[str] = []
        i = 0
        while i < len(lines):
            line = lines[i]
            nxt = lines[i + 1] if i + 1 < len(lines) else None
            if (
                nxt is not None
                and self._LABEL_ONLY_LINE_RE.match(line)
                and nxt.strip()
                and not self._is_title_line(nxt)
                and not self._LABEL_STARTS_LINE_RE.match(nxt)
            ):
                result.append(line.rstrip() + " " + nxt.strip())
                i += 2
                continue
            result.append(line)
            i += 1
        return "\n".join(result)

    def _normalize_whitespace_for_matching(self, text: str) -> str:
        """Привести горизонтальные пробелы к канону ДО сопоставления.

        Зачем. Разбор не должен зависеть от типографики: то же самое заявление,
        набранное в другом банке или прогнанное другим конвертером, отличается от
        нашего корпуса в первую очередь пробелами, а не смыслом. Замер
        (`tests/measure_format_robustness.py`) до этой правки: двойные пробелы
        меняли результат у **42 документов из 79**, причём у 28 из них менялся
        `documentType` — то есть ветка извлечения и рекомендованные акты; ещё 5
        документов ломал неразрывный пробел в суммах.

        Что делаем: любые горизонтальные пробелы (вкл. неразрывные и тонкие)
        обычный, серии один. ПЕРЕВОДЫ СТРОК НЕ ТРОГАЕМ: на структуре строк
        держатся label-anchored слои («Должник:» + следующая строка) и разбор
        таблиц. `raw_text` остаётся исходным — предпросмотр и построчная вставка
        правок в docx (§A) работают с ним.
        """
        normalized = self._HSPACE_ONE_RE.sub(" ", text)
        return self._HSPACE_RUN_RE.sub(" ", normalized)

    def _mask_attachment_list_amounts(self, text: str) -> str:
        """Убрать из сопоставления суммы, стоящие в ПЕРЕЧНЕ приложений.

        Перечень приложений — список ДОКУМЕНТОВ, а не денег должника (правило
        Андрея). «1. Квитанция о внесении денежных средств на депозитный счет АС в
        размере 25000» — это депозит на вознаграждение управляющего, и он утекал в
        `totalDebt`, а оттуда в акт как «основной долг в размере 25 000 руб.»
        (`Самобанкрот/Заявление РФЛ1.docx`).

         Режем ТОЛЬКО нумерованные пункты перечня. Блок приложений целиком трогать
        НЕЛЬЗЯ: у 7 документов корпуса за словом «Приложение» идёт приложенный расчёт
        задолженности («Просроченная ссудная задолженность: 8020.73»), и это
        единственный источник их финансов — слепая обрезка обнулила бы им суммы.

        Маскируем цифры пробелами, ДЛИНУ СОХРАНЯЕМ: смещения в тексте остаются
        валидными для остальных слоёв, а имена документов в перечне никому не нужны.
        """
        lines = text.split("\n")
        # Берём ПОСЛЕДНИЙ якорь, а не первый: слово «Приложение» встречается и в
        # ссылках по тексту («приложение № 1 к договору»), а перечень приложений —
        # в конце заявления. С первым якорем маска съедала нумерованные пункты
        # ПРОСИТЕЛЬНОЙ части («4. Включить требования … в размере 501 365 000,9 руб.»
        # в Заявл_Ликвидир_Оргтехника.docx) — то есть ровно те суммы, ради которых
        # всё и затевалось.
        anchors = [i for i, line in enumerate(lines) if self._ATTACH_ANCHOR_RE.match(line)]
        if not anchors:
            return text
        changed = False
        for i in range(anchors[-1], len(lines)):
            line = lines[i]
            if not self._ATTACH_ITEM_RE.match(line):
                continue
            if re.search(r"\d[\d\s  ]*[.,]?\d*\s*(?:руб|₽)|в\s+размере\s+\d", line, re.IGNORECASE):
                lines[i] = re.sub(r"\d", " ", line)
                changed = True
        return "\n".join(lines) if changed else text

    def _docx_parts_true_order(self, file_path: str) -> List[Tuple[str, str, int]]:
        """Части DOCX в РЕАЛЬНОМ порядке чтения документа + их ПЛОСКИЙ индекс.

        Возвращает `(text, origin, flat_index)`, где `flat_index` — позиция части
        в `extract_text` (плоский порядок «абзацы таблицы колонтитулы»).
        Порядок же самих элементов — истинный (таблицы стоят там, где они в теле,
        а не свалены в конец). Нужно секциям: в ФНС-сканах шапка (суд/ФНС) лежит
        в таблице В НАЧАЛЕ документа — по плоскому порядку она уезжала в конец и
        попадала в «Приложения». Классификация идёт по истинному порядку, а
        `flat_index` сохраняет инвариант пересборки (== extract_text).
        """
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        flat = self._docx_text_parts(file_path)
        pn = sum(1 for _text, origin in flat if origin == "body")   # абзацев тела
        cn = sum(1 for _text, origin in flat if origin == "table")  # ячеек таблиц

        doc = Document(file_path)
        result: List[Tuple[str, str, int]] = []
        p_i = 0
        c_i = 0
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                para = Paragraph(child, doc)
                if para.text.strip():
                    result.append((para.text.strip(), "body", p_i))
                    p_i += 1
            elif child.tag == qn("w:tbl"):
                table = Table(child, doc)
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            result.append((cell.text.strip(), "table", pn + c_i))
                            c_i += 1
        # Хвост плоского текста (колонтитулы/надписи/сноски) — с их плоскими
        # индексами; в теле их нет, порядок для классификации не важен.
        for k, (text, origin) in enumerate(flat[pn + cn:]):
            result.append((text, origin, pn + cn + k))
        return result

    def extract_sections(self, file_path: str) -> List[Dict[str, Any]]:
        """Тот же текст, что `extract_text`, разбитый на ЧЕТЫРЕ блока — ТОЛЬКО
        для предпросмотра/правки (`/docx-text`). Анализ и baseline «Скачать с
        правками» продолжают работать с плоским `extract_text`.

        Классификация — ПОСТРОЧНАЯ и по РЕАЛЬНОМУ порядку документа: части
        (`_docx_parts_true_order`) разворачиваются в строки по «\\n», каждая
        строка несёт ПЛОСКИЙ line-index (её позицию в `extract_text.split("\\n")`).
        Инвариант: соединение строк ВСЕХ блоков, отсортированных по `index`, через
        «\\n» байт-в-байт равно `extract_text(file)` (нулевой риск для golden).

        Строчная гранулярность нужна, т.к. в реальных заявлениях преамбула и
        маркер идут ОДНИМ абзацем («Руководствуясь ст…, ПРОШУ: …») — по абзацу
        преамбула уезжала в просительную. По строкам «ПРОШУ:» — своя строка.

        Блоки — непрерывные диапазоны по якорям (ищутся в теле И таблицах —
        в табличной вёрстке pdf2docx просьба/приложение лежат в ячейках):
          `header`      — вводная: всё до строки-титула (шапка-таблица ФНС в
                          реальном порядке стоит первой сюда);
          `body`        — описательно-мотивировочная: от титула до просьбы;
          `prayer`      — просительная: от строки-просьбы до «Приложение»;
          `attachments` — приложения: от «Приложение» и до конца.
        Якорь не найден граница схлопывается. Строки НИКОГДА не теряются.

        Не-DOCX (PDF-fallback) блоков не даёт — вызывающая сторона использует
        плоский `text`.
        """
        if Path(file_path).suffix.lower() not in (".docx", ".doc"):
            return []

        # Смещение line-index для каждой плоской части (часть может содержать
        # несколько строк через «\n»); line-index = позиция в extract_text-строках.
        flat = self._docx_text_parts(file_path)
        part_line_start: List[int] = []
        acc = 0
        for text, _origin in flat:
            part_line_start.append(acc)
            acc += text.count("\n") + 1

        # Развернуть части реального порядка в строки: (line_text, origin, line_idx).
        lines: List[Tuple[str, str, int]] = []
        for text, origin, fpi in self._docx_parts_true_order(file_path):
            base = part_line_start[fpi]
            for j, ln in enumerate(text.split("\n")):
                lines.append((ln, origin, base + j))

        # Якоря — по строкам тела и таблиц (колонтитулы/сноски не считаем).
        anchor_origins = ("body", "table")
        title_seq: Optional[int] = None
        prayer_seq: Optional[int] = None
        attach_seq: Optional[int] = None
        for seq, (text, origin, _idx) in enumerate(lines):
            if origin not in anchor_origins:
                continue
            if title_seq is None and self._is_title_line(text):
                title_seq = seq
                continue
            if prayer_seq is None and (
                self._PRAYER_COLON_RE.search(text) or self._PRAYER_START_RE.match(text)
            ):
                if title_seq is None or seq >= title_seq:
                    prayer_seq = seq
                continue
            if (
                attach_seq is None
                and prayer_seq is not None
                and seq >= prayer_seq
                and self._ATTACH_ANCHOR_RE.match(text)
            ):
                attach_seq = seq

        n = len(lines)
        header_end = title_seq if title_seq is not None else 0
        body_end = prayer_seq if prayer_seq is not None else n
        prayer_end = attach_seq if attach_seq is not None else n
        body_end = max(body_end, header_end)
        prayer_end = max(prayer_end, body_end)

        buckets: Dict[str, List[Dict[str, Any]]] = {
            "header": [], "body": [], "prayer": [], "attachments": []
        }
        for seq, (text, _origin, line_idx) in enumerate(lines):
            if seq < header_end:
                key = "header"
            elif seq < body_end:
                key = "body"
            elif seq < prayer_end:
                key = "prayer"
            else:
                key = "attachments"
            buckets[key].append({"text": text, "index": line_idx})

        titles = {
            "header": "Вводная часть",
            "body": "Описательно-мотивировочная часть",
            "prayer": "Просительная часть",
            "attachments": "Приложения",
        }
        return [
            {"id": key, "title": titles[key], "lines": buckets[key]}
            for key in ("header", "body", "prayer", "attachments")
            if buckets[key]
        ]

    def clean_extracted_value(self, value: str) -> str:
        """Очищает извлеченное значение от звездочек и других маскирующих символов"""
        if not value:
            return value

        cleaned = re.sub(r'\*+', '', value)
        cleaned = re.sub(r'\s+', ' ', cleaned)
        cleaned = cleaned.strip()

        return cleaned


    def _capitalize_full_name(self, full_name: str) -> str:
        if not full_name:
            return full_name
        tokens = [token for token in re.split(r"\s+", full_name.strip()) if token]
        capitalized_tokens = [self._capitalize_word(token) for token in tokens]
        return " ".join(capitalized_tokens)

    def _strip_ooo_prefix(self, name: Optional[str]) -> Optional[str]:
        """
        Убирает приставку ООО/Общество с ограниченной ответственностью, оставляя только название.
        """
        if not name:
            return name
        trimmed = name.strip(' «»"')
        match = re.match(
            r'^(?:ООО|Обществ[ао]\s+с\s+ограниченной\s+ответственностью)\s*[«"]?(.+?)[»"]?$',
            trimmed,
            re.IGNORECASE
        )
        if match:
            return self.clean_extracted_value(match.group(1).strip(' «»"'))
        return name

    def _convert_full_name_case(self, full_name: str, target_case: str) -> Optional[str]:
        if not full_name:
            return None
        tokens = [token for token in re.split(r"\s+", full_name.strip()) if token]
        if not tokens:
            return None
        inflected_tokens = [self._inflect_word(token, target_case) for token in tokens]
        result = " ".join(inflected_tokens).strip()
        return result or None

    def _extract_creditor_name_from_text(self, text: str) -> Optional[str]:
        if not text:
            return None
        # Приоритет: полное название организации после "Истец:"
        patterns = [
            r"Истец[:\s]*\n\s*((?:ООО|ОАО|ПАО|ЗАО|АО|Обществ[ао]\s+с\s+ограниченной\s+ответственностью|Публичное\s+акционерное\s+общество)[^,\n\\[\\]]+)",  # Истец:\nПАО Сбербанк
            r"Истец[:\s]+((?:ООО|ОАО|ПАО|ЗАО|АО|Общество\s+с\s+ограниченной\s+ответственностью|Публичное\s+акционерное\s+общество)[^,\n\\[\\]]+)",  # Истец: ПАО Сбербанк
                        r"Истец[:\s]*\n\s*([А-ЯЁ][А-ЯЁа-яё\s«»\"“”]{5,100}?)(?=\n|$|,|\[|ИНН|ОГРН|адрес|телефон|Место|Дата)",  # Истец:\nНазвание организации
                        r"Истец[:\s]+([А-ЯЁ][А-ЯЁа-яё\s«»\"“”]{5,100}?)(?=\n|$|,|\[|ИНН|ОГРН|адрес|телефон|Место|Дата)",  # Истец: Название организации
            r"Истец[:\s]+([^\n]+)",  # Общий паттерн
        ]
        for pattern_str in patterns:
            pattern = re.compile(pattern_str, re.IGNORECASE | re.MULTILINE)
            match = pattern.search(text)
            if match:
                creditor = match.group(1).strip()
                stop_words = ['Место нахождения', 'Дата государственной регистрации', 'ОГРН', 'ИНН', 'адрес']
                for stop_word in stop_words:
                    if stop_word in creditor:
                        creditor = creditor.split(stop_word)[0].strip()
                creditor = self.clean_extracted_value(creditor)
                if creditor and len(creditor) < 200 and not any(word in creditor.lower() for word in ['имеет право', 'потребовать', 'судебном порядке']):
                    return creditor


        from label_synonyms import CREDITOR_HEADER_LABELS, labels_alternation
        alt = labels_alternation(CREDITOR_HEADER_LABELS)
        m = re.search(
            rf"(?:{alt})\s*:?\s*"
            rf"([А-ЯЁ][^\n\r]{{3,150}}?)"
            rf"(?=\n|\r|$|\bИНН\b|\bОГРН|Дата\s+гос|Место\s+нахожд|Почтовый|[Аа]дрес|www\.|\bтел)",
            text, re.IGNORECASE,
        )
        if m:
            creditor = self.clean_extracted_value(m.group(1)).strip(" ,;")
            # «ФНС России в лице\nМежрайонной ИФНС …» — кредитор продолжается на
            # следующей строке: дописываем продолжение до реквизитов/адреса.
            if re.search(r"в\s+лице\s*$", creditor, re.IGNORECASE):
                cont = re.search(
                    r"в\s+лице\s*\n\s*([^\n]+?)\s*(?=\n|Адрес|ИНН|ОГРН|Почтов|$)",
                    text, re.IGNORECASE,
                )
                if cont:
                    creditor = (creditor + " " + self.clean_extracted_value(cont.group(1))).strip(" ,;")

            _paren = re.search(r"\(([^)]*)\)", creditor)
            _pure_opf = _paren and re.fullmatch(
                r"\s*(?:публичн\w+\s+|непубличн\w+\s+)?"
                r"(?:акционерн\w+\s+обществ\w+|обществ\w+\s+с\s+ограниченн\w+\s+ответственност\w+|"
                r"АО|ПАО|ООО|ОАО|ЗАО)\s*",
                _paren.group(1), re.IGNORECASE,
            )
            if not _pure_opf:
                creditor = re.split(r"\s*\(", creditor, maxsplit=1)[0].strip(" ,;")
            cl = creditor.lower()
            # Отсекаем мусор: маркеры шаблона [12]/[987], boilerplate-фразы, описания.
            has_marker = bool(re.search(r"\[\d", creditor))
            boilerplate = any(w in cl for w in [
                'имеет право', 'потребовать', 'судебном порядке', 'должник', 'ответчик',
                'требовани', 'утвердить', 'просит', 'в размере', 'из числа', ' руб',
            ])
            looks_creditor = bool(re.search(
                r'банк|общество|фнс|росси|инспекц|служб|\bАО\b|\bООО\b|\bПАО\b|\bОАО\b|\bЗАО\b',
                creditor, re.IGNORECASE))
            if creditor and 3 < len(creditor) < 150 and not has_marker and not boilerplate and looks_creditor:
                return creditor

        # Фолбэк 2: кредитор назван в просительной части — «(включить) требование
        # кредитора <Банк/Общество …>» (без шапки-метки). Берём название с ОПФ.
        m2 = re.search(
            r"(?:требовани\w*\s+)?кредитора\s+"
            r"((?:Банк\w*\s+)?(?:ВТБ|Сбербанк|[А-ЯЁ][А-ЯЁа-яё«»\"-]+)"
            r"[^,\n]{0,60}?(?:\([^)]*обществ[^)]*\)|«[^»]+»))",
            text, re.IGNORECASE,
        )
        if m2:
            creditor = self.clean_extracted_value(m2.group(1)).strip(" ,;")
            cl = creditor.lower()
            looks_cred = bool(re.search(r"банк|общество|\bАО\b|\bПАО\b|\bООО\b|\bОАО\b|\bЗАО\b", creditor, re.IGNORECASE))
            bad = any(w in cl for w in ("банкротств", "финансирован", "процедур", "должник", "требовани"))
            if creditor and 3 < len(creditor) < 150 and looks_cred and not bad:
                return creditor
        return None

    def _extract_creditor_block(self, text: str) -> Optional[str]:
        """Извлекает блок текста с реквизитами кредитора.

        Якорь — «Заявитель (кредитор)» / «Истец:» / «Кредитор:». Метка может стоять
        как на отдельной строке («Истец:\nООО …»), так и инлайн («Истец: ООО …»):
        в обоих случаях блок начинается со строки-метки и обрезается по началу блока
        должника, чтобы не захватить его реквизиты (ИНН/ОГРН/адрес должника).
        """
        if not text:
            return None
        header = text[:4500]
        for start_pattern in [
            r"Заявитель\s*\(кредитор\)\s*:?\s*",
            r"Истец\s*:\s*",
            r"Кредитор\s*:\s*",
            # «Заявитель Акционерное общество …» — без «(кредитор)» и двоеточия.
            r"Заявитель\s+(?=(?:Акционерн|Публичн|Общество|ООО|АО|ПАО|ЗАО|ОАО|ИП|ФНС|«))",
            # «Заявитель: <имя>» с двоеточием и именем без ОПФ-префикса («ББР Банк»,
            # «ООО ПКО …»). Последним, чтобы не перебивать более специфичные якоря.
            r"Заявитель\s*:\s*",
        ]:
            m = re.search(start_pattern, header, re.IGNORECASE)
            if m:
                block_start = m.start()
                block_end = min(block_start + 1200, len(header))
                block = header[block_start:block_end]
                # Обрезаем по началу блока должника / третьего лица / управляющего,
                # чтобы адрес/ИНН/ОГРН брались только из секции кредитора.
                cut = re.search(
                    r"\n\s*(?:Должник|Ответчик|Треть[ие]\s+лиц|"
                    r"(?:Финансов|Временн|Конкурсн)\w+\s+управляющ)",
                    block, re.IGNORECASE,
                )
                if cut:
                    block = block[: cut.start()]

                if re.search(r"ИНН|ОГРН|адрес|место\s+нахождения|\n\s*\d{6}[,\s]",
                             block, re.IGNORECASE):
                    return block
        return None


    def _clean_creditor_address(self, addr: str) -> str:
        """Очищает юр-адрес кредитора от постороннего: скобочных пометок и хвостов."""
        addr = re.sub(r"\s*\([^)]*\)", "", addr)
        addr = re.split(
            r"\s*(?:Почтов\w*\s*адрес|Фактическ\w*\s*адрес|Адрес\s*для|[Тт]елефон|[Тт]ел\.|"
            r"e-?mail|эл\.?\s*почт|ОГРН|ИНН|КПП|БИК|Дата\s+гос|р/с|к/с|корр)",
            addr, flags=re.IGNORECASE,
        )[0]
        addr = re.sub(r"\s+", " ", addr).strip().strip(",;. ")
        return addr

    def _extract_creditor_address(self, text: str) -> Optional[str]:
        """Извлекает ТОЛЬКО юридический адрес кредитора (без почтового и мусора)."""
        block = self._extract_creditor_block(text)
        if not block:
            return None
        # Приоритет — явные метки юр-адреса; «почтовый/фактический адрес» исключаем.
        for addr_pattern in [

            r"(?:место\s+нахождения|юридическ\w+\s+адрес)[:\s]*([0-9]{6}[,\s]+[^\n]+(?:\n[^\n]+)?)",
            r"(?:место\s+нахождения|юридическ\w+\s+адрес)[:\s]*([^\n]+)",
            # Плоская метка «Адрес:» в начале строки (индекс + продолжение на след. строке).
            r"(?:^|\n)\s*адрес[:\s]*([0-9]{6}[,\s]+[^\n]+(?:\n[^\n]+)?)",
        ]:
            addr_m = re.search(addr_pattern, block, re.IGNORECASE)
            if addr_m:
                addr = self._clean_creditor_address(self.clean_extracted_value(addr_m.group(1).strip()))
                if addr and 10 <= len(addr) <= 200 and re.search(r"\d{6}|город|\bг\.|ул\.|улиц|пр-?кт|проспект", addr, re.IGNORECASE):
                    return addr

        lines = [ln.strip() for ln in block.split("\n")]
        for i, line in enumerate(lines):
            if re.match(r"^\d{6}[,\s]", line) and not re.match(
                r"^\d{6}[,\s].*(?:почтов|фактическ|а/я|абонентск)", line, re.IGNORECASE
            ):
                parts = [line]

                for nxt in lines[i + 1:i + 5]:
                    if not nxt:
                        break
                    if re.match(r"^(?:Исх\b|Дата\b|№|тел|e-?mail|на\s+№)", nxt, re.IGNORECASE):
                        continue
                    if re.search(r"(?:\bнаб\b|\bул\b|улиц|\bд\.|\bстр\b|\bпер\b|пр-?кт|проспект|шоссе|корп|\bпом\b|\bкв\b|\bоф\b|\bзд\b|литер)", nxt, re.IGNORECASE):
                        parts.append(nxt)
                        continue
                    break
                addr = self._clean_creditor_address(self.clean_extracted_value(" ".join(parts)))
                if addr and 10 <= len(addr) <= 200:
                    return addr
        return None

    def _extend_address_from_text(self, text: str, address: Optional[str]) -> Optional[str]:
        if not text or not address:
            return None

        address_clean = address.strip()
        if not address_clean:
            return None

        pattern = re.compile(re.escape(address_clean), re.IGNORECASE)
        match = pattern.search(text)
        if not match:
            return None

        start_idx = text.rfind('\n', 0, match.start())
        line_start = start_idx + 1 if start_idx != -1 else match.start()
        end_idx = text.find('\n', match.end())
        line_end = end_idx if end_idx != -1 else min(len(text), match.end() + 150)
        extended_line = text[line_start:line_end].strip()
        if len(extended_line) > len(address_clean):
            extended = self.clean_extracted_value(extended_line)
            # Расширение до всей строки могло вернуть ведущий ярлык
            # («Юридический адрес: …»). Отрезаем его.
            extended = re.sub(
                r"^\s*(?:юридическ\w*\s*адрес|адрес\w*\s*регистрации|"
                r"адрес\w*\s*прописки|мест\w*\s*нахождени\w*|"
                r"мест\w*\s*жительства|адрес)\s*:?\s*",
                "", extended, flags=re.IGNORECASE
            ).strip(" ,;:")
            return extended or None
        return None

    def _normalize_mortgage_interface_fields(self, fields: Dict[str, Any], text: str):
        def _strip_procedure_phrase(name: Optional[str]) -> Optional[str]:
            if not name:
                return name
            cleaned = re.sub(r"\s*должника\s+введена\s+процедура.*$", "", name, flags=re.IGNORECASE)
            cleaned = re.sub(r"\s*введена\s+процедура\s+наблюдения.*$", "", cleaned, flags=re.IGNORECASE)
            return cleaned.strip()

        # Для competition_collateral специально извлекаем название должника из шапки после "Должник:" на следующей строке
        source_doc_type = fields.get("sourceDocumentType", "").lower()
        if source_doc_type == "competition_collateral":
            debtor_match = re.search(
                r"Должник[:\s]*\n\s*((?:ООО|ОАО|ПАО|ЗАО|АО|Обществ[ао]\s+с\s+ограниченной\s+ответственностью)[^\n]{0,200}?)(?=\n|$|ИНН|ОГРН|адрес|телефон)",
                text,
                re.IGNORECASE | re.MULTILINE
            )
            if debtor_match:
                debtor_name = debtor_match.group(1).strip()
                debtor_name = re.sub(r"\s+", " ", debtor_name)
                debtor_name = debtor_name[:200].strip()
                if debtor_name and len(debtor_name) > 5:
                    cleaned_name = self._strip_ooo_prefix(debtor_name)
                    if cleaned_name and "суд" not in cleaned_name.lower():
                        fields["applicantName"] = cleaned_name
                        fields["debtorName"] = cleaned_name
                        fields["mortgageDebtorName"] = cleaned_name
                        logger.info(f"Извлечено название должника для competition_collateral: {cleaned_name}")

        fields["mortgageDebtorName"] = _strip_procedure_phrase(self._strip_ooo_prefix(fields.get("mortgageDebtorName")))
        fields["debtorName"] = _strip_procedure_phrase(self._strip_ooo_prefix(fields.get("debtorName")))
        # Для applicantName сохраняем ОПФ (ООО, АО и т.д.) для маркеров [2], [2.1], [2.2]
        fields["applicantName"] = _strip_procedure_phrase(fields.get("applicantName"))

        name_candidates = [
            fields.get("mortgageDebtorName"),
            fields.get("debtorName"),
            fields.get("applicantName")
        ]
        proper_name = None
        for candidate in name_candidates:
            if candidate and "суд" not in candidate.lower():
                proper_name = self._capitalize_full_name(candidate)
                break

        if proper_name:
            fields["mortgageDebtorName"] = proper_name
            fields["debtorName"] = proper_name
            fields["applicantName"] = proper_name

            genitive = fields.get("applicantNameGenitive")
            if not genitive or "суд" in genitive.lower():
                genitive_name = self._convert_full_name_case(proper_name, "gent")
                if genitive_name:
                    fields["applicantNameGenitive"] = self._capitalize_full_name(genitive_name)

            dative = fields.get("mortgageDebtorNameDative")
            if not dative:
                dative_name = self._convert_full_name_case(proper_name, "datv")
                if dative_name:
                    fields["mortgageDebtorNameDative"] = self._capitalize_full_name(dative_name)

        # Пробрасываем общие кредитные поля в специальные ипотечные маркеры [111]-[114],
        # если они не были извлечены прямо.
        credit_amount = fields.get("mortgageCreditAmount111") or fields.get("creditAmount")
        if credit_amount and not fields.get("mortgageCreditAmount111"):
            fields["mortgageCreditAmount111"] = credit_amount

        credit_term = fields.get("mortgageCreditTerm112") or fields.get("creditTermMonths")
        if credit_term and not fields.get("mortgageCreditTerm112"):
            fields["mortgageCreditTerm112"] = credit_term

        interest_rate = fields.get("mortgageInterestRate113") or fields.get("creditInterestRate")
        if interest_rate and not fields.get("mortgageInterestRate113"):
            fields["mortgageInterestRate113"] = interest_rate

        penalty_rate = fields.get("mortgagePenaltyRate114") or fields.get("creditPenaltyRate")
        if penalty_rate and not fields.get("mortgagePenaltyRate114"):
            fields["mortgagePenaltyRate114"] = penalty_rate

        principal_amount = fields.get("mortgagePrincipalAmount11")
        if principal_amount:
            fields["principalDebt"] = principal_amount
            fields.setdefault("principalDebt13", principal_amount)

        interest_amount = fields.get("mortgageInterestAmount12")
        if interest_amount:
            fields["interest"] = interest_amount
            fields.setdefault("interest14", interest_amount)

        creditor_name = fields.get("creditorName", "")
        # Мусорный истец: пусто, длинно, нарратив («…исходит из положений…», из-за
        # опечатки метки «Истеп:» вместо «Истец:») или без признака организации.
        looks_bad = (
            not creditor_name
            or len(creditor_name) > 80
            or bool(re.search(r"имеет\s+право|досрочног|исходит|положени|\bсуд\b", creditor_name, re.IGNORECASE))
            or not re.search(r"\b(?:ООО|ОАО|ПАО|ЗАО|АО|Банк|Обществ\w+|Публичное|ФНС)\b", creditor_name, re.IGNORECASE)
        )
        if looks_bad:
            new_creditor = self._extract_creditor_name_from_text(text)
            if new_creditor and not re.search(r"исходит|положени|\bсуд\b", new_creditor, re.IGNORECASE):
                fields["creditorName"] = new_creditor
            elif re.search(r"Сбербанк", text, re.IGNORECASE):
                # Тело подтверждает «Публичное акционерное общество "Сбербанк России"».
                fields["creditorName"] = "ПАО Сбербанк"

        address = fields.get("applicantAddress")
        extended_address = self._extend_address_from_text(text, address)
        if extended_address:
            fields["applicantAddress"] = extended_address

        # Представитель истца: реальные заявления не содержат маркера [2.2], поэтому
        # паттерн mortgageRepresentative22 не срабатывает. Извлекаем ФИО из блока
        # «Представитель истца:», пропуская строку отделения/филиала банка.
        if not fields.get("mortgageRepresentative22"):
            rep = self._extract_representative_name(text)
            if rep:
                fields["mortgageRepresentative22"] = rep

        # Представитель ответчика: отдельная метка «Представитель ответчика:» в блоке
        # ответчика (в отличие от представителя истца). Заполняется вручную, но если
        # в заявлении есть — засеваем.
        if not fields.get("respondentRepresentativeName"):
            resp_rep = self._extract_respondent_representative_name(text)
            if resp_rep:
                fields["respondentRepresentativeName"] = resp_rep

    def _fill_mortgage_state_duty(self, fields: Dict[str, Any], text: str) -> None:
        """Ипотека: госпошлина из шапки «Госпошлина: X руб.», если финансовый каскад
        её не заполнил. Каскад берёт госпошлину из блока «ПРОСИТ СУД» и спотыкается на
        числах с пробелом-разрядом и точкой-десятичной («80 400.00»); шапка — надёжный
        якорь. Зовём ПОСЛЕ каскада, иначе значение затирается пересчётом финблока."""
        if fields.get("stateDuty16") or fields.get("stateDuty"):
            return
        m = re.search(
            r"Госпошлин[ауы]?\s*[:\-]\s*([0-9][0-9\s.,]*[0-9])\s*(?:руб|рублей)",
            text, re.IGNORECASE,
        )
        if m:
            duty = re.sub(r"\s+", " ", m.group(1)).strip()
            fields["stateDuty16"] = duty
            fields["stateDuty"] = duty

    def _extract_representative_name(self, text: str) -> Optional[str]:
        """ФИО представителя истца из блока «Представитель истца:».

        Первая строка записи часто — отделение/филиал банка (с номером), а ФИО идёт
        ниже. Берём первую строку-ФИО в пределах нескольких строк после метки,
        останавливаясь на следующем разделе/реквизите."""
        m = re.search(r"Представител\w*\s+истца\s*:?", text, re.IGNORECASE)
        if not m:
            return None
        from fio_detector import is_person_name, _normalize_fio
        for line in text[m.end():].split("\n")[:6]:
            s = line.strip().strip(",")
            if not s:
                continue
            if re.match(
                r"(?:Ответчик|Истец|СНИЛС|ИНН|ОГРН|Почтов\w+\s+адрес|Адрес|"
                r"Цена\s+иска|Госпошлин|Дата\s+рождения|Паспорт)\b",
                s, re.IGNORECASE,
            ):
                break
            if is_person_name(s):
                return _normalize_fio(s)
        return None

    def _extract_respondent_representative_name(self, text: str) -> Optional[str]:
        """ФИО представителя ответчика из блока «Представитель ответчика:».

        ФИО обычно идёт первой строкой после метки; берём первую строку-ФИО в
        пределах нескольких строк, останавливаясь на следующем разделе/реквизите."""
        m = re.search(r"Представител\w*\s+ответчик\w*\s*:?", text, re.IGNORECASE)
        if not m:
            return None
        from fio_detector import is_person_name, _normalize_fio
        for line in text[m.end():].split("\n")[:6]:
            s = line.strip().strip(",")
            if not s:
                continue
            if re.match(
                r"(?:Ответчик|Истец|Треть\w+\s+лиц|СНИЛС|ИНН|ОГРН|Почтов\w+\s+адрес|Адрес|"
                r"Цена\s+иска|Госпошлин|Дата\s+рождения|Паспорт|Контактн\w+\s+телефон)\b",
                s, re.IGNORECASE,
            ):
                break
            if is_person_name(s):
                return _normalize_fio(s)
        return None

    def _extract_mortgage_collateral_block(self, text: str) -> Optional[str]:
        """
        Пытается извлечь описание предмета залога для маркера [1221], даже если нет явных маркеров.
        Ищет текст после фраз о залоге или кредите.
        Возвращает строку с описанием (может содержать несколько предметов залога).
        """
        if not text:
            return None

        marker = "[1221]"
        marker_idx = text.find(marker)
        search_end = marker_idx if marker_idx != -1 else len(text)

        # Пытаемся найти предложение с описанием предмета залога
        # Вариант 1: После фразы о залоге "что подтверждается договором залога № ... от ... :"
        collateral_pattern1 = re.compile(
            r"что\s+подтверждается\s+договором\s+залога\s+№\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*:\s*",
            re.IGNORECASE | re.DOTALL
        )
        phrase_match = collateral_pattern1.search(text, 0, search_end)

        # Вариант 2: После фразы "В качестве обеспечения... что подтверждается договором залога"
        if not phrase_match:
            collateral_pattern2 = re.compile(
                r"В\s+качестве\s+обеспечения[^.]*?что\s+подтверждается\s+договором\s+залога\s+№\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*:\s*",
                re.IGNORECASE | re.DOTALL
            )
            phrase_match = collateral_pattern2.search(text, 0, search_end)

        # Вариант 3: После фразы "Кредит ... выдавался/предоставлялся ..."
        if not phrase_match:
            phrase_pattern = re.compile(
                r"кредит[^\n]{0,120}?(?:выдава[лс][ась]|предоставля[лс][ась]|предоставлен)[^:\n]*:\s*",
                re.IGNORECASE
            )
            phrase_match = phrase_pattern.search(text, 0, search_end)

        # Вариант 4: После фразы "предоставил в залог"
        if not phrase_match:
            collateral_pattern3 = re.compile(
                r"предоставил\s+в\s+залог[^:\n]*:\s*",
                re.IGNORECASE
            )
            phrase_match = collateral_pattern3.search(text, 0, search_end)

        if phrase_match:
            segment_start = phrase_match.end()
        else:
            # Если не нашли начало по фразам, ищем ближайший разрыв строки перед маркером
            prev_break = text.rfind("\n\n", 0, search_end)
            if prev_break == -1:
                prev_break = text.rfind('\n', 0, search_end)
            segment_start = prev_break if prev_break != -1 else max(0, search_end - 600)

        if segment_start >= search_end:
            return None

        segment = text[segment_start:search_end]

        # Ограничиваем блок следующими служебными абзацами/маркерами
        # Для залоговых заявлений также останавливаемся на "Наличие заложенного имущества"
        stop_pattern = re.search(
            r"\n{2,}(?=(?:По\s+состоянию|В\s+результате|ПРОСИТ|Просит|Суд|Определил|Установить|Взыскать|Сообщение|Кредитор|Должник|Наличие\s+заложенного))",
            segment,
            re.IGNORECASE
        )
        if stop_pattern:
            segment = segment[:stop_pattern.start()]

        # Также останавливаемся на фразе "Наличие заложенного имущества подтверждается выпиской из ЕГРН"
        stop_collateral = re.search(
            r"Наличие\s+заложенного\s+имущества\s+подтверждается\s+выписк\w+\s+из\s+ЕГРН",
            segment,
            re.IGNORECASE
        )
        if stop_collateral:
            segment = segment[:stop_collateral.start()]

        # ИНЛАЙНОВЫЕ стопы (без требования \n{2,}): OCR/pdf2docx часто склеивает всё
        # заявление в один абзац, и стопы по "\n{2,}" не срабатывают — захват уходил в
        # нарратив обязательств и просительную часть (баг «огромного текста в предмете
        # залога» на любом залоговом заявлении). Обрезаем на первой служебной фразе,
        # маркирующей конец описания предмета и начало текста об исполнении/долге/просьбе.
        inline_stop = re.search(
            r"(?:Банк\s+свои\s+обязательства|Должник\w*\s+в\s+настоящее\s+время|"
            r"По\s+состояни\w+|Обязательство\s*№|Наличие\s+заложенного\s+имущества|"
            r"выписк\w+\s+из\s+ЕГРН|ПРОСИТ\b|Нормативно\W|В\s+соответствии\s+с\s+п\.\s*\d)",
            segment,
            re.IGNORECASE,
        )
        if inline_stop:
            segment = segment[:inline_stop.start()]

        marker_inside = re.search(r"\[\d{1,4}\]", segment)
        if marker_inside:
            segment = segment[:marker_inside.start()]

        segment = segment.strip(" \n\r\t•-")
        if not segment or len(segment) < 10:
            return None

        # Не считаем залогом фразы вида "Кому выдана [ФИО]" — это указание получателя кредита/доверенности, не предмет залога
        if re.match(r"^Кому\s+выдана\s+[А-ЯЁа-яё\s]{5,80}", segment, re.IGNORECASE):
            return None
        if re.match(r"^[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:\s+\.\.\.)?\s*$", segment.strip()) and "залог" not in segment.lower() and "недвижим" not in segment.lower() and "авто" not in segment.lower():
            return None

        segment = re.sub(r'\s{3,}', '\n\n', segment)
        segment = re.sub(r'[ \t]+', ' ', segment)

        return segment.strip()

    def _split_collateral_items(self, description: str) -> List[str]:
        """
        Разделяет описание предметов залога на отдельные предметы.

        ВАЖНО: по вашему описанию каждый предмет залога
        оформлен отдельной строкой (визуально — отдельный абзац).
        Поэтому правило максимально простое и предсказуемое:

        - каждая непустая строка = отдельный предмет залога.
        """
        if not description:
            return []

        text = description.replace("\r\n", "\n").replace("\r", "\n")

        raw_lines = [line.strip() for line in text.split("\n")]

        items: List[str] = [line for line in raw_lines if line and len(line) >= 10]

        if not items:
            cleaned = text.strip()
            if cleaned and len(cleaned) >= 10:
                items = [cleaned]

        logger.info(f"Разделено предметов залога (по строкам): {len(items)}")
        for i, item in enumerate(items):
            logger.info(f"  Предмет {i + 1}: {item[:100]}...")

        return items

    def _clean_court_name(self, value: str) -> str:
        """Очищает название суда от лишнего текста"""
        if not value:
            return value

        stop_words = [
            "возникает",
            "у конкурсного",
            "по денежным",
            "обязательствам",
            "с даты",
            "вступления",
            "в законную",
            "силу",
            "решения",
            "суда",
            "арбитражного",
            "или",
            "третейского",
            "о взыскании",
            "с должника",
            "денежных",
            "средств"
        ]

        value_lower = value.lower()
        for stop_word in stop_words:
            if stop_word in value_lower:
                idx = value_lower.find(stop_word)
                value = value[:idx].strip()
                break

        if len(value) > 100:
            for delimiter in ['.', ',', '\n']:
                idx = value.find(delimiter)
                if 20 < idx < 100:
                    value = value[:idx].strip()
                    break
            else:
                value = value[:100].strip()

        return value.strip()



    def extract_legal_entity_short_name(self, text: Optional[str]) -> Optional[str]:
        """
        Извлекает краткое наименование юридического лица (без организационно-правовой формы).
        """
        if not text:
            return None

        normalized = self.clean_extracted_value(text)
        if not normalized:
            return None

        patterns = [
            r"(?:ООО|ОАО|ПАО|ЗАО|АО|Обществ[ао]\s+с\s+ограниченной\s+ответственностью)\s*[«\"]?([^«»\"\n,]+)",
            r"(?:компания\s+с\s+ограниченной\s+ответственностью)\s*[«\"]?([^«»\"\n,]+)",
        ]

        for pattern in patterns:
            match = re.search(pattern, normalized, re.IGNORECASE)
            if match:
                short = match.group(1)
                return self.clean_extracted_value(short.strip('«»" '))

        return None


    def _entity_from_text(self, text: str) -> Optional[str]:
        """Тип должника КФХ/ИП из ТЕКСТА (для определений суда без блока «Должник:»).

        Смотрит описание должника СРАЗУ после «несостоятельным (банкротом) …» —
        там тип указан словами («индивидуального предпринимателя Главы КФХ …»).
        Возвращает только kfh/ip (legal не определяем по тексту — рядом часто
        стоит ЮЛ-кредитор, можно ошибиться).
        """
        if not text:
            return None
        m = re.search(
            r"несостоятельн\w*\s*\(?\s*банкрот\w*\s*\)?\s*([^\n.]{0,70})",
            text, re.IGNORECASE,
        )
        if not m:
            return None
        ctx = m.group(1).lower()
        if ("глав" in ctx and "кфх" in ctx) or "крестьянск" in ctx or "к(ф)х" in ctx:
            return "kfh"
        if "индивидуальн" in ctx and "предпринимател" in ctx:
            return "ip"
        return None



    def _apply_stacked_party_layout(self, fields: Dict[str, Any], text: str) -> None:
        """Формат ВТБ «реестр Nл»: метки сторон идут СТОПКОЙ, значения — ниже.

        Шапка выглядит так:
            КРЕДИТОР:
            ДОЛЖНИК:
            ФИНАНСОВЫЙ
            УПРАВЛЯЮЩИЙ:
            Банк ВТБ (ПАО) … ОГРН … ИНН …            значение КРЕДИТОРА
            Горина Юлия Игоревна … ИНН … адрес …      значение ДОЛЖНИКА
            Удодов Сергей Александрович (ИНН …)        значение УПРАВЛЯЮЩЕГО

        Обычный построчный разбор берёт метку «ФИНАНСОВЫЙ» за ФИО должника и
        реквизиты банка — за должника. Здесь переустанавливаем стороны по порядку.
        """
        if not text:
            return
        sig = re.search(
            r"КРЕДИТОР\s*:\s*\n\s*ДОЛЖНИК\s*:\s*\n\s*ФИНАНСОВ\w*\s*\n?\s*УПРАВЛЯЮЩ\w*\s*:",
            text, re.IGNORECASE,
        )
        if not sig:
            return
        region = text[sig.end():]
        cut = re.search(r"РАЗМЕР\s+ТРЕБОВАНИЙ|\bдело\s*№|\bЗАЯВЛЕНИЕ\b", region, re.IGNORECASE)
        if cut:
            region = region[:cut.start()]


        fio_re = re.compile(
            r"(?m)^[ \t]*([А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+)[ \t]*(?=\(|\n|$)"
        )
        fios = list(fio_re.finditer(region))
        if len(fios) < 2:
            return

        cred_block = region[: fios[0].start()]
        debtor_block = region[fios[0].start(): fios[1].start()]
        manager_block = region[fios[1].start():]

        def _inn(s):
            m = re.search(r"ИНН[:\s]*([0-9]{10,12})\b", s)
            return m.group(1) if m else None

        def _ogrn(s):
            m = re.search(r"ОГРН[:\s]*([0-9]{13,15})\b", s)
            return m.group(1) if m else None

        # --- КРЕДИТОР (юрлицо) ---
        cred_name = ""
        for ln in cred_block.splitlines():
            ln = ln.strip()
            if ln:
                cred_name = re.split(r"\s*-{3,}|\bОГРН\b|\bИНН\b", ln)[0].strip()
                break
        if cred_name:
            fields["creditorName"] = cred_name
        if _inn(cred_block):
            fields["creditorInn"] = _inn(cred_block)
        if _ogrn(cred_block):
            fields["creditorOgrn"] = _ogrn(cred_block)
        ca = re.search(
            r"(?:Юридическ\w*\s+адрес|адрес)[:\s]*\n?\s*([0-9]{6}[^\n]+)",
            cred_block, re.IGNORECASE,
        )
        if ca:
            fields["creditorAddress"] = ca.group(1).strip()

        # --- ДОЛЖНИК (физлицо) ---
        fields["applicantName"] = fios[0].group(1).strip()
        fields["entityType"] = "individual"
        fields.pop("companyInn", None)
        fields.pop("ogrn", None)
        if _inn(debtor_block):
            fields["inn"] = _inn(debtor_block)
        bd = re.search(
            r"Дата\s+рождения[:\s]*([0-3]?\d[.,][01]?\d[.,]\d{4})",
            debtor_block, re.IGNORECASE,
        )
        if bd:
            fields["birthDate"] = bd.group(1).replace(",", ".")
        bp = re.search(r"Место\s+рождения[:\s]*([^\n]+)", debtor_block, re.IGNORECASE)
        if bp:
            fields["birthPlace"] = bp.group(1).strip().rstrip(" .,;")
        # Адрес должника может быть и без индекса («Ростовская обл., г. Шахты, …»),
        # и многострочным — берём всё после метки «Адрес …:» до конца блока.
        da = re.search(
            r"Адрес\s+(?:регистрации|проживания|места\s+жительства)?\s*:\s*([\s\S]+)",
            debtor_block, re.IGNORECASE,
        )
        if da:
            addr = re.sub(r"\s+", " ", da.group(1)).strip().rstrip(" ,;")
            addr = re.split(r"\b(?:ИНН|ОГРН|СНИЛС|КПП)\b", addr, flags=re.IGNORECASE)[0].strip().rstrip(" ,;")
            if re.search(r"[А-Яа-яЁё]{3}", addr):
                fields["applicantAddress"] = addr

        # --- ФИНАНСОВЫЙ УПРАВЛЯЮЩИЙ (физлицо) ---
        fields["managerName"] = fios[1].group(1).strip()
        if _inn(manager_block):
            fields["managerInn"] = _inn(manager_block)
        ma = re.search(r"(\b\d{6}\b\s*,[^\n]+(?:\n[^\n]+)?)", manager_block)
        if ma:
            addr = re.sub(r"\s*\n\s*", " ", ma.group(1)).strip().rstrip(" ,;")
            addr = re.sub(r"\s*-{3,}\s*", " ", addr).strip()
            fields["managerAddress"] = addr
        else:
            # Иначе остаётся мусор от общего разбора («…609391, ИНН 7702070139»).
            fields.pop("managerAddress", None)
        logger.info(
            "Раскладка ВТБ (метки стопкой): должник=%s, кредитор=%s, управляющий=%s"
            % (fields.get("applicantName"), fields.get("creditorName"), fields.get("managerName"))
        )

    _STACKED_SIG_RE = re.compile(
        r"КРЕДИТОР\s*:\s*\n\s*ДОЛЖНИК\s*:\s*\n\s*ФИНАНСОВ\w*\s*\n?\s*УПРАВЛЯЮЩ\w*\s*:",
        re.IGNORECASE,
    )

    def _fix_truncated_creditor_name(self, fields: Dict[str, Any], text: str) -> None:
        """Имя кредитора, обрезанное переносом строки внутри названия
        («Публичное Акционерное Общество "МТС-\\nБанк"»), дотягиваем по тексту до
        закрывающей кавычки."""
        name = fields.get("creditorName")
        if not name or not text:
            return
        base = name.rstrip()
        # Триггер: имя кончается дефисом или открытой (непарной) кавычкой.
        unbalanced = base.count('"') % 2 == 1 or base.count("«") > base.count("»")
        if not (base.endswith("-") or unbalanced):
            return
        m = re.search(re.escape(base) + r"\s*\n?\s*([A-Za-zА-Яа-яЁё][^\n]*?[»\"])", text)
        if not m:
            return
        tail = m.group(1).strip()
        joiner = "" if base.endswith("-") else " "
        fixed = re.sub(r"\s+", " ", (base + joiner + tail)).strip()
        fields["creditorName"] = fixed

    # Метки-реквизиты, приклеенные к имени кредитора из stacked-шапки банка
    # («…ПАО Сбербанк Место нахождения: …»): с любой из них начинается блок
    # реквизитов, а не название — режем имя по ней.
    _CREDITOR_NAME_TAIL_RE = re.compile(
        r"\s*(?:Мест\w*\s+нахожд\w*|Юридическ\w+\s+адрес|Почтов\w+(?:\s+адрес)?|"
        r"Адрес\s+для\s+корреспонденции|Адрес[:\s]|ОГРН|ИНН|Телефон|e-?mail|"
        r"Дата\s+(?:государственн\w+\s+)?регистрац\w+)",
        re.IGNORECASE,
    )
    # Исходящий номер письма, вклеенный pdf2docx в середину адреса
    # («…Муниципальный Округ Исх. №б/н от 29.04.2026 Замоскворечье…»).
    _DISPATCH_NUMBER_RE = re.compile(
        r"\s*Исх\.?\s*№\s*\S+\s+от\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*",
        re.IGNORECASE,
    )

    def _clean_creditor_artifacts(self, fields: Dict[str, Any], text: str) -> None:
        """Убирает из кредитора артефакты грязной вёрстки/stacked-шапки банка:
        1) хвост-реквизиты в creditorName («… Место нахождения: …» и т.п.);
        2) врезку исходящего номера «Исх. № … от ДД.ММ.ГГГГ» в creditorAddress.
        Реквизиты (ИНН/ОГРН/адрес) при этом не теряются — они берутся отдельно
        из блока кредитора/реестра, а сопоставление с реестром по обрезанному имени
        (напр. «ПАО Сбербанк») продолжает работать."""
        name = fields.get("creditorName")
        if isinstance(name, str) and name.strip():
            m = self._CREDITOR_NAME_TAIL_RE.search(name)
            if m and m.start() > 0:
                cut = name[: m.start()].strip(" ,;:-")
                if len(cut) >= 3:
                    fields["creditorName"] = cut

        addr = fields.get("creditorAddress")
        if isinstance(addr, str) and "исх" in addr.lower():
            cleaned = self._DISPATCH_NUMBER_RE.sub(" ", addr)
            cleaned = re.sub(r"\s+", " ", cleaned).strip(" ,;")
            if cleaned and cleaned != addr:
                fields["creditorAddress"] = cleaned

    def _apply_table_breakdown_finances(self, fields: Dict[str, Any], text: str) -> None:
        """Табличная разбивка задолженности: столбцы «Структура задолженности |
        Значение | RUR» по одному/нескольким договорам (значения суммируются).
        Гейт по сигнатуре таблицы; валидация суммы компонентов = «ОБЩАЯ
        ЗАДОЛЖЕННОСТЬ»/итог, иначе ничего не перекрываем."""
        if not text:
            return
        low = text.lower()
        if "структура" not in low or "rur" not in low:
            return
        tm = re.search(
            r"ОБЩАЯ\s+ЗАДОЛЖЕННОСТЬ[:\s]*(\d[\d   ]*(?:[.,]\d{1,2})?)",
            text, re.IGNORECASE,
        )
        if not tm:
            return
        total = self._fin_amount(tm.group(1))
        if total <= 0:
            return
        money = r"(\d[\d   ]*(?:[.,]\d{1,2})?)"
        pr = it = fo = pen = ld = 0.0
        n = 0
        seen = set()  # таблица в тексте бывает задвоена — дедуп по (категория, значение)
        brk_vals: Dict[str, list] = {}  # адденды по группам — для тултипа «откуда число»
        for m in re.finditer(
            r"(основн\w*\s+долг\w*|процент\w*|неустойк\w*|\bпени\b|штраф\w*|госпошлин\w*)"
            r"\s*\n\s*" + money + r"\s*\n?\s*(?:RUR|руб)",
            text, re.IGNORECASE,
        ):
            cat = m.group(1).lower()
            val = self._fin_amount(m.group(2))
            if val <= 0:
                continue
            _grp = ("forfeit" if ("неустой" in cat or "пени" in cat)
                    else "penalty" if "штраф" in cat
                    else "interest" if "процент" in cat
                    else "loan_duty" if "госпошл" in cat
                    else "principal")
            if (_grp, round(val, 2)) in seen:
                continue
            seen.add((_grp, round(val, 2)))
            brk_vals.setdefault(_grp, []).append(val)
            if "штраф" in cat:
                pen += val
            elif "неустой" in cat or "пени" in cat:
                fo += val
            elif "процент" in cat:
                it += val
            elif "госпошл" in cat:
                ld += val
            else:
                pr += val
            n += 1
        comp = pr + it + fo + pen + ld
        if n == 0 or abs(comp - total) > 1.5:
            return

        breakdown = {
            fkey: [self._fin_fmt(v) for v in brk_vals[grp]]
            for grp, fkey in (("principal", "principalDebt"), ("interest", "interest"),
                              ("forfeit", "forfeit"), ("penalty", "penalties"),
                              ("loan_duty", "loanStateDuty17"))
            if brk_vals.get(grp)
        }
        if breakdown:
            fields["financeBreakdown"] = breakdown
        else:
            fields.pop("financeBreakdown", None)
        fields["principalDebt"] = self._fin_fmt(pr)
        fields["principalDebt13"] = fields["principalDebt"]
        fields["interest"] = self._fin_fmt(it)
        fields["interest14"] = fields["interest"]
        fields["forfeit"] = self._fin_fmt(fo)
        fields["forfeit15"] = fields["forfeit"]
        fields["penalties"] = self._fin_fmt(pen)
        if ld > 0:
            fields["loanStateDuty17"] = self._fin_fmt(ld)
        fields["totalDebt"] = self._fin_fmt(total)
        fields["debtAmount"] = fields["totalDebt"]
        # Банкротная госпошлина — «Взыскать … в размере NN (прописью) рублей».
        gm = re.search(
            r"(?:гос)?пошлин\w*[^\d]{0,120}?в\s*размере\s*" + money + r"\s*(?:\([^)]*\))?\s*руб",
            text, re.IGNORECASE,
        )
        if gm:
            duty = self._fin_fmt(self._fin_amount(gm.group(1)))
            fields["stateDuty16"] = duty
            fields["stateDuty"] = duty
        logger.info(
            "Финансы из таблицы: осн=%s проц=%s неуст=%s ссуд.гп=%s итог=%s"
            % (fields["principalDebt"], fields["interest"], fields["forfeit"],
               fields.get("loanStateDuty17"), fields["totalDebt"])
        )

    def _clear_garbage_bankruptcy_duty(self, fields: Dict[str, Any]) -> None:
        """Банкротная госпошлина [16], совпавшая с итогом/осн.долгом — это мусор
        (отдельной банкротной госпошлины в документе нет). Удаляем поле."""
        amt = self._fin_amount
        bankr = amt(fields.get("stateDuty16") or fields.get("stateDuty"))
        if bankr <= 0:
            return
        total = amt(fields.get("totalDebt") or fields.get("debtAmount"))
        principal = amt(fields.get("principalDebt") or fields.get("principalDebt13"))
        if (total > 0 and abs(bankr - total) < 1) or (principal > 0 and abs(bankr - principal) < 1):
            fields.pop("stateDuty16", None)
            fields.pop("stateDuty", None)
            logger.info("Удалена мусорная банкротная госпошлина (совпала с итогом/осн.долгом)")

    def _fill_creditor_nlp(self, fields: Dict[str, Any], text: str) -> None:
        """Второстепенная NLP-сеть для кредитора: срабатывает ТОЛЬКО когда основной
        label-anchored парсер и реестр не дали имени. Берём из NER первую организацию
        по шапке, отсекая суд/должника и требуя признак организации (ОПФ/кавычки/
        налоговый орган). Роль (это кредитор) задаёт позиция-шапка, а не сам NER.
        Verbatim-срез обеспечивает `find_orgs`. Цель — не «сломаться» на незнакомой
        раскладке будущих документов, где нет привычной метки «Кредитор:/Заявитель:»."""
        if not text or (fields.get("creditorName") or "").strip():
            return
        for cand in _nlp.find_orgs(text[:1200]):
            low = cand.lower()
            # Отсев суда и явного не-кредитора.
            if re.search(r"\bсуд\b|арбитражн\w+\s+суд", low):
                continue
            # Признак организации: ОПФ / банк / налоговый орган / кавычки-название.
            looks_org = bool(re.search(
                r"\b(?:ООО|АО|ПАО|ЗАО|ОАО|ПКО|НАО|Банк)\b|банк|общество|"
                r"налогов|ифнс|\bфнс\b|инспекц|служб|[«\"]",
                cand, re.IGNORECASE))
            if looks_org and 3 < len(cand) < 150:
                fields["creditorName"] = cand
                return

    # Регион в любом регистре: «Ростовской области», «Республике Башкортостан»,
    # «ПО РОСТОВСКОЙ ОБЛАСТИ» и т.п. Нормализуется в _norm_fns_region.
    _FNS_REGION = (r"(?:Республик\w+\s+[А-Яа-яЁё][А-Яа-яЁё-]+|"
                   r"[А-Яа-яЁё]+(?:ой|ому)\s+(?:области|краю|округу|АО))")

    @staticmethod
    def _norm_fns_region(region: str) -> str:
        """ЗАГЛАВНЫЙ/смешанный регион канонический вид: «РОСТОВСКОЙ ОБЛАСТИ»
        «Ростовской области», «РЕСПУБЛИКЕ БАШКОРТОСТАН» «Республике Башкортостан».
        Гео-тип (области/краю/округу) со строчной, названия — с заглавной."""
        lower_words = {"области", "краю", "округу", "ао", "автономному"}
        out = []
        for w in re.sub(r"\s+", " ", region).strip().split(" "):
            out.append(w.lower() if w.lower() in lower_words else w[:1].upper() + w[1:].lower())
        return " ".join(out)

    def _build_fns_creditor(self, text: str) -> str:
        """Имя налогового органа по слоям шаблонов (приоритет — конкретной инспекции).

        Слои (первый сработавший): 1) «в лице Межрайонной ИФНС России № N по <регион>»;
        2) скобка «(Межрайонная ИФНС России № N по <регион>)»; 3) заглавная полная форма
        «МЕЖРАЙОННАЯ ИНСПЕКЦИЯ ФНС № N ПО <РЕГИОН>»; 4) управление «УФНС/УПРАВЛЕНИЕ ФНС
        по <регион>». Если ни один — общий «ФНС России» (безопасная деградация, не мусор).
        Возвращает всегда verbatim-нормализованное юр-имя.
        """
        rg = self._FNS_REGION
        # Слои с конкретной инспекцией: (№ инспекции, регион).
        insp_layers = [
            rf"в\s+лице\s+Межрайонн\w+\s+ИФНС\s+России\s+№?\s*(\d+)\s+по\s+({rg})",
            rf"\(\s*Межрайонн\w+\s+ИФНС\s+России\s+№?\s*(\d+)\s+по\s+({rg})\s*\)",
            rf"Межрайонн\w+\s+инспекци\w+\s+федеральн\w+\s+налогов\w+\s+служб\w+"
            rf"\s+№?\s*(\d+)\s+по\s+({rg})",
        ]
        for pat in insp_layers:
            m = re.search(pat, text, re.IGNORECASE)
            if m:
                return (f"ФНС России в лице Межрайонной ИФНС России № {m.group(1)} "
                        f"по {self._norm_fns_region(m.group(2))}")
        # Уровень управления по субъекту (без номера инспекции).
        for pat in (rf"УФНС\s+России\s+по\s+({rg})",
                    rf"управлени\w+\s+федеральн\w+\s+налогов\w+\s+служб\w+\s+по\s+({rg})"):
            m = re.search(pat, text, re.IGNORECASE)
            if m:
                return f"ФНС России в лице УФНС России по {self._norm_fns_region(m.group(1))}"
        return "ФНС России"


    # САМОБАНКРОТСТВО: структурный разбор шапки «должник первым, кредиторы списком»

    _SB_DEBTOR_LABEL_RE = re.compile(
        r"(?:от\s+)?должник\w*\s*(?:\(\s*заявител\w*\s*\))?\s*:?|(?<![а-яё])фио\s*:",
        re.IGNORECASE,
    )

    _SB_CREDITORS_RE = re.compile(r"(?:\d{1,2}[ \t]+)?кредитор\w*\s*\d{0,2}\s*[:.]", re.IGNORECASE)
    # Конец шапки — заголовок «Заявление …» (отдельной строкой, с продолжением-типом,
    # либо ЗАГЛАВНОЕ «ЗАЯВЛЕНИЕ» посреди строки перед «о признании…»).
    _SB_TITLE_RES = (
        re.compile(r"(?im)^\s*заявлени\w*\s*$"),
        re.compile(r"(?im)^\s*заявлени\w*\s+(?:должника|физическ|гражданина|о\s)"),
        re.compile(r"ЗАЯВЛЕНИЕ(?=\s*\n?\s*о\s)"),
    )
    # ФИО: Фамилия [( девичья )] Имя Отчество; допускаем перенос строки внутри
    # («Корсунов Вячеслав\nВалерьевич») — \s+ покрывает \n.
    _SB_FIO_RE = re.compile(
        r"([А-ЯЁ][А-ЯЁа-яё-]{1,25}(?:\s*\(\s*[А-ЯЁ][А-ЯЁа-яё-]{1,25}\s*\))?)"
        r"\s+([А-ЯЁ][а-яё-]{2,20})"
        r"\s+([А-ЯЁ][а-яё-]{2,25})"
    )
    # Метки-стопы внутри блока должника (конец значения очередного поля).
    _SB_FIELD_STOP = (
        r"(?:снилс|инн\b|огрн|кредитор|место\s+работы|телефон|контактн|адрес\s+для|"
        r"семейн|депозит|тел[.\s]|паспорт|выдан|дата\s+выдачи|код\s+подразделения|"
        r"дата\s+(?:и\s+место\s+)?рождения|место\s+рождения|$)"
    )

    def _apply_self_bankruptcy_layout(self, fields: Dict[str, Any], text: str):
        """Разводит должника и кредиторов в заявлении САМОБАНКРОТА.

        Шапка самобанкрота: должник (с паспортными реквизитами) идёт ПЕРВЫМ,
        затем СПИСОК кредиторов с их ИНН/ОГРН/адресами, затем третьи лица /
        уполномоченный орган. Общий label-парсер на такой раскладке путает
        стороны (берёт ОГРН/адрес кредитора как реквизиты должника) — здесь
        реквизиты должника извлекаются СТРОГО из его блока (до первого
        «Кредитор…»), а кредиторские поля очищаются (кредитора-заявителя нет).

        Возвращает список третьих лиц из шапки (может быть пустым) или None,
        если блок «Третьи лица:» не найден.
        """
        if not text:
            return None

        title_pos = len(text)
        for rx in self._SB_TITLE_RES:
            m = rx.search(text)
            if m and m.start() < title_pos:
                title_pos = m.start()
        header = text[: min(title_pos, 6000)]

        m_label = self._SB_DEBTOR_LABEL_RE.search(header)
        if m_label:
            blk_start = m_label.end()
        else:
            # Безметочная шапка (Федоренко): должник — первое ФИО после суда.
            m_fio0 = None
            for cand in self._SB_FIO_RE.finditer(header):
                fio_c = re.sub(r"\s+", " ", cand.group(0))
                if re.search(r"суд|област|район|граждан", fio_c, re.IGNORECASE):
                    continue
                if is_person_name(fio_c):
                    m_fio0 = cand
                    break
            if not m_fio0:
                return None
            blk_start = m_fio0.start()
        m_cred = self._SB_CREDITORS_RE.search(header, blk_start)
        blk = header[blk_start : m_cred.start() if m_cred else len(header)]
        if len(blk) < 30:
            return None

        fio = None
        for cand in self._SB_FIO_RE.finditer(blk[:300]):
            fio_c = re.sub(r"\s+", " ", cand.group(0)).strip()
            # Отсев ложных «ФИО» из служебных строк (название суда, дескриптор
            # гражданства «Гражданин Российской Федерации» и т.п.)
            if re.search(r"суд|банкрот|заявл|граждан", fio_c, re.IGNORECASE):
                continue
            if is_person_name(fio_c):
                fio = fio_c
                break
        if not fio and m_label:
            # OCR иногда переставляет «От Должника: <ФИО>» «<ФИО> от Должника:»:
            # ФИО оказывается ПЕРЕД меткой. Фолбэк — ищем персональное ФИО в участке
            # непосредственно перед меткой (исключая суд/регион), иначе в applicantName
            # протекал бы «Арбитражный суд …» из генерик-парсера.
            pre = header[max(0, m_label.start() - 150) : m_label.start()]
            for cand in self._SB_FIO_RE.finditer(pre):
                fio_c = re.sub(r"\s+", " ", cand.group(0)).strip()
                if re.search(r"суд|банкрот|заявл|област|район|граждан", fio_c, re.IGNORECASE):
                    continue
                if is_person_name(fio_c):
                    fio = fio_c
                    break
        if not fio:
            return None

        fields["applicantName"] = fio
        fields["debtorName"] = fio
        base = re.sub(self._NAME_PREFIX_RE, "", fio, flags=re.IGNORECASE).strip() or fio
        for case_key, conv in (
            ("applicantNameGenitive", self._convert_name_to_genitive),
            ("applicantNameDative", self._convert_name_to_dative),
            ("applicantNameAccusative", self._convert_name_to_accusative),
            ("applicantNameInstrumental", self._convert_name_to_instrumental),
        ):
            try:
                fields[case_key] = conv(base) or fio
            except Exception:
                fields[case_key] = fio


        pre_passport = re.split(r"паспорт|выдан", blk, flags=re.IGNORECASE)[0]
        birth = (
            re.search(r"дата\s*(?:и\s*место\s*)?рождения[:;\s]*(\d{2}\.\d{2}\.\d{4})", blk, re.IGNORECASE)
            or re.search(r"(\d{2}\.\d{2}\.\d{4})\s*(?:г\.?\s*р\.?|года\s+рождения)", blk, re.IGNORECASE)
            or re.search(
                r"(\d{1,2}\s+(?:январ|феврал|март|апрел|ма[яй]|июн|июл|август|сентябр|октябр|ноябр|декабр)[а-яё]*"
                r"\s+\d{4})\s+года\s+рождения",
                blk, re.IGNORECASE,
            )
            or re.search(r"\b(\d{2}\.\d{2}\.\d{4})\s*г?\b", pre_passport)
        )
        if birth:
            # Дата прописью («14 июня 1990») дд.мм.гггг: поле фронта — type=date.
            fields["birthDate"] = (
                self._normalize_obl_date(birth.group(1)) or birth.group(1).strip()
            )
        else:
            fields.pop("birthDate", None)

        bp = re.search(
            r"место\s+рождения[:;\s]*(?:\d{2}\.\d{2}\.\d{4}\s*)?(.{3,140}?)(?=\s*(?:[;]|" + self._SB_FIELD_STOP + r"))",
            blk, re.IGNORECASE | re.DOTALL,
        )
        if bp:
            place = re.sub(r"\s+", " ", bp.group(1)).strip(" ,;.")
            place = re.sub(r"(?<=[а-яё])-\s+(?=[а-яё])", "-", place)  # перенос «р-\nна» «р-на»
            if len(place) >= 3:
                fields["birthPlace"] = place
        elif fields.get("birthPlace"):
            fields.pop("birthPlace", None)

        snils = re.search(r"снилс[:\s]*([\d][\d\-\s]{9,15}\d)", blk, re.IGNORECASE)
        if snils:
            fields["snils"] = re.sub(r"\s+", " ", snils.group(1)).strip()
        else:
            fields.pop("snils", None)
        inn = re.search(r"\bинн\b[:\s]*(\d{10,12})", blk, re.IGNORECASE)
        if inn:
            fields["inn"] = inn.group(1)
            fields["companyInn"] = inn.group(1)
        else:
            fields.pop("inn", None)
            fields.pop("companyInn", None)

        addr = re.search(
            r"(?:адрес\s+(?:мест[аом]*\s+)?регистрации|место\s+жительства\s+по\s+регистрации|"
            r"мест[ао]\s+регистрации|адрес)\s*[:\s]\s*(.{5,240}?)(?=\s*" + self._SB_FIELD_STOP + r")",
            blk, re.IGNORECASE | re.DOTALL,
        )
        addr_val = None
        if addr:
            addr_val = re.sub(r"\s+", " ", addr.group(1)).strip(" ,;")
        else:
            # Фолбэк: первая индекс-строка блока.
            m6 = re.search(r"(\d{6}\s*,?\s*[^\n]{5,160}(?:\n[^\n]{2,80}){0,2})", blk)
            if m6:
                addr_val = re.sub(r"\s+", " ", m6.group(1)).strip(" ,;")
        if addr_val:
            fields["applicantAddress"] = addr_val
            fields["debtorAddress"] = addr_val
        else:
            fields.pop("applicantAddress", None)


        fields["entityType"] = "individual"
        for k in ("ogrn", "ogrnip", "companyOgrn",
                  "creditorName", "creditorAddress", "creditorInn", "creditorOgrn",
                  "thirdPartyName", "thirdPartyAddress", "thirdPartyInn",
                  "thirdPartyBirthDate", "thirdPartySnils"):
            fields.pop(k, None)

        logger.info(f" Самобанкротство: должник из шапки «{fio}», реквизиты строго из его блока")
        return self._extract_sb_third_parties(text, title_pos)

    def _extract_sb_third_parties(self, text: str, title_pos: int):
        """Третьи лица из шапки самобанкрота: «Третьи лица[, не заявляющие …]:».

        Формат блока: <название организации / ФИО> <адрес с индексом> …
        (повторяется). У третьего лица-персоны после адреса могут идти реквизиты
        (дата рождения, паспорт, СНИЛС, ИНН) — выделяем их из хвоста адреса.
        Возвращает список словарей или None (блока нет).
        """
        # Уточнение метки («, не заявляющие самостоятельных требований») может
        # быть разорвано переносом строки — допускаем \n до двоеточия.
        m = re.search(r"треть\w+\s+лиц[^:]{0,80}?:", text, re.IGNORECASE)
        if not m:
            return None
        end = m.end() + 1600
        # Стопы блока: уполномоченный орган, заголовок заявления, начало таблицы
        # обязательств (табличная шапка «№ п/п Содержание обязательства…»).
        for stop_rx in (
            re.compile(r"уполномоченн\w+\s+орган", re.IGNORECASE),
            re.compile(r"№\s*п/п|содержание\s+обязательства|итого\s*:", re.IGNORECASE),
        ):
            m_s = stop_rx.search(text, m.end())
            if m_s:
                end = min(end, m_s.start())
        for rx in self._SB_TITLE_RES:
            m_t = rx.search(text, m.end())
            if m_t:
                end = min(end, m_t.start())
        block = text[m.end(): end]


        if block.count("\n") < 2:
            block = re.sub(r"\s(?=\d{6}\b)", "\n", block)
            block = re.sub(
                r"\s(?=(?:УФНС|ИФНС|ФНС|Межрайонн|Управлен|Отдел|Инспекц|Отделени|Фонд|ООО|АО|ПАО|ЗАО|НАО)\b)"
                r"|\s(?=[А-ЯЁ][а-яё-]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:вич|вна|чна|ична)\b)",
                "\n", block,
            )


        _ORG_START = re.compile(
            r"^(?:межрайонн\w*|управлен\w*|отдел\w*|инспекц\w*|отделени\w*|фонд\w*|союз\w*)"
            r"|^(?:уфнс|ифнс|фнс|ооо|ао|пао|зао|нао)\b",
            re.IGNORECASE,
        )
        # Реквизитные строки персоны-третьего лица — продолжение записи, не новое имя.
        _REQ_LINE = re.compile(
            r"^(?:место\s+рождения|паспорт|выдан|снилс|инн|код\s+подразделения|"
            r"дата\s+рождения|\d{2}\.\d{2}\.\d{4})",
            re.IGNORECASE,
        )
        # Кредитные организации — это КРЕДИТОРЫ, в третьих лицах их не бывает
        # (случай слипшихся колонок таблицы «Кредиторы: | Третьи лица:»).
        _CREDIT_ORG = re.compile(r"банк|\bмкк\b|\bмфк\b|\bпко\b|\bкпк\b", re.IGNORECASE)

        entries = []
        cur_name: list = []
        cur_tail: list = []
        in_addr = False

        def flush():
            nonlocal cur_name, cur_tail, in_addr
            name = re.sub(r"\s+", " ", " ".join(cur_name)).strip(" ,;")
            tail = re.sub(r"\s+", " ", " ".join(cur_tail)).strip(" ,;")
            if name and len(name) >= 5 and not _CREDIT_ORG.search(name):
                entry = {"name": name}
                # Реквизиты персоны в хвосте адреса — отделяем от самого адреса.
                m_req = re.search(
                    r"\d{2}\.\d{2}\.\d{4}\s*(?:г\.?\s*р|года\s+рождения)|паспорт|снилс|инн",
                    tail, re.IGNORECASE,
                )
                addr = tail[: m_req.start()] if m_req else tail
                req = tail[m_req.start():] if m_req else ""
                if addr.strip(" ,;"):
                    entry["address"] = addr.strip(" ,;")
                bd = re.search(r"(\d{2}\.\d{2}\.\d{4})\s*(?:г\.?\s*р|года\s+рождения)", req, re.IGNORECASE)
                if bd:
                    entry["birthDate"] = bd.group(1)
                sn = re.search(r"снилс[:\s]*([\d][\d\-\s]{9,15}\d)", req, re.IGNORECASE)
                if sn:
                    entry["snils"] = re.sub(r"\s+", " ", sn.group(1)).strip()
                inn_m = re.search(r"\bинн\b[:\s]*(\d{10,12})", req, re.IGNORECASE)
                if inn_m:
                    entry["inn"] = inn_m.group(1)
                entries.append(entry)
            cur_name, cur_tail, in_addr = [], [], False

        for ln in (s.strip() for s in block.splitlines()):
            if not ln:
                continue
            if re.match(r"\d{6}\b", ln):
                in_addr = True
                cur_tail.append(ln)
                continue
            # «Имя и адрес одной строкой» («ПАО СБЕРБАНК 117312, г Москва…») —
            # начало НОВОЙ записи со встроенным адресом.
            m_inline = re.search(r"\s(\d{6}\b.*)$", ln)
            if m_inline and not _REQ_LINE.match(ln):
                flush()
                cur_name.append(ln[: m_inline.start()])
                cur_tail.append(m_inline.group(1))
                in_addr = True
                continue
            if in_addr:
                if _REQ_LINE.match(ln):
                    cur_tail.append(ln)  # реквизиты персоны — продолжение записи
                elif _ORG_START.match(ln) or (not re.search(r"\d", ln) and self._SB_FIO_RE.match(ln)):
                    flush()
                    cur_name.append(ln)
                elif re.search(r"\d", ln) or not re.match(r"[А-ЯЁ]{2,}", ln):
                    cur_tail.append(ln)  # хвост адреса
                else:
                    flush()
                    cur_name.append(ln)
            else:
                cur_name.append(ln)
        flush()
        return entries

    def _apply_fns_authority(self, fields: Dict[str, Any], text: str) -> None:
        """Заявления уполномоченного органа (ФНС) о банкротстве/включении в РТК.

        В таких заявлениях кредитор и заявитель — налоговый орган (ФНС России /
        Межрайонная ИФНС), указанный в шапке-бланке БЕЗ метки «Кредитор:/Заявитель:»,
        поэтому общий label-anchored парсер его не берёт, а заявителем ошибочно
        становится должник-физлицо (из «…требований кредитора … <ФИО должника>»).
        Детектор срабатывает ТОЛЬКО при бланке ФНС в шапке; должника не трогает.
        """
        if not text:
            return

        header_fns = bool(
            re.search(r"НАЛОГОВ\w+\s+СЛУЖБ", text[:400], re.IGNORECASE)
            or re.search(r"ЗАЯВЛЕНИ\w+\s+УПОЛНОМОЧЕНН\w+\s+ОРГАН", text[:700], re.IGNORECASE)
        )

        body_fns = bool(
            re.search(r"\b[ИУ]ФНС\b", text, re.IGNORECASE)
            and re.search(r"обязательн\w+\s+платеж", text, re.IGNORECASE)
        )
        if not (header_fns or body_fns):
            return

        _OPF = r"\b(?:ООО|АО|ПАО|ЗАО|ОАО|ПКО|НАО|Банк)\b"
        cur_cred = fields.get("creditorName") or ""
        if "фнс" not in cur_cred.lower() and re.search(_OPF, cur_cred, re.IGNORECASE):
            return

        # Имя налогового органа собираем слоями (см. _build_fns_creditor); если ни
        # один шаблон инспекции/управления не сработал — безопасный общий «ФНС России».
        creditor = self._build_fns_creditor(text)


        built_full = "в лице" in creditor.lower()
        if "фнс" not in cur_cred.lower() or (built_full and "в лице" not in cur_cred.lower()):
            fields["creditorName"] = creditor
        cred_now = fields.get("creditorName") or ""

        # НОВОЕ (справочник ФНС): выверенный юр-адрес КРЕДИТОРА в блок «Данные о
        # кредиторе» поле creditorAddress. Заполняем ТОЛЬКО его и ТОЛЬКО когда
        # кредитор — налоговый орган; адресов заявителя/должника не касаемся (иначе
        # юр-адрес инспекции подмешивается в адрес должника). Справочник fns_registry
        # находит адрес по имени органа — он инвариантен к тому, как оформлена шапка
        # конкретного заявления (и работает, даже если адреса в шапке нет).
        # Подсказка для дизамбигуации ТОРМ (один № инспекции обслуживает несколько
        # городов): шапка ДО блока «Должник», чтобы город/индекс инспекции не спутать
        # с городом должника — реестр выберет кандидата по городу/индексу из шапки.
        # Общий для должника-физлица и должника-ЮЛ — поэтому ДО развилки по типу лица.
        if "фнс" in (cred_now or "").lower():
            _dpos = re.search(r"Должник", text, re.IGNORECASE)
            _hint = text[: _dpos.start()] if _dpos else text[:1800]
            reg_addr = None
            try:
                reg_addr = _fns_reg.resolve_address(cred_now, _hint)
            except Exception as exc:  # справочник недоступен — тихо пропускаем
                logger.warning(f"Реестр ФНС не ответил: {exc}")
            if reg_addr:
                fields["creditorAddress"] = reg_addr
                # Адрес инспекции — из справочника, а не из разбора текста: это
                # независимое подтверждение, юристу его перепроверять не нужно.
                fields.setdefault(self._PROVENANCE_KEY, {})["creditorAddress"] = (
                    field_contract.SOURCE_REGISTRY
                )


        _PN = r"[А-ЯЁ][А-ЯЁа-яё]+(?:-[А-ЯЁ][А-ЯЁа-яё]+)?(?:\s+[А-ЯЁ][А-ЯЁа-яё]+){2}"

        # НАДЁЖНЫЙ УПРАВЛЯЮЩИЙ ФНС по якорю. Позиционный managerName-парсер спотыкается
        # об OCR-артефакт формы «Финансовый управляющий: На № <ФИО>» и о строчную
        # «утверждён» перед ФИО ФИО управляющего теряется (Мишунин). Якорь
        # «…управляющим утвержд(ён/ена) <ФИО>» либо «управляющий: [На №] <ФИО>» даёт
        # номинатив надёжно. Заполняем ТОЛЬКО если текущее значение не похоже на ФИО.
        if not is_person_name((fields.get("managerName") or "").strip()):
            _mm = re.search(r"управляющ\w*\s+утвержд[её]н\w*\s+(" + _PN + r")", text, re.IGNORECASE)
            if not _mm:
                _mm = re.search(r"управляющ\w*\s*:\s*(?:На\s*[№N]\s*)?(" + _PN + r")", text, re.IGNORECASE)
            if _mm:
                cand_m = _normalize_fio(_mm.group(1).strip())
                if is_person_name(cand_m) and "фнс" not in cand_m.lower():
                    fields["managerName"] = cand_m


        if is_person_name(fields.get("managerName") or ""):
            _ma = re.search(
                r"управляющ\w*[:\s][\s\S]{0,80}?Адрес(?:\s+для\s+корреспонденции)?\s*:\s*"
                r"(\d{6}[^\n]*?)"
                r"(?=\s*(?:Должник|ИНН|ОГРН|ОГРНИП|СНИЛС|КПП|Дело|Тел|e-?mail|www|№|$)|\n)",
                text, re.IGNORECASE,
            )
            if _ma:
                _madr = re.sub(r"\s+", " ", _ma.group(1)).strip().rstrip(" ,;")
                if re.search(r"[А-Яа-яЁё]{3}", _madr):
                    fields["managerAddress"] = _madr


        if self._apply_fns_legal_debtor(fields, text):
            return

        fns_debtor = None
        _dm = re.search(r"(" + _PN + r")\s+обратил", text)
        if not _dm:
            _dm = re.search(r"Должник\w*\s*:\s*(" + _PN + r")\s*(?:ИНН|ОГРН)", text, re.IGNORECASE)
        if _dm:
            cand = _normalize_fio(_dm.group(1).strip())  # каноничный Titlecase (снимает CAPS)
            if is_person_name(cand) and "фнс" not in cand.lower():
                fns_debtor = cand
        mgr_norm = _normalize_fio((fields.get("managerName") or "").strip())
        cur_debt = (fields.get("debtorName") or "").strip()
        cur_debt_norm = _normalize_fio(cur_debt)
        same_as_mgr = bool(mgr_norm) and cur_debt_norm == mgr_norm
        if fns_debtor and _normalize_fio(fns_debtor) != mgr_norm:
            # Якорь важнее позиционного значения: ставим его, если текущий должник —
            # дубль управляющего, не-ФИО (мусор) или просто отличается от якорного.
            if same_as_mgr or not is_person_name(cur_debt) or _normalize_fio(fns_debtor) != cur_debt_norm:
                fields["debtorName"] = fns_debtor
        elif same_as_mgr:
            # Якоря нет, но должник = управляющий — это заведомо ошибка: убираем дубль,
            # чтобы в акт не ушло ФИО управляющего как должника.
            fields.pop("debtorName", None)


        appl = fields.get("applicantName") or ""
        debt_name = (fields.get("debtorName") or "").strip()
        if (debt_name and debt_name != appl and "фнс" not in debt_name.lower()
                and is_person_name(debt_name)):
            fields["applicantName"] = debt_name
            base = re.sub(self._NAME_PREFIX_RE, "", debt_name, flags=re.IGNORECASE).strip() or debt_name
            for case_key, conv in (
                ("applicantNameGenitive", self._convert_name_to_genitive),
                ("applicantNameDative", self._convert_name_to_dative),
                ("applicantNameAccusative", self._convert_name_to_accusative),
                ("applicantNameInstrumental", self._convert_name_to_instrumental),
            ):
                try:
                    fields[case_key] = conv(base) or debt_name
                except Exception:
                    fields[case_key] = debt_name
        # Адреса ФНС-заявления. У заявителя-ФНС нет метки «Адрес:», поэтому общий
        # парсер кладёт в applicantAddress адрес ДОЛЖНИКА (или суда) — разводим их:
        #   • debtorAddress «Должник … Адрес: …», иначе — из applicantAddress (общий
        #     парсер часто кладёт туда именно адрес должника);
        #   • applicantAddress юр-адрес инспекции из шапки (перед Телефон/www.nalog).
        _STREET = r"ул|пр\.|просп|проспект|пер|переул|д\.|дом|улиц|ст-ца|стан|мкр|кв\."

        def _clean_addr(s: str) -> str:
            s = re.sub(r"\s+", " ", s).strip(" ,;")
            return re.split(r"\b(?:ИНН|ОГРН|ОГРНИП|СНИЛС|КПП)\b", s, flags=re.IGNORECASE)[0].strip(" ,;")


        _ADDR_MARK = (
            r"(?:\d{6}|\bг\.|\bгор\b|\bгород|\bобл\b|\bобласт|\bр-?н\b|\bрайон|\bул\.|"
            r"\bулиц|\bд\.|\bдом\b|\bкв\.|\bкорп|\bпос\b|\bпос[её]лок|\bст-ца|\bстаниц|"
            r"\bпер\.|\bпр-кт|\bпросп|\bмкр|\bхутор|\bс\.\s|\bсело|\bдеревн)"
        )
        _ADDR_CONT = r"(?:\n(?=[^\n]*" + _ADDR_MARK + r")(?![^\n]*:)[^\n]+){0,3}"
        debt_addr = None
        dm = re.search(
            r"Должник\w*[:\s][\s\S]{0,90}?\bАдрес[:\s]*([0-9А-ЯЁ][^\n]+" + _ADDR_CONT + r")",
            text, re.IGNORECASE,
        )

        dm_bare = None
        if not dm:
            dm_bare = re.search(
                r"Должник\w*\s*:\s*[\s\S]{0,200}?(\d{6}\s*,\s*[А-ЯЁ][^\n]+" + _ADDR_CONT + r")",
                text, re.IGNORECASE,
            )

        def _is_addr(s: str) -> bool:
            s = (s or "").strip()
            if not (8 <= len(s) <= 160) or not re.search(r"\d", s):
                return False
            # Стоп-слова заявления: если есть — это проза, а не адрес.
            if re.search(r"руководству|федеральн\w+\s+закон|уведомл|задолженност|приложен|"
                         r"направлен|уплач|несостоятельн|банкротств|\bстать\w+|\bст\.?\s*\d|"
                         r"в\s+лице|ликвидатор|сообщени|о\s+намерении|обратил|обратиться\s+в\s+суд",
                         s, re.IGNORECASE):
                return False
            # 1) Почтовый индекс — однозначный признак адреса (покрывает большинство).
            if re.search(r"\b\d{6}\b", s):
                return True
            # 2) Тип адресного объекта из ШИРОКОГО спектра (улицы/площади/наб/шоссе/туп/

            if re.search(r"\b(?:ул|улиц\w*|пер|переул\w*|пр-кт|пр-т|просп\w*|проспект|б-р|"
                         r"бульвар|наб|набережн\w*|пл|площад\w*|шоссе|туп|тупик|алле\w*|"
                         r"лини\w*|проезд|тракт|кв-л|квартал|мкр|микрорайон|городок|"
                         r"пос|пос[её]лок|село|сельсовет|деревн\w*|станиц\w*|ст-ца|хутор|"
                         r"аул|слобод\w*|город|гор|пгт|снт|днп)\b", s, re.IGNORECASE):
                return True
            # 3) Дом/строение/корпус/квартира с номером вплотную («д. 5», «дом 14»,
            #    «стр 3», «кв. 43») — адресный хвост без явного типа улицы/НП.
            return bool(re.search(r"\b(?:д|дом|влд|владени\w*|стр|строени\w*|корп|корпус|к|кв|"
                                  r"уч|участок)\b\.?\s*\d", s, re.IGNORECASE))

        cur_aa = (fields.get("applicantAddress") or "").strip()
        cand = None
        if dm:
            cand = _clean_addr(dm.group(1))
        elif dm_bare:
            cand = _clean_addr(dm_bare.group(1))
        if cand and _is_addr(cand):
            debt_addr = cand
        # ПОДСТРАХОВКА (гибрид): regex промахнулся ИЛИ вернул не-адрес достаём адрес
        # морфологически через Natasha AddrExtractor из окна блока «Должник:». Работает
        # независимо от раскладки (индекс в начале строки / после ОГРНИП / без метки),
        # т.е. страхует от невиданных макетов — чего конечным числом регулярок не выразить.
        if not debt_addr:
            _dblk = re.search(r"Должник\w*\s*:", text, re.IGNORECASE)
            if _dblk:
                nat = _nlp.complete_address(text[_dblk.end(): _dblk.end() + 300])
                if nat:
                    nat = _clean_addr(nat)
                    if _is_addr(nat):
                        debt_addr = nat
        if (not debt_addr and cur_aa and re.search(r"\b\d{6}\b", cur_aa)
                and re.search(_STREET, cur_aa, re.IGNORECASE)
                and not re.search(r"Неглинн|Станиславског", cur_aa, re.IGNORECASE)):
            # Метки «Должник Адрес:» нет, но в applicantAddress лежит адрес с улицей и
            # индексом (не центральный ФНС, не суд) — это фактически адрес должника.
            debt_addr = cur_aa
        if debt_addr and not (fields.get("debtorAddress") or "").strip():
            fields["debtorAddress"] = debt_addr


        if debt_addr:
            # Не затираем более полный уже найденный адрес (общий _collect_block_address
            # часто собирает многострочный адрес точнее): берём более длинный валидный.
            cur_aa2 = (fields.get("applicantAddress") or "").strip()
            if cur_aa2 and _is_addr(cur_aa2) and len(cur_aa2) > len(debt_addr) and debt_addr in cur_aa2:
                fields["applicantAddress"] = cur_aa2
            else:
                fields["applicantAddress"] = debt_addr
        else:
            aa = (fields.get("applicantAddress") or "").strip()
            creda = (fields.get("creditorAddress") or "").strip()
            if aa and (not _is_addr(aa) or (creda and aa == creda)
                       or re.search(r"Неглинн|Станиславског|www\.nalog|Адрес\s+для", aa, re.IGNORECASE)):
                fields.pop("applicantAddress", None)


    _FNS_ORG = (
        r"(?:Общество\s+с\s+ограниченной\s+ответственностью|Публичное\s+акционерное\s+общество|"
        r"Закрытое\s+акционерное\s+общество|Открытое\s+акционерное\s+общество|"
        r"Непубличное\s+акционерное\s+общество|Акционерное\s+общество|"
        r"ООО|ПАО|ЗАО|ОАО|НАО|АО)\s*«[^»]{1,80}»"
    )

    def _apply_fns_legal_debtor(self, fields: Dict[str, Any], text: str) -> bool:
        """Должник-ЮЛ в заявлении уполномоченного органа (ФНС). Возвращает True, если
        ветка отработала (тогда физлицо-логика вызывающего пропускается).

        В ФНС-заявлении против ЮЛ нет ни блока «Должник:», ни метки «Адрес:» — должник
        вводится оборотом «Общество … «X» ИНН … ОГРН … (далее – должник, ООО «X»)
        зарегистрировано <дата> по адресу: <адрес>.». Шапка как источник НЕНАДЁЖНА:
        она двухколоночная (бланк инспекции слева, адресат справа), и конвертер рвёт
        адрес должника между колонками («…г. Калуга, ул.» | «Труда, д. 29А»). Поэтому
        якоримся на предложение о регистрации — оно одноколоночное и есть всегда.
        """
        if not text:
            return False


        org = None
        for rx in (
            r"\(\s*далее\s*[–—-]\s*должник\w*\s*,\s*(" + self._FNS_ORG + r")\s*\)",
            r"призна(?:ть|нии)\s+(" + self._FNS_ORG + r")\s+несостоятельн",
        ):
            m = re.search(rx, text, re.IGNORECASE)
            if m:
                org = re.sub(r"\s+", " ", m.group(1)).strip()
                break
        if not org:
            return False


        qm = re.search(r"«([^»]{1,80})»", org)
        if not qm:
            return False
        quoted = re.escape(qm.group(1))

        req = re.search(
            r"«" + quoted + r"»\s*ИНН\s*(?P<inn>\d{10})(?:\s*КПП\s*\d{9})?\s*ОГРН\s*(?P<ogrn>\d{13})",
            text, re.IGNORECASE,
        )
        debtor_inn = req.group("inn") if req else ""
        debtor_ogrn = req.group("ogrn") if req else ""


        short = self._strip_ooo_prefix(org) or org
        fields["debtorName"] = short
        fields["legalShortName"] = short
        fields["applicantName"] = org

        for case_key in ("applicantNameGenitive", "applicantNameDative",
                         "applicantNameAccusative", "applicantNameInstrumental"):
            fields[case_key] = org

        if debtor_inn:
            fields["inn"] = debtor_inn
        if debtor_ogrn:
            fields["ogrn"] = debtor_ogrn

        for cred_key, debtor_val in (("creditorInn", debtor_inn), ("creditorOgrn", debtor_ogrn)):
            cur = re.sub(r"\D", "", str(fields.get(cred_key) or ""))
            if debtor_val and cur == debtor_val:
                fields.pop(cred_key, None)

        # Адрес — из предложения о регистрации, до конца абзаца. Хвост чистим
        # общими правилами (§J.3): при одностраничной PDF docx склейке к адресу
        # прилипает проза следующего предложения.
        addr_m = re.search(
            r"«" + quoted + r"»[^\n]{0,150}?зарегистрирован\w*\s+"
            r"(?:\d{1,2}[.,]\d{1,2}[.,]\d{4}\s+)?по\s+адресу\s*:\s*([^\n]{10,300})",
            text, re.IGNORECASE,
        )
        if addr_m:
            addr = re.sub(r"\s+", " ", addr_m.group(1)).strip()
            addr = re.split(r"\b(?:ИНН|ОГРН|КПП)\b", addr, flags=re.IGNORECASE)[0]
            addr = self._truncate_glued_address(addr).strip(" .,;")
            if 10 <= len(addr) <= 200 and re.search(r"\d", addr):
                fields["debtorAddress"] = addr
                fields["applicantAddress"] = addr

        # Тип лица — штатным детектором по обновлённым именам («ООО» в debtorName даёт
        # legal). Раньше он видел мусорного должника и решал individual в акт шли
        # рекомендации для физлица.
        entity = self.detect_entity_type(fields)
        if entity:
            fields["entityType"] = entity
        logger.info(f"ФНС против ЮЛ: должник={org}, тип лица={fields.get('entityType')}")
        return True


    _SRO_LEGAL_HEAD = r"(?:Союз\w*|Ассоциаци\w*|Некоммерческ\w*|НПС?|ААУ|СРО|Саморегулируем\w*|Крымск\w*|Российск\w*)"

    _SRO_ANCHOR_RE = re.compile(
        r"(?:из\s+числа\s+членов|членств[оа]|член(?:а|ом|ов)?|"
        r"саморегулируемо[йя]\s+организаци[ияй])"
        r"\s*[:\-–—.]?\s*"  # «.» терпит склейку «членов.НП» (pdf2docx съел пробел)
        r"(" + _SRO_LEGAL_HEAD + r"\b[^()]{2,200}?)"

        r"(?=\s*[()\d]|[;.]\s+[А-ЯЁ]|\n\s*\n|$)",
        re.IGNORECASE,
    )
    _SRO_HEAD_RE = re.compile(r"^\s*" + _SRO_LEGAL_HEAD + r"\b", re.IGNORECASE)

    @classmethod
    def _extract_sro_mention(cls, text: str) -> Optional[str]:
        """Первый фрагмент-упоминание СРО из текста (по якорям членства): начинается
        с правовой формы, закрыт кавычками. Многострочное имя схлопывается, хвост
        реквизитов обрезан. Берём первый, что опознан реестром либо закрыт кавычками."""
        from sro_registry import resolve_sro
        if not text:
            return None
        for m in cls._SRO_ANCHOR_RE.finditer(text):
            raw = re.sub(r"\s+", " ", m.group(1)).strip(" ,.;-—–")
            # Хвостовая метка реквизитов без значения («… «Созидание» ОГРН:») —
            # перезахват границы; обрезаем.
            raw = re.sub(r"[\s,;.]*(?:ОГРН|ИНН|КПП|дата\s+рег\w*|адрес)\s*:?\s*$", "",
                         raw, flags=re.IGNORECASE).strip(" ,.;-—–")
            # Дубль имени через дефис («… «Эгида» -Ассоциация … «Эгида»») —
            # оставляем до первого ЗАКРЫТОГО кавычками имени.
            mdup = re.match(r'(.*?[«"][^»"]{2,}[»"])\s*[-–]\s*(?=\w)', raw)
            if mdup:
                raw = mdup.group(1)
            if len(raw) >= 5 and (resolve_sro(raw) or re.search(r"[«\"][^»\"]{3,}[»\"]", raw)):
                return raw
        return None

    def _resolve_manager_sro(self, fields: Dict[str, Any], text: str) -> None:
        """Гарантирует каноничное СРО в `sroName`: находим упоминание, сверяем с
        реестром `sro_data`, подставляем NAIM_FULL.

        Порядок: (1) текущее значение основного парсера опознаётся реестром
        каноничное имя; (2) иначе (пусто/обрезок/мусор) переизвлекаем из текста по
        якорям членства канон или очищенное сырое; (3) если упоминания нет, а
        текущее не похоже на имя СРО (boilerplate) — чистим, чтобы не протекал мусор."""
        from sro_registry import resolve_sro

        cur = (fields.get("sroName") or "").strip()

        # 1) Текущее значение уже опознаётся реестром каноничное имя.
        full = resolve_sro(cur) if cur else None
        if full:
            fields["sroName"] = full
            return

        # 2) Переизвлекаем из текста (обрезок «Ассоциации «Межрегиональная», мусор).
        raw = self._extract_sro_mention(text)
        if raw:
            fields["sroName"] = resolve_sro(raw) or raw
            return

        # 3) Упоминания нет: сырое, не похожее на имя СРО, — не держим.
        if cur and not self._SRO_HEAD_RE.match(cur):
            fields.pop("sroName", None)

    def _fill_manager_address(self, fields: Dict[str, Any], text: str) -> None:
        """Фолбэк адреса управляющего: «… управляющий: <ФИО, м.б. в 2 строки>
        Адрес регистрации: <многострочный адрес>». Срабатывает только если адрес
        ещё не найден; схлопывает переносы строк и обрезает реквизиты."""
        if not text:
            return
        # Уже есть нормальный адрес (с индексом и без реквизитов) — не трогаем.
        _ex = fields.get("managerAddress")
        if _ex and re.search(r"\b\d{6}\b", _ex) and not re.search(r"\b(?:ИНН|СНИЛС|ОГРН|КПП)\b", _ex, re.IGNORECASE):
            return
        m = re.search(
            r"(?:финансов\w+|временн\w+|конкурсн\w+|арбитражн\w+)\s+управляющ\w+"
            r"[\s\S]{0,80}?адрес\s+регистрации[:\s]*([0-9]{6}[\s\S]{0,160}?)"
            r"(?=\n\s*(?:Дело|Тел|Исх|ЗАЯВЛЕНИЕ|ИНН|ОГРН|СНИЛС|№)|\n\s*\n|$)",
            text, re.IGNORECASE,
        )
        if not m:
            # Без метки «Адрес регистрации»: адрес-индекс сразу после реквизитов
            # управляющего («…управляющий: ФИО ИНН… СНИЛС…\n350012, …»).
            m = re.search(
                r"(?:финансов\w+|временн\w+|конкурсн\w+|арбитражн\w+)\s+управляющ\w+"
                r"[\s\S]{0,120}?СНИЛС[^\n]*\n\s*([0-9]{6}[\s\S]{0,140}?)"
                r"(?=\n\s*(?:Реквизиты|Дело|Тел|Исх|ЗАЯВЛЕНИЕ|ИНН|ОГРН|№|член)|\n\s*\n|$)",
                text, re.IGNORECASE,
            )
        if not m and (fields.get("managerName") or "").strip():

            m = re.search(
                r"(?:финансов\w+|временн\w+|конкурсн\w+|арбитражн\w+)\s+управляющ\w+"
                r"[:\s]*[\s\S]{0,120}?(?:^|\n)\s*адрес[^:\n]*:\s*([0-9]{6}[\s\S]{0,160}?)"
                r"(?=\n\s*(?:Дело|Тел|Исх|ЗАЯВЛЕНИЕ|ИНН|ОГРН|СНИЛС|№|Размер|Государственн)|\n\s*\n|$)",
                text, re.IGNORECASE | re.MULTILINE,
            )
        if not m:
            return
        addr = re.sub(r"\s+", " ", m.group(1)).strip()
        addr = re.split(r"\b(?:ИНН|ОГРН|ОГРНИП|СНИЛС|КПП)\b", addr, flags=re.IGNORECASE)[0]
        addr = addr.strip().rstrip(" ,;")
        if re.search(r"[А-Яа-яЁё]{3}", addr):
            fields["managerAddress"] = addr

    # Метка (якорь роли, стоп-метки других сторон) для NLP-достройки адреса.
    # Окно берём ПОСЛЕ якоря и ОБРЫВАЕМ на метке следующей стороны — иначе окно
    # перепрыгивает через контакты роли (напр. «Адрес для корреспонденции: email»)
    # в блок другой стороны и хватает чужой адрес (адрес должника вместо управляющего).
    _ADDR_ROLE_ANCHORS = (
        ("managerAddress",
         r"(?:финансов\w+|временн\w+|конкурсн\w+|арбитражн\w+)\s+управляющ\w+",
         r"Должник|Кредитор|Заявител|Ответчик|Взыскател|Третье\s+лицо|ЗАЯВЛЕНИЕ|Требование|Дело\s*№|В\s+производств"),
        ("applicantAddress",
         r"Заявител\w+\s*[:\-]",
         r"Должник|Кредитор|управляющ|Ответчик|Взыскател|Третье\s+лицо|ЗАЯВЛЕНИЕ|Требование|Дело\s*№|В\s+производств"),
    )

    def _refine_addresses_nlp(self, fields: Dict[str, Any], text: str) -> None:
        """Второстепенный слой: Natasha достраивает адрес по окну роли.

        Срабатывает КОНСЕРВАТИВНО, чтобы не портить рабочий regex:
        - поле пустое заполняем адресом из окна роли;
        - поле есть, но обрезано заменяем ТОЛЬКО если кандидат Natasha строго
          длиннее и текущее значение — его префикс (тогда не теряем нестандартный
          хвост «а/я»/литеру, который AddrExtractor склонен отбрасывать).
        Роль/принадлежность определяет якорь-метка, не NER. VERBATIM-срез исходного
        текста обеспечивает сам `complete_address`.
        """
        if not text:
            return

        def _norm(s: str) -> str:
            return re.sub(r"\s+", " ", s or "").strip().lower()

        for field, anchor, stop in self._ADDR_ROLE_ANCHORS:

            if field == "managerAddress" and not (fields.get("managerName") or "").strip():
                continue
            am = re.search(anchor, text, re.IGNORECASE)
            if not am:
                continue
            # Окно роли: до 400 символов после якоря, обрываем на пустой строке
            # и на метке следующей стороны (чтобы не захватить чужой адрес).
            win = text[am.end():am.end() + 400]
            win = re.split(r"\n\s*\n", win, maxsplit=1)[0]
            sm = re.search(stop, win, re.IGNORECASE)
            if sm:
                win = win[:sm.start()]
            cand = _nlp.complete_address(win)
            if not cand:
                continue
            cand = re.sub(r"\s+", " ", cand).strip()  # схлопываем переносы внутри адреса

            if not re.search(r"\b\d{6}\b", cand) and not re.search(r"\bд\.?\s*\d", cand, re.IGNORECASE):
                continue
            cur = (fields.get(field) or "").strip()
            if not cur:
                fields[field] = cand
            elif len(_norm(cand)) > len(_norm(cur)) and _norm(cand).startswith(_norm(cur)):
                fields[field] = cand


    _ADDR_COMPONENT_RES = (
        ("index", re.compile(r"\b(\d{6})\b")),
        ("region", re.compile(r"\bобл(?:асть|асти|\.)?(?=[\s,])", re.IGNORECASE)),
        ("city", re.compile(r"(?<!тер\.)\b(?:г|гор|город)\b\.?\s?(?=[А-ЯЁ])", re.IGNORECASE)),
        ("street", re.compile(r"\b(?:ул|улица)\b\.?", re.IGNORECASE)),
        ("house", re.compile(r"\b(?:д|дом)\.?\s*№?\s*(\d[\w/-]*)", re.IGNORECASE)),
    )

    @staticmethod
    def _addr_component_value(addr: str, m: "re.Match", kind: str) -> str:
        """Значение адресного компонента (какой именно город/улица/дом)."""
        if kind in ("index", "house"):
            return (m.group(1) or "").casefold()
        # region/city/street: имя — заглавное слово ПОСЛЕ маркера («г. Москва»)
        # либо ПЕРЕД ним («Ростовская обл,», «Ульяновская область»).
        ma = re.match(r"[\s.]*([А-ЯЁ][А-ЯЁа-яё-]{1,})", addr[m.end():])
        if ma:
            return ma.group(1).casefold()
        mb = re.search(r"([А-ЯЁ][А-ЯЁа-яё-]{1,})[\s,]*$", addr[: m.start()])
        return mb.group(1).casefold() if mb else ""

    # Ликвидатор: ФИО из 2-3 titlecase-токенов (Оленченко Олег Игоревич / Иванов Иван).
    _LIQ_PN = r"[А-ЯЁ][А-ЯЁа-яё]+(?:-[А-ЯЁ][А-ЯЁа-яё]+)?(?:\s+[А-ЯЁ][А-ЯЁа-яё]+){1,2}"

    def _detect_liquidation(self, text: str) -> bool:
        """True, если в заявлении есть сведения о ликвидации должника-ЮЛ.

        Сильные (однозначные) сигналы: «ликвидатор…», «ликвидационная комиссия»,
        «ликвидируемого должника». Слабые («решение/стадия/процесс … ликвидации»)
        засчитываются, только если это не «ликвидация задолженности/последствий/
        аварии» (иной смысл слова «ликвидация»). `\\s+` терпимо к переносам слов.
        """
        strong = (r"ликвидатор", r"ликвидационн\w*\s+комисси", r"ликвидируем\w*\s+должник")
        for rx in strong:
            if re.search(rx, text, re.IGNORECASE):
                return True
        weak = (r"(?:решени\w+|постановлени\w+)\s+о\s+ликвидаци",
                r"(?:стади\w+|процесс\w*)\s+ликвидаци")
        for rx in weak:
            for m in re.finditer(rx, text, re.IGNORECASE):
                tail = text[m.end():m.end() + 24]
                if not re.match(r"\s+(?:задолженност|последстви|авари)", tail, re.IGNORECASE):
                    return True
        return False


    _ABSENT_STRONG_RES = (
        # «ввести … процедуру конкурсного производства отсутствующего должника» —
        # ключевая формула просительной части (правило Андрея).
        re.compile(r"конкурсн\w*\s*производств\w*\s*отсутствующ\w*\s*должник", re.IGNORECASE),
        # «§ 2 Банкротство отсутствующего должника», «упрощённая процедура банкротства
        # отсутствующего должника».
        re.compile(r"банкротств\w*\s*отсутствующ\w*\s*должник", re.IGNORECASE),
        # «для ведения процедуры отсутствующего должника».
        re.compile(r"процедур\w*\s*отсутствующ\w*\s*должник", re.IGNORECASE),
        # «у должника имеются признаки отсутствующего должника», «отвечает критериям…».
        re.compile(r"(?:признак\w*|критери\w*)\s*отсутствующ\w*\s*должник", re.IGNORECASE),
        # «признать должника отсутствующим должником».
        re.compile(r"призна\w*\s*[\s\S]{0,40}?отсутствующ\w+\s*должник", re.IGNORECASE),
    )

    _ABSENT_ART230_RE = re.compile(
        r"(?:стать\w+|ст\.?)\s*230\s*(?:[\s\S]{0,60}?(?:банкротств|несостоятельн|127-ФЗ))",
        re.IGNORECASE,
    )

    def _detect_absent_debtor(self, text: str) -> bool:
        """True, если заявление подано в отношении ОТСУТСТВУЮЩЕГО должника-ЮЛ.

        Признак — упрощённая процедура § 2 гл. XI Закона о банкротстве: имущество
        должника заведомо не покрывает судебные расходы либо по счетам год нет
        операций. Проверено на корпусе (75 файлов): формулы срабатывают только на
        заявлениях этого вида.
        """
        if not text:
            return False
        if any(rx.search(text) for rx in self._ABSENT_STRONG_RES):
            return True
        return bool(self._ABSENT_ART230_RE.search(text))


    _DECEASED_STRONG_RES = (
        # «24.02.2019 Заемщик умер, что подтверждается свидетельством о смерти» —
        # типовая формула изложения факта смерти (оба референсных заявления).
        re.compile(r"умер(?:ла)?\s*,?\s*что\s*подтверждается\s*свидетельств\w*\s*о\s*смерти", re.IGNORECASE),
        # «о признании умершего должника несостоятельным (банкротом)».
        re.compile(r"призна\w*\s*[\s\S]{0,40}?умерш\w+\s*должник", re.IGNORECASE),
        # «заявление о признании умершего гражданина банкротом».
        re.compile(r"заявлени\w*\s*о\s*признании\s*умерш\w+", re.IGNORECASE),
    )

    _DECEASED_ART2231_RE = re.compile(
        r"(?:стать\w+|ст\.?)\s*223\s*\.?\s*1\s*(?:[\s\S]{0,60}?(?:банкротств|несостоятельн|127-ФЗ))",
        re.IGNORECASE,
    )

    def _detect_deceased(self, text: str) -> bool:
        """True, если заявление подано в отношении УМЕРШЕГО должника-физлица.

        Проверено на корпусе (77 файлов): формулы срабатывают только на заявлениях
        этого вида, ложных срабатываний нет.
        """
        if not text:
            return False
        if any(rx.search(text) for rx in self._DECEASED_STRONG_RES):
            return True
        return bool(self._DECEASED_ART2231_RE.search(text))


    _DEATH_DATE_BEFORE_RE = re.compile(
        r"(\d{1,2}[.,]\d{1,2}[.,]\d{4})\s*[^\n]{0,40}?\bумер(?:ла)?\b", re.IGNORECASE
    )
    # Обратный порядок: «Заемщик умер 24.02.2019», «умерла 07.11.2020 г.».
    _DEATH_DATE_AFTER_RE = re.compile(
        r"\bумер(?:ла)?\b\s*(?:\w+\s+){0,3}?(\d{1,2}[.,]\d{1,2}[.,]\d{4})", re.IGNORECASE
    )

    _DEATH_CERT_RE = re.compile(
        r"свидетельств\w*\s*о\s*смерти[\s\S]{0,80}?"
        r"\b([IVXLC]{1,7})\s*[-–—]?\s*([А-ЯЁ]{2})\s*(?:№|N)?\s*[-–—]?\s*(\d{6})\b",
        re.IGNORECASE,
    )
    def _extract_death_details(self, text: str) -> Dict[str, str]:
        """Сведения о смерти должника: дата смерти и свидетельство о смерти.

        Вызывается только для заявлений об умершем должнике. ФИО и адрес нотариуса
        здесь НЕ извлекаются — по решению Андрея эти два поля заполняет пользователь
        вручную. Серии свидетельства в референсных заявлениях тоже нет — извлекаем,
        только если банк её всё же указал.
        """
        out: Dict[str, str] = {}
        if not text:
            return out

        m = self._DEATH_DATE_BEFORE_RE.search(text) or self._DEATH_DATE_AFTER_RE.search(text)
        if m:
            out["deathDate"] = m.group(1).replace(",", ".")

        m = self._DEATH_CERT_RE.search(text)
        if m:
            out["deathCertificate"] = "%s-%s № %s" % (
                m.group(1).upper(), m.group(2).upper(), m.group(3)
            )
        return out

    def _extract_liquidator(self, text: str) -> Optional[str]:
        """Извлекает наименование ликвидатора (ФИО физлица или ОПФ организации).

        Слои с приоритетом: (1) ликвидатор-ЮЛ (управляющая организация) наименование =
        организация; (2) ликвидационная комиссия (председатель/руководитель/в составе);
        (3) единственный ликвидатор-физлицо (разные формулировки, вкл. официальную по
        ЕГРЮЛ). ФИО прогоняется через `_normalize_fio` + `is_person_name`.
        """
        pn = self._LIQ_PN
        # Слой 1 — ликвидатор-ЮЛ (управляющая организация): наименование = организация.
        m = re.search(
            r"(?:управляющ\w+\s+организаци\w+\s*\(ликвидатора\)|ликвидатора?)\s*[—–\-]?\s*"
            r"((?:ООО|ОАО|ЗАО|ПАО|НАО|АО)\s*[«\"][^»\"\n]{2,60}[»\"])\s+в\s+лице",
            text, re.IGNORECASE)
        if m:
            return re.sub(r"\s+", " ", m.group(1)).strip()
        # Слой 2 — ликвидационная комиссия.
        for rx in (
            r"(?:председател\w+|руководител\w+)\s+ликвидационн\w+\s+комисси\w+\s+(" + pn + r")",
            r"ликвидационн\w+\s+комисси\w+\s+в\s+составе\s+председател\w+\s+(" + pn + r")",
        ):
            cand = self._liquidator_person(text, rx)
            if cand:
                return cand
        # Слой 3 — единственный ликвидатор-физлицо.
        for rx in (
            r"в\s+лице\s+(?:единственного\s+)?ликвидатора\b\s*[:—–\-]*\s*(" + pn + r")",
            r"ликвидатором\s+назначен\w*\s+(" + pn + r")",
            r"полномочи\w+\s+ликвидатора\s+возложены\s+на\s+(" + pn + r")",
            r"лица,?\s*ответственного\s+за\s+ликвидацию,?\s*(" + pn + r")",
            r"органа,?\s*осуществляющего\s+ликвидацию,?\s*в\s+лице\s+(" + pn + r")",
            r"от\s+имени\s+юридического\s+лица\s*\(ликвидатора\)[,\s]*(" + pn + r")",
            r"ликвидатор\w*\s*:\s*(" + pn + r")",
        ):
            cand = self._liquidator_person(text, rx)
            if cand:
                return cand
        return None

    def _liquidator_person(self, text: str, rx: str) -> Optional[str]:
        """Матч ФИО ликвидатора по паттерну rx нормализация регистра + приведение к
        именительному падежу (грамматика «в лице ликвидатора <кого>» даёт косвенный
        падеж). Возвращает ФИО или None, если не похоже на имя."""
        m = re.search(rx, text, re.IGNORECASE)
        if not m:
            return None
        cand = self._fio_to_nominative(_normalize_fio(m.group(1).strip()))
        return cand if is_person_name(cand) else None


    _MGR_CAND_RE = re.compile(
        r"([А-ЯЁ][А-ЯЁа-яё]+(?:-[А-ЯЁ][А-ЯЁа-яё]+)?(?:\s+[А-ЯЁ][А-ЯЁа-яё]+){2})"
        r"\s*\(\s*ИНН\s*\d{12}\b",
        re.IGNORECASE,
    )

    def _extract_manager_candidate(self, text: str) -> Optional[str]:
        """ФИО кандидата в арбитражные управляющие из ИНИЦИИРУЮЩЕГО заявления.

        Существующие якоря ищут «управляющим утверждён <ФИО>» — прошедшее время, т.е.
        заявления о включении в РТК, где управляющий уже назначен судом. В заявлении о
        признании банкротом управляющего только ПРОСЯТ утвердить, и грамматика другая:
          • «конкурсным управляющим ООО «X» утвердить члена <СРО> (…) Белов Иван
            Алексеевич (ИНН 632127553358, регистрационный номер 22962)» — именительный;
          • «просит в качестве кандидатуры арбитражного управляющего утвердить Панина
            Александра Владимировича (ИНН 644923704326, СНИЛС …)» — родительный.
        Общее у обеих форм — ФИО вплотную перед скобкой с 12-значным ИНН; на него и
        якоримся, а слева требуем контекст утверждения (иначе это чьё-то другое ИНН).
        """
        if not text:
            return None
        for m in self._MGR_CAND_RE.finditer(text):
            left = text[max(0, m.start() - 250): m.start()]
            if not re.search(r"утвердит|кандидатур|управляющ", left, re.IGNORECASE):
                continue
            # Родительный именительный (см. §J.2): «Панина Александра Владимировича».
            cand = self._fio_to_nominative(_normalize_fio(m.group(1).strip()))
            if is_person_name(cand) and "фнс" not in cand.lower():
                return cand
        return None

    def _fio_to_nominative(self, fio: str) -> str:
        """Приводит ФИО к именительному падежу («Иванова Ивана Ивановича» «Иванов
        Иван Иванович»). Несклоняемые/уже-именительные токены не трогаются. Род —
        по подстроке отчества (вич masc, вна femn), чтобы фамилия-омоним склонялась
        верно («Иванова» masc «Иванов», femn «Иванова»)."""
        morph = self._ensure_morph()
        if not morph or not fio:
            return fio
        low = fio.lower()
        gender = "masc" if "вич" in low else ("femn" if "вна" in low else None)
        return " ".join(self._token_to_nominative(morph, t, gender) for t in fio.split())

    def _token_to_nominative(self, morph, tok: str, gender: Optional[str]) -> str:
        if "-" in tok:
            return "-".join(self._token_to_nominative(morph, p, gender) for p in tok.split("-"))
        parses = morph.parse(tok)
        if not parses:
            return tok
        # Предпочитаем разбор как имя собственное (фамилия/имя/отчество).
        p = next((x for x in parses if any(k in str(x.tag) for k in ("Surn", "Name", "Patr"))), parses[0])
        gramm = {"nomn"}
        if gender and "Surn" in str(p.tag):
            gramm.add(gender)
        inflected = p.inflect(gramm) or p.inflect({"nomn"})
        if inflected and inflected.word:
            return self._match_original_case(tok, inflected.word)
        return tok

    def _truncate_glued_address(self, addr: str) -> str:
        """Обрезает склейку двух адресов в одном поле.

        Пример (Корсунов): «347631, обл. Ростовская, …, кв. 61 Кредиторы ООО МКК
        Эквазайм 432071, Ульяновская область, …» — за адресом должника продолжен
        список кредиторов. Правила: (1) слово «кредитор» внутри адреса — обрезка
        по нему; (1a) «мягкий» маркер начала прозы (в лице/ликвидатор/сообщение/
        о намерении/обратиться в суд) — при однострочной PDF docx склейке к адресу
        приклеивается хвост «…ком. 314 В лице ликвидатора: … Сообщение №… о
        намерении обратиться в суд»; (2) повтор адресной категории (индекс/область/
        город/улица/дом) с ДРУГИМ значением — обрезка перед самым ранним повтором.
        Повтор с тем же значением («г. о. город Новочеркасск, г. Новочеркасск …») —
        ФИАС-стиль одного адреса, не склейка.
        """
        # Неразрывные пробелы (PDF docx) обычные, иначе `\s` местами промахивается.
        addr = addr.replace("\xa0", " ")
        m_cred = re.search(r"[\s,;]кредитор\w*", addr, re.IGNORECASE)
        if m_cred:
            addr = addr[: m_cred.start()]
        # Мягкие маркеры конца адреса / начала прозы заявления. «№ N» НЕ маркер —
        # в адресе бывает «дом № 5»; хвост «Сообщение №…» отсекает «сообщени».
        m_soft = re.search(
            r"\s+(?:в\s+лице\b|ликвидатор|председател\w+\s+ликвидационн|сообщени\w*\b|"
            r"о\s+намерении\b|обратил\w*\b|обратиться\s+в\s+суд)",
            addr, re.IGNORECASE)
        if m_soft:
            addr = addr[: m_soft.start()]
        cut = len(addr)
        for kind, rx in self._ADDR_COMPONENT_RES:
            first_val = None
            for m in rx.finditer(addr):
                val = self._addr_component_value(addr, m, kind)
                if first_val is None:
                    first_val = val
                    continue
                if val and first_val and val != first_val:
                    cut = min(cut, m.start())
                    break
        addr = addr[:cut].strip().rstrip(" ,;-")
        # Огрызок следующей записи: ОПФ-хвост («… д. 80 к. 9 ООО ПКО «РСВ»») —
        # в самом адресе организаций-ОПФ не бывает, это начало имени кредитора.
        addr = re.sub(r"\s+(?:ООО|ОАО|ЗАО|ПАО|НАО|АО)\b.*$", "", addr)
        return addr.strip().rstrip(" ,;-")


    _STRICT_CONTRACT = os.environ.get("SBERACT_STRICT_CONTRACT") == "1"

    def _contract_checkpoint(self, fields: Dict[str, Any], step: str) -> None:
        """Инвариант между шагами каскада: словарь всё ещё осмыслен?

        `analyze()` — длинная цепочка правок общего словаря, и шаг N может
        скопировать значение, которое шаг N+5 признает мусором и вычистит — копия
        при этом переживёт чистку (так «Арбитражный суд…» и попал в `loanDebt`
        через `amounts_mixin._reconcile_debt_amounts`). Чекпоинт показывает ШАГ,
        на котором словарь испортился, а не разбирательство через 200 строк.
        """
        issues = field_contract.find_issues(fields)
        if not issues:
            return
        for issue in issues:
            logger.warning(f"Контракт нарушен после шага «{step}»: {issue}")
        if self._STRICT_CONTRACT:
            raise AssertionError(f"Контракт нарушен после шага «{step}»: {issues}")

    def _sanitize_address_fields(self, fields: Dict[str, Any]) -> None:
        """Адресные поля не должны содержать реквизиты (ИНН/ОГРН/ОГРНИП/СНИЛС/КПП):
        обрезаем хвост по первому такому маркеру; если осмысленного адреса (букв)
        не осталось — поле было мусором, удаляем. Затем режем склейку двух
        адресов (повтор индекса/области/города/улицы/дома, «Кредиторы …» внутри)."""
        for k in ("address", "applicantAddress", "creditorAddress", "managerAddress",
                  "thirdPartyAddress", "debtorAddress"):
            v = fields.get(k)
            if not v or not isinstance(v, str):
                continue
            cleaned = re.split(
                r"\b(?:ИНН|ОГРНИП|ОГРН|СНИЛС|КПП)\b", v, flags=re.IGNORECASE
            )[0].strip().rstrip(" ,;-")
            cleaned = self._truncate_glued_address(cleaned)
            if not re.search(r"[А-Яа-яЁё]{3}", cleaned):
                fields.pop(k, None)
            elif cleaned != v:
                fields[k] = cleaned

    def _apply_stacked_finances(self, fields: Dict[str, Any], text: str) -> None:
        """Финансы формата ВТБ «реестр Nл» (сигнатура меток стопкой).

        Просительная: «…в общем размере TOTAL руб … из которых: – AMT руб – КАТ;
        – AMT руб – КАТ». Госпошлина — в шапке «ГОСПОШЛИНА: NN руб» (банкротная,
        в долг не входит). Срабатывает ТОЛЬКО при сигнатуре — не трогает общий парсер.
        """
        if not text or not self._STACKED_SIG_RE.search(text):
            return
        money = r"(\d[\d   ]*[.,]\d{2})"
        # Разбивку берём из ПРОСИТЕЛЬНОЙ части («ПРОШУ:/ПРОСИТ:»), а не из тела,
        # где есть отдельные разбивки по каждому обязательству.
        _pr = re.search(r"\bПРОШУ\s*:|\bПРОСИТ\s*:", text, re.IGNORECASE)
        search_text = text[_pr.start():] if _pr else text
        pm = re.search(
            r"в\s+(?:\w+\s+){0,2}размере\s*" + money + r"\s*руб[^\n]{0,80}?из\s+котор\w+\s*:?(.{0,500})",
            search_text, re.IGNORECASE | re.DOTALL,
        )
        if not pm:
            return
        total = self._fin_amount(pm.group(1))
        body = pm.group(2)
        pr = it = fo = 0.0
        found = False
        brk_vals: Dict[str, list] = {}  # адденды по полям — для тултипа «откуда число»
        for am in re.finditer(money + r"\s*руб[^–\-\n]*[–\-]\s*([^\n;]+)", body):
            val = self._fin_amount(am.group(1))
            lbl = am.group(2).lower()
            if val <= 0:
                continue
            # «основной долг (кредит, проценты)» — это долг (скобки не делают его процентами).
            if "неустой" in lbl or "пени" in lbl:
                fo += val
                brk_vals.setdefault("forfeit", []).append(val)
            elif "основн" in lbl or "долг" in lbl or "ссудн" in lbl:
                pr += val
                brk_vals.setdefault("principalDebt", []).append(val)
            elif "процент" in lbl:
                it += val
                brk_vals.setdefault("interest", []).append(val)
            else:
                continue
            found = True
        if not found:
            return
        comp_sum = pr + it + fo
        if total <= 0 or abs(comp_sum - total) > 1.5:
            # Разбивка не сошлась с общим размером — не перекрываем (страховка).
            return
        # Слой перекрывает значения полей — синхронно перекрываем и разбивку тултипа.
        if brk_vals:
            fields["financeBreakdown"] = {
                k: [self._fin_fmt(v) for v in vals] for k, vals in brk_vals.items()
            }
        else:
            fields.pop("financeBreakdown", None)
        fields["principalDebt"] = self._fin_fmt(pr)
        fields["principalDebt13"] = fields["principalDebt"]
        fields["interest"] = self._fin_fmt(it)
        fields["interest14"] = fields["interest"]
        fields["forfeit"] = self._fin_fmt(fo)
        fields["forfeit15"] = fields["forfeit"]
        fields["penalties"] = self._fin_fmt(0.0)
        fields["totalDebt"] = self._fin_fmt(total)
        fields["debtAmount"] = fields["totalDebt"]
        # Банкротная госпошлина из шапки «ГОСПОШЛИНА: NN руб».
        gm = re.search(r"ГОСПОШЛИНА[:\s]*" + money + r"\s*руб", text, re.IGNORECASE)
        if gm:
            duty = self._fin_fmt(self._fin_amount(gm.group(1)))
            fields["stateDuty16"] = duty
            fields["stateDuty"] = duty
        logger.info(
            "Финансы ВТБ (стопка): осн=%s проц=%s неуст=%s итог=%s пошлина=%s"
            % (fields["principalDebt"], fields["interest"], fields["forfeit"],
               fields["totalDebt"], fields.get("stateDuty16"))
        )

    def extract_fields(self, text: str, document_type: str) -> Dict[str, Any]:
        """
        Извлекает поля на основе паттернов
        """
        extracted_fields = {}
        logger.info(f"Начинаем извлечение полей для типа документа: {document_type}")
        logger.info(f"Длина текста: {len(text)} символов")

        if document_type not in self.patterns:
            logger.warning(f"Тип документа {document_type} не найден в паттернах")
            return extracted_fields

        patterns = self.patterns[document_type]
        logger.info(f"Найдено {len(patterns)} паттернов для извлечения")

        summable_fields = ['principalDebt', 'loanDebt', 'interest', 'penalties', 'forfeit', 'totalDebt']
        multiple_fields = ['contractNumber', 'contractDate', 'obligationType']
        obligations = []

        for pattern_info in patterns:
            field_name = pattern_info["name"]
            field_patterns = pattern_info["patterns"]
            logger.info(f"Обрабатываем поле: {field_name} с {len(field_patterns)} паттернами")

            if field_name in summable_fields:
                if self._extract_summable_field(extracted_fields, text, field_name, field_patterns):
                    continue
            elif field_name in multiple_fields:
                self._extract_multiple_field(extracted_fields, text, field_name, field_patterns)
            else:

                if field_name in ["inn", "ogrn", "ogrnip", "companyInn"]:
                    # ИНН/ОГРН/ОГРНИП должника из блока «Должник:»/«Ответчик:»
                    self._extract_party_inn_ogrn(extracted_fields, text, field_name)
                    continue


                if field_name == "applicantAddress":
                    # Адрес должника из блока «Должник:»/«Ответчик:»
                    if self._extract_party_address(extracted_fields, text, field_name):
                        continue

                # ВАЖНО: Для ИНН, ОГРН, ОГРНИП и companyInn мы уже обработали выше, пропускаем общие паттерны
                if field_name in ["inn", "ogrn", "ogrnip", "companyInn"]:
                    logger.info(f"Пропускаем общие паттерны для {field_name}, так как уже обработали в специальной логике")
                    continue

                for i, pattern in enumerate(field_patterns):
                    logger.info(f"  Паттерн {i+1} для {field_name}: {pattern}")

                    # Специальная обработка для courtName: приоритетно ищем в шапке (первые 2500 символов),
                    # чтобы не подхватить фрагмент из тела документа (юрлицо/ФЛ)
                    if field_name == "courtName":
                        text_header = text[:2500]
                        matches = re.findall(pattern, text_header, re.IGNORECASE)
                        if not matches:
                            matches = re.findall(pattern, text, re.IGNORECASE)
                        for match in matches:
                            if isinstance(match, tuple):
                                match_value = next((part for part in match if part), "")
                            else:
                                match_value = match

                            match_value = (match_value or "").strip()
                            if not match_value:
                                continue

                            cleaned_value = self._clean_court_name(match_value)
                            if not cleaned_value or len(cleaned_value) < 10:
                                continue
                            cl = cleaned_value.lower()
                            if "суд" not in cl and "арбитражн" not in cl:
                                continue
                            # Отбрасываем ссылки на закон: "Согласно п. 2 ст. 7...", "право на обращение в арбитражный суд"
                            if any(cl.startswith(p) or p in cl for p in ("согласно", "п. ", "ст. ", "закона право", "право на обращение")):
                                continue
                            region_words = ("области", "края", "республики", "города", "автономного округа", "автономной области")
                            if not any(r in cl for r in region_words):
                                continue
                            extracted_fields[field_name] = cleaned_value
                            logger.info(f"Extracted {field_name}: {cleaned_value}")
                            break

                        if field_name in extracted_fields:
                            break
                        continue
                    match = re.search(pattern, text, re.IGNORECASE)
                    if match:
                        if match.groups():
                            value = match.group(1).strip()
                        else:
                            value = match.group(0).strip()
                        logger.info(f"  Найдено совпадение: '{value}'")
                        # Значение ЭТОЙ итерации. Без сброса сюда доезжало
                        # cleaned_value от предыдущего поля/паттерна: пустой захват
                        # («требования:» без числа) не проходил фильтр ниже, а
                        # присваивание в конце цикла отрабатывало — и в
                        # requirementsSum уезжал номер дела «А47-7950/2011», который
                        # дальше нормализатор сумм превращал в 4 779 502 011,00.
                        cleaned_value = None
                        if value and len(value) > 2 and value.strip():  # Фильтруем слишком короткие значения и пустые строки
                            cleaned_value = self.clean_extracted_value(value)

                            if field_name == "courtName" and cleaned_value:
                                procedure_type = extracted_fields.get('procedureType', '').lower()
                                procedure_raw = extracted_fields.get('procedureTypeRaw', '').lower()
                                is_deceased = (procedure_type == 'deceased' or
                                             any(keyword in procedure_raw for keyword in ["умер", "умерший", "смерть", "смерти"]))

                                if is_deceased:
                                    cleaned_value = self._clean_court_name_deceased(cleaned_value)
                                else:
                                    cleaned_value = self._clean_court_name(cleaned_value)

                                if not cleaned_value:
                                    continue

                            if field_name == "applicantAddress" and cleaned_value:
                                cleaned_value = re.sub(r'\s*\[9\]\s*\.?', '', cleaned_value)
                                cleaned_value = re.sub(r'\s+', ' ', cleaned_value).strip()

                                cleaned_value = re.sub(r'\([^)]*родительный[^)]*\)', '', cleaned_value)
                                cleaned_value = re.sub(r'\[2\][^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'банкрот[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'процедур[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'реализац[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'имуществ[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'опубликов[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'сайт[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'Единого[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'федерального[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'реестр[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'сведений[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'банкротств[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'сообщение[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'№\s*[0-9]+[^0-9]*', '', cleaned_value)
                                cleaned_value = re.sub(r'\d{1,2}[.,]\d{1,2}[.,]\d{4}[^0-9]*', '', cleaned_value)

                                # Если в адресе есть почтовый индекс (6 цифр) — считаем, что он должен идти первым
                                if re.search(r'\b\d{6}\b', cleaned_value):
                                    cleaned_value = re.sub(r'^[^0-9]*', '', cleaned_value)
                                else:
                                    # Адрес без индекса (как в шапке: "Адрес регистрации: Ростовская область, ...") —
                                    # удаляем только служебные слова "Адрес", "адрес регистрации", "место жительства"
                                    cleaned_value = re.sub(
                                        r'^(Адрес|адрес|адрес\s+регистрации|место\s+жительства|место\s+регистрации)\s*[:\-–—]*\s*',
                                        '',
                                        cleaned_value,
                                        flags=re.IGNORECASE
                                    )
                                cleaned_value = re.sub(r'[^0-9,\s\-а-яёА-ЯЁ\.]+', '', cleaned_value)

                                cleaned_value = re.sub(r'\s+', ' ', cleaned_value).strip()

                                cleaned_value = re.sub(r'банкрот.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'процедур.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'реализац.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'имуществ.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'опубликов.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'сайт.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'Единого.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'федерального.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'реестр.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'сведений.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'банкротств.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'родительный.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'падеж.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'\[2\].*$', '', cleaned_value)
                                cleaned_value = re.sub(r'\[9\].*$', '', cleaned_value)
                                cleaned_value = re.sub(r'сообщение.*$', '', cleaned_value, flags=re.IGNORECASE)
                                cleaned_value = re.sub(r'01\.02\.1883.*$', '', cleaned_value)

                                cleaned_value = re.sub(r'\([^)]*\)', '', cleaned_value)  # Убираем скобки и их содержимое

                                cleaned_value = re.sub(r'^[^0-9]*', '', cleaned_value)

                                cleaned_value = re.sub(r'[^0-9,\s\-а-яёА-ЯЁ\.]', '', cleaned_value)
                                cleaned_value = re.sub(r'\s+', ' ', cleaned_value).strip()

                                if len(cleaned_value) < 20:
                                    address_patterns = [
                                        r'([0-9]{6}[,\s]+[^,\n]+(?:[,\s]+[^,\n]+)*?)(?:\n|$|[,\[]|(?:телефон|дата|огрн|инн|снилс|паспорт|серия|номер|банкрот|процедур|реализац|имуществ|опубликов|сайт|Единого|федерального|реестр|сведений|банкротств|родительный|падеж|\[2\]|\[9\]))',
                                        r'([0-9]{6}[,\s]+[А-ЯЁ][^,\n]+(?:[,\s]+[^,\n]+)*?)(?:\n|$|[,\[]|(?:телефон|дата|огрн|инн|снилс|паспорт|серия|номер))'
                                    ]

                                    for pattern in address_patterns:
                                        matches = re.findall(pattern, text, re.IGNORECASE | re.DOTALL)
                                        if matches:
                                            full_address = matches[0].strip()
                                            full_address = re.sub(r'банкрот.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'процедур.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'реализац.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'имуществ.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'опубликов.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'сайт.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'Единого.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'федерального.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'реестр.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'сведений.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'банкротств.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'родительный.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'падеж.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'\[2\].*$', '', full_address)
                                            full_address = re.sub(r'\[9\].*$', '', full_address)
                                            full_address = re.sub(r'сообщение.*$', '', full_address, flags=re.IGNORECASE)
                                            full_address = re.sub(r'01\.02\.1883.*$', '', full_address)
                                            full_address = re.sub(r'\([^)]*\)', '', full_address)
                                            full_address = re.sub(r'[^0-9,\s\-а-яёА-ЯЁ\.]', '', full_address)
                                            full_address = re.sub(r'\s+', ' ', full_address).strip()

                                        if len(full_address) > len(cleaned_value):
                                            cleaned_value = full_address
                                        break

                        # Не сохраняем creditorName, если подставился мусор ("кредиторов принимаются в установленном законом порядке")
                        if field_name == "creditorName" and cleaned_value and ("принимаются" in cleaned_value or "установленном законом" in cleaned_value):
                            logger.info(f"  Пропуск creditorName (мусорное значение): {cleaned_value[:60]}...")
                            continue
                        if not cleaned_value:
                            # Захват пустой или отбракован очисткой — пробуем следующий
                            # паттерн, а не сохраняем чужое значение.
                            continue
                        # Денежное поле обязано содержать денежный токен: иначе в
                        # сумму уезжают номера дел, ИНН и даты (см. money_token_from_capture).
                        if pattern_info.get("type") == "amount" and not self.money_token_from_capture(cleaned_value):
                            logger.info(f"  Пропуск {field_name}: '{cleaned_value[:40]}' не денежная сумма")
                            continue
                        extracted_fields[field_name] = cleaned_value
                        logger.info(f"Extracted {field_name}: {cleaned_value}")
                        break
                    else:
                        logger.info(f"  Совпадений не найдено")

        if "mortgageCollateralDescription1221" not in extracted_fields or not extracted_fields.get("mortgageCollateralDescription1221"):
            logger.info("Предмет залога [1221] не найден в основных паттернах, пробуем fallback метод...")
            collateral_block = self._extract_mortgage_collateral_block(text)
            if collateral_block:
                extracted_fields["mortgageCollateralDescription1221"] = collateral_block
                logger.info(f"Извлечено mortgageCollateralDescription1221 через fallback метод: {collateral_block[:150]}...")

        collateral_description = extracted_fields.get("mortgageCollateralDescription1221")
        if collateral_description and re.match(r"^Кому\s+выдана\s+", collateral_description.strip(), re.IGNORECASE):
            # "Кому выдана [ФИО]" — не описание залога, а указание получателя; убираем ложное значение
            del extracted_fields["mortgageCollateralDescription1221"]
            collateral_description = None
            logger.info("Удалено ложное описание залога (Кому выдана ...)")
        collaterals_list: List[Dict[str, Any]] = []
        if collateral_description:
            collateral_items = self._split_collateral_items(collateral_description)
            if collateral_items:
                collaterals_list = [self._make_collateral_obj(idx, item_desc)
                                    for idx, item_desc in enumerate(collateral_items)]

        # Довылавливаем залоговые авто/технику по VIN у ОСТАЛЬНЫХ «Обязательство №N»
        # (см. _extract_vin_collateral_items) — дедуп по VIN, ничего не перезаписывает.
        vin_items = self._extract_vin_collateral_items(text)
        if vin_items:
            existing_vins = {c.get("vin") for c in collaterals_list if c.get("vin")}
            for desc in vin_items:
                obj = self._make_collateral_obj(len(collaterals_list), desc)
                if obj.get("vin") and obj["vin"] in existing_vins:
                    continue
                collaterals_list.append(obj)
                if obj.get("vin"):
                    existing_vins.add(obj["vin"])

        if collaterals_list:
            # Сохраняем массив в extracted_fields для передачи во frontend
            extracted_fields["collaterals"] = collaterals_list
            logger.info(f"Создано {len(collaterals_list)} предметов залога")

        obligations = self.extract_obligations(text, extracted_fields)
        if obligations:
            extracted_fields['obligations'] = obligations
            logger.info(f"Добавлено {len(obligations)} обязательств")

            def _fmt_sum(value: float) -> str:
                return f"{value:,.2f}".replace(",", " ").replace(".", ",")

            # Сумма всех обязательств -> totalDebt ([12])
            total_obligation_sum = sum(o.get('amountValue', 0.0) or 0.0 for o in obligations)
            if total_obligation_sum > 0:
                if document_type == "rtk_application":
                    # Для РТК не переопределяем totalDebt, чтобы не ломать сумму из блока "ПРОСИТ СУД".
                    extracted_fields["obligationsTotalDebt"] = _fmt_sum(total_obligation_sum)
                    logger.info(
                        f"Сумма всех обязательств (obligationsTotalDebt, только для справки): {extracted_fields['obligationsTotalDebt']}"
                    )
                else:
                    extracted_fields['totalDebt'] = _fmt_sum(total_obligation_sum)
                    logger.info(f"Сумма всех обязательств (totalDebt/[12]): {extracted_fields['totalDebt']}")

            # Сумма залоговых обязательств -> [0007]
            collateral_sum = sum(o.get('amountValue', 0.0) or 0.0 for o in obligations if o.get('isCollateral'))
            if collateral_sum > 0:
                extracted_fields['ipCollateralClaimAmount'] = _fmt_sum(collateral_sum)
                logger.info(f"Сумма залоговых обязательств ([0007]): {extracted_fields['ipCollateralClaimAmount']}")

            # Сумма неустоек залоговых обязательств -> [0071]
            collateral_penalty_sum = sum(o.get('penalty0071Value', 0.0) or 0.0 for o in obligations if o.get('isCollateral'))
            if collateral_penalty_sum > 0:
                extracted_fields['penalty0071'] = _fmt_sum(collateral_penalty_sum)
                logger.info(f"Сумма неустоек залоговых обязательств ([0071]): {extracted_fields['penalty0071']}")

            ordered_contract_numbers = []
            ordered_contract_dates = []

            # Сохраняем исходный contractNumber, если он был извлечен и содержит буквы (более полный)
            original_contract_number = extracted_fields.get('contractNumber', '')
            has_letters_in_original = bool(original_contract_number and re.search(r'[A-Za-zА-ЯЁ]', original_contract_number))

            for obligation in obligations:
                number = obligation.get('contractNumber')
                date = obligation.get('contractDate')

                if number and number not in ordered_contract_numbers:
                    # Если исходный номер содержит буквы и текущий номер - только цифры,
                    # проверяем, не является ли текущий номер частью исходного
                    if has_letters_in_original and number and not re.search(r'[A-Za-zА-ЯЁ]', number):
                        # Если исходный номер содержит текущий номер как подстроку, используем исходный
                        if number in original_contract_number:
                            continue  # Пропускаем короткий номер
                    ordered_contract_numbers.append(number)
                if date and date not in ordered_contract_dates:
                    ordered_contract_dates.append(date)

            # Если исходный номер содержит буквы и более полный, используем его вместо номеров из обязательств
            if has_letters_in_original and original_contract_number:
                # Проверяем, есть ли в обязательствах номера, которые являются частью исходного
                all_obligation_numbers_are_substrings = all(
                    not re.search(r'[A-Za-zА-ЯЁ]', num) and num in original_contract_number
                    for num in ordered_contract_numbers if num
                ) if ordered_contract_numbers else False

                if all_obligation_numbers_are_substrings or not ordered_contract_numbers:
                    extracted_fields['contractNumber'] = original_contract_number
                    logger.info(f"Используем исходный полный номер договора: {extracted_fields['contractNumber']}")
                else:
                    all_numbers = [original_contract_number] + [n for n in ordered_contract_numbers if n not in original_contract_number]
                    extracted_fields['contractNumber'] = ", ".join(all_numbers)
                    logger.info(f"Сводный список номеров договоров (с приоритетом исходного): {extracted_fields['contractNumber']}")
            elif ordered_contract_numbers:
                extracted_fields['contractNumber'] = ", ".join(ordered_contract_numbers)
                logger.info(f"Сводный список номеров договоров: {extracted_fields['contractNumber']}")

            if ordered_contract_dates:
                extracted_fields['contractDate'] = ", ".join(ordered_contract_dates)
                logger.info(f"Сводный список дат договоров: {extracted_fields['contractDate']}")

            # Подставляем реальные номера договоров в обязательства с плейсхолдером "Договор_1", "Договор_2" и т.д.
            contract_number_raw = extracted_fields.get('contractNumber', '')
            if contract_number_raw and obligations:
                parts = [p.strip() for p in re.split(r'[,;]', contract_number_raw)]
                real_numbers = [
                    p for p in parts
                    if p and p.lower() not in ('заключен', 'заключенный', 'далее')
                    and not re.match(r'^Договор_\d+$', p, re.IGNORECASE)
                    and len(p) >= 2
                ]
                for idx, obl in enumerate(obligations):
                    if not isinstance(obl, dict):
                        continue
                    current = (obl.get('contractNumber') or '').strip()
                    if re.match(r'^Договор_\d+$', current, re.IGNORECASE) and idx < len(real_numbers):
                        obl['contractNumber'] = real_numbers[idx]
                        logger.info(f"Подставлен реальный номер договора в обязательство {idx + 1}: {real_numbers[idx]}")

        if 'requirementsSum' in extracted_fields and extracted_fields['requirementsSum']:
            requirements_sum_value = str(extracted_fields['requirementsSum']).strip()
            if requirements_sum_value and requirements_sum_value not in ["0", "0,00", "0.00", ""]:
                if 'totalDebt' not in extracted_fields or not extracted_fields['totalDebt'] or str(extracted_fields['totalDebt']).strip() in ["0", "0,00", "0.00", ""]:
                    extracted_fields['totalDebt'] = extracted_fields['requirementsSum']
                    logger.info(f"Используем сумму требований как общую сумму долга: {extracted_fields['requirementsSum']}")
        try:
            breakdown_pattern = re.compile(
                r"на\s+сумму\s+([0-9\s,]+)\s*руб[\s\S]{0,200}?из\s+которой[:\s]*"
                r"сумма\s+основного\s+долга\s+([0-9\s,]+)\s*руб"
                r"[\s\S]{0,120}?сумма\s+долга\s+по\s+процентам\s+([0-9\s,]+)\s*руб"
                r"[\s\S]{0,120}?сумма\s+задолженности\s+по\s+просроченным\s+процентам\s+([0-9\s,]+)\s*руб"
                r"[\s\S]{0,160}?расход[аов]*\s+по\s+оплате\s+госпошл[ины]{3}\s*[–—-]?\s*([0-9\s,]+)\s*руб"
                r"[\s\S]{0,80}?(?:ИТОГО\s*[–—-]?\s*([0-9\s,]+)\s*руб)?",
                re.IGNORECASE,
            )
            breakdown_match = breakdown_pattern.search(text)
            if breakdown_match:
                base_sum, principal_b, interest_b, overdue_b, duty_b, total_b = breakdown_match.groups()
                principal_b = self.normalize_amount_value(principal_b or "")
                interest_b = self.normalize_amount_value(interest_b or "")
                overdue_b = self.normalize_amount_value(overdue_b or "")
                duty_b = self.normalize_amount_value(duty_b or "")
                if principal_b:
                    extracted_fields["principalDebt"] = principal_b
                if interest_b:
                    extracted_fields["interest"] = interest_b
                if overdue_b:
                    extracted_fields["forfeit"] = overdue_b
                if duty_b:
                    extracted_fields["stateDuty"] = duty_b
                if total_b:
                    extracted_fields["totalDebt"] = self.normalize_amount_value(total_b)
                logger.info(
                    f"Разобрана сложная формулировка сумм: principal={principal_b}, interest={interest_b}, "
                    f"overdue={overdue_b}, duty={duty_b}, total={extracted_fields.get('totalDebt')}"
                )
        except Exception as e:
            logger.debug(f"Ошибка при разборе сложной формулировки сумм: {e}")


        try:
            breakdown_pattern = re.compile(
                r"на\s+сумму\s+([0-9\s,]+)\s*руб[^.]*?из\s+которой[:\s]*"
                r"сумма\s+основного\s+долга\s+([0-9\s,]+)\s*руб[^,]*,"
                r"\s*сумма\s+долга\s+по\s+процентам\s+([0-9\s,]+)\s*руб[^,]*,"
                r"\s*сумма\s+задолженности\s+по\s+просроченным\s+процентам\s+([0-9\s,]+)\s*руб[^,]*,"
                r"\s*расходы\s+по\s+оплате\s+госпошлин[ыа]\s*[–—-]?\s*([0-9\s,]+)\s*руб[^.]*?"
                r"(?:ИТОГО\s*[–—-]?\s*([0-9\s,]+)\s*руб)?",
                re.IGNORECASE | re.DOTALL,
            )
            breakdown_match = breakdown_pattern.search(text)
            if breakdown_match:
                base_sum, principal_b, interest_b, overdue_b, duty_b, total_b = breakdown_match.groups()
                principal_b = self.normalize_amount_value(principal_b or "")
                interest_b = self.normalize_amount_value(interest_b or "")
                overdue_b = self.normalize_amount_value(overdue_b or "")
                duty_b = self.normalize_amount_value(duty_b or "")
                if principal_b:
                    extracted_fields["principalDebt"] = principal_b
                if interest_b:
                    extracted_fields["interest"] = interest_b
                if overdue_b:
                    extracted_fields["forfeit"] = overdue_b
                if duty_b:
                    extracted_fields["stateDuty"] = duty_b
                if total_b:
                    extracted_fields["totalDebt"] = self.normalize_amount_value(total_b)
                else:
                    # если "ИТОГО" нет, общая сумма = базовая сумма + госпошлина
                    try:
                        base_f = float(self.normalize_amount_value(base_sum or "0").replace(" ", "").replace(",", "."))
                        duty_f = float(self.normalize_amount_value(duty_b or "0").replace(" ", "").replace(",", "."))
                        total_calc = base_f + duty_f
                        formatted_total = f"{total_calc:,.2f}".replace(",", " ").replace(".", ",")
                        extracted_fields["totalDebt"] = formatted_total
                    except Exception:
                        pass
                logger.info(
                    f"Разобрана сложная формулировка сумм: principal={principal_b}, interest={interest_b}, "
                    f"overdue={overdue_b}, duty={duty_b}, total={extracted_fields.get('totalDebt')}"
                )
        except Exception as e:
            logger.debug(f"Ошибка при разборе сложной формулировки сумм: {e}")

        # Попытка дополнительно вытащить суммы из "табличных" форматов,
        # когда название и сумма разбиты по разным строкам (как в таблице)
        try:
            if document_type == "rtk_application" and text:
                lines = text.splitlines()

                def _extract_amount_near_index(start_index: int) -> Optional[str]:
                    """
                    Ищет сумму в этой или ближайших строках ниже (типичный случай: название в одной ячейке, сумма в соседней).
                    """
                    n = len(lines)
                    for j in range(start_index, min(n, start_index + 3)):
                        line_j = lines[j]
                        m = re.search(r"([0-9][0-9\s\u00a0\u202f.,]+)", line_j)
                        if m:
                            return self.normalize_amount_value(m.group(1))
                    return None

                if not extracted_fields.get("principalDebt"):
                    for idx, line in enumerate(lines):
                        ll = line.lower()
                        if (
                            "сумма основного долга" in ll
                            or ("основной долг" in ll and "просроч" not in ll)
                            or "тело долга" in ll
                        ):
                            val = _extract_amount_near_index(idx)
                            if val:
                                extracted_fields["principalDebt"] = val
                                logger.info(f"principalDebt извлечён из табличного блока: {val}")
                                break

                if not extracted_fields.get("interest"):
                    for idx, line in enumerate(lines):
                        ll = line.lower()
                        if "сумма долга по процентам" in ll or "проценты" in ll:
                            val = _extract_amount_near_index(idx)
                            if val:
                                extracted_fields["interest"] = val
                                logger.info(f"interest извлечён из табличного блока: {val}")
                                break

                if not extracted_fields.get("forfeit"):
                    for idx, line in enumerate(lines):
                        ll = line.lower()
                        if "задолженности по просроченным процентам" in ll or "просроченным процентам" in ll:
                            val = _extract_amount_near_index(idx)
                            if val:
                                extracted_fields["forfeit"] = val
                                logger.info(f"forfeit извлечён из табличного блока: {val}")
                                break

                if not extracted_fields.get("stateDuty"):
                    for idx, line in enumerate(lines):
                        ll = line.lower()
                        if "госпошлин" in ll or "государственной пошл" in ll:
                            val = _extract_amount_near_index(idx)
                            if val:
                                extracted_fields["stateDuty"] = val
                                logger.info(f"stateDuty извлечена из табличного блока: {val}")
                                break
        except Exception as e:
            logger.debug(f"Ошибка при попытке извлечь суммы из табличного формата: {e}")

        # Если totalDebt отсутствует или равен 0, пытаемся собрать его из составных частей
        if 'totalDebt' not in extracted_fields or not extracted_fields['totalDebt'] or str(extracted_fields['totalDebt']).strip() in ["0", "0,00", "0.00", ""]:
            principal = self._safe_amount_field(extracted_fields.get("principalDebt"))
            interest = self._safe_amount_field(extracted_fields.get("interest"))
            forfeit_val = self._safe_amount_field(extracted_fields.get("forfeit"))
            loan_debt = self._safe_amount_field(extracted_fields.get("loanDebt"))

            composite_total = 0.0
            if loan_debt > 0:
                composite_total += loan_debt
            if principal > 0:
                composite_total += principal
            if interest > 0:
                composite_total += interest
            if forfeit_val > 0:
                composite_total += forfeit_val

            if composite_total > 0:
                formatted_total = f"{composite_total:,.2f}".replace(',', ' ').replace('.', ',')
                extracted_fields['totalDebt'] = formatted_total
                logger.info(f"Собрана общая сумма долга из частей: {formatted_total}")
            elif 'debtAmount' in extracted_fields and extracted_fields['debtAmount']:
                extracted_fields['totalDebt'] = extracted_fields['debtAmount']
                logger.info(f"Используем debtAmount как общую сумму долга: {extracted_fields['debtAmount']}")

        # Для rtk_application: если debtAmount явно равен только основному долгу,
        # а проценты и просроченные проценты уже найдены — корректируем суммы.
        if document_type == "rtk_application":
            try:
                principal = self._safe_amount_field(extracted_fields.get("principalDebt"))
                interest = self._safe_amount_field(extracted_fields.get("interest"))
                forfeit_val = self._safe_amount_field(extracted_fields.get("forfeit"))
                composite = principal + interest + forfeit_val
                current_debt = self._safe_amount_field(extracted_fields.get("debtAmount"))
                if composite > 0 and current_debt > 0:
                    # Если debtAmount практически совпадает с principalDebt и явно меньше суммы компонентов —
                    # считаем это ошибкой извлечения и подменяем.
                    if abs(current_debt - principal) < 1 and current_debt < composite:
                        formatted = f"{composite:,.2f}".replace(",", " ").replace(".", ",")
                        old_debt = extracted_fields.get("debtAmount")
                        extracted_fields["debtAmount"] = formatted
                        total_str = (extracted_fields.get("totalDebt") or "").strip()
                        total_val = self._safe_amount_field(total_str) if total_str else 0.0
                        if not total_val or abs(total_val - self._safe_amount_field(old_debt)) < 1:
                            extracted_fields["totalDebt"] = formatted
                        logger.info(f"Скорректирована сумма требований: debtAmount/totalDebt = {formatted} (principal+interest+forfeit)")
            except Exception as e:
                logger.debug(f"Ошибка при корректировке суммы требований: {e}")

        # Для rtk_application: суммы берём ИЗ БЛОКА "ПРОСИТ СУД"
        if document_type == "rtk_application":
            # Суммы/госпошлина/публикация из блока «ПРОСИТ СУД»
            self._extract_claim_block_amounts(extracted_fields, text)

        # Согласование сумм долга (principal/interest/forfeit/total/loanDebt)
        self._reconcile_debt_amounts(extracted_fields, text, document_type)

        # Согласование госпошлины/requirementsSum
        self._reconcile_state_duty(extracted_fields, text, document_type)


        if document_type != "rtk_application" and not extracted_fields.get("loanStateDuty17"):
            if extracted_fields.get("stateDuty16"):
                extracted_fields["loanStateDuty17"] = extracted_fields["stateDuty16"]
            elif extracted_fields.get("stateDuty"):
                extracted_fields["loanStateDuty17"] = extracted_fields["stateDuty"]


        # Реквизиты ЮЛ из блока «Должник:» (debtor_block)
        debtor_block = self._extract_legal_entity_requisites(extracted_fields, text)

        procedure_type, procedure_raw = self.determine_procedure_type(text)
        if procedure_type:
            extracted_fields['procedureType'] = procedure_type
            logger.info(f"Определен тип процедуры: {procedure_type}")
        if procedure_raw:
            extracted_fields['procedureTypeRaw'] = procedure_raw
            logger.info(f"Оригинальное описание процедуры: {procedure_raw}")

        # Пост-обработка процедуры и managerInn
        self._normalize_procedure_and_manager(extracted_fields, text, procedure_type)

        debtor_name_raw = extracted_fields.get("debtorName")
        applicant_name_raw = extracted_fields.get("applicantName")
        applicant_name_genitive_raw = extracted_fields.get("applicantNameGenitive")
        applicant_name_instrumental_raw = extracted_fields.get("applicantNameInstrumental")
        debtor_clean = None
        if debtor_name_raw:
            if re.search(r"^введена\s+процедура|процедура\s+наблюдения", debtor_name_raw, re.IGNORECASE):
                debtor_name_raw = None
        if debtor_name_raw:
            debtor_clean = re.sub(r"\[.*?\]", "", debtor_name_raw)
            debtor_clean = self.clean_extracted_value(debtor_clean)
            debtor_clean = debtor_clean.strip('«»" ')
            debtor_clean = re.sub(r"^\s*фио\s+", "", debtor_clean, flags=re.IGNORECASE)
            debtor_clean = re.sub(r"\s*должника\s+введена\s+процедура.*$", "", debtor_clean, flags=re.IGNORECASE)
            debtor_clean = re.sub(r"\s*введена\s+процедура\s+наблюдения.*$", "", debtor_clean, flags=re.IGNORECASE)
            debtor_clean = debtor_clean.strip()
            if debtor_clean and len(debtor_clean) > 3 and not re.match(r"^[а-яё]\s*\.?$", debtor_clean, re.IGNORECASE):
                extracted_fields["debtorName"] = debtor_clean
            else:
                extracted_fields.pop("debtorName", None)
        else:
            extracted_fields.pop("debtorName", None)


        applicant_clean = None
        if applicant_name_raw:
            if re.search(r"^[а-яё]\s+введена\s+процедура|^введена\s+процедура|процедура\s+наблюдения|соответствует\s+признак|признак\w*\s+банкрот", applicant_name_raw, re.IGNORECASE):
                applicant_name_raw = None
        if applicant_name_raw:
            applicant_clean = re.sub(r"\[.*?\]", "", applicant_name_raw)
            applicant_clean = self.clean_extracted_value(applicant_clean)
            applicant_clean = applicant_clean.strip('«»" ')
            applicant_clean = re.sub(r"^\s*фио\s+", "", applicant_clean, flags=re.IGNORECASE)
            applicant_clean = re.sub(r"\s*должника\s+введена\s+процедура.*$", "", applicant_clean, flags=re.IGNORECASE)
            applicant_clean = re.sub(r"\s*введена\s+процедура\s+наблюдения.*$", "", applicant_clean, flags=re.IGNORECASE)
            applicant_clean = applicant_clean.strip()
            if applicant_clean and len(applicant_clean) > 3 and not re.match(r"^[а-яё]\s*\.?$", applicant_clean, re.IGNORECASE):
                extracted_fields["applicantName"] = applicant_clean
            else:
                extracted_fields.pop("applicantName", None)
        else:
            extracted_fields.pop("applicantName", None)
        # Если не извлекли имя, пробуем короткое наименование по тексту/блоку должника
        if "applicantName" not in extracted_fields:
            legal_fallback_source = debtor_block or debtor_name_raw or applicant_name_raw or text[:1200]
            legal_short = self.extract_legal_entity_short_name(legal_fallback_source)
            if legal_short:
                extracted_fields["applicantName"] = legal_short

        if applicant_name_genitive_raw:
            if re.search(r"^введена\s+процедура|процедура\s+наблюдения", applicant_name_genitive_raw, re.IGNORECASE):
                applicant_name_genitive_raw = None
        if applicant_name_genitive_raw:
            genitive_clean = re.sub(r"\[.*?\]", "", applicant_name_genitive_raw)
            genitive_clean = self.clean_extracted_value(genitive_clean)
            genitive_clean = genitive_clean.strip('«»" ')
            genitive_clean = re.sub(r"^\s*фио\s+", "", genitive_clean, flags=re.IGNORECASE)
            # Убираем префиксы "ГЛАВА КФХ ИП" для родительного падежа
            genitive_clean = re.sub(
                r'^(ГЛАВА\s+КФХ\s+ИП\s+|ГЛАВА\s+КФХ\s+|КФХ\s+ИП\s+|ИП\s+ГЛАВА\s+КФХ\s+)',
                '',
                genitive_clean,
                flags=re.IGNORECASE
            ).strip()
            genitive_clean = re.sub(r"\s*введена\s+процедура\s+наблюдения.*$", "", genitive_clean, flags=re.IGNORECASE)
            genitive_clean = genitive_clean.strip()
            if genitive_clean and len(genitive_clean) > 3 and not re.match(r"^[а-яё]\s*\.?$", genitive_clean, re.IGNORECASE):
                extracted_fields["applicantNameGenitive"] = genitive_clean
            else:
                extracted_fields.pop("applicantNameGenitive", None)
        else:
            extracted_fields.pop("applicantNameGenitive", None)

        is_kfh = extracted_fields.get("isKfh", False)
        self._generate_debtor_case_forms(extracted_fields, text, applicant_clean, is_kfh, applicant_name_instrumental_raw)

        debtor_clean = self._tune_legal_entity_naming(extracted_fields, text, debtor_clean, applicant_clean, applicant_name_raw, debtor_block, debtor_name_raw)

        self._resolve_debtor_address(extracted_fields, text, debtor_block)

        if extracted_fields.get("ogrn"):
            extracted_fields["ogrn"] = re.sub(r"\D", "", extracted_fields["ogrn"])

        if extracted_fields.get("companyInn"):
            extracted_fields["companyInn"] = re.sub(r"\D", "", extracted_fields["companyInn"])

        self._sanitize_credit_params(extracted_fields, text)

        self._finalize_debtor_type(extracted_fields, text, debtor_clean)


        self._normalize_financial_block(extracted_fields, text)


        self._drop_law_context_dates(extracted_fields, text)

        logger.info(f"Итоговые извлеченные поля: {extracted_fields}")
        return extracted_fields

    _LAW_DATE_FIELDS = ("contractDate", "courtDecisionDate", "priorDecisionDate")

    def _extract_kommersant_publication(self, extracted_fields, text):
        """[67]/[68] — номер и дата публикации в газете «Коммерсантъ» (вынесено из extract_fields)."""
        # [68] — дата газеты «Коммерсантъ» (не путать с [11] ЕФРСБ)
        # Приоритет: типичная формулировка "в газете «Коммерсантъ» № 123 от 27.12.2025"
        kommersant_date_match = re.search(
            r"(?:газет[аы]\s+)?[«\"]?Коммерсант[ъ\"»']?[»\"]?\s*№\s*[0-9\-\/]+\s*от\s*(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
            text,
            re.IGNORECASE,
        )
        if not kommersant_date_match:
            # Вариант "в газете «Коммерсантъ» от 27.12.2025" — без номера
            kommersant_date_match = re.search(
                r"в[\s\u00a0\u202f]+газет[аы][\s\u00a0\u202f]+«?Коммерсант[\"ъ»']?»?\s*[^0-9]{0,40}?от[\s\u00a0\u202f]*(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
                text,
                re.IGNORECASE,
            )
        if not kommersant_date_match:
            # Вариант "27.12.2025 в газете «Коммерсантъ»" — дата стоит ПЕРЕД упоминанием газеты
            kommersant_date_match = re.search(
                r"(\d{1,2}[.,]\d{1,2}[.,]\d{4})[\s\u00a0\u202f]+в[\s\u00a0\u202f]+газет[аы][\s\u00a0\u202f]+«?Коммерсант[\"ъ»']?»?",
                text,
                re.IGNORECASE,
            )
        if not kommersant_date_match:

            kommersant_date_match = re.search(
                r"официальн[а-яё\s]+издан[а-яё\s]+газет[аы]\s+«?Коммерсант[\"ъ»']?»?[^0-9]{0,400}?от\s*(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
                text,
                re.IGNORECASE,
            )
        if not kommersant_date_match:

            kommersant_date_match = re.search(
                r"Коммерсант[ъ\"»']?\s*[^0-9]{0,200}?(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
                text,
                re.IGNORECASE,
            )
            if not kommersant_date_match and re.search(r"Коммерсант", text, re.IGNORECASE):
                date_pattern = re.compile(r"(\d{1,2}[.,]\d{1,2}[.,]\d{4})")
                best_match = None
                best_date = None  # сравниваем реальные даты, чтобы не брать старый закон 2002 года
                for m in date_pattern.finditer(text):
                    start, end = m.start(), m.end()
                    window_start = max(0, start - 150)
                    window_end = min(len(text), end + 150)
                    window = text[window_start:window_end]
                    if re.search(r"Коммерсант", window, re.IGNORECASE):
                        date_str = m.group(1).replace(",", ".")
                        try:
                            day, month, year = map(int, date_str.split("."))
                            # Грубая фильтрация нереалистичных годов
                            if year < 1990 or year > 2100:
                                continue
                            from datetime import date as _date
                            cur_date = _date(year, month, day)
                        except Exception:
                            continue
                        if best_date is None or cur_date > best_date:
                            best_date = cur_date
                            best_match = m
                if best_match:
                    kommersant_date_match = best_match
        if kommersant_date_match:
            kommersant_date = kommersant_date_match.group(1).strip().replace(",", ".")
            extracted_fields["kommersantDate"] = kommersant_date
            logger.info(f"Дата газеты «Коммерсантъ» [68]: {kommersant_date}")
        # [67] — номер газеты «Коммерсантъ» (только при явном «№ …», не путать с [9] ЕФРСБ)
        kommersant_number_match = re.search(
            r"(?:газет[аы]\s+)?[«\"]?Коммерсант[ъ\"»']?[»\"]?\s*№\s*([0-9\-\/]+)",
            text,
            re.IGNORECASE,
        )
        if kommersant_number_match:
            extracted_fields["kommersantNumber"] = kommersant_number_match.group(1).strip()
            logger.info(f"Номер газеты «Коммерсантъ» [67]: {extracted_fields['kommersantNumber']}")







    def _generate_debtor_case_forms(self, extracted_fields, text, applicant_clean, is_kfh, applicant_name_instrumental_raw):
        """Падежные формы имени должника [2.1] род., [2.3] твор., [2.4] вин., [2.2] дат.: очистка raw-значений и автогенерация склонений (КФХ — по ФИО без префиксов). Вынесено из extract_fields."""
        kfh_head_name = extracted_fields.get("kfhHeadName", "")
        name_for_inflection = applicant_clean

        if is_kfh and kfh_head_name:
            # Для КФХ используем только ФИО без префиксов
            name_for_inflection = kfh_head_name
            logger.info(f" КФХ: используем kfhHeadName '{kfh_head_name}' для генерации падежных форм")

        if "applicantNameGenitive" not in extracted_fields and name_for_inflection:
            genitive_auto = self._convert_name_to_genitive(name_for_inflection)
            if genitive_auto:
                extracted_fields["applicantNameGenitive"] = genitive_auto

        # Творительный падеж для должника (маркер [2.3])
        if applicant_name_instrumental_raw:
            instr_clean = re.sub(r"\[.*?\]", "", applicant_name_instrumental_raw)
            instr_clean = self.clean_extracted_value(instr_clean)
            instr_clean = instr_clean.strip('«»" ')
            instr_clean = re.sub(r"^\s*фио\s+", "", instr_clean, flags=re.IGNORECASE)
            # Убираем префиксы "ГЛАВА КФХ ИП" для творительного падежа
            instr_clean = re.sub(
                r'^(ГЛАВА\s+КФХ\s+ИП\s+|ГЛАВА\s+КФХ\s+|КФХ\s+ИП\s+|ИП\s+ГЛАВА\s+КФХ\s+)',
                '',
                instr_clean,
                flags=re.IGNORECASE
            ).strip()
            instr_clean = re.sub(r"\s*введена\s+процедура\s+наблюдения.*$", "", instr_clean, flags=re.IGNORECASE)
            instr_clean = instr_clean.strip()
            if instr_clean and len(instr_clean) > 3:
                extracted_fields["applicantNameInstrumental"] = instr_clean
            else:
                extracted_fields.pop("applicantNameInstrumental", None)

        if "applicantNameInstrumental" not in extracted_fields and name_for_inflection:
            instr_auto = self._convert_name_to_instrumental(name_for_inflection)
            if instr_auto:
                extracted_fields["applicantNameInstrumental"] = instr_auto

        # Обработка винительного падежа [2.4]
        applicant_name_accusative_raw = extracted_fields.get("applicantNameAccusative")
        if applicant_name_accusative_raw:
            accs_clean = re.sub(r"\[.*?\]", "", applicant_name_accusative_raw)
            accs_clean = self.clean_extracted_value(accs_clean)
            accs_clean = accs_clean.strip('«»" ')
            accs_clean = re.sub(r"^\s*фио\s+", "", accs_clean, flags=re.IGNORECASE)
            # Убираем префиксы "ГЛАВА КФХ ИП" для винительного падежа
            accs_clean = re.sub(
                r'^(ГЛАВА\s+КФХ\s+ИП\s+|ГЛАВА\s+КФХ\s+|КФХ\s+ИП\s+|ИП\s+ГЛАВА\s+КФХ\s+)',
                '',
                accs_clean,
                flags=re.IGNORECASE
            ).strip()
            accs_clean = re.sub(r"\s*введена\s+процедура\s+наблюдения.*$", "", accs_clean, flags=re.IGNORECASE)
            accs_clean = accs_clean.strip()
            if accs_clean and len(accs_clean) > 3:
                if not re.search(r"^введена\s+процедура|процедура\s+наблюдения", accs_clean, re.IGNORECASE):
                    extracted_fields["applicantNameAccusative"] = accs_clean
            else:
                extracted_fields.pop("applicantNameAccusative", None)

        if "applicantNameAccusative" not in extracted_fields and name_for_inflection:
            # Для КФХ используем правильное склонение с одушевленностью (кого?)
            accs_auto = self._convert_name_to_accusative(name_for_inflection)
            if accs_auto:
                extracted_fields["applicantNameAccusative"] = accs_auto

        # Дательный падеж для должника (маркер [2.2]) - везде это имя должника в дательном падеже, кроме ипотеки
        # НЕ генерируем дательный падеж, если name_for_inflection содержит "суд" (это название суда, а не имя должника)
        if "applicantNameDative" not in extracted_fields and name_for_inflection:
            if "суд" not in name_for_inflection.lower():
                # Убираем префикс "ИП" перед склонением (как для других падежей)
                name_for_dative = re.sub(
                    r'^(ИП\s+|ГЛАВА\s+КФХ\s+ИП\s+|ГЛАВА\s+КФХ\s+|КФХ\s+ИП\s+)',
                    '',
                    name_for_inflection,
                    flags=re.IGNORECASE
                ).strip()
                if not name_for_dative:
                    name_for_dative = name_for_inflection

                dative_auto = self._convert_name_to_dative(name_for_dative)
                if dative_auto:
                    extracted_fields["applicantNameDative"] = dative_auto
                    logger.info(f"Сгенерирован дательный падеж для [2.2]: {dative_auto}")
            else:
                logger.warning(f"Пропущена генерация дательного падежа - name_for_inflection содержит 'суд': {name_for_inflection}")

    def _normalize_procedure_and_manager(self, extracted_fields, text, procedure_type):
        """Пост-обработка процедуры и управляющего: чистка названия суда для «умерший», фолбэк процедуры «наблюдение» по raw-описанию, валидация managerInn по контрольной сумме. Вынесено из extract_fields."""
        if procedure_type == 'deceased' and extracted_fields.get('courtName'):
            court_name = extracted_fields['courtName']
            cleaned_court_name = self._clean_court_name_deceased(court_name)
            if cleaned_court_name != court_name:
                extracted_fields['courtName'] = cleaned_court_name
                logger.info(f"Очищено название суда для процедуры 'умерший': '{court_name}' -> '{cleaned_court_name}'")

        # Фолбэк: если тип процедуры явно не нормализовался, но по "сырому" описанию видно "наблюдение" —
        # считаем, что это процедура наблюдения.
        if not extracted_fields.get("procedureType"):
            raw_lower = (extracted_fields.get("procedureTypeRaw") or "").lower()
            if any(token in raw_lower for token in ["наблюден", "наблюдения", "наблюдение"]):
                extracted_fields["procedureType"] = "observation"
                logger.info("Процедура по умолчанию установлена как 'observation' на основании procedureTypeRaw")

        manager_inn = extracted_fields.get("managerInn")
        if manager_inn:
            normalized_inn = re.sub(r"\D", "", manager_inn)
            # Отсеиваем мусор, не проходящий контрольную сумму (например, 689768767676),
            # чтобы в поле не подставлялся заведомо ложный ИНН.
            if normalized_inn and is_valid_inn(normalized_inn):
                extracted_fields["managerInn"] = normalized_inn
            else:
                logger.debug(f"managerInn '{manager_inn}' отброшен: не прошёл контрольную сумму")
                extracted_fields.pop("managerInn", None)







    def _extract_multiple_field(self, extracted_fields, text, field_name, field_patterns):
        """Поля с множественными значениями (contractNumber/contractDate/obligationType): сбор всех вхождений по паттернам, защита C (закон != договор), фильтры мусора, дедуп с сохранением порядка. Вынесено из pattern-цикла."""
        found_values = []
        for pattern in field_patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            for match in matches:
                value = match.strip()
                if value and len(value) > 2:
                    cleaned_value = self.clean_extracted_value(value)
                    if not cleaned_value:  # Если после очистки ничего не осталось
                        continue

                    if field_name == 'contractNumber':
                        # Защита C: «№ 353-ФЗ» — ссылка на закон, не номер договора.
                        if looks_like_law_ref(cleaned_value):
                            logger.info(f"Пропуск contractNumber (ссылка на закон): {cleaned_value}")
                        elif not any(word in cleaned_value.lower() for word in ['считается', 'поручительства', 'возникшим', 'которые', 'согласно', 'установленную', 'принятые', 'договор', 'кредитный', 'путем', 'подписания', 'далее']) and len(cleaned_value.strip()) >= 2:
                            found_values.append(cleaned_value)
                            logger.info(f"Found {field_name}: {cleaned_value}")
                    elif field_name == 'obligationType':
                        if not any(word in cleaned_value.lower() for word in ['возникшим', 'которые', 'согласно', 'установленную', 'принятые']):
                            found_values.append(cleaned_value)
                            logger.info(f"Found {field_name}: {cleaned_value}")
                    else:
                        found_values.append(cleaned_value)
                        logger.info(f"Found {field_name}: {cleaned_value}")

        if found_values:
            unique_values = []
            seen_values = set()
            for value in found_values:
                normalized = value.lower()
                if normalized in seen_values:
                    continue
                seen_values.add(normalized)
                unique_values.append(value)

            if unique_values:
                extracted_fields[field_name] = ", ".join(unique_values)
                logger.info(f"All {field_name}: {extracted_fields[field_name]}")

    def _detect_kfh_head(self, text):
        """Ранняя детекция КФХ по тексту («ГЛАВА КФХ ИП ФИО»). Возвращает (is_kfh_detected, kfh_head_name). Вынесено из analyze."""
        # Ищем в любом месте документа, не только в блоке должника
        kfh_patterns = [
            r"глава\s+кфх\s+ип\s+([А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+){1,2})",  # "ГЛАВА КФХ ИП ФИО"
            r"кфх\s+ип\s+([А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+){1,2})",  # "КФХ ИП ФИО"
            r"глава\s+кфх\s+([А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+){1,2})",  # "ГЛАВА КФХ ФИО"
        ]

        is_kfh_detected = False
        kfh_head_name = None

        for pattern in kfh_patterns:
            kfh_match = re.search(pattern, text, re.IGNORECASE)
            if kfh_match:
                is_kfh_detected = True
                kfh_head_name = kfh_match.group(1).strip()
                logger.info(f"КФХ обнаружено по паттерну '{pattern}': глава КФХ - {kfh_head_name}")
                break
        return is_kfh_detected, kfh_head_name

    def _regenerate_ip_dative(self, extracted_fields, text, ip_specific_fields):
        """Перегенерация дательного падежа [2.2] для должника-ИП: выбор корректного ФИО (приоритет ip_specific/extracted, фильтр «не суд»), снятие префиксов, склонение. Вынесено из analyze."""
        # Перегенерируем дательный падеж [2.2] с правильным именем должника
        # Используем правильное имя из ip_specific_fields или extracted_fields
        correct_name_for_dative = None

        def is_valid_name(name):
            if not name or "суд" in name.lower():
                return False
            words = name.strip().split()
            if len(words) < 2 or len(words) > 4:
                return False
            valid_words = [w for w in words if w.upper() not in ["ИП", "ГЛАВА", "КФХ"]]
            if len(valid_words) < 2:
                return False
            return True

        # Приоритет 1: applicantName из ip_specific_fields
        if ip_specific_fields.get("applicantName") and is_valid_name(ip_specific_fields.get("applicantName")):
            correct_name_for_dative = ip_specific_fields.get("applicantName")
            logger.info(f"Используем applicantName из ip_specific_fields для [2.2]: {correct_name_for_dative}")

        # Приоритет 2: debtorName из ip_specific_fields
        if not correct_name_for_dative and ip_specific_fields.get("debtorName") and is_valid_name(ip_specific_fields.get("debtorName")):
            correct_name_for_dative = ip_specific_fields.get("debtorName")
            logger.info(f"Используем debtorName из ip_specific_fields для [2.2]: {correct_name_for_dative}")

        # Приоритет 3: applicantName из extracted_fields
        if not correct_name_for_dative and extracted_fields.get("applicantName") and is_valid_name(extracted_fields.get("applicantName")):
            correct_name_for_dative = extracted_fields.get("applicantName")
            logger.info(f"Используем applicantName из extracted_fields для [2.2]: {correct_name_for_dative}")

        # Приоритет 4: debtorName из extracted_fields
        if not correct_name_for_dative and extracted_fields.get("debtorName") and is_valid_name(extracted_fields.get("debtorName")):
            correct_name_for_dative = extracted_fields.get("debtorName")
            logger.info(f"Используем debtorName из extracted_fields для [2.2]: {correct_name_for_dative}")

        if correct_name_for_dative:
            # Убираем префикс "ИП" перед склонением
            name_for_dative = re.sub(
                r'^(ИП\s+|ГЛАВА\s+КФХ\s+ИП\s+|ГЛАВА\s+КФХ\s+|КФХ\s+ИП\s+)',
                '',
                correct_name_for_dative,
                flags=re.IGNORECASE
            ).strip()
            if name_for_dative:
                dative_auto = self._convert_name_to_dative(name_for_dative)
                if dative_auto:
                    extracted_fields["applicantNameDative"] = dative_auto
                    logger.info(f"Перегенерирован дательный падеж для [2.2]: {dative_auto} (из имени: {correct_name_for_dative})")
                else:
                    logger.warning(f"Не удалось сгенерировать дательный падеж для: {name_for_dative}")
            else:
                logger.warning(f"После удаления префиксов имя стало пустым: {correct_name_for_dative}")
        else:
            logger.warning(f"Не найдено правильное имя должника для генерации дательного падежа [2.2]")
            logger.warning(f"  applicantName (ip_specific): {ip_specific_fields.get('applicantName')}")
            logger.warning(f"  debtorName (ip_specific): {ip_specific_fields.get('debtorName')}")
            logger.warning(f"  applicantName (extracted): {extracted_fields.get('applicantName')}")
            logger.warning(f"  debtorName (extracted): {extracted_fields.get('debtorName')}")



    def _maybe_upgrade_to_observation_collateral(self, extracted_fields, text, document_type, text_lower):
        """Доп. проверка после извлечения: ЮЛ + залог + наблюдение -> тип observation_collateral (+ поля залога). Возвращает (возможно обновлённый) document_type. Вынесено из analyze."""
        if document_type == "rtk_application":
            entity_type = extracted_fields.get("entityType", "").lower()

            # ОБЯЗАТЕЛЬНАЯ ПРОВЕРКА: документ считается залоговым ТОЛЬКО если есть формулировка "обеспеченное залогом"
            required_collateral_phrases = [
                r"обеспеченное\s+залогом",
                r"обеспечено\s+залогом",
                r"обеспечен\s+залогом",
                r"обеспечена\s+залогом",
                r"обеспечены\s+залогом",
                r"В\s+качестве\s+обеспечения\s+исполнения\s+обязательств\s+кредитному\s+договору\s+Заемщик\s+предоставил\s+в\s+залог\s+Банку",
                r"обязательство.*обеспечен.*залогом",
                r"обязательства.*обеспечен.*залогом"
            ]
            has_required_collateral_phrase = any(re.search(pattern, text, re.IGNORECASE | re.DOTALL) for pattern in required_collateral_phrases)

            # Дополнительные индикаторы залога (только если есть обязательная формулировка)
            has_collateral = False
            if has_required_collateral_phrase:
                has_collateral = any(re.search(pattern, text, re.IGNORECASE | re.DOTALL) for pattern in [
                    r"В\s+качестве\s+обеспечения\s+исполнения\s+обязательств\s+кредитному\s+договору\s+Заемщик\s+предоставил\s+в\s+залог\s+Банку",
                    r"что\s+подтверждается\s+договором\s+залога",
                    r"договор\s+залога\s+№"
                ])

            is_observation = any(keyword in text_lower for keyword in ["наблюден", "наблюдения", "процедура наблюдения"]) or \
                            entity_type == "legal" or \
                            (extracted_fields.get("procedureType") or "").lower() == "observation"

            # Не повышаем до observation_collateral, если в тексте есть явное «ИП ФИО»
            # в имени должника (та же проверка, что и при первичной классификации).
            has_ip_name = self._detect_ip_debtor_name(text)

            if entity_type == "legal" and has_collateral and is_observation and not has_ip_name:
                document_type = "observation_collateral"
                logger.info("Определен тип документа: observation_collateral (Наблюдение с залогом для ЮЛ после извлечения полей)")
                collateral_fields = self.extract_physical_collateral_fields(text)
                for key, value in collateral_fields.items():
                    if value:
                        extracted_fields[key] = value

                if "mortgageCollateralDescription1221" not in extracted_fields or not extracted_fields.get("mortgageCollateralDescription1221"):
                    logger.info("Предмет залога [1221] не найден для observation_collateral, пробуем fallback метод...")
                    collateral_block = self._extract_mortgage_collateral_block(text)
                    if collateral_block:
                        extracted_fields["mortgageCollateralDescription1221"] = collateral_block
                        logger.info(f"Извлечено mortgageCollateralDescription1221 через fallback метод: {collateral_block[:150]}...")

                extracted_fields["observationHasCollateral"] = "true"
        return document_type

    def _postprocess_mortgage_and_case_number(self, extracted_fields, text, document_type):
        """Постобработка ипотеки/залога (корректировка интерфейсных значений) и валидация номера дела («/ГОД» в конце, отсев доверенностей/договоров). Вынесено из analyze."""
        if document_type == "mortgage_claim" or document_type == "competition_collateral":
            self._normalize_mortgage_interface_fields(extracted_fields, text)
            # Для конкурсного производства с залогом явно считаем должника юридическим лицом
            if document_type == "competition_collateral":
                extracted_fields["entityType"] = "legal"
            # Если есть ИНН финуправляющего — добавляем его в отображаемое имя
            manager_name_cc = extracted_fields.get("managerName")
            manager_inn_cc = extracted_fields.get("managerInn")
            if manager_name_cc and manager_inn_cc and "инн" not in manager_name_cc.lower():
                extracted_fields["managerName"] = f"{manager_name_cc.strip()} (ИНН {manager_inn_cc})"


        if extracted_fields.get("entityType") not in ("kfh", "ip", "legal"):
            text_entity = self._entity_from_text(text)
            if text_entity:
                extracted_fields["entityType"] = text_entity
                if text_entity == "kfh":
                    extracted_fields["isKfh"] = True


        _cn = extracted_fields.get("caseNumber")
        if _cn and not self._is_valid_case_number(_cn):
            _cn_norm = re.sub(r"\s+", "", str(_cn)).upper().replace("Ё", "Е")
            if self._MAGISTRATE_CASE_NUMBER_RE.match(_cn_norm):
                # Дело мирового участка («по делу № 2-7-3490/2025» — ссылка на
                # судебный приказ в теле заявления). Номером арбитражного дела оно
                # не является, а спасательная регулярка ниже отрезала бы ему голову
                # и выдала правдоподобный, но чужой «7-3490/2025».
                extracted_fields.pop("caseNumber", None)
                _cn = None
        if _cn and not self._is_valid_case_number(_cn):
            # Хвост мог прилипнуть без пробела («…/2026Исх. Док.») — жадный класс
            # захватил заглавную букву после года. Спасаем канонический номер.
            _core = re.search(
                r"[А-ЯA-Z]?\d{1,4}[А-ЯA-Z]?[-–]\d{1,15}/(?:19|20)\d{2}",
                str(_cn).upper().replace("Ё", "Е"),
            )
            if _core:
                extracted_fields["caseNumber"] = _core.group(0)
            else:
                extracted_fields.pop("caseNumber", None)







    def _analyze_ip_collection(self, text, document_type):
        """Ветка анализа ИП-взыскания: extract_fields + слияние ip_specific_fields (без перезаписи ИНН/сумм), залог авто [1221], дательный падеж, обязательства. Возвращает extracted_fields. Вынесено из analyze."""
        extracted_fields = self.extract_fields(text, "rtk_application")  # Используем rtk_application как базовый тип

        ip_specific_fields = self.extract_ip_enforcement_fields(text)
        logger.info(f"Извлеченные специфичные поля для ИП (взыскание): {list(ip_specific_fields.keys())}")
        logger.info(f" Значения полей [1000], [1001], [1004]:")
        logger.info(f"  creditAmount [1000]: {ip_specific_fields.get('creditAmount')}")
        logger.info(f"  creditTermMonths [1001]: {ip_specific_fields.get('creditTermMonths')}")
        logger.info(f"  debtSnapshotDate [1004]: {ip_specific_fields.get('debtSnapshotDate')}")

        for key, value in ip_specific_fields.items():
            if value:
                # ВАЖНО: Не перезаписываем ИНН и ОГРН, если они уже были извлечены из блока "Ответчик:" в основном цикле
                if key in ["inn", "ogrnip", "ogrn"]:
                    if key in extracted_fields and extracted_fields[key]:
                        logger.info(f"Пропускаем перезапись {key} из ip_specific_fields (уже извлечен из блока Ответчик: {extracted_fields[key]})")
                        continue
                # Не перезаписываем уже найденные денежные поля основного извлечения,
                # чтобы не заносить шум из fallback-паттернов ip_specific_fields.
                if key in ["principalDebt13", "interest14", "forfeit15", "principalDebt", "interest", "forfeit", "totalDebt", "stateDuty16", "stateDuty"]:
                    if key in extracted_fields and extracted_fields[key]:
                        # Для госпошлины разрешаем обновить 0,00 -> ненулевое значение.
                        if key in ["stateDuty16", "stateDuty"]:
                            existing_amount = self._safe_amount_field(extracted_fields.get(key))
                            incoming_amount = self._safe_amount_field(value)
                            if existing_amount <= 0 < incoming_amount:
                                logger.info(
                                    f"Обновляем {key}: текущее значение {extracted_fields[key]} заменяется на ненулевое {value}"
                                )
                            else:
                                logger.info(f"Пропускаем перезапись {key} из ip_specific_fields (уже есть корректное значение: {extracted_fields[key]})")
                                continue
                        else:
                            logger.info(f"Пропускаем перезапись {key} из ip_specific_fields (уже есть корректное значение: {extracted_fields[key]})")
                            continue
                extracted_fields[key] = value
                logger.info(f"Установлено поле ИП {key}: {value}")

        # Для «Взыскание ИП залог авто» [1221] — описание авто (марка, модель, год, VIN и т.д.)
        if document_type == "ip_collection_collateral_auto":
            car_1221 = self._extract_car_collateral_1221(text)
            if car_1221:
                extracted_fields["mortgageCollateralDescription1221"] = car_1221
                logger.info(f"Установлено mortgageCollateralDescription1221 (залог авто): {car_1221[:80]}...")

        self._regenerate_ip_dative(extracted_fields, text, ip_specific_fields)

        obligations = self.extract_obligations(text, extracted_fields)
        if obligations:
            extracted_fields['obligations'] = obligations
            logger.info(f"Добавлено {len(obligations)} обязательств для ИП (взыскание)")
        return extracted_fields

    def _analyze_legal_collection(self, text, document_type):
        """Ветка анализа взыскания с ЮЛ: extract_fields + данные о залоге (без перезаписи сумм), дательный падеж, обязательства. Возвращает extracted_fields. Вынесено из analyze."""
        # Используем основной extract_fields для ЮЛ, чтобы получить все поля
        extracted_fields = self.extract_fields(text, "rtk_application")  # Используем rtk_application как базовый тип

        # Для типов с залогом извлекаем только данные о залоге (НЕ перезаписываем суммы долга/процентов)
        if document_type in ("legal_collection_collateral", "legal_collection_collateral_auto"):
            # Извлекаем данные о залоге (аналогично ИП), но используем только нужные поля
            ip_specific_fields = self.extract_ip_enforcement_fields(text)
            logger.info(f"Извлеченные специфичные поля для ЮЛ с залогом: {list(ip_specific_fields.keys())}")

            # Только поля, связанные с залогом, которые нужно добавить для ЮЛ
            collateral_fields = [
                "ipCollateralContractNumber", "ipCollateralContractDate", "ipCollateralClaimAmount",
                "mortgageCollateralDescription1221", "contractNumber", "contractDate",
                "creditAmount", "creditTermMonths", "creditInterestRate", "creditPenaltyRate",
                "debtSnapshotDate", "courtName", "creditorName", "debtAmount"
            ]

            # Поля, которые НИКОГДА не перезаписываем для ЮЛ (они уже правильно извлечены в основном цикле)
            fields_to_preserve = [
                "inn", "ogrnip", "ogrn", "principalDebt13", "interest14", "forfeit15",
                "interest", "principalDebt", "totalDebt", "stateDuty16"
            ]

            for key, value in ip_specific_fields.items():
                if value:
                    # ВАЖНО: Не перезаписываем важные поля сумм, если они уже были извлечены в основном цикле
                    if key in fields_to_preserve:
                        if key in extracted_fields and extracted_fields[key]:
                            # Для госпошлины разрешаем обновление 0,00 -> ненулевое значение.
                            if key == "stateDuty16":
                                existing_amount = self._safe_amount_field(extracted_fields.get(key))
                                incoming_amount = self._safe_amount_field(value)
                                if existing_amount <= 0 < incoming_amount:
                                    logger.info(
                                        f" Обновляем {key} для ЮЛ с залогом: текущее значение "
                                        f"{extracted_fields[key]} заменяется на ненулевое {value}"
                                    )
                                else:
                                    logger.info(
                                        f"Пропускаем перезапись {key} из ip_specific_fields "
                                        f"(уже извлечен в основном цикле: {extracted_fields[key]})"
                                    )
                                    continue
                            else:
                                logger.info(f"Пропускаем перезапись {key} из ip_specific_fields (уже извлечен в основном цикле: {extracted_fields[key]})")
                                continue
                    # Госпошлину для ЮЛ с залогом сохраняем отдельно (может приходить из ip_specific_fields
                    # как корректное ненулевое значение, даже если в основном цикле было 0,00).
                    if key == "stateDuty16":
                        extracted_fields["stateDuty16"] = value
                        # Синхронизируем общее поле госпошлины с [16], чтобы генератор не получил 0,00.
                        extracted_fields["stateDuty"] = value
                        logger.info(f"Установлено поле ЮЛ с залогом {key}: {value}")
                    elif key in collateral_fields:
                        extracted_fields[key] = value
                        logger.info(f"Установлено поле ЮЛ с залогом {key}: {value}")
                    else:
                        logger.debug(f"Пропущено поле {key} (не относится к залогу для ЮЛ)")

            # Для «Взыскание ЮЛ залог авто» [1221] — описание авто (марка, модель, год, VIN и т.д.)
            if document_type == "legal_collection_collateral_auto":
                car_1221 = self._extract_car_collateral_1221(text)
                if car_1221:
                    extracted_fields["mortgageCollateralDescription1221"] = car_1221
                    logger.info(f"Установлено mortgageCollateralDescription1221 (залог авто): {car_1221[:80]}...")

            # Специальное извлечение процентов для ЮЛ с залогом: "просроченные проценты" или "проценты" после тире или без него
            interest_patterns = [
                r"просроченные\s+проценты\s*[–—-]\s*([0-9\s\u00a0\u202f,]+(?:[.,][0-9]+)?)\s*(?:руб(?:\.|лей)?|₽|р\.?)?",  # просроченные проценты – X
                r"просроченные\s+проценты[:\s]+([0-9\s,]+(?:[.,][0-9]+)?)\s*(?:руб|рублей|₽|р\.?)?",  # просроченные проценты: X или просроченные проценты X
                r"проценты\s*[–—-]\s*([0-9\s\u00a0\u202f,]+(?:[.,][0-9]+)?)\s*(?:руб(?:\.|лей)?|₽|р\.?)?",  # проценты – X
                r"проценты[:\s]+([0-9\s,]+(?:[.,][0-9]+)?)\s*(?:руб|рублей|₽|р\.?)?",  # проценты: X или проценты X
                r"([0-9\s,]+(?:[.,][0-9]+)?)\s*(?:руб|рублей|₽|р\.?)?\s*[–—-]\s*просроченные\s+проценты",  # X – просроченные проценты
                r"([0-9\s,]+(?:[.,][0-9]+)?)\s*(?:руб|рублей|₽|р\.?)?\s*[–—-]\s*проценты",  # X – проценты
            ]

            for pattern in interest_patterns:
                match = re.search(pattern, text, re.IGNORECASE)
                if match:
                    interest_value = match.group(1).strip()
                    interest_value = re.sub(r'\s+', ' ', interest_value)
                    if interest_value:
                        extracted_fields["interest14"] = interest_value
                        extracted_fields["interest"] = interest_value
                        logger.info(f"Перезаписано interest14 для ЮЛ с залогом из паттерна '{pattern[:50]}...': {interest_value}")
                        break

        obligations = self.extract_obligations(text, extracted_fields)
        if obligations:
            extracted_fields['obligations'] = obligations
            logger.info(f"Добавлено {len(obligations)} обязательств для ЮЛ (взыскание)")
        return extracted_fields

    def _analyze_ip_enforcement(self, text, document_type):
        """Ветка анализа ИП-исполнения (взыскание/реализация/реструктуризация): extract_fields + ip_specific, умерший, дательный падеж, обязательства. Возвращает extracted_fields. Вынесено из analyze."""
        extracted_fields = self.extract_fields(text, "rtk_application")  # Используем rtk_application как базовый тип

        ip_specific_fields = self.extract_ip_enforcement_fields(text)
        logger.info(f"Извлеченные специфичные поля для ИП: {list(ip_specific_fields.keys())}")

        text_lower = text.lower()
        is_deceased = any(keyword in text_lower for keyword in ["умер", "умерший", "смерть", "смерти"])

        for key, value in ip_specific_fields.items():
            if value:
                # ВАЖНО: Не перезаписываем ИНН и ОГРН, если они уже были извлечены из блока "Ответчик:" в основном цикле
                if key in ["inn", "ogrnip", "ogrn"]:
                    if key in extracted_fields and extracted_fields[key]:
                        logger.info(f"Пропускаем перезапись {key} из ip_specific_fields (уже извлечен из блока Ответчик: {extracted_fields[key]})")
                        continue
                if key == 'courtName' and is_deceased:
                    cleaned_value = self._clean_court_name_deceased(value)
                    if cleaned_value:
                        extracted_fields[key] = cleaned_value
                        logger.info(f"Установлено поле ИП {key} (очищено для умершего): {cleaned_value}")
                    else:
                        logger.info(f"Поле ИП {key} очищено до пустого значения, пропускаем")
                else:
                    # Всегда перезаписываем специфичные поля для ИП, даже если они уже есть
                    extracted_fields[key] = value
                    logger.info(f"Установлено поле ИП {key}: {value}")

        if document_type in ["ip_enforcement_statement_collateral", "ip_enforcement_realization_collateral",
                             "ip_enforcement_restructuring_collateral"]:
            extracted_fields["ipHasCollateral"] = "true"
        else:
            collateral_detected = self.detect_ip_collateral(text)
            extracted_fields["ipHasCollateral"] = "true" if collateral_detected else "false"

        obligations = self.extract_obligations(text, extracted_fields)
        if obligations:
            extracted_fields['obligations'] = obligations
            logger.info(f"Добавлено {len(obligations)} обязательств для ИП")
        return extracted_fields

    def _analyze_physical_collateral(self, text, document_type):
        """Ветка анализа ФЛ/ЮЛ с залогом (реализация/реструктуризация/наблюдение/конкурс): extract_fields + поля залога, обязательства. Возвращает extracted_fields. Вынесено из analyze."""
        extracted_fields = self.extract_fields(text, "rtk_application")

        # Извлекаем поля залога (аналогично ИП, но без префикса ИП)
        collateral_fields = self.extract_physical_collateral_fields(text)
        logger.info(f"Извлеченные поля залога: {list(collateral_fields.keys())}")
        for key, value in collateral_fields.items():
            if value:
                extracted_fields[key] = value
                logger.info(f"Установлено поле залога {key}: {value}")

        if "mortgageCollateralDescription1221" not in extracted_fields or not extracted_fields.get("mortgageCollateralDescription1221"):
            logger.info(" Предмет залога [1221] не найден в extract_physical_collateral_fields, пробуем fallback метод...")
            collateral_block = self._extract_mortgage_collateral_block(text)
            if collateral_block:
                extracted_fields["mortgageCollateralDescription1221"] = collateral_block
                logger.info(f"Извлечено mortgageCollateralDescription1221 через fallback метод: {collateral_block[:150]}...")

        if document_type == "observation_collateral":
            extracted_fields["observationHasCollateral"] = "true"
        else:
            extracted_fields["physicalHasCollateral"] = "true"

        obligations = self.extract_obligations(text, extracted_fields)
        if obligations:
            extracted_fields['obligations'] = obligations
            logger.info(f"Добавлено {len(obligations)} обязательств для {document_type}")
        return extracted_fields


    def _drop_law_context_dates(self, fields: Dict[str, Any], text: str) -> None:
        """Удаляет даты, которые в исходном тексте стоят ТОЛЬКО в ссылках на закон/Пленум.

        Дата рождения и иные поля не затрагиваются. Если у даты есть хотя бы одно
        вхождение вне законного контекста — она сохраняется.
        """
        if not text:
            return
        for fname in self._LAW_DATE_FIELDS:
            val = fields.get(fname)
            if not val or not isinstance(val, str):
                continue
            dates = [d.strip() for d in val.split(",") if d.strip()]
            if not dates:
                continue
            kept = []
            for d in dates:
                occ = [m.start() for m in re.finditer(re.escape(d), text)]
                if occ and all(date_in_law_context(text, s, s + len(d)) for s in occ):
                    logger.info(f"Пропуск {fname} (дата в ссылке на закон): {d}")
                    continue
                kept.append(d)
            if kept:
                fields[fname] = ", ".join(kept)
            else:
                fields.pop(fname, None)

    # Типы ТС (многословные — раньше однословных, чтобы «грузовой тягач» не срезался
    # до «грузовой»). «автомобиль» тут НЕ тип, а метка предмета — не берём.
    _VEHICLE_TYPE_RE = (
        r"седельн\w+\s+тягач\w*|грузов\w+\s+тягач\w*|бортов\w+\s+грузов\w+|"
        r"грузов\w+|легков\w+|полуприцеп\w*|прицеп\w*|тягач\w*|автобус\w*|"
        r"самосвал\w*|фургон\w*|мотоцикл\w*|трактор\w*|экскаватор\w*|погрузчик\w*"
    )

    def _extract_vehicle_type(self, desc: str) -> Optional[str]:
        """Вид ТС (грузовой/полуприцеп/прицеп/тягач/…) для наименования залога-авто."""
        m = re.search(r"(" + self._VEHICLE_TYPE_RE + r")", desc, re.IGNORECASE)
        if m:
            t = re.sub(r"\s+", " ", m.group(1).strip().lower())
            return t[0].upper() + t[1:]
        return None

    def _extract_collateral_brand_model(self, desc: str) -> Optional[str]:
        """Марка+модель авто. Поддержка меток «Марка/Модель [ТС]:» и OCR «Мо дель»."""
        d = re.sub(r"мо\s+дель", "модель", desc, flags=re.IGNORECASE)  # OCR-склейка
        mk = re.search(r"марк[аиуе](?:\s*тс)?\s*[:：]\s*([A-Za-zА-Яа-яЁё0-9\- ]{1,30}?)\s*(?=[,;.]|модел|год|vin|кузов|$)", d, re.IGNORECASE)
        md = re.search(r"модел[ьи](?:\s*тс)?\s*[:：]\s*([A-Za-zА-Яа-яЁё0-9\- ]{1,30}?)\s*(?=[,;.]|год|vin|кузов|рама|$)", d, re.IGNORECASE)
        parts = []
        if mk and mk.group(1).strip():
            parts.append(mk.group(1).strip())
        if md and md.group(1).strip():
            parts.append(md.group(1).strip())
        return " ".join(parts) or None

    def _extract_freeform_make_model(self, desc: str, vtype: Optional[str], year: Optional[str]) -> Optional[str]:
        """Марка/модель без меток: «…, Грузовой ИЖ 27175 2009.» «ИЖ 27175».
        Берём текст сразу после вида ТС до года/знака препинания."""
        if not vtype:
            return None
        m = re.search(
            re.escape(vtype) + r"\s+([A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё0-9\- ]{1,40}?)"
            r"(?=\s*(?:19\d{2}|20[0-3]\d)\b|[,;.]|$)",
            desc, re.IGNORECASE,
        )
        if not m:
            return None
        s = re.sub(r"\s+", " ", m.group(1)).strip()
        # Не марка, если это метка или сам вид ТС.
        if not s or re.match(r"(?i)марк|модел|vin|год|тс\b", s):
            return None
        return s

    def _extract_collateral_address(self, desc: str) -> Optional[str]:
        """Извлекает ПОЛНЫЙ адрес объекта недвижимости из описания залога.

        Берёт всё после метки адреса и обрезает хвост (кадастр, стоимость,
        год постройки, площадь и пр.), чтобы в адресе не было постороннего.
        """
        m = re.search(
            r"(?:по\s+адресу|адрес[уе]?|расположен\w*\s+по\s+адресу|местонахожд\w*)\s*[:：,]?\s*([^\n]+)",
            desc, re.IGNORECASE,
        )
        if not m:
            return None
        addr = re.sub(r"\s+", " ", m.group(1)).strip()
        # Обрезаем хвост: кадастр, стоимость, год постройки, площадь, этажность, VIN.
        addr = re.split(
            r"\s*[,(]?\s*(?:\d{2}:\d{2}:\d{6,7}:\d+|кадастров\w+|залогов\w+\s+стоим|оценочн\w+\s+стоим|"
            r"рыночн\w+\s+стоим|начальн\w+\s+(?:продажн\w+\s+)?цен|стоимост\w+|год\s+постройки|"
            r"площад\w+|\d+[\s-]*этажн|\bVIN\b)",
            addr, flags=re.IGNORECASE,
        )[0]
        # Хвостовая «(» от «(кадастровый номер …)» остаётся, если скобку не съел split.
        addr = addr.strip().rstrip(",;.( ")
        return addr if len(addr) >= 6 else None

    def _extract_car_year(self, desc: str) -> Optional[str]:
        """Год выпуска авто: «год выпуска: 2011», «Год: 2020», «Год выпуска ТС: 1992»,
        «2011 г.в.» / «2011 года выпуска» либо хвостовой год «ИЖ 27175 2009.»."""
        m = re.search(r"год\w*(?:\s+выпуска)?(?:\s*тс)?\s*[:：]?\s*(19\d{2}|20[0-3]\d)", desc, re.IGNORECASE)
        if m:
            return m.group(1)
        m = re.search(r"\b(19\d{2}|20[0-3]\d)\s*(?:г\.?\s*в\.?|года?\s+выпуска)", desc, re.IGNORECASE)
        if m:
            return m.group(1)
        # Хвостовой год в свободном описании: «…ИЖ 27175 2009.» (без метки/«г.в.»).
        m = re.search(r"(?<!\d)(19\d{2}|20[0-3]\d)(?!\d)\s*[.;,]?\s*$", desc.strip())
        if m:
            return m.group(1)
        return None

    def _extract_plate(self, desc: str) -> Optional[str]:
        """Гос. рег. знак авто. Метка в любом виде: грз / гсз / г/н / гос (рег.) знак /
        государственный (регистрационный) знак / гос. номер / рег. знак."""
        m = re.search(
            r"(?:\bгрз\b|\bгсз\b|\bг\.?\s*/?\s*н\.?|"
            r"гос(?:ударственн\w+)?\.?\s*(?:рег(?:истрационн\w+)?\.?\s*)?(?:знак|номер)|"
            r"рег(?:истрационн\w+)?\.?\s*знак)"
            r"\s*[:：№]?\s*([А-ЯЁA-Z]{1,3}\s?\d{2,4}\s?[А-ЯЁA-Z]{0,3}\s?\d{0,3})",
            desc, re.IGNORECASE,
        )
        if m and re.search(r"\d", m.group(1)) and re.search(r"[А-ЯЁA-Z]", m.group(1)):
            return re.sub(r"\s+", "", m.group(1)).upper()
        return None

    def _extract_collateral_value(self, desc: str) -> Optional[str]:
        """Извлекает залоговую стоимость предмета залога (число)."""
        m = re.search(
            r"(?:залогов\w+\s+стоимост\w+|оценочн\w+\s+стоимост\w+|рыночн\w+\s+стоимост\w+|"
            r"стоимост\w+\s+(?:предмета\s+)?залога|начальн\w+\s+(?:продажн\w+\s+)?цен\w+|стоимост\w+)"
            r"\s*(?:залога|объекта)?\s*[:：]?\s*(?:в\s+размере\s+)?"
            r"([0-9][0-9\s  .,]*[0-9]|[0-9])\s*(?:руб|₽|р\.)",
            desc, re.IGNORECASE,
        )
        if m:
            num = re.sub(r"[\s  ]", "", m.group(1))
            return num
        return None

    # Фраза залоговой стоимости (для вырезания из описания «иного» залога).
    _VALUE_PHRASE_RE = (
        r"[,;.]?\s*(?:общ\w+\s+)?(?:залогов\w+\s+стоимост\w+|оценочн\w+\s+стоимост\w+|рыночн\w+\s+стоимост\w+|"
        r"стоимост\w+\s+(?:предмета\s+)?залога|начальн\w+\s+(?:продажн\w+\s+)?цен\w+|стоимост\w+)"
        r"\s*(?:залога|объекта)?\s*[:：]?\s*(?:в\s+размере\s+)?[0-9][0-9\s.,]*[0-9]?\s*(?:руб\w*|₽|р\.)"
    )

    def _strip_value_phrase(self, text: str) -> str:
        """Убирает фразу залоговой стоимости из текста (чтобы не дублировать в описании)."""
        return re.sub(self._VALUE_PHRASE_RE, "", text, flags=re.IGNORECASE).strip(" ,;.-—–")

    def _extract_other_object_name(self, desc: str) -> Optional[str]:
        """Наименование «иного» предмета залога: до первой запятой/двоеточия (там
        обычно начинаются атрибуты — «, страна изготовления:», «, находящееся по
        адресу:») либо по ключевому слову вида."""
        m = re.match(r"\s*([^:：,\n]{2,100}?)\s*[:：,]", desc)
        if m:
            name = m.group(1)
        else:
            km = re.search(
                r"(ценн\w+\s+бумаг\w*|акци\w+|облигаци\w+|вексел\w+|"
                r"дол[яюи]\s+в\s+уставн\w+\s+капитал\w*|дол[яюи]\s+в\s+праве\w*|"
                r"оборудовани\w*|товар\w*\s+в\s+оборот\w*|имуществ\w+\s+прав\w*\s+требовани\w*|"
                r"имуществ\w+\s+прав\w*|прав\w*\s+требовани\w*|\bпа[йи]\b)",
                desc, re.IGNORECASE,
            )
            name = km.group(1) if km else desc.split(",")[0][:40]
        name = self._strip_value_phrase(name).strip(" -—–,;.")
        return (name[0].upper() + name[1:]) if name else None

    def _extract_collateral_object_name(self, desc: str) -> Optional[str]:
        """Извлекает вид объекта недвижимости (дом / земельный участок / квартира …).

        Здание/помещение/дом и т.п. — если есть, это САМ объект (земельный участок
        под ним фигурирует лишь как «расположенное на земельном участке …»). Поэтому
        строения ищем первым проходом, землю — вторым."""
        m = re.search(
            r"(жил\w+\s+дом|нежил\w+\s+(?:помещени\w+|здани\w+)|"
            r"квартир\w+|комнат\w+|машино-?мест\w*|гараж\w*|нежил\w+\s+здани\w+|"
            r"здани\w+|строени\w+|сооружени\w+|помещени\w+|\bдом\b)",
            desc, re.IGNORECASE,
        )
        if not m:
            m = re.search(r"(земельн\w+\s+участ\w+)", desc, re.IGNORECASE)
        if m:
            name = m.group(1).strip()
            return name[0].upper() + name[1:]
        return None

    # ------------------------------------------------------------------ #
    # Предмет ипотеки: структурированные объекты залога для формы.        #
    # Отдельный top-level ключ result["mortgageProperties"] — вне fields  #
    # и вне collaterals, поэтому golden-снимок его не фиксирует.          #
    # ------------------------------------------------------------------ #

    @staticmethod
    def _mp_num(s: str) -> str:
        """Число суммы → канон без разрядных пробелов, десятичная запятая как есть."""
        s = (s or "").replace(" ", " ").replace(" ", " ")
        return re.sub(r"\s+", "", s).strip(" .,")

    @staticmethod
    def _mp_type(name: str) -> Optional[str]:
        """Тип объекта из его наименования/описания (для сопоставления «в том числе»)."""
        n = (name or "").lower()
        if "участ" in n:
            return "участок"
        if "дом" in n:
            return "дом"
        if "квартир" in n:
            return "квартира"
        if "машино" in n:
            return "машиноместо"
        if "гараж" in n:
            return "гараж"
        if "помещ" in n:
            return "помещение"
        return None

    def _mp_description(self, raw: str) -> str:
        """Основное описание объекта: до адреса/кадастра/записи ЕГРН."""
        d = (raw or "").strip().lstrip("-–—• \t").strip()
        d = re.split(
            r",?\s*(?:расположен\w*|находящ\w*|по\s+адресу|кадастров\w+|Запис\w+\s+в\s+ЕГРН)",
            d, maxsplit=1, flags=re.IGNORECASE,
        )[0].strip(" ,;.")
        return (d[0].upper() + d[1:]) if d else d

    def _mp_egrn(self, desc: str):
        """(номер ЕГРН, дата) из «Запись в ЕГРН: …, № … от …» описания объекта.

        Долевая собственность даёт несколько записей (2/4, 1/4, 1/4) — собираем все
        с метками долей: «2/4 <№>; 1/4 <№>; …», дата — первой записи."""
        m = re.search(r"Запис\w+\s+в\s+ЕГРН\s*:?\s*(.+?)(?:Ипотека|$)", desc or "", re.IGNORECASE | re.DOTALL)
        if not m:
            return "", ""
        seg = m.group(1)
        recs = re.findall(
            r"(?:(\d\s*/\s*\d)\s*,\s*)?(?:№|N)\s*([0-9A-Za-zА-Яа-я:/\-]+?)\s+от\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
            seg,
        )
        if not recs:
            m2 = re.search(r"(?:№|N)\s*([0-9A-Za-zА-Яа-я:/\-]+)", seg)
            return (m2.group(1).strip(" .") if m2 else ""), ""
        if len(recs) == 1:
            _share, num, date = recs[0]
            return num.strip(" ."), date.replace(",", ".")
        parts = []
        for share, num, _date in recs:
            num = num.strip(" .")
            parts.append(f"{share.replace(' ', '')} {num}" if share else num)
        return "; ".join(parts), recs[0][2].replace(",", ".")

    def _mp_breakdown(self, region: str) -> Dict[str, str]:
        """Разбивка «в том числе <тип> — <сумма> руб.» → {тип: сумма}."""
        bd: Dict[str, str] = {}
        for m in re.finditer(
            r"(жил\w+\s+дом\w*|земельн\w+\s+участ\w*|квартир\w+|нежил\w+\s+помещени\w*|"
            r"помещени\w+|гараж\w*|машино-?мест\w*)\s*[-–—:]?\s*"
            r"([\d][\d\s  .,]*\d|\d)\s*руб",
            region, re.IGNORECASE,
        ):
            typ = self._mp_type(m.group(1))
            if typ and typ not in bd:
                bd[typ] = self._mp_num(m.group(2))
        return bd

    def _mp_amounts(self, text: str, kind: str):
        """(итог, разбивка) для оценки (kind='value') или начальной цены (kind='start')."""
        if kind == "value":
            pats = [
                r"котор\w+\s+составил\w*\s+([\d\s  .,]+?)\s*руб",
                r"рыночн\w+\s+стоимост\w+[^.]{0,80}?составля\w+\s+([\d\s  .,]+?)\s*руб",
            ]
        else:
            pats = [
                r"начальн\w+\s+продажн\w+\s+цен\w+[^.]{0,80}?составля\w+\s+([\d\s  .,]+?)\s*руб",
                r"Установить\s+начальн\w+\s+цен\w+\s+продажи[^.]{0,90}?в\s+размере\s+([\d\s  .,]+?)\s*руб",
                r"начальн\w+\s+продажн\w+\s+цен\w+\s+должна\s+быть\s+установлен\w*\s+в\s+размере\s+([\d\s  .,]+?)\s*руб",
            ]
        stops = r"(?:Таким\s+образом|Следовательно|Определить|Согласно|В\s+соответствии|\n\n)"
        for pat in pats:
            m = re.search(pat, text, re.IGNORECASE)
            if m:
                total = self._mp_num(m.group(1))
                region = re.split(stops, text[m.end(): m.end() + 300], maxsplit=1, flags=re.IGNORECASE)[0]
                return total, self._mp_breakdown(region)
        return "", {}

    def _mp_npc_strategy(self, text: str) -> str:
        """Стратегия НПЦ (%): приоритет «залоговая стоимость … в размере NN%»,
        фолбэк — словесное «восьмидесяти процентам» = 80%."""
        m = re.search(r"залогов\w+\s+стоимост\w+[^\n]{0,80}?в\s+размере\s+(\d{1,3})\s*%", text, re.IGNORECASE)
        if m:
            return m.group(1) + "%"
        if re.search(r"восьмидесяти\s+процент", text, re.IGNORECASE):
            return "80%"
        return ""

    def _mp_appraisal_report(self, text: str) -> str:
        """Отчёт об оценке: «№<номер> от <дата>» (общий для всех объектов)."""
        m = re.search(
            r"отчет\w*\s+об\s+оценк\w+[^\n№N]{0,40}?(?:№|N)\s*([^\s,()\n]+)"
            r"(?:\s+(?:от|г\.?)\s*(\d{1,2}[.,]\d{1,2}[.,]\d{4}))?",
            text, re.IGNORECASE,
        )
        if m:
            num = m.group(1).strip(" .,")
            date = (m.group(2) or "").replace(",", ".")
            return f"№{num} от {date}" if date else f"№{num}"
        # ДДУ и часть исков: оценка ссылается на «заключение № <номер> от <дата>»
        # (без слов «об оценке»/«о стоимости имущества»).
        m3 = re.search(
            r"заключени\w+[^\n№N]{0,20}?(?:№|N)\s*([^\s,()\n]+)"
            r"\s+от\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
            text, re.IGNORECASE,
        )
        if m3:
            num = m3.group(1).strip(" .,")
            date = m3.group(2).replace(",", ".")
            return f"№{num} от {date}"
        m2 = re.search(
            r"заключени\w+\s+о\s+стоимости\s+имуществ\w*\s+от\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
            text, re.IGNORECASE,
        )
        if m2:
            return "от " + m2.group(1).replace(",", ".")
        return ""

    # Начало записи объекта залога (в начале строки, после снятого буллета).
    _MP_OBJ_START_RE = re.compile(
        r"^(?:жил\w+\s+дом|квартир\w+|земельн\w+\s+участ\w+|нежил\w+\s+(?:помещени\w+|здани\w+)|"
        r"помещени\w+|гараж\w*|машино-?мест\w*|здани\w+|строени\w+|комнат\w+)",
        re.IGNORECASE,
    )

    def _mp_collateral_chunks(self, text: str) -> List[str]:
        """Абзацы-предметы из блока «…залог… а именно: …» до «В силу п.1 ст. 77…».

        Объекты идут отдельными строками, но буллет «-» бывает только у первого
        (грязная вёрстка). Поэтому режем не по буллетам, а по строкам, начинающимся
        с ключевого слова вида объекта; строки-продолжения (перенос адреса) клеим
        к текущему объекту."""
        m = re.search(
            r"а\s+именно\s*:?\s*(.+?)(?:В\s+силу\s+(?:п\.?\s*1\s+)?ст\.?\s*77|В\s+силу\s+ст\.|"
            r"Банк\s+исполнил|Право\s+собственности\s+на\s+вышеуказанн|\n\s*\n|$)",
            text, re.IGNORECASE | re.DOTALL,
        )
        if not m:
            return []
        block = m.group(1)
        chunks: List[str] = []
        cur: Optional[str] = None
        for raw in block.split("\n"):
            ln = raw.strip(" -–—•\t")
            if not ln:
                continue
            if self._MP_OBJ_START_RE.match(ln):
                if cur:
                    chunks.append(cur)
                cur = ln
            elif cur is not None:
                cur += " " + ln
        if cur:
            chunks.append(cur)
        return chunks

    @staticmethod
    def _mp_cadastral(chunk: str) -> str:
        # «кадастровый номер: 23:…» и «кадастровый номер земельного участка № 11:…»
        # (между «номер» и числом могут стоять слова «земельного участка №»).
        m = re.search(r"кадастров\w+\s+номер[^\d\n]{0,40}?(\d{2}:\d{2}:\d{5,7}:\d+)", chunk, re.IGNORECASE)
        return m.group(1) if m else ""

    @staticmethod
    def _mp_address(chunk: str) -> str:
        m = re.search(r"по\s+адресу\s*:?\s*(.+?)(?:,?\s*кадастров\w+\s+номер|\.\s*Запис\w+|$)", chunk, re.IGNORECASE)
        return m.group(1).strip(" ,;.") if m else ""

    def _detect_mortgage_kind(self, text: str) -> str:
        """Вид ипотеки: 'military' | 'ddu' | 'civil'.

        Военная (приоритет): третье лицо — ФГКУ «Росвоенипотека» (накопительно-
        ипотечная система жилищного обеспечения военнослужащих) или продукт «Военная
        ипотека».
        ДДУ: долевое строительство — «участник(а) долевого строительства», «договор
        участия в долевом строительстве», «инвестирование строительства», застройщик
        (кредит на строящееся жильё, предмет залога — права требования по ДДУ)."""
        if re.search(
            r"накопительно-?ипотечн\w+\s+систем\w+|Росвоенипотек\w*|Военн\w+\s+ипотек\w+",
            text, re.IGNORECASE,
        ):
            return "military"
        if re.search(
            r"долев\w+\s+строительств\w*|договор\w*\s+участия\s+в\s+строительстве|"
            r"инвестировани\w+\s+строительств\w*|специализированн\w+\s+застройщик\w*|"
            r"участник\w*\s+долевого\s+строительства",
            text, re.IGNORECASE,
        ):
            return "ddu"
        return "civil"

    def _mp_ddu_contract(self, text: str):
        """(номер договора ДДУ, дата) из предмета залога/тела: «по ДОГОВОРУ № <c>
        УЧАСТИЯ В ДОЛЕВОМ СТРОИТЕЛЬСТВЕ от <d>» либо «договор участия в долевом
        строительстве № <c> … от <d>»."""
        m = re.search(
            r"ДОГОВОР\w*\s*№?\s*(.+?)\s+УЧАСТИЯ\s+В\s+ДОЛЕВОМ\s+СТРОИТЕЛЬСТВЕ\s+от\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
            text, re.IGNORECASE,
        )
        if not m:
            m = re.search(
                r"договор\w*\s+участия\s+в\s+долевом\s+строительстве\s*№?\s*(.+?)\s+от\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
                text, re.IGNORECASE,
            )
        if not m:
            return "", ""
        num = re.sub(r"\s+", " ", m.group(1)).strip(" .,№")
        return num, m.group(2).replace(",", ".")

    @staticmethod
    def _mp_ddu_egrn(seg: str):
        """(номер записи ЕГРН, дата) для ДДУ: «Запись в ЕГРН от <дата> … номер … № <номер>»
        (дата ПЕРЕД номером, в отличие от гражданской ипотеки)."""
        m = re.search(
            r"Запис\w*\s+в\s+ЕГРН\s+от\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})[^№N\n]{0,80}?(?:№|N)\s*([0-9:/\-]+)",
            seg, re.IGNORECASE,
        )
        if m:
            return m.group(2).strip(" ."), m.group(1).replace(",", ".")
        return "", ""

    def _build_ddu_property(self, text: str) -> List[Dict[str, Any]]:
        """Предмет ипотеки для ДДУ: одна карточка — имущественные права требования
        участника долевого строительства по договору (а не физический объект)."""
        # Блок предмета залога из просительной части: «Обратить взыскание … на предмет
        # залога: - имущественные права требования … по ДОГОВОРУ …».
        m = re.search(
            r"имуществен\w+\s+прав\w+\s+требован\w+(.+?)(?:Установить\s+начальн|Определить\s+способ|"
            r"\n\s*\d+\.\s|$)",
            text, re.IGNORECASE | re.DOTALL,
        )
        seg = m.group(0) if m else text
        contract, ddate = self._mp_ddu_contract(seg if m else text)
        egrn, egrn_date = self._mp_ddu_egrn(seg)
        val_total, _ = self._mp_amounts(text, "value")
        st_total, _ = self._mp_amounts(text, "start")
        addr = self._mp_address(seg)
        return [{
            "id": "mortgageProperty-0",
            "description": "Имущественные права требования участника долевого строительства",
            "cadastralNumber": self._mp_cadastral(seg),
            "address": addr,
            "value": val_total or "",
            "startingPrice": st_total or "",
            "npcStrategy": self._mp_npc_strategy(text),
            "appraisalReport": self._mp_appraisal_report(text),
            "egrnRecord": egrn,
            "egrnRecordDate": egrn_date,
            "dduContract": contract,
            "dduDate": ddate,
        }]

    def _build_mortgage_properties(
        self, text: str, collaterals: List[Dict[str, Any]], kind: Optional[str] = None,
    ) -> List[Dict[str, Any]]:
        """Структурированные предметы ипотеки для формы.

        Для ДДУ — отдельная ветка (права требования по договору долевого участия).
        Для гражданской/военной: объекты (вид/описание/кадастр/адрес/ЕГРН) разбираем
        прямо из блока залога — надёжнее общего `collaterals`, который на грязной
        вёрстке (буллет только у первого объекта) теряет второй. Стоимость/начальную
        цену — из абзаца оценки (разбивка «в том числе <тип>») по типу объекта;
        стратегию НПЦ и отчёт об оценке — общие для всех объектов."""
        if kind == "ddu":
            return self._build_ddu_property(text)

        chunks = self._mp_collateral_chunks(text)
        # Фолбэк: если блок не найден, но общий парсер что-то дал — используем его.
        if not chunks and collaterals:
            chunks = [c.get("description") or c.get("objectName") or "" for c in collaterals]
        chunks = [c for c in chunks if c and len(c) >= 10]
        if not chunks:
            return []

        npc = self._mp_npc_strategy(text)
        report = self._mp_appraisal_report(text)
        val_total, val_bd = self._mp_amounts(text, "value")
        st_total, st_bd = self._mp_amounts(text, "start")
        out: List[Dict[str, Any]] = []
        for i, chunk in enumerate(chunks):
            typ = self._mp_type(chunk)
            # Есть разбивка «в том числе <тип>» — берём сумму объекта; иначе общий
            # итог распространяется на всё заложенное имущество (ставим его на каждый).
            value = val_bd.get(typ, "") if val_bd else (val_total or "")
            start = st_bd.get(typ, "") if st_bd else (st_total or "")
            egrn, egrn_date = self._mp_egrn(chunk)
            out.append({
                "id": f"mortgageProperty-{i}",
                "description": self._mp_description(chunk),
                "cadastralNumber": self._mp_cadastral(chunk),
                "address": self._mp_address(chunk),
                "value": value or "",
                "startingPrice": start or "",
                "npcStrategy": npc,
                "appraisalReport": report,
                "egrnRecord": egrn,
                "egrnRecordDate": egrn_date,
                "dduContract": "",
                "dduDate": "",
            })
        return out

    _VIN_COLLATERAL_RE = re.compile(
        r"(?:[-–—]\s*)?(?:движимое\s+имущество|автомобиль)\s*[:：]\s*VIN\s*[:：]?\s*"
        r"[A-ZА-Я0-9]{5,20}[^\n]*?\.(?=\s*(?:\n|$))",
        re.IGNORECASE,
    )

    def _extract_vin_collateral_items(self, text: str) -> List[str]:
        """Залоговые авто/спецтехника по VIN — глобальный скан по ВСЕМУ документу.

        Заявления с несколькими «Обязательство №N» дают у КАЖДОГО своё
        описание предмета залога («движимое имущество: VIN: … ; Год: …» /
        «автомобиль: VIN: …»), но `mortgageCollateralDescription1221` берёт
        только первое совпадение в тексте — остальные обязательства теряют
        свой залог. Дедуп по VIN — один физический предмет может обеспечивать
        сразу несколько обязательств (повторяется в тексте дословно).
        """
        if not text:
            return []
        norm = text.replace("\xa0", " ")
        items: List[str] = []
        seen_vins = set()
        for m in self._VIN_COLLATERAL_RE.finditer(norm):
            desc = re.sub(r"^\s*[-–—]\s*", "", m.group(0)).strip()
            # Дедуп-ключ — сам код VIN как он есть в документе (не привязываемся к
            # стандартной длине 17: в исходниках попадаются урезанные/опечатанные VIN).
            vin_m = re.search(r"VIN\s*[:：]?\s*([A-ZА-Я0-9]{5,20})", desc, re.IGNORECASE)
            key = vin_m.group(1).upper() if vin_m else desc
            if key in seen_vins:
                continue
            seen_vins.add(key)
            items.append(desc)
        return items

    def _extract_all_collateral_items(self, text: str) -> List[str]:
        """Сканирует ВЕСЬ документ и собирает описания всех предметов залога.

        Залоги бывают у нескольких обязательств. Берём:
          - маркированные строки «- <дом/квартира/земельный участок/Автомобиль/марка …>»;
          - инлайн «…автотранспортное средство: <описание>».
        Каждое описание — до конца строки (со всеми атрибутами: VIN, год, грз, стоимость).
        """
        if not text:
            return []
        norm = text.replace("\xa0", " ").replace(" ", " ")
        items: List[str] = []

        _inventory_re = re.compile(
            r"(?:имеется|зарегистрирован\w*)\s+следующ\w+\s+(?:движим\w+\s+|недвижим\w+\s+)?имуществ\w*"
            r"|опис\w+\s+имуществ"

            r"|не\s+зарегистрирован\w*\s+(?:движим\w+|недвижим\w+)(?:\s+и\s+(?:движим\w+|недвижим\w+))?\s+имуществ\w*",
            re.IGNORECASE,
        )

        def _is_inventory_item(pos: int) -> bool:

            ctx = norm[max(0, pos - 1500):pos]
            last = None
            for im in _inventory_re.finditer(ctx):
                last = im
            return bool(last) and "залог" not in ctx[last.end():].lower()

        # Маркер пункта перечня: буллет ЛИБО нумерация «1)», «2.», «3 )».
        # Нумерация обязательна, иначе списки вида «1) Объект недвижимости-Жилой
        # дом … 3) Земельный участок …» распознавались лишь частично: буллетом
        # ошибочно работал дефис внутри «недвижимости-Жилой», а у пунктов без
        # такого дефиса (земельный участок) маркера не находилось вовсе.
        bullet_re = re.compile(
            r"(?:[-–—•]|(?:^|\n)[ \t]*\d{1,2}\s*[).])\s*"
            r"((?:Автомобил\w*|марк[аи]\s*[:：]|жил\w*\s*дом|\bдом\b|квартир\w*|"
            r"земельн\w+\s+участ\w*|нежил\w*|помещени\w*|здани\w*|гараж\w*|машино-?мест\w*|"
            r"комнат\w*|строени\w*|сооружени\w*|"
            # «иное»: ценные бумаги, доли, оборудование, товары, имущественные права и т.п.
            r"ценн\w+\s+бумаг\w*|акци\w+|облигаци\w+|вексел\w+|дол[яюи]\s+в\s+(?:уставн|праве)|"
            r"оборудовани\w*|товар\w*\s+в\s+оборот\w*|имуществ\w+\s+прав\w*|прав\w*\s+требовани\w*|"
            r"\bпа[йи]\b)[^\n]*)",
            re.IGNORECASE,
        )
        for m in bullet_re.finditer(norm):
            s = self._join_bullet_continuation(norm, m.start(1)).rstrip(" .,;")
            if len(s) >= 8 and not _is_inventory_item(m.start()):
                items.append(s)
        for m in re.finditer(r"(?:авто)?транспортн\w+\s+средств\w*\s*[:：]\s*([^\n]+)", norm, re.IGNORECASE):
            s = m.group(1).strip().rstrip(" .,;")
            if len(s) >= 8 and not _is_inventory_item(m.start()):
                items.append(s)

        items.extend(self._extract_pledge_block_items(norm))

        # «движимое имущество: VIN: …»/«автомобиль: VIN: …» без буллета из
        # списка выше (метка «движимое имущество» в него не входит) — отдельные
        # предметы залога у КАЖДОГО «Обязательство №N» в документах со
        # Сбербанковским форматом кредитных линий. Дедуп по VIN — уже найденные
        # тем же кодом (если предмет попал через bullet_re/pledge-выше) не дублируем.
        seen_vins = {vm.group(1).upper() for it in items
                     for vm in [re.search(r"VIN\s*[:：]?\s*([A-ZА-Я0-9]{5,20})", it, re.IGNORECASE)]
                     if vm}
        for desc in self._extract_vin_collateral_items(norm):
            vm = re.search(r"VIN\s*[:：]?\s*([A-ZА-Я0-9]{5,20})", desc, re.IGNORECASE)
            key = vm.group(1).upper() if vm else desc
            if key in seen_vins:
                continue
            seen_vins.add(key)
            items.append(desc)
        return items

    def _extract_pledge_block_items(self, norm: str) -> List[str]:
        """Предметы залога из блоков «в залог передано следующее имущество: …».
        Возвращает описания предметов (каждое оканчивается на «стоимостью N руб.»)."""
        _val = r"стоимост\w+\s+[0-9][0-9\s.,]*[0-9]\s*руб"
        items: List[str] = []
        for bm in re.finditer(r"в\s+залог\s+переда\w+\s+", norm, re.IGNORECASE):
            head = norm[bm.end(): bm.end() + 5000]
            # Итемизированный список: «следующее [движимое] имущество: <предметы>».
            lm = re.match(r"следующее\s+[^:\n]{0,40}?имуществ\w*\s*[:：]", head, re.IGNORECASE)
            if lm:
                block = head[lm.end():]
                # Конец блока — итоговая строка «Общая стоимость … составляет … руб».
                endm = re.search(r"Общая\s+стоимост\w+.*?составляет.*?руб", block,
                                 re.IGNORECASE | re.DOTALL)
                if endm:
                    block = block[:endm.start()]
                buf = ""
                for line in block.split("\n"):
                    buf += " " + line.strip()
                    if re.search(_val, buf, re.IGNORECASE):
                        s = re.sub(r"\s+", " ", buf).strip(" .,;")
                        if len(s) >= 12 and "общая стоимост" not in s.lower():
                            items.append(s)
                        buf = ""
                continue
            # Сводная форма без перечня: «…[движимое] имущество … общей стоимостью N
            # руб.» (предметы вынесены в Приложение) — одна карточка «Иное».
            sm = re.match(
                r"([^:\n]{0,25}?имуществ\w*.{0,240}?общ\w+\s+" + _val + r")",
                head, re.IGNORECASE | re.DOTALL,
            )
            if sm:
                s = re.sub(r"\s+", " ", sm.group(1)).strip(" .,;")
                if len(s) >= 12:
                    items.append(s)
        return items

    def _join_bullet_continuation(self, norm: str, obj_start: int) -> str:
        """Собирает предмет залога, перенесённый pdf2docx на несколько строк:
        адрес и «(кадастровый номер …)» часто уезжают на следующие строки, а
        `bullet_re` берёт только первую (до `\\n`). Продолжение приклеиваем ТОЛЬКО
        пока первая строка НЕ завершена терминатором перечня (`;`/`.`) — т.е. это
        реальный перенос, а не следующий пункт/предложение. Стоп — на терминаторе,
        следующем буллете или пустой строке. Дефис-переносы («об-\\nласти») склеиваем."""
        window = norm[obj_start: obj_start + 600]
        lines = window.split("\n")
        acc = (lines[0] if lines else "").strip()
        # Первая строка уже завершена — предмет однострочный, ничего не тянем.
        if re.search(r"[;.]\s*$", acc):
            return acc
        for cur in (ln.strip() for ln in lines[1:]):
            # Стоп на следующем пункте перечня. Нумерация («2) Объект…») здесь так
            # же обязательна, как буллет: без неё адрес первого предмета вбирал в
            # себя начало второго — «…ул. Ленина, 65 2) Объект недвижимости-Жилой
            # дом, общей». Пункт без завершающей точки — обычное дело в заявлениях.
            if not cur or re.match(r"(?:[-–—•]|\d{1,2}\s*[).])\s*\S", cur):
                break
            # Последний пункт перечня не имеет следующего маркера, и склейка
            # утекала в текст за перечнем («…ул. Ленина, 65 ПАО Сбербанк
            # обязательства по предоставлению кредита исполнены…»). Адрес,
            # оканчивающийся номером дома, считаем завершённым — продолжаем только
            # ради кадастрового номера, который дописывают отдельной строкой.
            if (re.search(r"по\s+адресу", acc, re.IGNORECASE) and re.search(r"\d\s*$", acc)
                    and not re.match(r"[(\[]|кадастр", cur, re.IGNORECASE)):
                break
            acc = (acc[:-1] + cur) if acc.endswith("-") else (acc + " " + cur)
            if re.search(r"[;.]\s*$", acc):
                break
        return acc.strip()

    def _collateral_dedupe_key(self, obj: Dict[str, Any]):
        """Ключ уникальности предмета залога (VIN для авто, кадастр/адрес для недвижимости)."""
        t = obj.get("collateralType")
        if t == "auto":
            return ("auto", (obj.get("vin") or obj.get("brandModel") or obj.get("description", "")[:60]).upper())
        if t == "real_estate":
            # Кадастровый номер уникален — им и различаем. Без него адреса мало:
            # на одном участке стоят два жилых дома с одним почтовым адресом, и
            # ключ по адресу схлопывал их в один предмет. Добавляем название с
            # площадью («Жилой дом, площадь 114,8 кв.м»), которое их и различает;
            # для настоящего дубля оно совпадает, так что дедуп продолжает работать.
            cadastral = (obj.get("cadastralNumber") or "").strip()
            if cadastral:
                return ("re", cadastral.lower())
            address = (obj.get("address") or "").strip().lower()
            name = (obj.get("objectName") or obj.get("description", "")[:60]).strip().lower()
            return ("re", address, name)
        # «Иное»: описания однотипных предметов (линия № 6 / № 7) совпадают в первых
        # символах — добавляем стоимость, чтобы разные предметы не схлопнулись.
        return ("other", obj.get("description", "")[:80].lower(), obj.get("collateralValue", ""))

    def _dedupe_collaterals(self, objs: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Убирает дубли предметов залога, оставляя наиболее заполненный."""
        def filled(o):
            return sum(1 for k in ("objectName", "collateralValue", "cadastralNumber",
                                   "address", "vin", "brandModel") if o.get(k))
        by_key: Dict[Any, Dict[str, Any]] = {}
        order: List[Any] = []
        for o in objs:
            key = self._collateral_dedupe_key(o)
            if key not in by_key:
                by_key[key] = o
                order.append(key)
            elif filled(o) > filled(by_key[key]):
                by_key[key] = o
        result = [by_key[k] for k in order]
        for i, o in enumerate(result):
            o["id"] = f"collateral-{i}"
        return result

    def _collateral_has_substance(self, obj: Dict[str, Any]) -> bool:
        """Проверяет, что предмет залога реальный (а не шаблонный «Предметом залога является …»).

        Недвижимость/транспорт — всегда реальны (по ключевым словам классификации).
        «Иное» — только если содержит реальный предмет: ценные бумаги, долю, оборудование,
        товары в обороте, имущественные права и т.п.
        """
        if obj.get("collateralType") in ("auto", "real_estate"):
            return True

        if not obj.get("collateralValue"):
            return False
        name = (obj.get("objectName") or "").strip()
        if len(name) >= 8 and name.lower() != "иное":
            return True
        desc = (obj.get("description") or "").lower()
        return bool(re.search(
            r"ценн\w+\s+бумаг|акци\w|облигаци|вексел|дол[яюи]\s+в\s+(?:уставн|праве)|"
            r"оборудовани|товар\w*\s+в\s+оборот|имуществ\w+\s+прав|прав\w*\s+требовани|"
            r"\bпа[йи]\b|спецтехник|самоходн\w+\s+машин",
            desc,
        ))

    def _make_collateral_obj(self, idx: int, desc: str) -> Dict[str, Any]:
        """Создаёт объект залога с определённым типом и извлечёнными полями."""
        ctype = self._classify_collateral(desc)
        clip = desc if len(desc) < 200 else desc[:200]
        obj = {
            "id": f"collateral-{idx}", "description": desc, "collateralType": ctype,
            "objectName": "", "collateralValue": "", "cadastralNumber": "", "address": "",
            "vin": "", "brandModel": "", "otherDescription": "",
        }
        # Залоговая стоимость (для любого типа).
        val = self._extract_collateral_value(desc)
        if val:
            obj["collateralValue"] = val

        if ctype == "auto":
            vin = re.search(r"(?:VIN\s*[:：]?\s*)?\b([A-HJ-NPR-Z0-9]{17})\b", desc, re.IGNORECASE)
            if vin:
                obj["vin"] = vin.group(1).upper()
            vtype = self._extract_vehicle_type(desc)
            year = self._extract_car_year(desc)
            bm = self._extract_collateral_brand_model(desc)
            if not bm:  # марка/модель без меток: «…Грузовой ИЖ 27175 2009.»
                bm = self._extract_freeform_make_model(desc, vtype, year)
            if bm:
                obj["brandModel"] = bm
            # Наименование авто: «<Вид ТС> <Марка Модель>, <год> г.в., гос. знак…,
            # цвет…, кузов…». Вид ТС и марка/модель — «голова», остальное — атрибуты.
            head = " ".join(p for p in (vtype, bm) if p).strip()
            attrs = []
            if year:
                attrs.append(f"{year} г.в.")
            plate = self._extract_plate(desc)
            if plate:
                attrs.append(f"гос. знак {plate}")
            cm = re.search(r"цвет\s*[:：]?\s*([А-ЯЁа-яё\-]+)", desc, re.IGNORECASE)
            if cm:
                attrs.append(f"цвет {cm.group(1).lower()}")
            km = re.search(r"кузов\s*[№:：]?\s*([A-ZА-ЯЁ0-9\-]{4,})", desc, re.IGNORECASE)
            if km:
                attrs.append(f"кузов {km.group(1)}")
            name = (head + ", " if head and attrs else head) + ", ".join(attrs)
            obj["objectName"] = name.strip(" ,") or "Автомобиль"
        elif ctype == "real_estate":
            cad = re.search(r"(\d{2}:\d{2}:\d{6,7}:\d{1,6})", desc)
            if cad:
                obj["cadastralNumber"] = cad.group(1)
            addr = self._extract_collateral_address(desc)
            if addr:
                obj["address"] = addr
            # Наименование — вид объекта + доступные атрибуты (год постройки, площадь).
            parts = [self._extract_collateral_object_name(desc) or "Недвижимость"]
            yb = re.search(r"год\s+(?:постройки|возведения|строительства)\s*[:：]?\s*(\d{4})", desc, re.IGNORECASE)
            if yb:
                parts.append(f"год постройки {yb.group(1)}")
            am = re.search(r"площад\w*\s*[:：]?\s*([\d.,]+)\s*(?:кв\.?\s*м|м2|м²|кв\.?\s*метр\w*)", desc, re.IGNORECASE)
            if am:
                parts.append(f"площадь {am.group(1)} кв.м")
            fm = re.search(r"(\d+)[\s-]*этажн\w+", desc, re.IGNORECASE)
            if fm:
                parts.append(f"{fm.group(1)}-этажный")
            obj["objectName"] = ", ".join(parts)
        else:
            # «Иное»: наименование и стоимость — отдельными полями; в описание кладём
            # остаток (без дублирования наименования и стоимости).
            name = self._extract_other_object_name(desc)
            obj["objectName"] = name or "Иное"
            rest = re.sub(r"^\s*[-–—•]\s*", "", desc)  # срезаем ведущий маркер
            # Убираем ведущее наименование с двоеточием («Ценные бумаги: …» «…»).
            mcolon = re.match(r"\s*[^:：\n]{2,60}[:：]\s*", rest)
            if mcolon:
                rest = rest[mcolon.end():]
            elif name and rest.lower().startswith(name.lower()):
                rest = rest[len(name):].lstrip(" :,-—–")
            obj["otherDescription"] = self._strip_value_phrase(rest)
        return obj

    def _is_car_collateral_document(self, text: str) -> bool:
        """Определяет наличие залога авто: связка «марка … (модель/VIN/год/кузов)»
        в любом месте текста (в т.ч. инлайн «…автотранспортное средство: Автомобиль, марка: …»)."""
        if not text:
            return False
        return bool(re.search(
            r"марк[аи]\s*[:：][^\n]{0,150}?(?:модель|vin|год\s+выпуска|кузов|рама)",
            text, re.IGNORECASE,
        ))

    def _extract_car_collateral_1221(self, text: str) -> Optional[str]:
        """
        Извлекает описание предмета залога [1221] для залога авто в формате:
        марка: Opel, модель: Astra, год выпуска: 2021, VIN: WP6AB2A45BLA52999
        Ищет предложение, начинающееся с «марка» и содержащее модель, год, VIN, кузов, рама.
        """
        if not text:
            return None
        normalized = text.replace('\u202f', ' ').replace('\xa0', ' ')
        # Связка «(автотранспортное средство:) (Автомобиль,) марка: …» до конца строки —
        # чтобы захватить и год выпуска, и VIN, и залоговую стоимость.
        prefix = r"(?:(?:авто)?транспортн\w+\s+средств\w*\s*[:：]?\s*)?(?:Автомобил\w*[,:\s]*)?"
        candidates = []
        for m in re.finditer(prefix + r"марк[аи]\s*[:：][^\n]+", normalized, re.IGNORECASE):
            block = re.sub(r'\s+', ' ', m.group(0)).strip().rstrip('.,; ')
            if len(block) >= 12 and re.search(r"модель|vin|год\s+выпуска|кузов|рама", block, re.IGNORECASE):
                candidates.append(block)
        if candidates:
            # Предпочитаем описание, где есть залоговая стоимость.
            for c in candidates:
                if re.search(r"стоимост", c, re.IGNORECASE):
                    return c
            return candidates[0]
        return None

    def detect_ip_collateral(self, text: str) -> bool:
        """
        Определяет, содержит ли заявление указание на залоговое имущество.
        """
        if not text:
            return False

        text_lower = text.lower()
        keywords = [
            "предмет залога",
            "заложенн",
            "залоговое имущество",
            "обращении взыскания на заложенное",
            "ипотек",
            "договор залога",
            "в залог"
        ]

        return any(keyword in text_lower for keyword in keywords)

    def extract_physical_collateral_fields(self, text: str) -> Dict[str, str]:
        """
        Извлекает поля залога для физических лиц (аналогично ИП, но без префикса ИП).
        """
        fields: Dict[str, str] = {}

        if not text:
            return fields

        normalized_text = text.replace('\u202f', ' ').replace('\xa0', ' ')

        def clean_numeric(value: str) -> str:
            cleaned = value.replace('\u202f', ' ').replace('\xa0', ' ')
            cleaned = re.sub(r'\s*(?:руб\.?|рублей|₽)\.?$', '', cleaned, flags=re.IGNORECASE)
            cleaned = cleaned.strip()
            return cleaned

        def set_field(key: str, patterns, postprocess=None, flags=re.IGNORECASE | re.MULTILINE):
            for pattern in patterns:
                match = re.search(pattern, normalized_text, flags)
                if match:
                    value = match.group(1) if match.groups() else match.group(0)
                    value = value.strip()
                    value = self.clean_extracted_value(value)
                    if postprocess:
                        value = postprocess(value)
                    if value:
                        fields[key] = value
                        logger.info(f"Extracted {key}: {value}")
                        break

        def filter_contract_number(value):
            """Фильтрует некорректные значения для номера договора"""
            if not value:
                return None
            value = value.strip()
            if value.lower() in ["от", "none", "null", ""]:
                return None
            value = re.sub(r'^\s*от\s+', '', value, flags=re.IGNORECASE)
            value = re.sub(r'\s+от\s*$', '', value, flags=re.IGNORECASE)
            value = re.sub(r'от\s*\d+\s*$', '', value, flags=re.IGNORECASE)
            value = re.sub(r'от\s*\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*$', '', value, flags=re.IGNORECASE)
            return value.strip() if value.strip() else None

        # Номер договора залога [0005] - для ФЛ (используем те же поля, что и для ИП)
        set_field(
            "ipCollateralContractNumber",  # Используем то же поле, что и для ИП
            [
                r"что\s+подтверждается\s+договором\s+залога\s+№\s*([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)(?=\s+от\s+\d|\s+от\s*$|\s*\[|\s*$|\s*\.)",
                r"что\s+подтверждается\s+договором\s+залога[,\s]+([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)\s+от\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}",
                r"договором\s+залога\s+№\s*([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)(?=\s+от\s+\d|\s+от\s*$|\s*\[|\s*$|\s*\.)",
                r"договором\s+залога[,\s]+([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)\s+от\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}",
                r"В\s+качестве\s+обеспечения[^.]*?договором\s+залога\s+№\s*([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)(?=\s+от\s+\d|\s+от\s*$|\s*\[|\s*$|\s*\.)",
                r"В\s+качестве\s+обеспечения[^.]*?договором\s+залога[,\s]+([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)\s+от\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}",
                r"предоставил\s+в\s+залог[^.]*?договором\s+залога\s+№\s*([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)(?=\s+от\s+\d|\s+от\s*$|\s*\[|\s*$|\s*\.)",
                r"предоставил\s+в\s+залог[^.]*?договором\s+залога[,\s]+([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)\s+от\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}",
                r"договор\s+залога[:\s]*№\s*([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)(?=\s+от\s+\d|\s+от\s*$|\s*$|\s*\.)",
                r"договор\s+залога[^.]*?№\s*([А-ЯЁ0-9/-]+(?:[,\s\-]+[А-ЯЁ0-9/-]+)*?)(?=\s+от\s+\d|\s+от\s*$|\s*$|\s*\.)",
            ],
            postprocess=filter_contract_number
        )

        # Дата договора залога [0006]
        set_field(
            "ipCollateralContractDate",  # Используем то же поле, что и для ИП
            [
                r"что\s+подтверждается\s+договором\s+залога[,\s]+[А-ЯЁ0-9/-]+(?:\s+от|от)\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
                r"договором\s+залога[,\s]+[А-ЯЁ0-9/-]+(?:\s+от|от)\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
                r"договор\s+залога[:\s]*№?\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})",
                r"что\s+подтверждается\s+договором\s+залога[^.]*?от\s+(\d{1,2}[.,]\d{1,2}[.,]\d{4})"
            ]
        )

        # Сумма требований в реестре [0007]
        set_field(
            "ipCollateralClaimAmount",  # Используем то же поле, что и для ИП
            [
                r"Установить\s+требования[^.]*?в\s+размере\s+([0-9\s.,]+)\s*(?:руб|рублей|₽|р\.?)",
                r"требования[^.]*?кредиторов[^.]*?в\s+размере\s+([0-9\s.,]+)\s*(?:руб|рублей|₽|р\.?)",
                r"в\s+реестре\s+требований\s+кредиторов[^.]*?в\s+размере\s+([0-9\s.,]+)\s*(?:руб|рублей|₽|р\.?)",
                r"требования[^.]*?Публичного\s+акционерного\s+общества[^.]*?в\s+размере\s+([0-9\s.,]+)\s*(?:руб|рублей|₽|р\.?)"
            ],
            postprocess=clean_numeric
        )

        # Извлекаем описание предмета залога [1221] для ФЛ с залогом.
        # Только при наличии признаков залога — иначе ложный warning.
        has_collateral_markers = bool(re.search(r"договор\w*\s+залога|ипотек|предмет\s+залога|\[1221\]", text, re.IGNORECASE))
        if has_collateral_markers and "mortgageCollateralDescription1221" not in fields:
            logger.info(f" Ищем mortgageCollateralDescription1221 для ФЛ с залогом...")
            collateral_patterns = [
                r"что\s+подтверждается\s+договором\s+залога\s+№\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*:\s*([\s\S]+?)(?=Наличие\s+заложенного\s+имущества\s+подтверждается\s+выпиской\s+из\s+ЕГРН|Наличие\s+заложенного|\[1221\]|\.\s+[А-ЯЁ]|\n\s*\n|ПРОСИТ|По\s+состоянию|Сумма\s+к|В\s+результате|$)",
                r"подтверждается\s+договором\s+залога\s+№\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*:\s*([\s\S]+?)(?=Наличие\s+заложенного\s+имущества\s+подтверждается\s+выпиской\s+из\s+ЕГРН|Наличие\s+заложенного|\[1221\]|\.\s+[А-ЯЁ]|\n\s*\n|ПРОСИТ|По\s+состоянию|Сумма\s+к|В\s+результате|$)",
                r"В\s+качестве\s+обеспечения[^.]*?что\s+подтверждается\s+договором\s+залога\s+№\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*:\s*([\s\S]+?)(?=Наличие\s+заложенного\s+имущества\s+подтверждается\s+выпиской\s+из\s+ЕГРН|Наличие\s+заложенного|\[1221\]|\.\s+[А-ЯЁ]|\n\s*\n|ПРОСИТ|По\s+состоянию|Сумма\s+к|В\s+результате|$)",
                r"договором\s+залога\s+№\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*:\s*([\s\S]+?)(?=Наличие\s+заложенного|\[1221\]|\.\s+[А-ЯЁ]|\n\s*\n|ПРОСИТ|По\s+состоянию|Сумма\s+к|В\s+результате|$)",
                r"залога\s+№\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*:\s*([^\n]+(?:\n[^\n]+)*?)(?=Наличие\s+заложенного|\[1221\]|\.\s+[А-ЯЁ]|\n\s*\n|ПРОСИТ|По\s+состоянию|Сумма\s+к|В\s+результате|$)",
                r"залога\s+№\s*[А-ЯЁ0-9/-]+(?:\s+от|от)\s+\d{1,2}[.,]\d{1,2}[.,]\d{4}\s*:\s*([^\n\.]+(?:\.[^\n]*)?)(?=Наличие|\[1221\]|\n\s*\n|ПРОСИТ|По\s+состоянию|Сумма\s+к|В\s+результате|$)"
            ]
            set_field("mortgageCollateralDescription1221", collateral_patterns)
            if "mortgageCollateralDescription1221" in fields:
                extracted_value = fields['mortgageCollateralDescription1221']
                logger.info(f"Извлечено mortgageCollateralDescription1221 ({len(extracted_value)} символов): {extracted_value[:150]}...")
            else:
                logger.warning(f"mortgageCollateralDescription1221 НЕ найдено!")

        return fields






    def _get_recommended_acts(self, document_type: str, extracted_fields: Dict[str, Any], text: str) -> Dict[str, Any]:
        """
        Определяет рекомендуемые акты на основе типа документа, типа лица и наличия залога.

        Returns:
            Словарь с рекомендациями: {
                "entityType": "individual" | "legal" | "ip" | "kfh",
                "collateralOption": "collateral" | "collateral_auto" | "no_collateral",
                "recommendedActIds": ["act_id1", "act_id2", ...]
            }
        """
        # Определяем тип лица. Приоритет: КФХ ЮЛ (ООО/Общество) ИП (должник/ответчик ИП ФИО) ФЛ
        entity_type_raw = (extracted_fields.get("entityType") or "").lower()
        is_kfh = extracted_fields.get("isKfh", False)
        text_lower_for_entity = text.lower()[:5000]  # блок должника обычно в начале

        if is_kfh or entity_type_raw == "kfh":
            recommended_entity_type = "kfh"
        elif entity_type_raw in ("legal", "юридическое лицо", "юрлицо"):
            recommended_entity_type = "legal"
        elif entity_type_raw == "ip":
            recommended_entity_type = "ip"
        elif entity_type_raw == "individual":
            # Доверяем извлечённому типу ФЛ — не переопределяем по тексту (в тексте может быть ООО/ПАО кредитора)
            recommended_entity_type = "individual"
        else:
            # Тип не определён при извлечении — проверяем по тексту только блок должника (не кредитора)
            debtor_name = (extracted_fields.get("debtorName") or extracted_fields.get("applicantName") or "").lower()
            legal_in_debtor_name = (
                "ооо" in debtor_name or "оао" in debtor_name or "пао" in debtor_name or "зао" in debtor_name or
                "общество с ограниченной ответственностью" in debtor_name or "ограниченной ответственностью" in debtor_name or
                bool(re.search(r"\b(ооо|оао|пао|зао)\b", debtor_name))
            )
            legal_in_text = (
                "ооо" in text_lower_for_entity or
                "общество с ограниченной ответственностью" in text_lower_for_entity or
                "ограниченной ответственностью" in text_lower_for_entity or
                bool(re.search(r"\bооо\b", text_lower_for_entity)) or
                bool(re.search(r"\b(оао|пао|зао)\b", text_lower_for_entity))
            )
            if legal_in_debtor_name or (legal_in_text and not re.search(r"[а-яё]{2,}\s+[а-яё]{2,}\s+[а-яё]{2,}", debtor_name)):
                recommended_entity_type = "legal"
            elif re.search(r"(?:должник|ответчик)[:\s]*\n?\s*ип\s+[а-яё]", text_lower_for_entity) or re.search(r"индивидуальный предприниматель\s+[а-яё]", text_lower_for_entity):
                recommended_entity_type = "ip"
            else:
                recommended_entity_type = "individual"

        has_collateral = False
        has_auto_collateral = False

        # ОБЯЗАТЕЛЬНАЯ ПРОВЕРКА: документ считается залоговым ТОЛЬКО если есть формулировка "обеспеченное залогом"
        # или похожие формулировки, которые явно указывают на залог (ограничиваем расстояние, чтобы не ловить случайные совпадения)
        required_collateral_phrases = [
            r"обеспеченное\s+залогом",
            r"обеспечено\s+залогом",
            r"обеспечен\s+залогом",
            r"обеспечена\s+залогом",
            r"обеспечены\s+залогом",
            r"В\s+качестве\s+обеспечения\s+исполнения\s+обязательств\s+кредитному\s+договору\s+Заемщик\s+предоставил\s+в\s+залог\s+Банку",
            r"обязательство[^.]{0,120}?обеспечен[^.]{0,30}?залогом",
            r"обязательства[^.]{0,120}?обеспечен[^.]{0,30}?залогом"
        ]

        has_required_collateral_phrase = any(re.search(pattern, text, re.IGNORECASE | re.DOTALL) for pattern in required_collateral_phrases)

        if not has_required_collateral_phrase:
            logger.info("Документ НЕ содержит обязательной формулировки 'обеспеченное залогом' - залог не определяется")
            has_collateral = False
            has_collateral_fields = False
            has_collateral_text = False
        else:
            # Проверяем поля залога в извлеченных данных (не считаем залогом "Кому выдана [ФИО]" — это не описание залога)
            _collateral_desc = (extracted_fields.get("mortgageCollateralDescription1221") or "").strip()
            _valid_collateral_desc = bool(_collateral_desc and not re.match(r"^Кому\s+выдана\s+", _collateral_desc, re.IGNORECASE))
            has_collateral_fields = bool(
                extracted_fields.get("ipCollateralContractNumber") or
                extracted_fields.get("ipCollateralContractDate") or
                _valid_collateral_desc or
                extracted_fields.get("physicalHasCollateral") == "true" or
                extracted_fields.get("observationHasCollateral") == "true" or
                extracted_fields.get("ipHasCollateral") == "true"
            )

            collateral_indicators = [
                r"В\s+качестве\s+обеспечения\s+исполнения\s+обязательств\s+кредитному\s+договору\s+Заемщик\s+предоставил\s+в\s+залог\s+Банку",
                r"что\s+подтверждается\s+договором\s+залога",
                r"договор\s+залога\s+№",
                r"предоставил\s+в\s+залог",
                r"предоставляет\s+в\s+залог"
            ]

            has_collateral_text = any(re.search(pattern, text, re.IGNORECASE | re.DOTALL) for pattern in collateral_indicators)

            has_collateral = has_collateral_fields or has_collateral_text

        auto_indicators = [
            r"автомобил",
            r"транспортное\s+средство",
            r"vin",
            r"марка.*модель",
            r"гос\.?\s*номер"
        ]

        has_auto_fields = bool(
            extracted_fields.get("vin") or
            extracted_fields.get("brandModel")
        )

        if has_collateral:
            has_auto_collateral = has_auto_fields or any(re.search(pattern, text, re.IGNORECASE | re.DOTALL) for pattern in auto_indicators)

        if has_auto_collateral:
            recommended_collateral_option = "collateral_auto"
        elif has_collateral:
            recommended_collateral_option = "collateral"
        else:
            recommended_collateral_option = "no_collateral"

        recommended_act_ids = []
        text_lower = text.lower()

        is_realization = "realization" in document_type or "реализац" in text_lower
        is_restructuring = "restructuring" in document_type or "реструктур" in text_lower
        is_observation = "observation" in document_type or "наблюден" in text_lower
        is_competition = "competition" in document_type or "конкурсн" in text_lower


        is_rtk_inclusion = document_type == "rtk_application" and not self._detect_self_bankruptcy(text)

        if is_rtk_inclusion:
            recommended_act_ids.append("final_rtk_inclusion")
            recommended_act_ids.append("acceptance_definition")

        if is_realization and not is_rtk_inclusion:
            recommended_act_ids.append("final_realization")
            if recommended_collateral_option != "no_collateral":
                recommended_act_ids.append("acceptance_no_motion_no_duty_collateral")
            else:
                recommended_act_ids.append("acceptance_definition")

        if is_restructuring and not is_rtk_inclusion:
            recommended_act_ids.append("final_restructuring")
            if recommended_collateral_option != "no_collateral":
                recommended_act_ids.append("acceptance_no_motion_no_duty_collateral")
            else:
                recommended_act_ids.append("acceptance_definition")

        # Для наблюдения (только не для физлиц — ФЛ не бывает процедура наблюдение/конкурс)
        if is_observation and recommended_entity_type != "individual" and not is_rtk_inclusion:
            recommended_act_ids.append("final_observation")
            if recommended_collateral_option != "no_collateral":
                recommended_act_ids.append("acceptance_no_motion_no_duty_collateral")
            else:
                recommended_act_ids.append("acceptance_definition")

        # Для конкурсного производства (только не для физлиц)
        if is_competition and recommended_entity_type != "individual" and not is_rtk_inclusion:
            recommended_act_ids.append("final_competition")
            if recommended_collateral_option != "no_collateral":
                recommended_act_ids.append("acceptance_no_motion_no_duty_collateral")
            else:
                recommended_act_ids.append("acceptance_definition")

        if recommended_entity_type == "ip" and recommended_collateral_option != "no_collateral" and not is_rtk_inclusion:
            if "realization" in document_type:
                recommended_act_ids.append("final_realization")
            elif "restructuring" in document_type:
                recommended_act_ids.append("final_restructuring")
            recommended_act_ids.append("acceptance_no_motion_no_duty_collateral")

        if recommended_entity_type == "ip" and recommended_collateral_option == "no_collateral" and not is_rtk_inclusion:
            if "realization" in document_type:
                recommended_act_ids.append("final_realization")
            elif "restructuring" in document_type:
                recommended_act_ids.append("final_restructuring")
            recommended_act_ids.append("acceptance_definition")

        # Для КФХ (всегда наблюдение)
        if recommended_entity_type == "kfh" and not is_rtk_inclusion:
            recommended_act_ids.append("final_observation")
            if recommended_collateral_option != "no_collateral":
                recommended_act_ids.append("acceptance_no_motion_no_duty_collateral")
            else:
                recommended_act_ids.append("acceptance_definition")

        # Для ЮЛ без залога (наблюдение по умолчанию)
        if recommended_entity_type == "legal" and recommended_collateral_option == "no_collateral" and not is_realization and not is_restructuring and not is_rtk_inclusion:
            recommended_act_ids.append("final_observation")
            recommended_act_ids.append("acceptance_definition")

        if recommended_entity_type == "ip":
            recommended_act_ids = [act_id for act_id in recommended_act_ids if act_id != "final_competition"]

        if recommended_entity_type == "individual":
            recommended_act_ids = [act_id for act_id in recommended_act_ids if act_id not in ("final_observation", "final_competition")]

        recommended_act_ids = list(dict.fromkeys(recommended_act_ids))

        logger.info(f"Рекомендации: entityType={recommended_entity_type}, collateralOption={recommended_collateral_option}, acts={recommended_act_ids}")

        return {
            "entityType": recommended_entity_type,
            "collateralOption": recommended_collateral_option,
            "recommendedActIds": recommended_act_ids
        }

    def find_date_near_contract(self, text: str, contract_number: str) -> str:
        """
        Ищет дату рядом с номером договора
        """
        def _is_plausible_date(date_str: str) -> bool:
            m = re.match(r'^(\d{1,2})[.,](\d{1,2})[.,](\d{4})$', date_str or "")
            if not m:
                return False
            day = int(m.group(1))
            month = int(m.group(2))
            year = int(m.group(3))
            return 1 <= day <= 31 and 1 <= month <= 12 and 1900 <= year <= 2100

        # 1) Приоритет: конструкции "DD.MM.YYYY ... заключен договор № CONTRACT"
        direct_pattern = rf'(?<!\d)(\d{{1,2}}[.,]\d{{1,2}}[.,]\d{{4}})(?!\d)[\s\S]{{0,220}}?(?:заключ[её]н[ао]?|заключили)[\s\S]{{0,120}}?договор[^\n]{{0,80}}?№\s*{re.escape(contract_number)}'
        direct_match = re.search(direct_pattern, text, re.IGNORECASE)
        if direct_match:
            direct_date = direct_match.group(1)
            if _is_plausible_date(direct_date):
                return direct_date

        # 2) Фолбэк: ищем дату рядом с каждым вхождением номера договора
        for contract_match in re.finditer(re.escape(contract_number), text, re.IGNORECASE):
            start = max(0, contract_match.start() - 300)
            end = min(len(text), contract_match.end() + 300)
            context = text[start:end]
            date_pattern = r'(?<!\d)(\d{1,2}[.,]\d{1,2}[.,]\d{4})(?!\d)'
            for date_match in re.finditer(date_pattern, context):
                candidate = date_match.group(1)
                if _is_plausible_date(candidate):
                    return candidate

        return ''


    def detect_obligation_type(self, text: str, contract_number: str) -> str:
        """
        Определяет тип обязательства по контексту
        """
        context_pattern = rf'.{{0,100}}{re.escape(contract_number)}.{{0,100}}'
        context_match = re.search(context_pattern, text, re.IGNORECASE)

        local_type = 'Договор'
        if context_match:
            context = context_match.group(0).lower()

            if 'карт' in context and ('кредит' in context or 'эмисси' in context):
                local_type = 'Кредитная карта'
            elif 'эмисси' in context:  # эмиссионный контракт = выпуск кредитной карты
                local_type = 'Кредитная карта'
            elif 'кредитн' in context:
                local_type = 'Кредитный договор'
            elif 'займ' in context:
                local_type = 'Договор займа'
            elif 'поручительств' in context:
                local_type = 'Договор поручительства'
            elif 'договор залога' in context or 'договора залога' in context or 'договором залога' in context:
                local_type = 'Договор залога'
            elif 'ссуд' in context:
                local_type = 'Договор ссуды'


        if local_type == 'Договор':
            vm = re.search(
                r'вид\s+обязательств\w*[:\s]*\n?\s*([A-Za-zА-Яа-яЁё]+)',
                text, re.IGNORECASE)
            if vm:
                v = vm.group(1).lower()
                if 'кредит' in v:
                    local_type = 'Кредитный договор'
                elif 'залог' in v:
                    local_type = 'Договор залога'
                elif v.startswith('за'):  # заём / займ
                    local_type = 'Договор займа'
                elif 'поручит' in v:
                    local_type = 'Договор поручительства'

        return local_type

    def calculate_confidence(self, document_type: Optional[str], extracted_fields: Dict[str, str], text: str) -> float:
        """
        Рассчитывает уверенность в результатах анализа
        """
        if not extracted_fields:
            return 0.0

        if document_type == "ip_enforcement_statement" or document_type == "ip_enforcement_statement_collateral":
            required_fields = [
                "applicantName", "inn", "ogrnip", "creditAmount", "principalDebt13"
            ]
        else:
            required_fields = [
                "applicantName", "courtName", "caseNumber", "debtAmount"
            ]

        filled_required = sum(1 for field in required_fields if field in extracted_fields)

        base_confidence = filled_required / len(required_fields)

        additional_factors = 0.0

        if "applicantName" in extracted_fields:
            name = extracted_fields["applicantName"]
            if len(name.split()) >= 3:  # ФИО должно содержать минимум 3 части
                additional_factors += 0.1

        if document_type == "ip_enforcement_statement" or document_type == "ip_enforcement_statement_collateral":
            amount = (
                extracted_fields.get("totalDebt")
                or extracted_fields.get("creditAmount")
                or extracted_fields.get("principalDebt13")
            )
            if amount and re.match(r'^\d+(?:[.,]\d+)?$', amount.replace(' ', '')):
                additional_factors += 0.1

            contract_date = extracted_fields.get("contractDate")
            if contract_date and re.match(r'^\d{1,2}[.,]\d{1,2}[.,]\d{4}$', contract_date):
                additional_factors += 0.1
        else:
            if "debtAmount" in extracted_fields:
                amount = extracted_fields["debtAmount"]
                if re.match(r'^\d+$', amount.replace(' ', '')):
                    additional_factors += 0.1

            if "applicationDate" in extracted_fields:
                date = extracted_fields["applicationDate"]
                if re.match(r'^\d{1,2}[.,]\d{1,2}[.,]\d{4}$', date):
                    additional_factors += 0.1

        total_confidence = min(base_confidence + additional_factors, 1.0)

        return round(total_confidence, 2)

    def get_page_count(self, file_path: str) -> int:
        """
        Получает количество страниц в документе (.docx или .pdf).
        """
        try:
            suffix = Path(file_path).suffix.lower()
            if suffix == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(file_path)
                return max(1, len(reader.pages))
            doc = Document(file_path)
            return max(1, len(doc.paragraphs) // 20)
        except Exception:
            return 1

    def validate_field(self, field_name: str, value: str, field_type: str) -> bool:
        """
        Валидирует извлеченное поле
        """
        if not value or len(value.strip()) < 2:
            return False

        if field_type == "date":
            date_pattern = r'^\d{1,2}[.,]\d{1,2}[.,]\d{4}$'
            return bool(re.match(date_pattern, value))

        elif field_type == "amount":
            amount_pattern = r'^\d+$'
            return bool(re.match(amount_pattern, value.replace(' ', '')))

        elif field_type == "person_name":
            return len(value.split()) >= 2

        return True

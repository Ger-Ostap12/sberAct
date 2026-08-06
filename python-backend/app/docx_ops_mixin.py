# -*- coding: utf-8 -*-
"""Низкоуровневые docx-операции для генератора (вынос без изменения поведения).

Подстановка/замена плейсхолдеров и regex в параграфах, таблицах и колонтитулах,
сбор и удаление маркеров с контекстом, форматирование дат и сумм. Группа
самодостаточна (self-вызовы только внутри группы, атрибутов host нет).
Поведение 1-в-1 под gen-golden.
"""
import logging
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Set

from docx import Document
from docx.oxml.shared import qn

logger = logging.getLogger(__name__)


class _ParagraphView:
    """Абзац по его XML-элементу.

    `doc.paragraphs` не видит абзацы внутри врезок (`w:txbxContent`), а в
    ипотечных актах там лежат шапка суда и подпись секретаря — маркеры оттуда
    тоже надо чистить. Полноценный `docx.text.paragraph.Paragraph` требует
    родительскую часть, которой у нас на руках нет, поэтому обёртка минимальна.
    """

    __slots__ = ("_element",)

    def __init__(self, element):
        self._element = element


class DocxOpsMixin:

    def _remove_placeholder_with_context(self, doc: Document, placeholder: str):
        """
        Удаляет маркер вместе с контекстом вокруг него (например, "Дело№ [1]" удаляется полностью).

        Args:
            doc: Документ для обработки
            placeholder: Маркер для удаления (например, "[1]", "[4]", "[990]")
        """
        # Экранируем маркер для использования в регулярном выражении
        escaped_placeholder = re.escape(placeholder)

        # Определяем паттерны контекста для разных маркеров
        context_patterns = self._get_context_patterns_for_marker(placeholder)

        # Пробуем удалить маркер с контекстом по каждому паттерну
        removed = False
        for pattern in context_patterns:
            # Заменяем {MARKER} на экранированный маркер
            regex_pattern = pattern.replace("{MARKER}", escaped_placeholder)
            # Добавляем опциональные пробелы и знаки препинания после контекста
            # Это позволяет удалить контекст даже если после него есть запятая, точка и т.д.
            regex_pattern_with_punctuation = regex_pattern + r"(?:\s*[,\s\.;:]*)?"
            # Используем регулярное выражение для поиска и удаления
            if self._replace_regex_in_doc(doc, regex_pattern_with_punctuation, ""):
                logger.debug(f"Удален маркер {placeholder} с контекстом по паттерну: {pattern}")
                removed = True
                break

        # Если не удалось удалить с контекстом — удаляем маркер и при необходимости весь абзац
        if not removed:
            self._replace_placeholder_in_doc(doc, placeholder, "")
            # Дополнительно: очищаем абзацы, где после удаления маркера остаётся только подпись/контекст
            self._clear_paragraphs_containing_only_placeholder(doc, placeholder)

    # Граница предложения: точка/восклицательный/вопросительный/точка с запятой,
    # за которыми пробел. Двоеточие границей НЕ считаем — «Дата извещения: [1310]»
    # это одна фраза, и подпись должна уйти вместе с маркером.
    # Границы КЛАУЗЫ (не всего предложения) для удаления пустого маркера: добавлена
    # запятая. Иначе пустой [1233] (площадь) сносил весь абзац предмета залога, а
    # точки внутри дат/сокращений («13.09.2022», «г.») рвали фразу в произвольном месте.
    # Двоеточие НЕ разделитель — метка «адрес:» уходит вместе со своим значением.
    _SENTENCE_END = re.compile(r"[.!?;,]")

    def _remove_placeholder_phrase(self, doc: Document, placeholder: str) -> None:
        """Удаляет предложение с незаполненным маркером целиком, пустой абзац — вместе с абзацем.

        Отличие от `_remove_placeholder_with_context`: та снимает маркер и
        прилегающую пунктуацию, оставляя висеть подпись («Дата извещения:»).
        В ипотечных актах подпись без значения — брак, поэтому убираем фразу.

        Правим текст по run'ам, а не через `p.text = ...`: присваивание схлопывает
        абзац в один run и теряет полужирный/курсив внутри него.

        ТОЛЬКО для ипотеки. Банкротные акты ходят прежним путём — их поведение
        зафиксировано golden-эталонами и меняться не должно.
        """
        # Список материализуем ДО правок: удаление абзаца на ходу сбивает
        # ленивый обход lxml, и часть абзацев проскакивала бы необработанной.
        emptied = []
        for par in list(self._iter_all_paragraphs(doc)):
            touched = False
            # Маркер может быть разорван по run'ам — сверяемся по склейке абзаца.
            while placeholder in self._paragraph_text(par):
                text = self._paragraph_text(par)
                pos = text.index(placeholder)
                start, end = self._sentence_bounds(text, pos, len(placeholder))
                if not self._delete_span_in_paragraph(par, start, end):
                    break  # защита от зацикливания, если удалить не удалось
                touched = True
            # Удаляем ТОЛЬКО те абзацы, которые опустошили сами: пустые строки
            # в шаблоне держат вёрстку, сносить их нельзя.
            if touched and not self._paragraph_text(par).strip():
                emptied.append(par._element)

        for element in emptied:
            parent = element.getparent()
            # Последний абзац в ячейке/врезке удалять нельзя — Word считает
            # такой контейнер повреждённым и отказывается открывать документ.
            if parent is not None and len(parent.findall(qn("w:p"))) > 1:
                parent.remove(element)

    def _remove_placeholder_token(self, doc: Document, placeholder: str) -> None:
        """Мягкое удаление: стираем ТОЛЬКО сам маркер, окружающий текст и абзац
        оставляем нетронутыми.

        Для boilerplate-актов (извещение), где каждый абзац — обязательный
        шаблонный текст, а маркеры вроде [1302] могут не иметь источника данных.
        Удалять предложение/абзац здесь нельзя: теряются шапка суда и половина
        текста. Незаполненный маркер оставляет пустое место — вёрстка и текст
        совпадают с шаблоном. ТОЛЬКО для ипотеки.
        """
        for par in list(self._iter_all_paragraphs(doc)):
            while placeholder in self._paragraph_text(par):
                text = self._paragraph_text(par)
                pos = text.index(placeholder)
                end = pos + len(placeholder)
                # Схлопываем один прилегающий пробел, чтобы «поручением [X] [Y]»
                # не превратилось в тройной пробел при пустом маркере.
                if end < len(text) and text[end] == " " and pos > 0 and text[pos - 1] == " ":
                    end += 1
                if not self._delete_span_in_paragraph(par, pos, end):
                    break  # защита от зацикливания

    def _is_clause_boundary(self, text: str, idx: int) -> bool:
        """Настоящая ли граница клаузы символ `text[idx]`.

        Точка внутри даты или числа границей НЕ является: «15.03.2026» — это одно
        значение, а не три предложения. Без этой проверки зачистка соседнего
        пустого маркера отрезала кусок даты и в акт уезжало «15.03.» вместо
        «15.03.2026» (и «01.2025» вместо «01.01.2025»).

        Точка в сокращении («г.Ростов», «ул.Мира») тоже не граница — после неё
        нет пробела, фраза продолжается.
        """
        char = text[idx]
        if char != ".":
            return True
        prev_char = text[idx - 1] if idx > 0 else ""
        next_char = text[idx + 1] if idx + 1 < len(text) else ""
        if prev_char.isdigit() and next_char.isdigit():
            return False  # 15.03.2026
        if next_char and not next_char.isspace():
            return False  # г.Ростов, ул.Мира
        return True

    def _sentence_bounds(self, text: str, pos: int, length: int) -> tuple:
        """Границы предложения вокруг маркера: [начало, конец) в координатах текста."""
        left = 0
        for match in self._SENTENCE_END.finditer(text, 0, pos):
            if self._is_clause_boundary(text, match.start()):
                left = match.end()
        right = len(text)
        for match in self._SENTENCE_END.finditer(text, pos + length):
            if self._is_clause_boundary(text, match.start()):
                right = match.end()
                break
        # Съедаем пробелы по краям, чтобы не оставить двойной пробел в абзаце.
        while left < pos and text[left].isspace():
            left += 1
        while right < len(text) and text[right].isspace():
            right += 1
        return left, right

    def _paragraph_text(self, par) -> str:
        return "".join(node.text or "" for node in par._element.iter(qn("w:t")))

    def _delete_span_in_paragraph(self, par, start: int, end: int) -> bool:
        """Вырезает срез [start, end) из абзаца, не трогая форматирование остальных run'ов."""
        if end <= start:
            return False
        offset = 0
        changed = False
        for node in par._element.iter(qn("w:t")):
            value = node.text or ""
            node_start, node_end = offset, offset + len(value)
            offset = node_end
            if node_end <= start or node_start >= end:
                continue
            head = value[: max(0, start - node_start)]
            tail = value[max(0, end - node_start):] if end < node_end else ""
            node.text = head + tail
            node.set(qn("xml:space"), "preserve")
            changed = True
        return changed

    def _iter_all_paragraphs(self, doc: Document):
        """Абзацы тела, таблиц, врезок и колонтитулов — один проход без дублей."""
        seen = set()
        roots = [doc.element.body]
        for section in doc.sections:
            for part in (section.header, section.footer,
                         section.first_page_header, section.first_page_footer,
                         section.even_page_header, section.even_page_footer):
                try:
                    roots.append(part._element)
                except Exception:  # у секции может не быть отдельного колонтитула
                    continue
        for root in roots:
            for element in root.iter(qn("w:p")):
                if id(element) in seen:
                    continue
                seen.add(id(element))
                yield _ParagraphView(element)

    def _clear_paragraphs_containing_only_placeholder(self, doc: Document, placeholder: str) -> None:
        """Очищает абзацы, которые состоят только из подписи и пустого маркера (контекст удаляется)."""
        escaped = re.escape(placeholder)
        context_patterns = self._get_context_patterns_for_marker(placeholder)

        def clear_in_paragraphs(paragraphs):
            for p in paragraphs:
                if not p.text or placeholder not in p.text:
                    continue
                text_stripped = p.text.strip()
                # Абзац целиком совпадает с одним из контекстных паттернов (подпись + маркер) — очищаем
                for pattern in context_patterns:
                    regex = pattern.replace("{MARKER}", escaped)
                    full_para = r"^\s*" + regex + r"\s*[.,;:\s]*$"
                    if re.match(full_para, text_stripped, re.IGNORECASE):
                        p.text = ""
                        logger.debug(f"Очищен абзац (подпись+маркер {placeholder})")
                        break
                else:
                    # Иначе: если после удаления маркера остаётся только мусор/пунктуация — тоже очищаем
                    after_remove = re.sub(escaped, "", p.text, count=1).strip()
                    if not after_remove or re.match(r"^[\s.,;:–—\-]+$", after_remove) or len(after_remove) < 3:
                        p.text = ""
                        logger.debug(f"Очищен абзац с пустым маркером {placeholder}")

        clear_in_paragraphs(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    clear_in_paragraphs(cell.paragraphs)
        for section in doc.sections:
            clear_in_paragraphs(section.header.paragraphs)
            clear_in_paragraphs(section.footer.paragraphs)
            for table in section.header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        clear_in_paragraphs(cell.paragraphs)
            for table in section.footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        clear_in_paragraphs(cell.paragraphs)

    def _get_context_patterns_for_marker(self, placeholder: str) -> List[str]:
        """
        Возвращает список паттернов контекста для указанного маркера.

        Args:
            placeholder: Маркер (например, "[1]", "[4]", "[990]")

        Returns:
            Список регулярных выражений для поиска маркера с контекстом

        Про «широкие» паттерны вида «ключевое слово ... {MARKER}»: промежуток
        задан классом [^\[\],.;()]{0,40} — без запятых, точек и скобок и не
        длиннее 40 символов. Раньше там стоял [^\[\]]*, и это ломалось на
        абзацах, где ключевое слово встречается дважды: чистка идёт ПОСЛЕ
        подстановки значений, скобок между сущностями уже нет, и паттерн для
        пустого [4] (ИНН должника) захватывал всё от ИНН кредитора — «ИНН
        7707083893) о банкротстве Иванов И.И. (дата рождения: 01.01.1970,» —
        отчего адрес должника оказывался в скобках рядом с банком.
        """
        # Извлекаем номер маркера
        match = re.match(r'\[([^\]]+)\]', placeholder)
        if not match:
            return [f"{re.escape(placeholder)}"]  # Если не удалось распарсить, возвращаем только маркер

        marker_number = match.group(1)

        # Паттерны контекста для разных маркеров
        context_patterns_map = {
            # [1] - Номер дела (широкий паттерн убирает контекст до маркера)
            "1": [
                r"(?:(?:Дело|дело|номер\s+дела|по\s+делу)[^\[\],.;()]{0,40})?{MARKER}",
                r"(?:Дело|дело)[\s№]*{MARKER}",
                r"(?:Дело|дело)[\s:]*№[\s]*{MARKER}",
                r"номер\s+дела[\s:]*№?[\s]*{MARKER}",
                r"по\s+делу[\s:]*№[\s]*{MARKER}",
                r"№[\s]*{MARKER}(?=\s|$|,|\.|;|:|\n)",
                r"(?:Дело|дело)\s*№\s*{MARKER}",
                r"(?:Дело|дело)\s*{MARKER}",
            ],
            # [2], [2.1], [2.2], [2.3], [2.4] - ФИО должника (обычно без префикса)
            "2": [],
            "2.1": [],
            "2.2": [],
            "2.3": [],
            "2.4": [],
            # [3] - Дата рождения (широкий паттерн: любой текст между ключевой фразой и маркером)
            "3": [
                r"(?:дата\s+рождения|Дата\s+рождения|дата\s+рожд\.|Дата\s+рожд\.)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:дата\s+рождения|Дата\s+рождения)[\s:]*{MARKER}",
                r"(?:дата\s+рожд\.|Дата\s+рожд\.)[\s:]*{MARKER}",
            ],
            # [3.1] - Город/место рождения
            "3.1": [
                r"(?:город\s*/\s*место\s+рождения|место\s+рождения|город\s+рождения)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:город\s*/\s*место\s+рождения|Город\s*/\s*место\s+рождения)[\s:]*{MARKER}",
                r"(?:место\s+рождения|Место\s+рождения)[\s:]*{MARKER}",
                r"(?:город\s+рождения|Город\s+рождения)[\s:]*{MARKER}",
            ],
            # [4] - ИНН
            "4": [
                r"(?:ИНН|инн)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:ИНН|инн)[\s:]*{MARKER}",
                r"(?:ИНН|инн)[\s№]*{MARKER}",
            ],
            # [4.1] - ОГРНИП
            "4.1": [
                r"(?:ОГРНИП|огрнип)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:ОГРНИП|огрнип)[\s:]*{MARKER}",
                r"(?:ОГРНИП|огрнип)[\s№]*{MARKER}",
            ],
            # [5] - СНИЛС
            "5": [
                r"(?:СНИЛС|снилс)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:СНИЛС|снилс)[\s:]*{MARKER}",
                r"(?:СНИЛС|снилс)[\s№]*{MARKER}",
            ],
            # [6] - Адрес регистрации
            "6": [
                r"(?:адрес\s+регистрации|адрес|Адрес)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:адрес|Адрес)[\s:]*{MARKER}",
                r"(?:адрес\s+регистрации|Адрес\s+регистрации)[\s:]*{MARKER}",
            ],
            # [7] - Дата решения суда
            "7": [
                r"(?:дата\s+решения\s+суда|дата\s+решения|Дата\s+решения)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:дата\s+решения|Дата\s+решения)[\s:]*{MARKER}",
                r"(?:дата\s+решения\s+суда|Дата\s+решения\s+суда)[\s:]*{MARKER}",
            ],
            # [8] - ФИО финансового управляющего
            "8": [
                r"(?:финансовый\s+управляющий|ФИО\s+финансового\s+управляющего)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:финансовый\s+управляющий|Финансовый\s+управляющий)[\s:]*{MARKER}",
                r"(?:ФИО\s+финансового\s+управляющего|ФИО\s+Финансового\s+управляющего)[\s:]*{MARKER}",
            ],
            # [9] - Номер сообщения
            "9": [
                r"(?:номер\s+сообщения|сообщение|Номер\s+сообщения|Сообщение)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:номер\s+сообщения|Номер\s+сообщения)[\s:]*№?[\s]*{MARKER}",
                r"(?:сообщение|Сообщение)[\s:]*№?[\s]*{MARKER}",
                # Полный блок "[на сайте ЕФРСБ №[9] от [11]]" — удаляем целиком, если [9] пустой
                # Не завязываемся на наличие [11], чтобы не оставался хвост "от ]"
                r"\[\s*на\s+сайте\s+ЕФРСБ[^\]]*{MARKER}[^\]]*\]",
            ],
            # [11] - Дата публикации ЕФРСБ
            "11": [
                r"(?:дата\s+публикации|Дата\s+публикации)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:дата\s+публикации|Дата\s+публикации)[\s:]*{MARKER}",
                r"(?:дата\s+публикации\s+на\s+сайте|Дата\s+публикации\s+на\s+сайте)[\s:]*{MARKER}",
                # Если даты нет и маркер пустой, убираем хвост вида "от ]"
                r"(?:от\s*\[{MARKER}\]|от\s*])",
            ],
            # [67] - Номер газеты «Коммерсантъ». При пустом номере удаляем весь блок с газетой.
            "67": [
                r"№\s*{MARKER}",
                r"(?:газет[аы]\s+)?«?Коммерсант[ъ\"»']?»?\s*№\s*{MARKER}",
                # Полный блок "в газете «Коммерсантъ» №[67] от [68]" — удаляем целиком, если [67] пустой
                r"в\s+газет[аы]\s+«?Коммерсант[ъ\"»']?»?\s*№\s*{MARKER}\s*от\s*\[68\]",
            ],
            # [68] - Дата газеты «Коммерсантъ»
            "68": [
                r"(?:газет[аы]\s+)?«?Коммерсант[ъ\"»']?»?[^0-9]{0,40}{MARKER}",
                r"(?:от\s+)?{MARKER}",
                # Если по какой-то причине остался блок без "в газете", удаляем "от [68]" после номера
                r"от\s*{MARKER}",
            ],
            # [12] - Общая сумма долга
            "12": [
                r"(?:общая\s+сумма|сумма\s+долга|Общая\s+сумма|Сумма\s+долга)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:общая\s+сумма|Общая\s+сумма)[\s:]*{MARKER}",
                r"(?:сумма\s+долга|Сумма\s+долга)[\s:]*{MARKER}",
            ],
            # [13] - Основной долг
            "13": [
                # Разбивка целого в родительном падеже: «…в размере [12] руб., из них
                # основного долга в размере [13] руб.». Ставим первой — общий паттерн
                # ниже («основной долг») именительный и такую фразу не берёт, отчего в
                # акте самобанкрота (разбивки по кредиторам нет) висело «основного
                # долга в размере  руб.». Хвост «руб.» снимаем здесь же.
                r",?\s*(?:из\s+них\s+|из\s+которых\s*:?\s*)?основного\s+долга\s+в\s+размере\s*{MARKER}\s*руб(?:\.|лей)?",
                r"(?:основной\s+долг|Основной\s+долг)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:основной\s+долг|Основной\s+долг)[\s:]*{MARKER}",
            ],
            # [14] - Проценты
            "14": [
                # См. [13]: «просроченных процентов в размере [14] руб.».
                r",?\s*(?:из\s+них\s+|из\s+которых\s*:?\s*)?(?:просроченных\s+)?процентов\s+в\s+размере\s*{MARKER}\s*руб(?:\.|лей)?",
                r"(?:проценты|Проценты)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:проценты|Проценты)[\s:]*{MARKER}",
            ],
            # [15] - Неустойка
            "15": [
                r"(?:неустойка|штрафные\s+санкции|Неустойка|Штрафные\s+санкции)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:неустойка|Неустойка)[\s:]*{MARKER}",
                r"(?:штрафные\s+санкции|Штрафные\s+санкции)[\s:]*{MARKER}",
            ],
            # [16] - Банкротная госпошлина
            "16": [
                r"(?:госпошлина|государственная\s+пошлина|банкротная\s+госпошлина)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:госпошлина|Госпошлина)[\s:]*{MARKER}",
                r"(?:государственная\s+пошлина|Государственная\s+пошлина)[\s:]*{MARKER}",
                r"(?:банкротная\s+госпошлина|Банкротная\s+госпошлина)[\s:]*{MARKER}",
            ],
            # [17] - Ссудная госпошлина (вторая госпошлина, в шаблонах иногда тоже называется просто "государственная пошлина")
            "17": [
                # Частый шаблон из актов:
                # "[17] – государственная пошлина"
                r"{MARKER}\s*[–—-]\s*государственная\s+пошлина[^\n]*",
                r"(?:ссудная\s+госпошлина|Ссудная\s+госпошлина|госпошлина|Госпошлина|государственная\s+пошлина|Государственная\s+пошлина)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:ссудная\s+госпошлина|Ссудная\s+госпошлина|госпошлина|Госпошлина|государственная\s+пошлина|Государственная\s+пошлина)[\s:]*{MARKER}",
            ],
            # [18] - Установка срока на предоставление возражений
            "18": [
                r"(?:установка\s+срока\s+на\s+предоставление\s+возражений|срок\s+на\s+предоставление\s+возражений)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:установка\s+срока\s+на\s+предоставление\s+возражений|Установка\s+срока\s+на\s+предоставление\s+возражений)[\s:]*{MARKER}",
                r"(?:срок\s+на\s+предоставление\s+возражений|Срок\s+на\s+предоставление\s+возражений)[\s:]*{MARKER}",
            ],
            # [19] - На рассмотрение заявления в срок
            "19": [
                r"(?:на\s+рассмотрение\s+заявления\s+в\s+срок|На\s+рассмотрение\s+заявления\s+в\s+срок)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:на\s+рассмотрение\s+заявления\s+в\s+срок|На\s+рассмотрение\s+заявления\s+в\s+срок)[\s:]*{MARKER}",
            ],
            # [20] - Срок для оставления без движения
            "20": [
                r"(?:срок\s+для\s+оставления\s+без\s+движения|Срок\s+для\s+оставления\s+без\s+движения)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:срок\s+для\s+оставления\s+без\s+движения|Срок\s+для\s+оставления\s+без\s+движения)[\s:]*{MARKER}",
            ],
            # [22] - Номер обособленного спора
            "22": [
                r"(?:номер\s+обособленного\s+спора|Номер\s+обособленного\s+спора)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:номер\s+обособленного\s+спора|Номер\s+обособленного\s+спора)[\s:]*{MARKER}",
            ],
            # [23] - Дата поступления заявления в суд (согласно штампу) — широкий паттерн убирает всю фразу
            "23": [
                r"(?:дата\s+поступления\s+заявления\s+в\s+суд|Дата\s+поступления\s+заявления\s+в\s+суд)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:согласно\s+штампу|Согласно\s+штампу)[^\[\],.;()]{0,40}{MARKER}",
                r"(?:дата\s+поступления\s+заявления\s+в\s+суд|Дата\s+поступления\s+заявления\s+в\s+суд)[\s:]*{MARKER}",
                r"(?:согласно\s+штампу|Согласно\s+штампу)[\s:]*{MARKER}",
            ],
            # [24] - Дата направления в суд
            "24": [
                r"(?:дата\s+направления\s+в\s+суд|Дата\s+направления\s+в\s+суд)[\s:]*{MARKER}",
            ],
            # [99] - Дата и время судебного заседания
            "99": [
                r"(?:дата\s+и\s+время\s+судебного\s+заседания|Дата\s+и\s+время\s+судебного\s+заседания)[\s:]*{MARKER}",
                r"(?:судебное\s+заседание|Судебное\s+заседание)[\s:]*{MARKER}",
            ],
            # [80] - Дата ПП депозит
            "80": [
                r"(?:дата\s+ПП\s+депозит|Дата\s+ПП\s+депозит|дата\s+пп\s+депозит)[\s:]*{MARKER}",
                r"(?:ПП\s+депозит|пп\s+депозит)[\s:]*{MARKER}",
            ],
            # [81] - Дата ПП ГП
            "81": [
                r"(?:дата\s+ПП\s+ГП|Дата\s+ПП\s+ГП|дата\s+пп\s+гп)[\s:]*{MARKER}",
                r"(?:ПП\s+ГП|пп\s+гп)[\s:]*{MARKER}",
            ],
            # [25] - Название/ФИО третьего лица
            "25": [
                r"(?:Название\s*/\s*ФИО\s+третьего\s+лица|название\s*/\s*фИО\s+третьего\s+лица|ФИО\s+третьего\s+лица|фИО\s+третьего\s+лица|третье\s+лицо|Третье\s+лицо)[\s:]*{MARKER}",
                r"(?:поручитель|Поручитель)[\s:]*{MARKER}",
            ],
            # [88] - Дата состояния задолженности
            "88": [
                r"(?:дата\s+состояния|Дата\s+состояния)[\s:]*{MARKER}",
                r"(?:дата\s+состояния\s+задолженности|Дата\s+состояния\s+задолженности)[\s:]*{MARKER}",
            ],
            # [989] - Название кредитора
            "989": [
                r"(?:кредитор|Кредитор)[\s:]*{MARKER}",
                r"(?:название\s+кредитора|Название\s+кредитора)[\s:]*{MARKER}",
            ],
            # [990] - ОГРН кредитора
            "990": [
                r"(?:ОГРН|огрн)[\s:]*{MARKER}",
                r"(?:ОГРН|огрн)[\s№]*{MARKER}",
            ],
            # [991] - ИНН кредитора
            "991": [
                r"(?:ИНН|инн)[\s:]*{MARKER}",
                r"(?:ИНН|инн)[\s№]*{MARKER}",
            ],
            # [1000] - Сумма кредита
            "1000": [
                r"(?:сумма\s+кредита|Сумма\s+кредита)[\s:]*{MARKER}",
            ],
            # [1001] - Срок кредита
            "1001": [
                r"(?:срок\s+кредита|Срок\s+кредита)[\s:]*{MARKER}",
            ],
            # [1002] - Процентная ставка
            "1002": [
                r"(?:процентная\s+ставка|Процентная\s+ставка)[\s:]*{MARKER}",
            ],
            # [1003] - Ставка неустойки
            "1003": [
                r"(?:ставка\s+неустойки|Ставка\s+неустойки)[\s:]*{MARKER}",
            ],
            # [1004] - Дата расчета задолженности
            "1004": [
                r"(?:дата\s+расчета|Дата\s+расчета)[\s:]*{MARKER}",
                r"(?:дата\s+расчета\s+задолженности|Дата\s+расчета\s+задолженности)[\s:]*{MARKER}",
            ],
            # [1221] - Описание предмета залога
            "1221": [
                r"(?:описание\s+предмета\s+залога|Описание\s+предмета\s+залога)[\s:]*{MARKER}",
            ],
            # [35] - Иное
            "35": [
                r"(?:иное|Иное)[\s:]*{MARKER}",
                r"(?:иное|Иное)[\s–—-]*{MARKER}",
            ],
            # [36] - Срочные проценты на основной долг
            "36": [
                r"(?:срочные\s+проценты\s+на\s+основной\s+долг|Срочные\s+проценты\s+на\s+основной\s+долг)[\s:]*{MARKER}",
                r"(?:срочные\s+проценты\s+на\s+основной\s+долг|Срочные\s+проценты\s+на\s+основной\s+долг)[\s–—-]*{MARKER}",
            ],
            # [37] - Срочные проценты на просроченный основной долг
            "37": [
                r"(?:срочные\s+проценты\s+на\s+просроченный\s+основной\s+долг|Срочные\s+проценты\s+на\s+просроченный\s+основной\s+долг)[\s:]*{MARKER}",
                r"(?:срочные\s+проценты\s+на\s+просроченный\s+основной\s+долг|Срочные\s+проценты\s+на\s+просроченный\s+основной\s+долг)[\s–—-]*{MARKER}",
            ],
            # [66] - Дата определения о принятии заявления к производству
            "66": [
                r"(?:дата\s+определения\s+о\s+принятии\s+заявления\s+к\s+производству|Дата\s+определения\s+о\s+принятии\s+заявления\s+к\s+производству)[\s:]*{MARKER}",
                r"(?:дата\s+определения|Дата\s+определения)[\s:]*{MARKER}",
            ],
            # [100]-[104] — даты договоров: удаляем фразу «кредитный договор от [10X] № [11X]» или «договор поручительства от [104] № [114]»
            # Улучшенные паттерны: удаляют всю фразу, включая номер договора, если дата пустая
            "100": [
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}\s+№\s+\[110\][^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}\s+№\s+[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}[^,.\n]*",
            ],
            "101": [
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}\s+№\s+\[111\][^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}\s+№\s+[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}[^,.\n]*",
            ],
            "102": [
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}\s+№\s+\[112\][^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}\s+№\s+[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}[^,.\n]*",
            ],
            "103": [
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}\s+№\s+\[113\][^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}\s+№\s+[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+{MARKER}[^,.\n]*",
            ],
            "104": [
                r"(?:,\s*)?договор\s+поручительства\s+от\s+{MARKER}\s+№\s+\[114\][^,.\n]*",
                r"(?:,\s*)?договор\s+поручительства\s+от\s+{MARKER}\s+№\s+[^,.\n]*",
                r"(?:,\s*)?договор\s+поручительства\s+от\s+{MARKER}[^,.\n]*",
            ],
            # [110]-[114] — номера договоров: удаляем фразу «кредитный договор от [10X] № [11X]» или «договор поручительства от [104] № [114]»
            # Улучшенные паттерны: удаляют всю фразу, включая дату договора, если номер пустой
            "110": [
                r"(?:,\s*)?кредитный\s+договор\s+от\s+\[100\]\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+[^,]+?\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор[^,]*?\s+№\s+{MARKER}[^,.\n]*",
            ],
            "111": [
                r"(?:,\s*)?кредитный\s+договор\s+от\s+\[101\]\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+[^,]+?\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор[^,]*?\s+№\s+{MARKER}[^,.\n]*",
            ],
            "112": [
                r"(?:,\s*)?кредитный\s+договор\s+от\s+\[102\]\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+[^,]+?\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор[^,]*?\s+№\s+{MARKER}[^,.\n]*",
            ],
            "113": [
                r"(?:,\s*)?кредитный\s+договор\s+от\s+\[103\]\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор\s+от\s+[^,]+?\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?кредитный\s+договор[^,]*?\s+№\s+{MARKER}[^,.\n]*",
            ],
            "114": [
                r"(?:,\s*)?договор\s+поручительства\s+от\s+\[104\]\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?договор\s+поручительства\s+от\s+[^,]+?\s+№\s+{MARKER}[^,.\n]*",
                r"(?:,\s*)?договор\s+поручительства[^,]*?\s+№\s+{MARKER}[^,.\n]*",
            ],
        }

        # Получаем паттерны для данного маркера
        patterns = context_patterns_map.get(marker_number, [])

        # Если паттернов нет, возвращаем только маркер
        if not patterns:
            return [f"{re.escape(placeholder)}"]

        return patterns

    def _collect_placeholders(self, doc: Document) -> Set[str]:
        """
        Собирает все плейсхолдеры вида [123], [2.1], [DATE] и т.п. из параграфов
        и таблиц документа, включая колонтитулы.
        Игнорирует произвольные блоки в квадратных скобках вроде
        "[на сайте ЕФРСБ №20899106 от 19.12.2025]", чтобы не удалять уже
        подставленный текст как "неизвестный маркер".
        """
        # Истинный маркер: либо чисто цифровой с точками (2, 2.1, 415 и т.п.),
        # либо специальные вроде [DATE].
        marker_pattern = re.compile(r'\[((?:\d+(?:\.\d+)*|DATE))\]')
        raw_bracket_pattern = re.compile(r'\[[^\]]+\]')
        placeholders: Set[str] = set()

        def collect_from_text(text: str):
            if not text:
                return
            for raw in raw_bracket_pattern.findall(text):
                if marker_pattern.match(raw):
                    placeholders.add(raw)

        # Собираем из основных параграфов
        for paragraph in doc.paragraphs:
            collect_from_text(paragraph.text)

        # Собираем из таблиц основного документа
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    collect_from_text(cell.text)

        # Собираем из колонтитулов
        for section in doc.sections:
            # Заголовки
            for paragraph in section.header.paragraphs:
                collect_from_text(paragraph.text)
            for table in section.header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        collect_from_text(cell.text)

            # Подвалы
            for paragraph in section.footer.paragraphs:
                collect_from_text(paragraph.text)
            for table in section.footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        collect_from_text(cell.text)

        return placeholders

    def _format_date_66(self, date_value: str) -> str:
        """
        Форматирует дату для маркера [66] в вид: «день» месяц год года.
        Пример: 15.03.2025 -> «15» марта 2025 года.
        Принимает дату в форматах DD.MM.YYYY или YYYY-MM-DD.
        """
        if not date_value or not str(date_value).strip():
            return ""
        date_str = str(date_value).strip()
        day, month, year = None, None, None
        if re.match(r"^\d{4}-\d{2}-\d{2}$", date_str):
            try:
                dt = datetime.strptime(date_str, "%Y-%m-%d")
                day, month, year = dt.day, dt.month, dt.year
            except ValueError:
                return ""
        elif re.match(r"^\d{1,2}[.,]\d{1,2}[.,]\d{4}$", date_str):
            normalized = date_str.replace(",", ".")
            parts = normalized.split(".")
            if len(parts) == 3:
                try:
                    day = int(parts[0])
                    month = int(parts[1])
                    year = int(parts[2])
                    if day < 1 or day > 31 or month < 1 or month > 12 or year < 1900 or year > 2100:
                        return ""
                except (ValueError, IndexError):
                    return ""
        if day is None or month is None or year is None:
            return ""
        months_genitive = [
            "", "января", "февраля", "марта", "апреля", "мая", "июня",
            "июля", "августа", "сентября", "октября", "ноября", "декабря"
        ]
        if month < 1 or month > 12:
            return ""
        month_name = months_genitive[month]
        # Форматируем день с ведущим нулем (05 вместо 5)
        day_formatted = f"{day:02d}"
        return f"«{day_formatted}» {month_name} {year} года"

    def _normalize_date_format(self, date_value: str) -> str:
        """
        Нормализует формат даты из YYYY.MM.DD или YYYY-MM-DD в DD.MM.YYYY.
        Пример: 2026.04.02 -> 02.04.2026
        """
        if not date_value or not str(date_value).strip():
            return date_value

        date_str = str(date_value).strip()

        # Проверяем формат YYYY.MM.DD или YYYY-MM-DD
        pattern1 = r'^(\d{4})[.-](\d{1,2})[.-](\d{1,2})$'
        match1 = re.match(pattern1, date_str)
        if match1:
            year = match1.group(1)
            month = match1.group(2).zfill(2)  # Добавляем ведущий ноль
            day = match1.group(3).zfill(2)    # Добавляем ведущий ноль
            return f"{day}.{month}.{year}"

        # Если формат уже DD.MM.YYYY, просто нормализуем (добавляем ведущие нули)
        pattern2 = r'^(\d{1,2})[.,](\d{1,2})[.,](\d{4})$'
        match2 = re.match(pattern2, date_str)
        if match2:
            day = match2.group(1).zfill(2)
            month = match2.group(2).zfill(2)
            year = match2.group(3)
            return f"{day}.{month}.{year}"

        # Datetime-local из формы: YYYY-MM-DDTHH:MM(:SS) -> DD.MM.YYYY HH:MM
        # (иначе «T» из ISO-разделителя протекает в акт: «08.08.2026T20:42»).
        m3 = re.match(r'^(\d{4})[.-](\d{1,2})[.-](\d{1,2})[T ](\d{1,2}):(\d{2})', date_str)
        if m3:
            return (f"{m3.group(3).zfill(2)}.{m3.group(2).zfill(2)}.{m3.group(1)} "
                    f"{m3.group(4).zfill(2)}:{m3.group(5)}")

        # Если формат не распознан, возвращаем как есть
        return date_str

    def _replace_placeholder_in_doc(self, doc: Document, placeholder: str, value: str) -> bool:
        """
        Заменяет указанный плейсхолдер во всём документе, включая колонтитулы.

        Returns:
            bool: True, если были выполнены замены.
        """
        replaced = False

        def replace_in_paragraphs(paragraphs):
            nonlocal replaced
            for paragraph in paragraphs:
                if placeholder in paragraph.text:
                    # Абзац со встроенной фигурой (w:pict/w:drawing — напр. рамка-шапка
                    # извещения): сеттер `paragraph.text` пересобирает run'ы абзаца и
                    # УНИЧТОЖАЕТ фигуру. Такие абзацы пропускаем — их w:t заполнит
                    # общий xml-проход ниже, не трогая структуру фигуры.
                    el = paragraph._p
                    if next(el.iter(qn("w:pict")), None) is not None or next(el.iter(qn("w:drawing")), None) is not None:
                        continue
                    paragraph.text = paragraph.text.replace(placeholder, value)
                    replaced = True

        def replace_in_tables(tables):
            for table in tables:
                for row in table.rows:
                    for cell in row.cells:
                        replace_in_paragraphs(cell.paragraphs)

        def replace_in_xml_text_nodes(elements):
            """
            Дополнительный проход по XML-текстам (w:t), чтобы заменить маркеры
            внутри фигур/текстовых блоков (textbox), которые python-docx не
            всегда отдает через paragraph/table API.
            """
            nonlocal replaced
            for element in elements:
                for text_node in element.iter(qn("w:t")):
                    if text_node.text and placeholder in text_node.text:
                        text_node.text = text_node.text.replace(placeholder, value)
                        replaced = True

        replace_in_paragraphs(doc.paragraphs)
        replace_in_tables(doc.tables)

        for section in doc.sections:
            replace_in_paragraphs(section.header.paragraphs)
            replace_in_tables(section.header.tables)
            replace_in_paragraphs(section.footer.paragraphs)
            replace_in_tables(section.footer.tables)

        # Включаем замену в текстбоксах и других контейнерах, не доступных как paragraphs/tables.
        xml_elements = [doc.part.element]
        for section in doc.sections:
            xml_elements.append(section.header.part.element)
            xml_elements.append(section.footer.part.element)
        replace_in_xml_text_nodes(xml_elements)

        return replaced

    def _replace_regex_in_doc(self, doc: Document, pattern: str, replacement: str) -> bool:
        """
        Заменяет текст по регулярному выражению во всём документе, включая колонтитулы.
        """
        replaced = False
        regex = re.compile(pattern, re.IGNORECASE)

        def replace_in_paragraphs(paragraphs):
            nonlocal replaced
            for paragraph in paragraphs:
                if regex.search(paragraph.text):
                    paragraph.text = regex.sub(replacement, paragraph.text)
                    replaced = True

        def replace_in_tables(tables):
            for table in tables:
                for row in table.rows:
                    for cell in row.cells:
                        replace_in_paragraphs(cell.paragraphs)

        def replace_in_xml_text_nodes(elements):
            nonlocal replaced
            for element in elements:
                for text_node in element.iter(qn("w:t")):
                    if text_node.text and regex.search(text_node.text):
                        text_node.text = regex.sub(replacement, text_node.text)
                        replaced = True

        replace_in_paragraphs(doc.paragraphs)
        replace_in_tables(doc.tables)

        for section in doc.sections:
            replace_in_paragraphs(section.header.paragraphs)
            replace_in_tables(section.header.tables)
            replace_in_paragraphs(section.footer.paragraphs)
            replace_in_tables(section.footer.tables)

        xml_elements = [doc.part.element]
        for section in doc.sections:
            xml_elements.append(section.header.part.element)
            xml_elements.append(section.footer.part.element)
        replace_in_xml_text_nodes(xml_elements)

        return replaced

    def _format_amount_value(self, value: float) -> str:
        """
        Форматирует число как денежную сумму:
        - 1234567.0  -> '1 234 567'
        - 1234567.8  -> '1 234 567,80'
        """
        try:
            num = float(value)
            # Если нет копеек — показываем только целые рубли без ",00"
            if abs(num - round(num)) < 1e-9:
                return f"{num:,.0f}".replace(",", " ")
            # Иначе показываем две копейки
            return f"{num:,.2f}".replace(",", " ").replace(".", ",")
        except Exception:
            return str(value)


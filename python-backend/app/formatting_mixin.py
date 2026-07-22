# -*- coding: utf-8 -*-
"""Форматирование/постобработка готового документа (вынос без изменения поведения).

Постобработка текста (дата прописью, номер дела, ISO-даты), единый шрифт Times
New Roman 11, применение стилей. Группа независима (нет self-вызовов).
Поведение 1-в-1 под gen-golden.
"""
import logging
import re
from typing import Any, Dict

from docx import Document
from docx.shared import Pt

logger = logging.getLogger(__name__)


class FormattingMixin:

    def _postprocess_document_formatting(self, doc: Document, cleaned_data: Dict[str, Any]) -> None:
        """
        Финальный проход по документу:
        - добавляем ведущий ноль в датах вида «5» месяца (делаем «05»);
        - убираем дублирование в номере дела вида А53-2345-4/2025-4/2025;
        - приводим даты формата 2026.04.02 / 2026-04-02 к виду 02.04.2026;
        - удаляем любые оставшиеся маркеры [XXX], если они вдруг не были очищены.
        """

        # Глобальные флаги наличия структурированных данных по ЕФРСБ и Коммерсанту.
        # Используем их совместно с локальным анализом текста каждого абзаца.
        has_efrsb_data = bool(
            str(cleaned_data.get("messageNumber") or "").strip()
            or str(cleaned_data.get("efirsbPublicationDate") or "").strip()
        )
        has_kommersant_data = bool(
            str(cleaned_data.get("kommersantNumber") or "").strip()
            or str(cleaned_data.get("kommersantDate") or "").strip()
        )

        def process_text(text: str) -> str:
            if not text:
                return text

            # Локальные флаги по фактическому тексту абзаца
            text_has_efrsb = bool(re.search(r"ЕФРСБ|Единого\s+федерального\s+реестра\s+сведений\s+о\s+банкротстве", text, re.IGNORECASE))
            text_has_kommersant = bool(re.search(r"Коммерсант", text, re.IGNORECASE))

            has_efrsb = has_efrsb_data or text_has_efrsb
            has_kommersant = has_kommersant_data or text_has_kommersant

            # 1) «5» февраля 2025 года -> «05» февраля 2025 года
            def _fix_quoted_day(match: re.Match) -> str:
                day_str = match.group(1)
                month_name = match.group(2)
                year_str = match.group(3)
                try:
                    day_int = int(day_str)
                except ValueError:
                    return match.group(0)
                day_formatted = f"{day_int:02d}"
                return f"«{day_formatted}» {month_name} {year_str} года"

            months_pattern = r"(января|февраля|марта|апреля|мая|июня|июля|августа|сентября|октября|ноября|декабря)"
            text = re.sub(r"«(\d{1,2})»\s+" + months_pattern + r"\s+(\d{4})\s+года", _fix_quoted_day, text)

            # 2) Дело №А53-2345-4/2025-4/2025 -> Дело №А53-2345-4/2025
            def _fix_case_number(match: re.Match) -> str:
                prefix = match.group(1)   # А53-2345-4
                year1 = match.group(2)    # 2025
                suffix = match.group(3)   # 4
                year2 = match.group(4)    # 2025
                if year1 == year2 and suffix in prefix:
                    return f"Дело №{prefix}/{year1}"
                return match.group(0)

            text = re.sub(
                r"Дело\s*№\s*([А-ЯЁA-Z0-9-]+)/(\d{4})-([А-ЯЁA-Z0-9-]+)/(\d{4})",
                _fix_case_number,
                text
            )
            # 2.1) Дело № А44-1233-4/2025-4 -> Дело № А44-1233-4/2025 (лишний суффикс после года)
            text = re.sub(r"(Дело\s*№\s*[А-ЯЁA-Z0-9-]+/\d{4})-\d+\b", r"\1", text)

            # 3) Даты формата 2026.04.02 / 2026-04-02 -> 02.04.2026
            def _fix_iso_date(match: re.Match) -> str:
                year = match.group(1)
                month = match.group(2).zfill(2)
                day = match.group(3).zfill(2)
                return f"{day}.{month}.{year}"

            text = re.sub(r"(\d{4})[.-](\d{1,2})[.-](\d{1,2})", _fix_iso_date, text)

            # 3.5) Подчищаем хвосты вида "от ]" / "от ] -", которые могли остаться
            # после удаления маркеров ЕФРСБ ([9], [11]) и т.п.
            text = re.sub(r"\s*от\s*]\s*[–—-]?", "", text)

            # 3.6) Фиксим слипание "опубликованы в" -> "опубликованыв" после удаления блока ЕФРСБ
            text = re.sub(r"(опубликованы)\s*в", r"\1 в", text, flags=re.IGNORECASE)

            # 3.7) Обработка блока ЕФРСБ:
            # Варианты поведения:
            # - если в этом абзаце есть и ЕФРСБ, и Коммерсант -> полностью удаляем блок ЕФРСБ: "[на сайте ЕФРСБ №... от ...]"
            # - если есть только ЕФРСБ -> убираем только квадратные скобки, сам текст оставляем
            # - если ЕФРСБ нет -> ничего тут не делаем (блок с [9]/[11] уже удалён раньше по маркерам)
            if has_efrsb:
                if has_kommersant:
                    # Оставляем только Коммерсант, убираем целиком блок "[на сайте ЕФРСБ ...]"
                    text = re.sub(
                        r"\s*\[\s*на\s+сайте\s+ЕФРСБ[^\]]*\]\s*",
                        " ",
                        text,
                        flags=re.IGNORECASE,
                    )
                    # На случай варианта без скобок "на сайте ЕФРСБ ..." после предыдущих проходов:
                    text = re.sub(
                        r"\s*на\s+сайте\s+ЕФРСБ[^,.\n]*[, ]*",
                        " ",
                        text,
                        flags=re.IGNORECASE,
                    )
                else:
                    # ЕФРСБ есть, Коммерсанта нет — оставляем текст без квадратных скобок
                    text = re.sub(
                        r"\s*\[\s*(на\s+сайте\s+ЕФРСБ[^\]]*?)\s*\]\s*",
                        r" \1 ",
                        text,
                        flags=re.IGNORECASE,
                    )

            # 3.8) Хвост "в газете «Коммерсантъ»":
            # - если в этом абзаце есть только ЕФРСБ (Коммерсанта нет) — убираем "в газете «Коммерсантъ» ..."
            # - если есть Коммерсант (с ЕФРСБ или без него) — хвост оставляем, чтобы данные Коммерсанта попали в акт
            if has_efrsb and not has_kommersant:
                text = re.sub(
                    r"[, ]*в\s+газет[аы]\s+«?Коммерсант[\"ъ»']?»?[^.\n]*",
                    "",
                    text,
                    flags=re.IGNORECASE,
                )

            # 4) На всякий случай убираем оставшиеся маркеры вида [123], [2.2] и т.п.
            text = re.sub(r"\[[0-9\.]+\]", "", text)

            # ВАЖНО: не сжимаем последовательности пробелов/табов до одного,
            # чтобы не ломать ручное выравнивание (например, номер дела и дата по краям строки).
            # Если где-то останутся «двойные» пробелы — это лучше, чем съехавшая в одну сторону шапка.
            return text

        # Применяем ко всем параграфам и таблицам, включая колонтитулы
        def process_paragraphs(paragraphs):
            for p in paragraphs:
                if p.text:
                    new_text = process_text(p.text)
                    if new_text != p.text:
                        p.text = new_text

        def process_tables(tables):
            for table in tables:
                for row in table.rows:
                    for cell in row.cells:
                        process_paragraphs(cell.paragraphs)

        process_paragraphs(doc.paragraphs)
        process_tables(doc.tables)

        for section in doc.sections:
            process_paragraphs(section.header.paragraphs)
            process_tables(section.header.tables)
            process_paragraphs(section.footer.paragraphs)
            process_tables(section.footer.tables)

    def set_times_new_roman_11(self, doc: Document):
        """
        Устанавливает шрифт Times New Roman 11 для всего документа
        """
        logger.info(" Устанавливаем шрифт Times New Roman 11 для всего документа")

        # Устанавливаем шрифт для всех стилей
        for style in doc.styles:
            if hasattr(style, 'font'):
                style.font.name = 'Times New Roman'
                style.font.size = Pt(11)
                logger.info(f"Установлен шрифт для стиля: {style.name}")

        # Устанавливаем шрифт для всех параграфов
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)

        # Устанавливаем шрифт для всех таблиц
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(11)

        logger.info(" Шрифт Times New Roman 11 установлен для всего документа")

    def apply_document_styles(self, doc: Document):
        """
        Применяет базовые стили к документу
        """
        # Настройка стилей параграфов
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(11)

        # Настройка отступов
        paragraph_format = style.paragraph_format
        paragraph_format.line_spacing = 1.5
        paragraph_format.space_after = Pt(12)


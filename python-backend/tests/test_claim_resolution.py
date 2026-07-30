# -*- coding: utf-8 -*-
"""Исход по иску (ипотека): маркеры [888] и [889].

Проверяем три ветки радиогруппы «Удовлетворение иска» на настоящем .docx —
подстановка идёт общим маппингом, и мимо документа её не проверить.
"""
import os
import sys

import pytest
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

from claim_resolution import CLAIM_RESOLUTION_TEXTS, build_partial_denial  # noqa: E402
from document_generator import DocumentGenerator  # noqa: E402

CREDITOR = "ПАО Сбербанк"


def _doc(tmp_path):
    """Мини-шаблон с обоими маркерами и с [989] снаружи абзаца [889]."""
    d = Document()
    d.add_paragraph("Исковые требования [989] – [888].")
    d.add_paragraph("[889]")
    path = tmp_path / "act.docx"
    d.save(str(path))
    return Document(str(path))


def _text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def _generate(tmp_path, resolution, creditor=CREDITOR):
    doc = _doc(tmp_path)
    DocumentGenerator().replace_document_data(
        doc,
        {"claimResolution": resolution, "creditorName": creditor, "sourceDocumentType": "mortgage_claim"},
    )
    return _text(doc)


@pytest.mark.parametrize("resolution,expected", sorted(CLAIM_RESOLUTION_TEXTS.items()))
def test_resolution_text_substituted(tmp_path, resolution, expected):
    """[888] получает формулировку, соответствующую выбранной радиокнопке."""
    text = _generate(tmp_path, resolution)
    assert expected in text
    assert "[888]" not in text


def test_partial_fills_denial_paragraph(tmp_path):
    """При «частично» абзац [889] появляется, а истец в нём уже подставлен."""
    text = _generate(tmp_path, "partial")
    assert "В удовлетворении остальной части заявленных исковых требований" in text
    assert CREDITOR in text.split("\n")[1]
    assert "[989]" not in text
    assert "[889]" not in text


@pytest.mark.parametrize("resolution", ["full", "deny"])
def test_non_partial_drops_denial_paragraph(tmp_path, resolution):
    """При «полностью» и «отказать» отказывать в остальной части не в чем."""
    text = _generate(tmp_path, resolution)
    assert "остальной части" not in text
    assert "[889]" not in text


def test_unknown_resolution_leaves_markers_empty(tmp_path):
    """Незнакомое значение радиогруппы не должно молча подставить чужой текст."""
    text = _generate(tmp_path, "whatever")
    assert not any(v in text for v in CLAIM_RESOLUTION_TEXTS.values())


def test_empty_mortgage_markers_are_removed(tmp_path):
    """Незаполненный маркер уносит фразу целиком, а опустевший абзац — исчезает.

    [1310] есть в маппинге, [1443] пока нет — проверяем обе ветки зачистки.
    """
    d = Document()
    d.add_paragraph("Дата извещения: [1310].")
    d.add_paragraph("Адрес филиала истца: [1443].")
    d.add_paragraph("Текст без маркеров.")
    path = tmp_path / "empty.docx"
    d.save(str(path))
    doc = Document(str(path))

    DocumentGenerator().replace_document_data(doc, {"sourceDocumentType": "mortgage_claim"})

    text = _text(doc)
    assert "[1310]" not in text and "[1443]" not in text
    # Подпись не должна висеть без значения.
    assert "Дата извещения" not in text
    assert "Адрес филиала истца" not in text
    assert "Текст без маркеров." in text


def test_partial_denial_keeps_marker_without_creditor():
    """Без наименования истца в абзаце остаётся видимый [989], а не пустое место."""
    assert build_partial_denial(None).count("[989]") == 1
    assert build_partial_denial("   ").count("[989]") == 1
    assert "[989]" not in build_partial_denial(CREDITOR)

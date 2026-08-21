# -*- coding: utf-8 -*-
"""Суммы берутся из таблицы расчёта задолженности, а не из плоского текста.

Заявление приходит в приложение как DOCX после конвертера, и таблица расчёта в
нём — настоящая таблица Word с верными парами «метка → значение». Но текст для
анализа собирается по ячейкам подряд, так что метка и значение становятся просто
соседними строками. Регулярки ловили первую подходящую подпись:

    в т.ч. на просроченные проценты
    7 559,11

читалось как «проценты по кредиту», и 7 559,11 (часть неустойки) уезжала во ВСЕ
акты, включая общую сумму долга.

Подстановка из таблицы включается только когда сходится арифметика
(ссудная + проценты + неустойка + госпошлина = ИТОГО) — это страхует от таблиц,
где конвертер сдвинул ячейки.

Второй круг (03.08): по конвертерному пути файла нет вовсе. PDF конвертируется,
пользователь правит текст в предпросмотре, и в анализ уходит POST /analyze-text —
табличная починка там не вызывалась, и во всех актах снова стояла 7 559,11.
Поэтому те же пары восстанавливаются из ПЛОСКОГО текста (метка, следом значение)
с тем же гейтом на сходимость арифметики.
"""
import os
import sys

import pytest
from docx import Document

_THIS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(_THIS, "..", "app")))

from document_analyzer import DocumentAnalyzer  # noqa: E402

# Таблица из реального заявления (Манукян, ПАО Сбербанк).
REAL_ROWS = [
    ("Госпошлина", "22 568,82"),
    ("Задолженность по неустойке", "407 123,69"),
    ("в т.ч. на просроченные проценты", "7 559,11"),
    ("в т.ч.на просроченную ссудную задолженность", "399 564,58"),
    ("в т.ч. неустойка за неисполнение условий договора", "0,00"),
    ("Проценты за кредит", "0,00"),
    ("в т.ч.срочнные", "0,00"),
    ("в т.ч.просроченные", "0,00"),
    ("Ссудная задолженность", "2 438 262,70"),
    ("в т.ч.срочная", "0,00"),
    ("в т.ч.просроченная", "2 438 262,70"),
    ("ИТОГО задолженность по состоянию на 22.04.2021", "2 867 955,21"),
]


def _make_docx(tmp_path, rows, body_text="Заявление кредитора о признании банкротом."):
    doc = Document()
    doc.add_paragraph(body_text)
    table = doc.add_table(rows=0, cols=2)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    path = tmp_path / "zayavlenie.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture(scope="module")
def analyzer():
    return DocumentAnalyzer()


def _apply(analyzer, tmp_path, rows, fields=None):
    """Прогоняет только разбор таблицы, без полного анализа (он медленный)."""
    path = _make_docx(tmp_path, rows)
    result = {"fields": dict(fields or {})}
    analyzer._apply_table_amounts(path, result)
    return result["fields"]


def test_amounts_taken_from_table(analyzer, tmp_path):
    fields = _apply(analyzer, tmp_path, REAL_ROWS)
    assert fields["loanDebt"] == "2 438 262,70"
    assert fields["interest"] == "0,00"
    assert fields["forfeit"] == "407 123,69"
    # Госпошлина из таблицы расчёта — ССУДНАЯ ([17]): она входит в ИТОГО
    # требований. В stateDuty ([16]) стоит банкротная, платится отдельно за
    # подачу заявления, и одним и тем же рублём эти маркеры быть не могут.
    assert fields["loanStateDuty17"] == "22 568,82"
    assert "stateDuty" not in fields
    assert fields["totalDebt"] == "2 867 955,21"


def test_subtotal_rows_ignored(analyzer, tmp_path):
    """«в т.ч. …» — подпункты, они не должны попадать в суммы."""
    fields = _apply(analyzer, tmp_path, REAL_ROWS)
    assert "7 559,11" not in fields.values(), fields
    assert "399 564,58" not in fields.values(), fields


def test_principal_filled_from_loan_debt(analyzer, tmp_path):
    """[13] читает и principalDebt, и loanDebt — заполняем оба."""
    fields = _apply(analyzer, tmp_path, REAL_ROWS)
    assert fields["principalDebt"] == "2 438 262,70"


def test_total_used_for_requirements(analyzer, tmp_path):
    """Банк просит включить ИТОГО целиком — эта же сумма идёт в требования."""
    fields = _apply(analyzer, tmp_path, REAL_ROWS)
    assert fields["debtAmount"] == "2 867 955,21"
    assert fields["requirementsSum"] == "2 867 955,21"


def test_previous_garbage_overwritten(analyzer, tmp_path):
    """Мусор, добытый регулярками из плоского текста, вытесняется таблицей."""
    fields = _apply(analyzer, tmp_path, REAL_ROWS,
                    {"totalDebt": "7 559,11", "interest": "7 559,11", "debtAmount": "22,04"})
    assert fields["totalDebt"] == "2 867 955,21"
    assert fields["interest"] == "0,00"
    assert fields["debtAmount"] == "2 867 955,21"


def test_broken_table_left_alone(analyzer, tmp_path):
    """Если арифметика не сходится, таблице не верим — поля не трогаем."""
    rows = [
        ("Госпошлина", "22 568,82"),
        ("Ссудная задолженность", "100 000,00"),
        ("Проценты за кредит", "1 000,00"),
        ("Задолженность по неустойке", "500,00"),
        ("ИТОГО задолженность по состоянию на 22.04.2021", "999 999,99"),  # не сходится
    ]
    before = {"totalDebt": "111 111,11", "interest": "222,22"}
    fields = _apply(analyzer, tmp_path, rows, before)
    assert fields == before


def test_consistent_table_applied(analyzer, tmp_path):
    """Сошлась арифметика — суммы применяются."""
    rows = [
        ("Госпошлина", "1 000,00"),
        ("Ссудная задолженность", "100 000,00"),
        ("Проценты за кредит", "20 000,00"),
        ("Задолженность по неустойке", "5 000,00"),
        ("ИТОГО задолженность по состоянию на 01.01.2026", "126 000,00"),
    ]
    fields = _apply(analyzer, tmp_path, rows)
    assert fields["totalDebt"] == "126 000,00"
    assert fields["loanDebt"] == "100 000,00"


def test_document_without_table_untouched(analyzer, tmp_path):
    """Документ без таблицы расчёта — ничего не меняем."""
    rows = [("Наименование", "значение"), ("Ещё строка", "текст")]
    before = {"totalDebt": "1,00"}
    fields = _apply(analyzer, tmp_path, rows, before)
    assert fields == before


# ─────────────── Плоский текст: путь /analyze-text (после конвертера) ───────────────

# Ровно то, что отдаёт extract_text на DOCX от конвертера для этого заявления:
# таблица развёрнута по ячейкам, метка и значение — соседние строки.
REAL_FLAT_TEXT = """\
ПРОШУ:
6. Включить требования ПАО Сбербанк в третью очередь реестра требований
кредиторов должника в размере 2 867 955,21 рублей.
состоянию на 22.04.2021 составляет:
Госпошлина
22 568,82
Задолженность по неустойке
407 123,69
в т.ч. на просроченные проценты
7 559,11
в т.ч.на просроченную ссудную задолженность
399 564,58
в т.ч. неустойка за неисполнение условий договора
0,00
Проценты за кредит
0,00
в т.ч.срочнные
0,00
в т.ч.просроченные
0,00
Ссудная задолженность
2 438 262,70
в т.ч.срочная
0,00
в т.ч.просроченная
2 438 262,70
ИТОГО задолженность по состоянию на 22.04.2021
2 867 955,21
"""

# Тот же документ, но текстовым слоем PDF (pypdf): порядок ячеек перемешан,
# метки стоят рядом с чужими значениями. Верить такой раскладке нельзя.
SCRAMBLED_PDF_TEXT = """\
Госпошлина 22 568,82
в т.ч.срочнные
Задолженность по неустойке
0,00
0,00
2 867 955,21
в т.ч.срочная
2 438 262,70
ИТОГО задолженность по состоянию на 22.04.2021
в т.ч.просроченная
0,00
в т.ч.на просроченную ссудную задолженность
7 559,11
399 564,58
Ссудная задолженность
407 123,69
в т.ч. на просроченные проценты
2 438 262,70
Проценты за кредит 0,00
"""


def _apply_text(analyzer, text, fields=None):
    result = dict(fields or {})
    analyzer._apply_text_debt_table(result, text)
    return result


def test_flat_text_amounts_taken_from_table(analyzer):
    """Конвертерный путь: пары метка/значение восстанавливаются из строк."""
    fields = _apply_text(analyzer, REAL_FLAT_TEXT)
    assert fields["loanDebt"] == "2 438 262,70"
    assert fields["interest"] == "0,00"
    assert fields["forfeit"] == "407 123,69"
    assert fields["loanStateDuty17"] == "22 568,82"
    assert "stateDuty" not in fields
    assert fields["totalDebt"] == "2 867 955,21"
    assert fields["principalDebt"] == "2 438 262,70"
    assert fields["debtAmount"] == "2 867 955,21"


def test_flat_text_overwrites_regex_garbage(analyzer):
    """Ровно тот дефект: 7 559,11 в процентах и в общей сумме долга."""
    before = {"totalDebt": "7 559,11", "interest": "7 559,11",
              "principalDebt": "0,00", "debtAmount": "22,04"}
    fields = _apply_text(analyzer, REAL_FLAT_TEXT, before)
    assert fields["totalDebt"] == "2 867 955,21"
    assert fields["interest"] == "0,00"
    assert fields["principalDebt"] == "2 438 262,70"
    assert "7 559,11" not in fields.values(), fields


def test_flat_text_label_and_value_on_one_line(analyzer):
    """Сжатая вёрстка: «метка значение» одной строкой — тоже пара."""
    text = (
        "Ссудная задолженность 100 000,00\n"
        "Проценты за кредит 20 000,00\n"
        "Задолженность по неустойке 5 000,00\n"
        "Госпошлина 1 000,00\n"
        "ИТОГО задолженность по состоянию на 01.01.2026 126 000,00\n"
    )
    fields = _apply_text(analyzer, text)
    assert fields["loanDebt"] == "100 000,00"
    assert fields["interest"] == "20 000,00"
    assert fields["totalDebt"] == "126 000,00"


def test_scrambled_pdf_text_left_alone(analyzer):
    """Перемешанный текстовый слой PDF: арифметика не сойдётся — не трогаем."""
    before = {"totalDebt": "2 845 386,39", "interest": "2 438 262,70"}
    fields = _apply_text(analyzer, SCRAMBLED_PDF_TEXT, before)
    assert fields == before


def test_plain_document_text_left_alone(analyzer):
    """Обычный текст с суммами, без таблицы расчёта, — поля не трогаем."""
    text = (
        "Задолженность по кредитному договору составляет 500 000,00 руб.\n"
        "Государственная пошлина уплачена платёжным поручением.\n"
        "В размере 25 000,00 рублей внесено на депозит суда.\n"
    )
    before = {"totalDebt": "500 000,00"}
    fields = _apply_text(analyzer, text, before)
    assert fields == before


def test_one_label_does_not_take_two_values(analyzer):
    """Метка потребляет ровно одно следующее значение, а не серию."""
    rows = analyzer._debt_rows_from_text("Ссудная задолженность\n1,00\n2,00\n3,00\n")
    assert rows == [("Ссудная задолженность", "1,00")]


def test_dates_are_not_read_as_amounts(analyzer):
    """«…на 22.04.2021» — дата, а не сумма: пара не создаётся."""
    rows = analyzer._debt_rows_from_text(
        "ИТОГО задолженность по состоянию на 22.04.2021\nв т.ч.просроченная\n"
    )
    assert rows == []


def test_analyze_from_text_uses_debt_table(analyzer):
    """Главный страж: разбор таблицы включён в каскад analyze_from_text.

    Дефект был именно в проводке: сам разбор работал, но по конвертерному пути
    (POST /analyze-text) его никто не вызывал.
    """
    text = (
        "Арбитражный суд Ростовской области\n"
        "Заявитель (кредитор): Публичное акционерное общество «Сбербанк России»\n"
        "ИНН 7707083893, ОГРН 1027700132195\n"
        "Должник: Манукян Сос Жораевич, 30.01.1962 года рождения\n"
        "Заявление о признании гражданина несостоятельным (банкротом).\n"
        "Задолженность по кредитному договору №716994 от 21.03.2008 по\n"
        "состоянию на 22.04.2021 составляет:\n"
        + REAL_FLAT_TEXT.split("составляет:\n", 1)[1]
    )
    fields = analyzer.analyze_from_text(text)["fields"]
    assert fields["totalDebt"] == "2 867 955,21"
    assert fields["interest"] == "0,00"
    assert fields["loanDebt"] == "2 438 262,70"
    assert fields["forfeit"] == "407 123,69"
    assert "7 559,11" not in fields.values(), fields

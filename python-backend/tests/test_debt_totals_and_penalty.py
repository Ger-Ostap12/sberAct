# -*- coding: utf-8 -*-
"""Согласование сумм долга и остаток слова «неустойки» без суммы.

Оба дефекта видны на одном акте (заявление Манукян, ПАО Сбербанк):

  «…в размере 2 438 262,70 руб., из них основного долга 407 123,69 руб.,
   просроченных процентов 2 438 262,70 руб.»   — целое равно одной из частей;
  «…2 438 262,70 руб. процентов, неустойки в третью очередь…»  — слово осталось,
   хотя неустойки в заявлении нет.

Причины: (1) «сумма частей» считалась только по principalDebt, а основной долг
попал в loanDebt («Ссудная задолженность» из таблицы расчёта); (2) шаблон пишет
«[15] неустойки» без «руб.», а зачистка требовала «руб.» обязательно.
"""
import os
import sys

import pytest
from docx import Document

_THIS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(_THIS, "..", "app")))

from document_analyzer import DocumentAnalyzer  # noqa: E402
from document_generator import DocumentGenerator  # noqa: E402


# --- 1. Общая сумма долга = основной долг + проценты + неустойка ---------------

@pytest.fixture
def analyzer():
    return DocumentAnalyzer()


def _reconcile(analyzer, fields):
    analyzer._reconcile_debt_amounts(fields, "", "initiation_physical")
    return fields


def test_total_uses_loan_debt_as_principal(analyzer):
    """Основной долг из loanDebt участвует в сумме наравне с principalDebt."""
    fields = _reconcile(analyzer, {
        "loanDebt": "407 123,69",
        "interest": "2 438 262,70",
        "totalDebt": "2 867 955,21",   # ИТОГО из таблицы, включает госпошлину
    })
    assert fields["totalDebt"] == "2 845 386,39", fields["totalDebt"]


def test_total_not_replaced_by_interest_alone(analyzer):
    """Целое не должно схлопываться до одной из частей."""
    fields = _reconcile(analyzer, {
        "loanDebt": "407 123,69",
        "interest": "2 438 262,70",
        "totalDebt": "2 438 262,70",
    })
    assert fields["totalDebt"] != fields["interest"], fields
    assert fields["totalDebt"] == "2 845 386,39", fields["totalDebt"]


def test_principal_debt_still_preferred(analyzer):
    """При наличии principalDebt берём его, а не loanDebt."""
    fields = _reconcile(analyzer, {
        "principalDebt": "100 000,00",
        "loanDebt": "999 999,99",
        "interest": "50 000,00",
        "totalDebt": "900 000,00",
    })
    assert fields["totalDebt"] == "150 000,00", fields["totalDebt"]


def test_forfeit_included(analyzer):
    fields = _reconcile(analyzer, {
        "loanDebt": "100 000,00",
        "interest": "20 000,00",
        "forfeit": "5 000,00",
        "totalDebt": "500 000,00",
    })
    assert fields["totalDebt"] == "125 000,00", fields["totalDebt"]


def test_matching_total_left_alone(analyzer):
    """Согласованные суммы не трогаем."""
    fields = _reconcile(analyzer, {
        "loanDebt": "100 000,00",
        "interest": "20 000,00",
        "totalDebt": "120 000,00",
    })
    assert fields["totalDebt"] == "120 000,00"


def test_interest_only_keeps_old_behaviour(analyzer):
    """Известны не оба компонента — работает прежнее правило «больше в 1.15 раза»."""
    fields = _reconcile(analyzer, {
        "interest": "100 000,00",
        "totalDebt": "110 000,00",  # меньше порога 1.15 — не трогаем
    })
    assert fields["totalDebt"] == "110 000,00"


# --- 2. Слово «неустойки» не остаётся без суммы -------------------------------

def _apply_penalty(text: str, data=None):
    doc = Document()
    doc.add_paragraph(text)
    DocumentGenerator()._apply_penalty_removal(doc, data or {})
    return doc.paragraphs[0].text


def test_penalty_word_removed_without_rub():
    """Шаблонная форма «[15] неустойки» — без «руб.» между маркером и словом."""
    out = _apply_penalty("Включить требование в размере [12] руб., из них [13] руб. основного "
                         "долга, [14] руб. процентов, [15] неустойки в третью очередь реестра.")
    assert "неустойк" not in out.lower(), out
    assert "[15]" not in out, out
    assert "[14] руб. процентов" in out, out


def test_penalty_with_rub_still_removed():
    out = _apply_penalty("из них [13] руб. основного долга, [14] руб. процентов, "
                         "[15] руб. неустойки следует включить в третью очередь.")
    assert "неустойк" not in out.lower(), out


def test_penalty_kept_when_amount_present():
    """Есть данные о неустойке — ничего не вырезаем."""
    out = _apply_penalty("из них [13] руб. основного долга, [14] руб. процентов, "
                         "[15] неустойки в третью очередь реестра.",
                         {"forfeit15": "1 000,00"})
    assert "[15]" in out and "неустойки" in out, out


def test_standalone_penalty_sentence_untouched():
    """Самостоятельное предложение про неустойку (без [13]/[14] рядом) не трогаем."""
    src = "Требования об установлении [15] руб. неустойки учесть отдельно в реестре требований."
    assert _apply_penalty(src) == src


# --- 3. В шаблонах нет слипшихся «[NN]руб.» -----------------------------------

TEMPLATES = os.path.abspath(os.path.join(_THIS, "..", "..", "Templates"))


@pytest.mark.skipif(not os.path.isdir(TEMPLATES), reason="папка Templates недоступна")
def test_no_glued_rub_in_templates():
    """«[14]руб.» давало в акте «2 438 262,70руб.» — пробел должен быть в шаблоне."""
    import re
    from pathlib import Path

    glued = []
    for path in sorted(Path(TEMPLATES).rglob("*.docx")):
        if path.name.startswith("~$"):
            continue
        for i, p in enumerate(Document(str(path)).paragraphs):
            if re.search(r"\](руб|рублей)", p.text):
                glued.append(f"{path.name} para {i}")
    assert not glued, "слипшиеся маркер и единица измерения: " + "; ".join(glued)

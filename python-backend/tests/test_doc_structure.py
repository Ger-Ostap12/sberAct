# -*- coding: utf-8 -*-
"""Структурный слой: провенанс частей DOCX без сдвига плоского текста (§S.2.B, шаг B1).

Смысл шага — начать хранить МЕСТО каждой части документа (таблица, строка,
колонка, номер абзаца), ничего при этом не меняя в разборе. Поэтому главная
проверка здесь одна и она про НЕИЗМЕННОСТЬ: плоский текст, который уходит в
анализ, обязан остаться прежним до символа на всём корпусе.

Пока этот инвариант держится, структурный слой не может сдвинуть ни golden
разбора, ни golden генерации — по построению, а не по удаче.

Вторая группа проверок — про объединённые ячейки. `row.cells` в python-docx
отдаёт такую ячейку по разу на каждую перекрытую колонку. Структурный слой
схлопывает повторы в `col_span`, а `flat_pairs` разворачивает их обратно:
убрать дубликаты из плоского текста значит сдвинуть эталон, и это отдельная
осознанная правка (§S.3.2), а не побочный эффект провенанса.
"""
import os
import sys

import pytest
from docx import Document

_THIS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(_THIS, "..", "app")))
sys.path.insert(0, os.path.abspath(os.path.join(_THIS, "golden")))

from doc_structure import (  # noqa: E402
    DocPart,
    cells_of,
    docx_parts,
    flat_pairs,
    index_cells,
    paragraph_after_label,
    structural_value,
    value_below_label,
    value_inline_in_cell,
    value_right_of_label,
)
from document_analyzer import DocumentAnalyzer  # noqa: E402

try:
    from _snapshot import CORPUS_DIR, corpus_files
except Exception:  # корпус недоступен — структурные проверки всё равно идут
    CORPUS_DIR, corpus_files = None, lambda: []


@pytest.fixture(scope="module")
def analyzer():
    return DocumentAnalyzer()


# --- Инвариант: плоский текст не изменился -----------------------------------

_FILES = corpus_files()


def _reference_flat_parts(analyzer, file_path):
    """НЕЗАВИСИМАЯ копия прежнего плоского обхода — эталон для сверки.

    Сравнивать новый обход с `extract_text` бесполезно: тот теперь сам построен
    на `doc_structure`, и проверка выродилась бы в «сломанное равно сломанному»
    (ровно так и вышло на первом заходе — сломанный разворот повторов эту
    проверку не уронил). Поэтому старый код воспроизведён здесь дословно и
    правиться вместе с продовым НЕ ДОЛЖЕН: он и есть договор.
    """
    doc = Document(file_path)
    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append((paragraph.text.strip(), "body"))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append((cell.text.strip(), "table"))

    seen_headers = set()
    for section in doc.sections:
        for container in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if container is None:
                continue
            chunk_parts = []
            for paragraph in container.paragraphs:
                if paragraph.text.strip():
                    chunk_parts.append(paragraph.text.strip())
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            chunk_parts.append(cell.text.strip())
            chunk = "\n".join(chunk_parts)
            if chunk and chunk not in seen_headers:
                seen_headers.add(chunk)
                parts.append((chunk, "colophon"))

    for chunk in analyzer._extract_docx_raw_xml_text(file_path):
        parts.append((chunk, "raw"))
    return parts


@pytest.mark.skipif(not _FILES, reason="корпус Заявления/ недоступен")
@pytest.mark.parametrize("rel_path", _FILES or [])
def test_плоский_текст_совпадает_с_прежним_обходом(analyzer, rel_path):
    """Новый обход == прежний, part за part, на каждом документе корпуса.

    Это и есть договор совместимости: пока он выполняется, любой структурный
    запрос — надстройка, а не подмена.
    """
    abs_path = os.path.join(CORPUS_DIR, rel_path)
    assert flat_pairs(analyzer._docx_structure(abs_path)) == _reference_flat_parts(
        analyzer, abs_path)


@pytest.mark.skipif(not _FILES, reason="корпус Заявления/ недоступен")
@pytest.mark.parametrize("rel_path", (_FILES or [])[:12])
def test_extract_text_не_изменился(analyzer, rel_path):
    """Вход анализа — плоский текст; он обязан совпасть с прежним до символа."""
    abs_path = os.path.join(CORPUS_DIR, rel_path)
    expected = "\n".join(t for t, _ in _reference_flat_parts(analyzer, abs_path))
    assert analyzer.extract_text(abs_path) == expected


# --- Координаты ---------------------------------------------------------------

def _make_docx(tmp_path, rows):
    """DOCX с одним абзацем и одной таблицей — минимальный носитель разметки."""
    doc = Document()
    doc.add_paragraph("Должник: Иванов Иван Иванович")
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    path = str(tmp_path / "t.docx")
    doc.save(path)
    return path


def test_ячейки_получают_координаты(tmp_path):
    path = _make_docx(tmp_path, [["ИНН", "7707083893"], ["ОГРН", "1027700132195"]])
    cells = cells_of(docx_parts(path))
    by_coord = {(c.row, c.col): c.text for c in cells}
    assert by_coord == {
        (0, 0): "ИНН", (0, 1): "7707083893",
        (1, 0): "ОГРН", (1, 1): "1027700132195",
    }
    assert all(c.table_id == 0 for c in cells)


def test_абзацы_получают_номер_и_стиль(tmp_path):
    path = _make_docx(tmp_path, [["a", "b"]])
    body = [p for p in docx_parts(path) if p.origin == "body"]
    assert body[0].para_index == 0
    assert body[0].style  # имя стиля есть; конкретное значение зависит от шаблона


def test_индекс_ячеек_даёт_соседа_справа(tmp_path):
    """Ради этого запроса всё и затевалось: «значение = ячейка справа от метки»."""
    path = _make_docx(tmp_path, [["ИНН", "7707083893"]])
    idx = index_cells(docx_parts(path))
    label = next(c for c in idx.values() if c.text == "ИНН")
    right = idx[(label.table_id, label.row, label.col + 1)]
    assert right.text == "7707083893"


# --- Объединённые ячейки -------------------------------------------------------

def test_объединённая_ячейка_не_повторяется_в_структуре(tmp_path):
    """Повтор перекрытой колонки схлопывается в col_span, а не в дубль."""
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(0, 0).text = "Расчёт задолженности"
    for c, value in enumerate(("Основной долг", "Проценты", "Неустойка")):
        table.cell(1, c).text = value
    path = str(tmp_path / "merged.docx")
    doc.save(path)

    cells = cells_of(docx_parts(path))
    header = [c for c in cells if c.text == "Расчёт задолженности"]
    assert len(header) == 1, "объединённая ячейка должна попасть в структуру один раз"
    assert header[0].col_span == 3


def test_плоский_текст_сохраняет_повторы_объединённой_ячейки(tmp_path):
    """Дедуп в плоском тексте сдвинул бы эталон — здесь его быть не должно.

    Проверка идёт в ОБЕ стороны: структура схлопывает, плоский вид разворачивает.
    """
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(0, 0).text = "Шапка"
    path = str(tmp_path / "merged2.docx")
    doc.save(path)

    parts = docx_parts(path)
    assert len([p for p in parts if p.text == "Шапка"]) == 1
    assert len([t for t, _ in flat_pairs(parts) if t == "Шапка"]) == 3


def test_соседние_ячейки_с_одинаковым_текстом_не_считаются_объединёнными(tmp_path):
    """Разделяющий признак — тождество XML-узла, а не совпадение текста.

    В таблицах расчёта задолженности одинаковые значения в соседних колонках —
    норма («0,00» и «0,00»), и схлопывать их нельзя.
    """
    path = _make_docx(tmp_path, [["0,00", "0,00", "0,00"]])
    cells = [c for c in cells_of(docx_parts(path)) if c.text == "0,00"]
    assert len(cells) == 3
    assert all(c.col_span == 1 for c in cells)
    assert [c.col for c in cells] == [0, 1, 2]


# --- Устойчивость --------------------------------------------------------------

# --- Структурные запросы (шаг B2) ---------------------------------------------
#
# Ради них весь слой и затевался: одна структурная формулировка вместо семи
# регулярок на поле. Отрицательные проверки тут важнее положительных — они
# описывают, где запрос ОБЯЗАН промолчать, а не выдать соседнее значение.

def _cell(text, row, col, span=1, table_id=0):
    return DocPart(text, "table", table_id=table_id, row=row, col=col, col_span=span)


def _para(text, index):
    return DocPart(text, "body", para_index=index)


def test_значение_справа_от_метки():
    parts = [_cell("ИНН", 0, 0), _cell("7707083893", 0, 1)]
    assert value_right_of_label(parts, "ИНН") == "7707083893"


def test_пустая_колонка_между_меткой_и_значением_перешагивается():
    """Разделительная колонка в вёрстке — обычное дело, значение за ней."""
    parts = [_cell("ИНН", 0, 0), _cell("7707083893", 0, 2)]
    assert value_right_of_label(parts, "ИНН") == "7707083893"


def test_объединённая_метка_перешагивается_по_своей_ширине():
    """Метка на три колонки: значение в четвёртой, а не во второй.

    Без учёта `col_span` запрос вернул бы пустоту или чужую ячейку.
    """
    parts = [_cell("ИНН", 0, 0, span=3), _cell("7707083893", 0, 3)]
    assert value_right_of_label(parts, "ИНН") == "7707083893"


def test_значение_не_берётся_из_соседней_строки():
    """Сосед — только по горизонтали в той же строке таблицы."""
    parts = [_cell("ИНН", 0, 0), _cell("7707083893", 1, 1)]
    assert value_right_of_label(parts, "ИНН") is None


def test_значение_не_берётся_из_другой_таблицы():
    parts = [_cell("ИНН", 0, 0, table_id=0), _cell("7707083893", 0, 1, table_id=1)]
    assert value_right_of_label(parts, "ИНН") is None


def test_далёкая_ячейка_не_считается_значением():
    """За пределом max_gap начинается уже соседний блок реквизитов."""
    parts = [_cell("ИНН", 0, 0), _cell("7707083893", 0, 9)]
    assert value_right_of_label(parts, "ИНН") is None


def test_значение_под_меткой():
    parts = [_cell("ИНН", 0, 0), _cell("7707083893", 1, 0)]
    assert value_below_label(parts, "ИНН") == "7707083893"


def test_значение_в_той_же_ячейке_после_двоеточия():
    parts = [_cell("ИНН: 7707083893", 0, 0)]
    assert value_inline_in_cell(parts, "ИНН") == "7707083893"


def test_подпись_без_разделителя_не_становится_значением():
    """«ИНН организации» — это ПОДПИСЬ. Без разделителя брать нечего.

    Это главный риск послабления: убери требование двоеточия — и поле
    заполнится словом «организации».
    """
    assert value_inline_in_cell([_cell("ИНН организации", 0, 0)], "ИНН") is None


def test_метка_целым_словом():
    """«ИННОВАЦИИ» не метка «ИНН» — иначе поле наберёт мусора из шапки."""
    parts = [_cell("ИННОВАЦИИ", 0, 0), _cell("что-то", 0, 1)]
    assert value_right_of_label(parts, "ИНН") is None


def test_простыня_в_ячейке_не_считается_меткой():
    """Ячейка длиной с абзац — это текст заявления, а не подпись поля."""
    long_cell = "ИНН " + "и прочий текст мотивировочной части заявления " * 3
    parts = [_cell(long_cell, 0, 0), _cell("7707083893", 0, 1)]
    assert value_right_of_label(parts, "ИНН") is None


def test_следующий_абзац_после_абзаца_метки():
    """Узкая шапка: подпись отдельной строкой, значение — следующей."""
    parts = [_para("Дата рождения", 0), _para("27.11.1961", 1)]
    assert paragraph_after_label(parts, "Дата рождения") == "27.11.1961"


def test_абзацный_запрос_не_смотрит_в_таблицы():
    parts = [_para("ИНН", 0), _cell("7707083893", 0, 1)]
    assert paragraph_after_label(parts, "ИНН") is None


def test_абзац_со_значением_не_считается_подписью():
    """Найдено теневым отчётом, а не придумано.

    Первый заход требовал лишь НАЧАЛА с метки, и абзац «Адрес регистрации:
    413105, …» объявлялся подписью, а значением становился следующий абзац —
    «Дата государственной регистрации: 20.06.1991». Из 69 расхождений отчёта
    почти все были ровно этим.
    """
    parts = [
        _para("Адрес регистрации: 413105, САРАТОВСКАЯ ОБЛАСТЬ, г. ЭНГЕЛЬС", 0),
        _para("Дата государственной регистрации: 20.06.1991", 1),
    ]
    assert paragraph_after_label(parts, "Адрес регистрации") is None
    assert value_inline_in_cell(parts, "Адрес регистрации") == (
        "413105, САРАТОВСКАЯ ОБЛАСТЬ, г. ЭНГЕЛЬС")


def test_подпись_с_двоеточием_и_без_равнозначны():
    """«Адрес» и «Адрес:» отдельной строкой — одна и та же подпись."""
    for label_text in ("Дата рождения", "Дата рождения:", "Дата рождения —"):
        parts = [_para(label_text, 0), _para("27.11.1961", 1)]
        assert paragraph_after_label(parts, "Дата рождения") == "27.11.1961"


def test_ячейка_с_меткой_и_значением_не_отдаёт_соседа():
    """Тот же дефект в таблице: «ИНН: 7707083893» — не подпись для соседа."""
    parts = [_cell("ИНН: 7707083893", 0, 0), _cell("770701001", 0, 1)]
    assert value_right_of_label(parts, "ИНН") is None
    assert value_inline_in_cell(parts, "ИНН") == "7707083893"


def test_значение_после_метки_в_абзаце_берётся_на_месте():
    """Самая частая форма в заявлениях — «метка: значение» одной строкой."""
    parts = [_para("ИНН: 7707083893", 0)]
    assert value_inline_in_cell(parts, "ИНН") == "7707083893"


def test_значение_обрезается_на_подписи_следующего_поля():
    """Одна ячейка часто несёт несколько полей подряд.

    Найдено теневым отчётом: по `claimAmount` на 18 документах значение уезжало
    вместе с чужой подписью — «545 418,01 рублей Государственная пошлина:
    16 136 рублей».
    """
    parts = [_cell("Размер требований: 545 418,01 рублей "
                   "Государственная пошлина: 16 136 рублей", 0, 0)]
    assert value_inline_in_cell(parts, "Размер требований") == "545 418,01 рублей"


def test_длинная_подпись_режет_раньше_короткой():
    """«Адрес» не должен съедать «Адрес регистрации» — реестр отдаёт длинные первыми."""
    parts = [_cell("ИНН: 7707083893 Адрес регистрации: г. Ростов", 0, 0)]
    assert value_inline_in_cell(parts, "ИНН") == "7707083893"


# --- Сведение запросов по реестру меток ---------------------------------------

def test_поле_берётся_по_реестру_меток_а_не_по_строке_в_коде():
    """`address` в реестре имеет шесть написаний; сработать должно любое.

    Метки живут в `label_synonyms`, и это принципиально: как только они
    расползутся по коду запросов, реестр снова начнёт врать (§S.2.D).
    """
    parts = [_cell("Юридический адрес", 0, 0), _cell("344002, г. Ростов-на-Дону", 0, 1)]
    assert structural_value(parts, "address") == ("344002, г. Ростов-на-Дону", "cell_right")


def test_уровень_запроса_возвращается_наружу():
    """Спорные случаи должны быть ВИДНЫ в теневом отчёте, а не проглочены."""
    assert structural_value([_cell("ИНН: 7707083893", 0, 0)], "inn") == (
        "7707083893", "cell_inline")


def test_неизвестное_поле_молчит():
    assert structural_value([_cell("ИНН", 0, 0), _cell("7707083893", 0, 1)],
                            "поля-такого-нет") is None


# --- Устойчивость --------------------------------------------------------------

def test_стиль_недоступен_не_роняет_разбор(tmp_path):
    """Битые DOCX от конвертера — норма входа; стиль это подсказка, не требование."""
    class _Broken:
        text = "Строка"

        @property
        def style(self):
            raise RuntimeError("стиль недоступен")

    class _Doc:
        paragraphs = [_Broken()]
        tables: list = []
        sections: list = []

    parts = docx_parts("несуществующий.docx", doc=_Doc())
    assert parts == [DocPart("Строка", "body", para_index=0, style=None)]

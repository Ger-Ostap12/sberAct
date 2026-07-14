# -*- coding: utf-8 -*-
"""Инвариант посекционного предпросмотра (`extract_sections`).

Ключевая гарантия нулевого риска для анализа/golden: секции — это лишь
ПРЕДСТАВЛЕНИЕ строк `extract_text`. Соединение строк всех секций в порядке их
глобального `index` должно быть БАЙТ-В-БАЙТ равно `extract_text`, а сами индексы
— полной партицией частей (без потерь и дублей).
"""
import glob
import os

import pytest

from document_analyzer import DocumentAnalyzer

_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))


def _corpus(limit=20):
    files = []
    for sub in ("Templates", "Заявления", "python-backend/generated"):
        base = os.path.join(_ROOT, sub)
        if os.path.isdir(base):
            files += [f for f in glob.glob(os.path.join(base, "**", "*.docx"), recursive=True)
                      if "~$" not in f]
    return sorted(files)[:limit]


CORPUS = _corpus()


def _reassemble(sections):
    lines = [ln for sec in sections for ln in sec["lines"]]
    lines.sort(key=lambda ln: ln["index"])
    return "\n".join(ln["text"] for ln in lines), [ln["index"] for ln in lines]


@pytest.mark.skipif(not CORPUS, reason="нет DOCX для проверки инварианта")
@pytest.mark.parametrize("path", CORPUS)
def test_sections_reassemble_to_extract_text(path):
    da = DocumentAnalyzer()
    text = da.extract_text(path)
    sections = da.extract_sections(path)
    n_parts = len(da._docx_text_parts(path))

    reassembled, indices = _reassemble(sections)
    assert reassembled == text, f"секции не пересобираются в extract_text: {path}"
    # Партиция: каждый индекс ровно один раз, покрыт весь диапазон частей.
    assert sorted(indices) == list(range(n_parts)), f"индексы не образуют партицию: {path}"


def _build_docx(tmp_path):
    """Синтетический DOCX: интро, блок должника, таблица, блок кредитора, просьба."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("В Арбитражный суд Ростовской области")
    doc.add_paragraph("Должник: ИП Иванов Иван Иванович")
    doc.add_paragraph("ИНН 6100000000")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Договор №1"
    table.rows[0].cells[1].text = "1 000 000 руб."
    doc.add_paragraph("Кредитор: ПАО Сбербанк")
    doc.add_paragraph("Прошу суд признать должника банкротом")
    out = os.path.join(tmp_path, "synthetic.docx")
    doc.save(out)
    return out


def test_sections_segment_debtor_creditor_tables(tmp_path):
    da = DocumentAnalyzer()
    path = _build_docx(tmp_path)

    text = da.extract_text(path)
    sections = da.extract_sections(path)
    n_parts = len(da._docx_text_parts(path))
    reassembled, indices = _reassemble(sections)

    assert reassembled == text
    assert sorted(indices) == list(range(n_parts))

    by_id = {s["id"]: " ".join(ln["text"] for ln in s["lines"]) for s in sections}
    assert "Иванов" in by_id.get("debtor", "")
    assert "Сбербанк" in by_id.get("creditor", "")
    assert "Договор" in by_id.get("tables", "")
    assert "банкротом" in by_id.get("finances", "")
    # Должник и кредитор — РАЗНЫЕ секции (не слиты).
    assert "Сбербанк" not in by_id.get("debtor", "")

"""Дата и время заседания [99] словами."""

import json
import os
import re
import sys

import pytest
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

from document_generator import DocumentGenerator  # noqa: E402
from hearing_datetime import format_hearing_datetime  # noqa: E402


class TestОсновнойФормат:
    def test_пример_андрея(self):
        assert format_hearing_datetime("14.10.2026 14:58") == (
            "14 октября 2026 года в 14 часов 58 минут")

    def test_вход_из_datetime_local(self):
        assert format_hearing_datetime("2026-10-14T14:58") == (
            "14 октября 2026 года в 14 часов 58 минут")

    def test_ведущий_ноль_в_дне_сохраняется(self):
        assert format_hearing_datetime("02.10.2026 11:00") == (
            "02 октября 2026 года в 11 часов 00 минут")

    def test_ведущий_ноль_в_часах_и_минутах(self):
        assert format_hearing_datetime("02.10.2026 09:05") == (
            "02 октября 2026 года в 09 часов 05 минут")

    def test_день_без_ведущего_нуля_на_входе(self):
        assert format_hearing_datetime("2.10.2026 09:05") == (
            "02 октября 2026 года в 09 часов 05 минут")

    def test_дата_без_времени(self):
        assert format_hearing_datetime("14.10.2026") == "14 октября 2026 года"

    def test_время_через_дефис(self):
        # datetime-local иногда доезжает с дефисом вместо двоеточия.
        assert format_hearing_datetime("2026-10-14T14-58") == (
            "14 октября 2026 года в 14 часов 58 минут")

    def test_секунды_отбрасываются(self):
        assert format_hearing_datetime("14.10.2026 14:58:30") == (
            "14 октября 2026 года в 14 часов 58 минут")


class TestСогласование:
    @pytest.mark.parametrize("время,ждём", [
        ("01:00", "в 01 час 00 минут"),
        ("02:00", "в 02 часа 00 минут"),
        ("03:00", "в 03 часа 00 минут"),
        ("04:00", "в 04 часа 00 минут"),
        ("05:00", "в 05 часов 00 минут"),
        ("11:00", "в 11 часов 00 минут"),
        ("12:00", "в 12 часов 00 минут"),
        ("13:00", "в 13 часов 00 минут"),
        ("14:00", "в 14 часов 00 минут"),
        ("20:00", "в 20 часов 00 минут"),
        ("21:00", "в 21 час 00 минут"),
        ("22:00", "в 22 часа 00 минут"),
        ("23:00", "в 23 часа 00 минут"),
        ("00:00", "в 00 часов 00 минут"),
    ])
    def test_часы(self, время, ждём):
        assert ждём in format_hearing_datetime("14.10.2026 " + время)

    @pytest.mark.parametrize("время,ждём", [
        ("10:01", "01 минуту"),
        ("10:02", "02 минуты"),
        ("10:04", "04 минуты"),
        ("10:05", "05 минут"),
        ("10:11", "11 минут"),
        ("10:12", "12 минут"),
        ("10:14", "14 минут"),
        ("10:21", "21 минуту"),
        ("10:22", "22 минуты"),
        ("10:25", "25 минут"),
        ("10:31", "31 минуту"),
        ("10:59", "59 минут"),
        ("10:00", "00 минут"),
    ])
    def test_минуты(self, время, ждём):
        assert format_hearing_datetime("14.10.2026 " + время).endswith(ждём)


class TestМесяцы:
    @pytest.mark.parametrize("месяц,имя", [
        (1, "января"), (2, "февраля"), (3, "марта"), (4, "апреля"),
        (5, "мая"), (6, "июня"), (7, "июля"), (8, "августа"),
        (9, "сентября"), (10, "октября"), (11, "ноября"), (12, "декабря"),
    ])
    def test_все_месяцы_в_родительном(self, месяц, имя):
        assert format_hearing_datetime("05.%02d.2026 10:00" % месяц) == (
            "05 %s 2026 года в 10 часов 00 минут" % имя)


class TestНеПортитВход:
    @pytest.mark.parametrize("значение", [
        "", None, "   ", "не указано", "скоро",
        "32.10.2026 10:00",   # дня не существует
        "14.13.2026 10:00",   # месяца не существует
        "14.10.1800 10:00",   # год вне диапазона
        "14.10.2026 25:00",   # часа не существует
        "14.10.2026 10:70",   # минуты не существует
        "14.10.26 10:00",     # двузначный год не принимаем
    ])
    def test_неразобранное_возвращается_как_есть(self, значение):
        assert format_hearing_datetime(значение) == значение


БАЗА = {
    "caseNumber": "А53-1234/2026",
    "debtorName": "Иванов Иван Иванович",
    "applicantName": "Иванов Иван Иванович",
    "creditorName": "ПАО Сбербанк",
    "judge": "Ерёмин Ф.Ф.",
    "entityType": "individual",
    "date": "17.07.2026",
    "obligations": [],
    "selectedEntityType": "individual",
    "selectedCollateralOption": "no_collateral",
    # Инициирование физлица без залога — шаблон принятия с маркером [99].
    "selectedApplicationKind": "other",
    "selectedActsIds": "acceptance_definition",
    "selectedActsData": json.dumps([{"id": "acceptance_definition", "selected": True}]),
}


class TestДоАкта:
    def _текст(self, **поля):
        данные = dict(БАЗА)
        данные.update(поля)
        результат = DocumentGenerator().generate("initiation_individual", данные)
        assert результат["success"], результат
        куски = []
        for инфо in результат["documents"].values():
            d = Document(инфо["file_path"])
            куски.append("\n".join(p.text for p in d.paragraphs))
        return re.sub(r"\s+", " ", "\n".join(куски))

    def test_заседание_печатается_словами(self):
        текст = self._текст(courtHearingDateTime99="2026-10-02T09:05")
        assert "на 02 октября 2026 года в 09 часов 05 минут" in текст
        assert "02.10.2026" not in текст
        assert "09:05" not in текст

    def test_пример_андрея_в_акте(self):
        текст = self._текст(courtHearingDateTime99="14.10.2026 14:58")
        assert "на 14 октября 2026 года в 14 часов 58 минут" in текст

    def test_пустое_поле_не_оставляет_маркер(self):
        assert "[99]" not in self._текст()

    def test_срок_конкурсного_остаётся_цифрами(self):
        # Срок конкурсного считается из того же поля [99], но ему нужна дата
        # цифрами: формат меняем в точке подстановки, cleaned_data не трогаем.
        данные = dict(БАЗА)
        данные.update({
            "debtorName": "ООО «Сириус»",
            "applicantName": "ООО «Сириус»",
            "legalShortName": "Сириус",
            "entityType": "legal",
            "selectedEntityType": "legal",
            "debtorStatus": "liquidation",
            "courtHearingDateTime99": "2026-10-02T09:05",
            "selectedActsIds": "acceptance_definition,final_competition",
            "selectedActsData": json.dumps([
                {"id": "acceptance_definition", "selected": True},
                {"id": "final_competition", "selected": True},
            ]),
        })
        результат = DocumentGenerator().generate(
            "initiation_legal_competition_liquidation", данные)
        assert результат["success"], результат
        куски = []
        for инфо in результат["documents"].values():
            d = Document(инфо["file_path"])
            куски.append("\n".join(p.text for p in d.paragraphs))
        текст = re.sub(r"\s+", " ", "\n".join(куски))
        assert "конкурсное производство сроком до 02.10.2026" in текст
        # Предлог перед маркером у этого шаблона свой («по делу»), берём без него.
        assert "02 октября 2026 года в 09 часов 05 минут" in текст

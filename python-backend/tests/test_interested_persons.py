# -*- coding: utf-8 -*-
"""Тесты заинтересованных лиц в генерации: заинтересованное лицо = наследник = третье лицо.

Решение от 16.07: это одна категория, лица равны — в маркеры [25.x]/[52.x]/[54.x]
подставляется любой из них, порядок пула «наследники, затем третьи лица».
Место рождения [53.x] в пул не входит: поля нет ни у наследника, ни у третьего лица,
маркер вычищается из текста вместе с подписью («уроженка [53.1], »).

Покрывают дефекты:
  • heirs/thirdParties не доезжали до генерации — [25.x] оставались пустыми;
  • [52]/[53] без данных оставляли в тексте «( года рождения, уроженка , адрес…)»;
  • лиц больше, чем слотов в шаблоне — лишние терялись молча.
"""
import pytest
from docx import Document

from document_generator import DocumentGenerator


@pytest.fixture(scope="module")
def gen():
    return DocumentGenerator()


HEIR = {"name": "Иванова Мария Ивановна", "address": "г. Ростов-на-Дону, ул. Ленина, д. 1"}
THIRD = {"name": "Сидоров Пётр Петрович", "birthDate": "05.05.1980",
         "address": "г. Москва, ул. Тверская, д. 2"}


def _doc(text: str) -> Document:
    doc = Document()
    doc.add_paragraph(text)
    return doc


# ── Пул лиц ──

def test_pool_order_heirs_then_third_parties(gen):
    """Наследники идут первыми, третьи лица следом — лица равны, но порядок фиксирован."""
    pool = gen._collect_interested_persons({"heirs": [HEIR], "thirdParties": [THIRD]})
    assert [p["name"] for p in pool] == [HEIR["name"], THIRD["name"]]


def test_pool_accepts_either_source_alone(gen):
    """В маркер подставляется любой из них — наследник или третье лицо, безразлично."""
    assert gen._collect_interested_persons({"heirs": [HEIR]})[0]["name"] == HEIR["name"]
    assert gen._collect_interested_persons({"thirdParties": [THIRD]})[0]["name"] == THIRD["name"]


def test_pool_skips_nameless_and_falls_back_to_legacy_flat_fields(gen):
    assert gen._collect_interested_persons({"heirs": [{"name": "  ", "address": "X"}]}) == []
    legacy = gen._collect_interested_persons(
        {"thirdPartyName": "Петров П.П.", "thirdPartyAddress": "г. Сочи"}
    )
    assert legacy[0]["name"] == "Петров П.П." and legacy[0]["address"] == "г. Сочи"


# ── Подстановка в маркеры ──

def test_heir_fills_markers_and_birth_place_is_cleaned(gen):
    """Наследник заполняет [25.1]/[54.1]; [52.1] (даты рождения у наследника нет)
    и [53.1] (место рождения — данных нет нигде) вычищаются вместе с подписью."""
    doc = _doc("заинтересованное лицо: [25.1] ([52.1] года рождения, "
               "уроженка [53.1], адрес регистрации: [54.1])")
    gen.replace_document_data(doc, {"heirs": [HEIR]})

    text = doc.paragraphs[0].text
    assert text == f"заинтересованное лицо: {HEIR['name']} (адрес регистрации: {HEIR['address']})"
    assert "[" not in text and "уроженка" not in text


def test_third_party_birth_date_is_kept(gen):
    """У третьего лица дата рождения есть — [52.1] заполняется, вычищается только [53.1]."""
    doc = _doc("заинтересованное лицо: [25.1] ([52.1] года рождения, "
               "уроженка [53.1], адрес регистрации: [54.1])")
    gen.replace_document_data(doc, {"thirdParties": [THIRD]})

    text = doc.paragraphs[0].text
    assert f"{THIRD['birthDate']} года рождения" in text
    assert "уроженка" not in text and "[" not in text


def test_person_without_details_leaves_no_empty_parentheses(gen):
    doc = _doc("заинтересованное лицо: [25.1] ([52.1] года рождения, "
               "уроженка [53.1], адрес регистрации: [54.1])")
    gen.replace_document_data(doc, {"heirs": [{"name": "Иванова М.И.", "address": ""}]})

    assert doc.paragraphs[0].text == "заинтересованное лицо: Иванова М.И."


# ── Лиц больше, чем слотов в шаблоне ──

def test_persons_over_slots_are_appended_inline_with_details(gen):
    """Слот в шаблоне один, лиц трое — лишние дописываются построчно, без новых маркеров
    (тот же приём, что для обязательств сверх слотов)."""
    doc = _doc("заинтересованное лицо: [25.1] ([52.1] года рождения, "
               "уроженка [53.1], адрес регистрации: [54.1])")
    gen.replace_document_data(doc, {"heirs": [HEIR, THIRD, {"name": "Кузнецов К.К.", "address": "г. Тула"}]})

    text = doc.paragraphs[0].text
    for person in (HEIR, THIRD):
        assert person["name"] in text
    assert "Кузнецов К.К. (адрес регистрации: г. Тула)" in text
    assert f"{THIRD['birthDate']} года рождения" in text


def test_persons_over_slots_appended_as_names_when_slot_is_bare(gen):
    """Слот без скобок с деталями («наследник [25.1] принял») — дописываем только ФИО."""
    doc = _doc("его наследник [25.1] принял(а) наследство")
    gen.replace_document_data(doc, {"heirs": [HEIR, {"name": "Кузнецов К.К.", "address": "г. Тула"}]})

    assert doc.paragraphs[0].text == (
        f"его наследник {HEIR['name']}, Кузнецов К.К. принял(а) наследство"
    )


def test_marker_25_uses_first_person_of_pool(gen):
    doc = _doc("третье лицо: [25]")
    gen.replace_document_data(doc, {"heirs": [HEIR], "thirdParties": [THIRD]})
    assert doc.paragraphs[0].text == f"третье лицо: {HEIR['name']}"

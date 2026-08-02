# -*- coding: utf-8 -*-
"""Вёрстка промежуточных актов: дата и «Дело №», «Судья» и ФИО — по краям строки.

Части строки были разведены ПРОБЕЛАМИ, подогнанными под длину маркеров. После
подстановки значений длина менялась, и части слипались: «15» августа 2026 года
Дело № 2-1142/2010». Теперь между ними один символ табуляции и табстоп по
правому краю текстовой области — правая часть прижата к полю при любой длине
левой.

Тест схемного порядка pPr обязателен: <w:tabs> должен стоять после <w:pStyle>
и до <w:jc>, иначе Word отвергает файл (на это уже наступали 01.08.2026 при
такой же правке ипотечных решений).
"""
import os

import pytest
from docx import Document

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

_THIS = os.path.dirname(os.path.abspath(__file__))
TEMPLATES = os.path.abspath(os.path.join(_THIS, "..", "..", "Templates"))

FILES = [
    "промежуточные_особые/внести отложка/продление упрощенка.docx",
    "промежуточные_особые/внести Назначение после упрощенки/Переход из упрощенки в основное производство.docx",
    "промежуточные_особые/внести отложка/отложение.docx",
    "промежуточные_особые/внести Возврат/возврат ртк ГП.docx",
    # Акты Б/Д — вторая волна тех же жалоб.
    "промежуточные_особые/продление БД.docx",
    "промежуточные_особые/Определение БД иное.docx",
    "промежуточные_особые/принятие после БД.docx",
    "промежуточные_особые/внести Обездвижка/Определение БД нет ГП.docx",
    "промежуточные_особые/внести Обездвижка/Определение БД не гп залог.docx",
    "промежуточные_особые/внести Обездвижка/бд ртк правопреемство -.docx",
]

# Порядок дочерних элементов pPr по схеме OOXML (нужная часть).
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
    "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]


def _path(rel):
    return os.path.join(TEMPLATES, rel.replace("/", os.sep))


pytestmark = pytest.mark.skipif(
    not os.path.isdir(TEMPLATES), reason="папка Templates недоступна"
)


@pytest.mark.parametrize("rel", FILES)
def test_case_line_uses_tab(rel):
    """Строка «дата ⇥ Дело №» разделена табуляцией, а не пробелами."""
    doc = Document(_path(rel))
    matches = [p for p in doc.paragraphs if "Дело №" in p.text and "[66]" in p.text]
    assert matches, "не найден абзац с датой и номером дела"
    for p in matches:
        assert "\t" in p.text, repr(p.text)
        assert "     " not in p.text, f"остались пробелы-разделители: {p.text!r}"


@pytest.mark.parametrize("rel", FILES)
def test_judge_line_uses_tab(rel):
    """Строка «Судья ⇥ ФИО» разделена табуляцией."""
    doc = Document(_path(rel))
    matches = [p for p in doc.paragraphs
               if p.text.strip().startswith("Судья") and "[415]" in p.text]
    assert matches, "не найден абзац подписи судьи"
    for p in matches:
        assert "\t" in p.text, repr(p.text)
        assert "     " not in p.text, f"остались пробелы-разделители: {p.text!r}"


@pytest.mark.parametrize("rel", FILES)
def test_right_tab_stop_present(rel):
    """У обеих строк есть табстоп, выровненный по правому краю."""
    doc = Document(_path(rel))
    section = doc.sections[0]
    expected = int(
        (section.page_width - section.left_margin - section.right_margin) / 914400 * 1440
    )
    checked = 0
    for p in doc.paragraphs:
        if "\t" not in p.text or not ("Дело №" in p.text or "[415]" in p.text):
            continue
        stops = list(p.paragraph_format.tab_stops)
        assert stops, f"нет табстопа: {p.text!r}"
        positions = [ts.position.twips for ts in stops]
        alignments = [str(ts.alignment) for ts in stops]
        assert any("RIGHT" in a for a in alignments), f"{p.text!r}: {alignments}"
        # Правый край с точностью до пары twips (округление EMU→twips).
        assert any(abs(pos - expected) <= 2 for pos in positions), \
            f"{p.text!r}: {positions}, ожидалось ~{expected}"
        checked += 1
    assert checked >= 2, f"проверено строк: {checked}"


def test_no_space_separated_columns_anywhere():
    """Страж по ВСЕМ шаблонам: части строки не разводятся пробелами.

    Жалоба приходила дважды — сперва по четырём актам, потом по всем Б/Д.
    Проверяем разом весь каталог, чтобы третьего раза не было.
    """
    import re
    from pathlib import Path

    glued = []
    for path in sorted(Path(TEMPLATES).rglob("*.docx")):
        if path.name.startswith("~$"):
            continue
        for i, p in enumerate(Document(str(path)).paragraphs):
            text = p.text
            if not re.search(r"\S {5,}\S", text):
                continue
            if ("Дело" in text and "[" in text) or text.strip().startswith("Судья"):
                glued.append(f"{path.name} para {i}: {text.strip()[:40]!r}")
    assert not glued, "части строки разведены пробелами: " + "; ".join(glued)


def test_judge_is_marker_not_name():
    """Подпись судьи — маркер [415], а не вписанное ФИО.

    В «продление БД» и «принятие иниц залог недвига» стояло «Судья К.О. Панова»,
    и в акт всегда попадала она, кого бы ни выбрали в интерфейсе.
    """
    import re
    from pathlib import Path

    hardcoded = []
    name_re = re.compile(
        r"^\s*Судья\b[\s\.]*"
        r"(?:[А-ЯЁ]\.\s*[А-ЯЁ]\.\s*[А-ЯЁ][а-яё]+|[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.)\s*$"
    )
    for path in sorted(Path(TEMPLATES).rglob("*.docx")):
        if path.name.startswith("~$"):
            continue
        for i, p in enumerate(Document(str(path)).paragraphs):
            if name_re.match(p.text) and "[415]" not in p.text:
                hardcoded.append(f"{path.name} para {i}: {p.text.strip()!r}")
    assert not hardcoded, "ФИО судьи вписано в шаблон: " + "; ".join(hardcoded)


@pytest.mark.parametrize("rel", FILES)
def test_ppr_schema_order(rel):
    """Порядок элементов внутри pPr соответствует схеме OOXML."""
    doc = Document(_path(rel))
    for i, p in enumerate(doc.paragraphs):
        pPr = p._p.find(f"{W}pPr")
        if pPr is None:
            continue
        names = [c.tag.replace(W, "") for c in pPr if c.tag.startswith(W)]
        idx = [PPR_ORDER.index(n) for n in names if n in PPR_ORDER]
        assert idx == sorted(idx), f"para {i}: нарушен порядок pPr: {names}"

# -*- coding: utf-8 -*-
"""Две правки на одном реальном акте («Решение Умерший», дело 2-1142/2010).

1. Удаление пустого маркера с контекстом съедало соседние сущности. Паттерн
   «ИНН …{MARKER}» с промежутком [^\\[\\]]* задумывался как «ключевое слово плюс
   немного контекста», но чистка идёт ПОСЛЕ подстановки значений: скобок в тексте
   уже нет, и пустой [4] (ИНН должника) уносил всё от ИНН кредитора —
   «ИНН 7707083893) о банкротстве Иванов И.И. (дата рождения: 01.01.1970,».
   В акте это выглядело так, будто у банка указан адрес должника.

2. Хвост незаполненного слота обязательства убирался только перед знаком
   препинания. В «…договорам от 21.03.2008 № 716994, от № в размере 2 438 262,70
   руб.» после хвоста идёт продолжение фразы, и «от №» оставалось в тексте.
"""
import os
import sys

import pytest
from docx import Document

_THIS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(_THIS, "..", "app")))

from document_generator import DocumentGenerator  # noqa: E402


def _doc(text: str) -> Document:
    d = Document()
    d.add_paragraph(text)
    return d


def _text(d: Document) -> str:
    return d.paragraphs[0].text


# --- 1. Чистка пустого маркера не выходит за пределы своей клаузы --------------

HEADER = ("ознакомившись с заявлением публичного акционерного общества "
          "«Сбербанк России» (ОГРН 1027700132195, ИНН 7707083893) о банкротстве "
          "Манукян Соса Жораевича (дата рождения: 30.01.1962, ИНН [4], СНИЛС [5], "
          "место регистрации: 347810, Ростовская область, ул. Ленина, д.65), "
          "установил следующее.")


@pytest.fixture
def gen():
    return DocumentGenerator()


def test_empty_inn_keeps_creditor_requisites(gen):
    d = _doc(HEADER)
    gen._remove_placeholder_with_context(d, "[4]")
    out = _text(d)
    assert "ИНН 7707083893" in out, out          # ИНН банка на месте
    assert "1027700132195" in out, out


def test_empty_inn_keeps_debtor_name_and_birth(gen):
    d = _doc(HEADER)
    gen._remove_placeholder_with_context(d, "[4]")
    out = _text(d)
    assert "о банкротстве Манукян Соса Жораевича" in out, out
    assert "дата рождения: 30.01.1962" in out, out


def test_empty_inn_marker_itself_removed(gen):
    d = _doc(HEADER)
    gen._remove_placeholder_with_context(d, "[4]")
    assert "[4]" not in _text(d)


def test_address_stays_with_debtor(gen):
    """Адрес принадлежит должнику и должен остаться в его скобках."""
    d = _doc(HEADER)
    for ph in ("[4]", "[5]"):
        gen._remove_placeholder_with_context(d, ph)
    out = _text(d)
    assert "место регистрации: 347810" in out, out
    # Между банком и адресом сохраняется ФИО должника — адрес не «переехал» к банку.
    assert out.index("Манукян") < out.index("место регистрации"), out


def test_own_context_still_removed(gen):
    """Собственный контекст маркера по-прежнему убирается вместе с ним."""
    d = _doc("Должник (ИНН [4]) признан банкротом.")
    gen._remove_placeholder_with_context(d, "[4]")
    out = _text(d)
    assert "[4]" not in out and "ИНН" not in out, out


# --- 2. Хвост незаполненного слота обязательств -------------------------------

@pytest.fixture
def render():
    """Мини-объект с нужными методами миксина рендеринга обязательств."""
    return DocumentGenerator()


def _clean_tails(gen, text: str) -> str:
    """Повторяет зачистку хвостов так, как она идёт в replace_obligations_data."""
    d = _doc(text)
    for _ in range(8):
        gen._replace_regex_in_doc(d, r"(?:,\s*)?от\s+№\s*(?=,|\.|\)|;|$)", "")
        gen._replace_regex_in_doc(d, r"(?:,\s*)?№\s*от\s*(?=,|\.|\)|;|$)", "")
        gen._replace_regex_in_doc(d, r"(?:,\s*)?от\s+№\s*(?=[а-яё])", " ")
        gen._replace_regex_in_doc(d, r"(?:,\s*)?№\s*от\s*(?=[а-яё])", " ")
    return _text(d)


def test_tail_removed_before_continuation(render):
    out = _clean_tails(render, "по кредитным договорам от 21.03.2008 № 716994, от № "
                               "в размере 2 438 262,70 руб.")
    assert "от №" not in out, out
    assert "716994 в размере" in out, out  # слова не склеились


def test_tail_removed_before_punctuation(render):
    out = _clean_tails(render, "по договорам от 21.03.2008 № 716994, от № .")
    assert "от №" not in out, out


def test_real_contract_not_touched(render):
    """Заполненные реквизиты договора зачистка не трогает."""
    src = "по кредитному договору от 21.03.2008 № 716994 в размере 100 руб."
    assert _clean_tails(render, src) == src


def test_mortgage_order_tail_removed(render):
    """Ипотечный порядок маркеров обратный — «№ от» тоже чистим."""
    out = _clean_tails(render, "по договору № от в размере 100 руб.")
    assert "№ от" not in out, out

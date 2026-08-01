# -*- coding: utf-8 -*-
"""Полнота ипотечного маппинга и безопасность зачистки.

Незаведённый маркер в ипотеке — не безобидная дыра: зачистка удаляет его
ВМЕСТЕ с окружающей фразой, поэтому пропущенный в маппинге номер уносит из акта
соседние суммы и даты. Тест стережёт, чтобы каждый маркер шаблонов был известен
генератору.
"""
import os
import re
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

from document_generator import DocumentGenerator  # noqa: E402

TEMPLATES = Path(__file__).resolve().parents[2] / "Templates" / "ипотека"
PART_RE = re.compile(r"word/(document|header\d*|footer\d*)\.xml$")

# Обрабатываются отдельными механизмами, а не через field_mapping.
SPECIAL = {"DATE", "66", "777", "992", "002.1", "100", "110"}


def _markers(path: Path) -> set:
    z = zipfile.ZipFile(path)
    text = ""
    for part in z.namelist():
        if PART_RE.match(part):
            text += " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", z.read(part).decode("utf-8")))
    return set(re.findall(r"\[([0-9][0-9.]*)\]", text))


def _known_markers() -> set:
    gen = DocumentGenerator()
    mapping = gen._apply_debtor_name_field_mapping(
        {"sourceDocumentType": "mortgage_claim", "fields": {}}, gen._base_field_mapping(), True
    )
    known = set(mapping.values()) | SPECIAL
    # Слотовые маркеры описаны в маппинге базой без суффикса.
    for base in list(known):
        for i in range(1, 6):
            known.add(f"{base}.{i}")
    return known


TEMPLATE_FILES = sorted(p for p in TEMPLATES.rglob("*.docx") if not p.name.startswith("~"))


@pytest.mark.parametrize("path", TEMPLATE_FILES, ids=[p.name for p in TEMPLATE_FILES])
def test_every_template_marker_is_known(path):
    missing = sorted(_markers(path) - _known_markers())
    assert not missing, f"маркеры без маппинга (зачистка съест соседний текст): {missing}"


# --- зачистка не должна резать ВНУТРИ дат ------------------------------------
# Баг, который это стережёт: точка внутри «15.03.2026» считалась концом фразы,
# удаление пустого соседнего маркера начиналось с середины даты, и в акт уезжало
# «15.03.» вместо «15.03.2026» (и «01.2025» вместо «01.01.2025»).
TRUNCATED_DATE = re.compile(r"\d{1,2}\.\d{1,2}\.(?!\d)|(?<!\d)\d{1,2}\.\d{4}\b")


def test_property_area_taken_from_field_not_only_description():
    """[1233] берётся из поля area объекта, а не только из текста описания.

    Раньше поле игнорировалось: при заполненной площади маркер оставался пустым
    и уносил из резолютивки «общей площадью …». Разворачивание mortgageProperties
    живёт в generate(), поэтому проверяем на настоящей генерации пакета.
    """
    fields = {
        "sourceDocumentType": "mortgage_claim",
        "selectedMortgageKind": "ddu",
        "courtName": "Октябрьский районный суд г. Ростова-на-Дону",
        "caseNumber": "2-1/2026",
        "creditorName": "ПАО Сбербанк",
        "contractNumber": "1", "contractDate": "01.01.2020",
    }
    data = dict(fields)
    data["fields"] = dict(fields)
    data["mortgageProperties"] = [{
        "description": "Квартира",          # площади в описании НЕТ
        "area": "54,3 кв. м",
        "address": "г. Ростов-на-Дону, ул. Мира, д. 5",
    }]

    result = DocumentGenerator().generate("mortgage", data)
    assert result.get("success"), result.get("error")

    decision = Document(result["documents"]["mortgage_decision"]["file_path"])
    text = "\n".join(p.text for p in decision.paragraphs)
    assert "54,3 кв. м" in text, text
    assert "общей площадью" in text


@pytest.mark.parametrize(
    "text",
    [
        "Дата 15.03.2026 [1229] город вынесения.",
        "за период с 01.01.2025 по 31.12.2025 [1229] (включительно), в размере 100.",
        "адрес г.Ростов-на-Дону [1229], дом 5.",
    ],
)
def test_cleanup_never_truncates_dates(tmp_path, text):
    d = Document()
    d.add_paragraph(text)
    path = tmp_path / "d.docx"
    d.save(str(path))
    doc = Document(str(path))

    DocumentGenerator().replace_document_data(doc, {"sourceDocumentType": "mortgage_claim"})

    got = doc.paragraphs[0].text if doc.paragraphs else ""
    assert "[1229]" not in got
    assert not TRUNCATED_DATE.search(got), f"дата разорвана зачисткой: {got!r}"


# --- падежи суда -------------------------------------------------------------
# [0] именительный, [002.1] родительный, [1302] вышестоящая в родительном,
# [1302.1] вышестоящая в именительном. В шаблонах стояли не те маркеры, и в акт
# уезжало «Принять к производству Октябрьский районный суд» и «направлено в
# Ростовского областного суда».
CASE_EXPECTATIONS = [
    ("mortgage_acceptance_long", "Принять к производству Октябрьского районного суда"),
    ("mortgage_acceptance_long", "по почтовому адресу Октябрьского районного суда"),
    ("mortgage_notice", "на решение Октябрьского районного суда"),
    ("mortgage_notice", "направлено в Ростовский областной суд"),
]


@pytest.fixture(scope="module")
def generated_package():
    fields = {
        "sourceDocumentType": "mortgage_claim",
        "selectedMortgageKind": "civil",
        "courtName": "Октябрьский районный суд г. Ростова-на-Дону",
        "higherCourt": "Ростовский областной суд",
        "caseNumber": "2-1/2026",
        "creditorName": "ПАО Сбербанк",
        "date": "15.03.2026",
        "courtDecisionDate": "15.03.2026",
        "courtSubmissionDate24": "25.03.2026",
        "contractNumber": "1",
        "contractDate": "01.01.2020",
    }
    data = dict(fields)
    data["fields"] = dict(fields)
    data["debtors"] = [{"name": "Фонарева Инна Викторовна"}]
    result = DocumentGenerator().generate("mortgage", data)
    assert result.get("success"), result.get("error")
    return result


@pytest.mark.parametrize("doc_key,expected", CASE_EXPECTATIONS)
def test_court_name_case_in_acts(generated_package, doc_key, expected):
    doc = Document(generated_package["documents"][doc_key]["file_path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert expected in text, f"нет ожидаемого падежа: {expected}"


def test_decision_date_falls_back_to_common_date_field():
    """[7] берётся из поля date, если courtDecisionDate не заполнен.

    Форма пишет «Дату принятия решения» в общее поле date; без бэкфилла [7]
    оставался пустым, а зачистка уносила вместе с ним «г.» перед городом
    («14.08.2026 г. Новороссийск» превращалось в «Новороссийск»).
    """
    fields = {
        "sourceDocumentType": "mortgage_claim",
        "selectedMortgageKind": "civil",
        "courtName": "Приморский районный суд г. Новороссийска",
        "caseNumber": "2-777/2026",
        "creditorName": "ПАО Сбербанк",
        "contractNumber": "1", "contractDate": "01.01.2020",
        "date": "14.08.2026",   # courtDecisionDate НЕ заполнен — как приходит с формы
    }
    data = dict(fields)
    data["fields"] = dict(fields)
    data["debtors"] = [{"name": "Фонарева Инна Викторовна"}]

    result = DocumentGenerator().generate("mortgage", data)
    assert result.get("success"), result.get("error")

    doc = Document(result["documents"]["mortgage_decision"]["file_path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "14.08.2026" in text
    assert "г. Новороссийск" in text, "потерян префикс «г.» перед городом"


@pytest.mark.parametrize("path", TEMPLATE_FILES, ids=[p.name for p in TEMPLATE_FILES])
def test_ppr_children_follow_schema_order(path):
    """Порядок детей <w:pPr> обязан соответствовать схеме OOXML.

    Word отвергает файл, если, например, <w:jc> стоит раньше <w:pStyle>.
    Тест стережёт правки вёрстки, которые делаются скриптами.
    """
    import xml.etree.ElementTree as ET

    order = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
             "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
             "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
             "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
    pos = {n: i for i, n in enumerate(order)}
    z = zipfile.ZipFile(path)
    for part in z.namelist():
        if not PART_RE.match(part):
            continue
        root = ET.fromstring(z.read(part))
        for ppr in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr"):
            seen = [pos[t] for t in (c.tag.split("}")[-1] for c in ppr) if t in pos]
            assert seen == sorted(seen), f"{path.name}/{part}: порядок pPr нарушен"

# -*- coding: utf-8 -*-
"""Ширина врезки-шапки в извещении.

Врезка позиционирована от края страницы (margin-left:120.35pt) и при ширине
270pt доходила до 390pt, поджимая правую колонку с адресами получателей
(«Копия: ПАО Сбербанк …») — та обрезалась. Ширину уменьшили до 220pt.

Высоту не фиксируем: у формы стоит mso-fit-shape-to-text, она подстраивается,
если адрес суда переносится на вторую строку.
"""
import os
import re
import zipfile

import pytest

_THIS = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.abspath(os.path.join(
    _THIS, "..", "..", "Templates", "ипотека", "Извещение_МАРКИРОВАННЫЙ.docx"))

# Врезка начинается на 120.35pt от края страницы; правое поле — на 552pt
# (страница 595pt минус 42.9pt поля). Запас под правую колонку должен остаться
# не меньше 200pt, иначе адреса получателей снова начнут обрезаться.
BOX_LEFT_PT = 120.35
RIGHT_MARGIN_PT = 552.1
MIN_RIGHT_COLUMN_PT = 200.0

pytestmark = pytest.mark.skipif(
    not os.path.exists(TEMPLATE), reason="шаблон извещения недоступен")


def _box_style() -> str:
    xml = zipfile.ZipFile(TEMPLATE).read("word/document.xml").decode("utf-8", "replace")
    m = re.search(r'style="([^"]*width:[^"]*)"', xml)
    assert m, "во врезке извещения не найден стиль с шириной"
    return m.group(1)


def test_box_width_leaves_room_for_recipients():
    style = _box_style()
    m = re.search(r"width:([\d.]+)pt", style)
    assert m, style
    width = float(m.group(1))
    free = RIGHT_MARGIN_PT - (BOX_LEFT_PT + width)
    assert free >= MIN_RIGHT_COLUMN_PT, (
        f"врезка шириной {width}pt оставляет правой колонке {free:.1f}pt, "
        f"нужно не меньше {MIN_RIGHT_COLUMN_PT}pt")


def test_box_autofits_height():
    """Ширина уменьшена — текст может перенестись, высота должна подстраиваться."""
    assert "mso-fit-shape-to-text:t" in _box_style()


def test_docx_still_valid():
    """Файл перепаковывался вручную — проверяем, что архив цел."""
    z = zipfile.ZipFile(TEMPLATE)
    assert z.testzip() is None
    assert "word/document.xml" in z.namelist()


def test_recipients_column_starts_right_of_the_box():
    """Адреса получателей не должны заходить левее врезки.

    Без левого отступа их длинные строки переносились под рамку и левее неё:
    «д.29» оказывалось слева от рамки, адрес ФГКУ — во всю ширину под ней.
    Шапка получателей идёт до первого абзаца письма (выключка по ширине).
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu

    doc = Document(TEMPLATE)
    left_margin_pt = Emu(doc.sections[0].left_margin).pt
    box_right_from_margin = 340.35 - left_margin_pt  # правый край врезки

    checked = 0
    for i, p in enumerate(doc.paragraphs):
        if p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
            break  # начался текст письма
        if not p.text.strip():
            continue
        indent = p.paragraph_format.left_indent
        assert indent is not None, f"para {i} без отступа: {p.text.strip()[:40]!r}"
        assert Emu(indent).pt >= box_right_from_margin, (
            f"para {i}: отступ {Emu(indent).pt:.0f}pt меньше правого края врезки "
            f"{box_right_from_margin:.0f}pt — текст заедет под рамку")
        checked += 1
    assert checked >= 5, f"проверено абзацев шапки: {checked}"


def test_letter_body_keeps_full_width():
    """Сам текст письма и подпись остаются на всю ширину."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu

    doc = Document(TEMPLATE)
    body = [p for p in doc.paragraphs if p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY]
    assert body, "не найден текст письма"
    for p in body:
        indent = p.paragraph_format.left_indent
        assert indent is None or Emu(indent).pt < 50, (
            f"текст письма получил отступ шапки: {p.text.strip()[:40]!r}")

# -*- coding: utf-8 -*-
"""Резолютивка «Включить требование …» — управление у суммы и пустая разбивка.

1. **«в размере».** В шаблонах ВКЛ в РТК фраза была написана как
   «…адрес регистрации: [988]) [12] руб., из которых: …», и в готовом акте сумма
   повисала без управления: «…д. 19) 180 532,78 руб., из которых». В теле того же
   акта («образовалась просроченная задолженность в размере [12] руб.») оборот
   правильный — расходились именно шаблоны. Страж следит, чтобы «в размере» не
   пропало снова при очередном пересохранении шаблона в Р7-Офисе.

2. **Пустая разбивка.** В заявлении самобанкрота разбивки долга нет (суммы
   разложены по кредиторам), и в акте оставалось «…в размере 824 831,33 руб.,
   из них основного долга в размере  руб., просроченных процентов в размере
   руб.». Чистка пустых маркеров такую фразу не брала: её паттерны для [13]/[14]
   именительного падежа («основной долг»), а в шаблоне родительный.
"""
import os
import re
import sys

import pytest
from docx import Document

_THIS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(_THIS, "..", "app")))

TEMPLATES = os.path.abspath(os.path.join(_THIS, "..", "..", "Templates"))

pytestmark = pytest.mark.skipif(
    not os.path.isdir(TEMPLATES), reason="папка Templates недоступна")

# «Включить требование …(реквизиты)… [12] руб.» — резолютивная фраза ВКЛ в РТК.
INCLUDE_RE = re.compile(r"Включить\s+требование")
AMOUNT_RE = re.compile(r"(?:\)|\[988\])\s*(?!в\s+размере)\[12\]")


def _paragraphs():
    for root, _dirs, files in os.walk(TEMPLATES):
        for name in sorted(files):
            if not name.endswith(".docx") or name.startswith("~$"):
                continue
            path = os.path.join(root, name)
            for i, paragraph in enumerate(Document(path).paragraphs):
                yield os.path.relpath(path, TEMPLATES), i, paragraph.text


def test_include_claim_has_v_razmere():
    offenders = [
        f"{rel} [P{i}]: {text[:110]}"
        for rel, i, text in _paragraphs()
        if INCLUDE_RE.search(text) and AMOUNT_RE.search(text)
    ]
    assert not offenders, "сумма без «в размере»:\n" + "\n".join(offenders)


def test_include_claim_amount_has_comma_before_breakdown():
    """«руб.из которых» — слипшийся хвост из того же абзаца."""
    offenders = [
        f"{rel} [P{i}]: {text[:110]}"
        for rel, i, text in _paragraphs()
        if INCLUDE_RE.search(text) and re.search(r"руб\.\s*из\s+которых", text)
    ]
    assert not offenders, "нет запятой перед разбивкой:\n" + "\n".join(offenders)


# --- Чистка пустой разбивки в родительном падеже -----------------------------

def _act_with_empty_breakdown(tmp_path):
    doc = Document()
    doc.add_paragraph(
        "В обоснование заявления заявитель указывает, что должник имеет "
        "неисполненные обязательства в размере [12] руб.,"
        "из них основного долга в размере [13] руб., "
        "просроченных процентов в размере [14] руб."
    )
    path = tmp_path / "act.docx"
    doc.save(str(path))
    return str(path)


def test_empty_breakdown_phrase_removed(tmp_path):
    from document_generator import DocumentGenerator

    gen = DocumentGenerator()
    path = _act_with_empty_breakdown(tmp_path)
    doc = Document(path)
    for marker in ("[13]", "[14]"):
        gen._remove_placeholder_with_context(doc, marker)
    text = doc.paragraphs[0].text
    assert "основного долга" not in text, text
    assert "просроченных процентов" not in text, text
    # Целое остаётся нетронутым.
    assert "[12]" in text, text

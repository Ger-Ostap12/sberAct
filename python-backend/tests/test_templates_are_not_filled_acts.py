# -*- coding: utf-8 -*-
"""Страж: в Templates не должно быть готовых актов вместо шаблонов.

Дважды находили файлы с данными реальных дел вместо маркеров:
  • «продление БД.docx» — дело А53-41727-2/2025, судья Панова К.О., должник
    Белозорева В.Н. (отсюда жалоба «судья всегда Панова»);
  • «принятие иниц залог недвига.docx» — дело А53-39019-4/2025, должник
    Болдырева О.В., пять договоров и залоговая квартира; маркер на весь файл
    был ровно один.

Такой файл проходит генерацию молча и выдаёт юристу чужие персональные данные,
поэтому проверяем весь каталог, а не отдельные имена.
"""
import os
import re

import pytest
from docx import Document

_THIS = os.path.dirname(os.path.abspath(__file__))
TEMPLATES = os.path.abspath(os.path.join(_THIS, "..", "..", "Templates"))

pytestmark = pytest.mark.skipif(
    not os.path.isdir(TEMPLATES), reason="папка Templates недоступна")

# Номер конкретного дела: «Дело № А53-39019-4/2025».
CASE_NO = re.compile(r"Дело\s*№\s*А\d{2}-\d{3,}")
# Дата в шапке акта прописью: «26» февраля 2026 года.
HEADER_DATE = re.compile(
    r"«\d{2}»\s+(?:январ|феврал|март|апрел|ма|июн|июл|август|сентябр|октябр|ноябр|декабр)\w*\s+20\d\d")
# «в составе судьи Пановой К.О.»
JUDGE_IN_TEXT = re.compile(r"в\s+составе\s+судьи\s+[А-ЯЁ][а-яё]+(?:ой|ым|а|ы)?\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.")
# Подпись «Судья К.О. Панова»
JUDGE_SIGN = re.compile(
    r"^\s*Судья\b[\s\.]*"
    r"(?:[А-ЯЁ]\.\s*[А-ЯЁ]\.\s*[А-ЯЁ][а-яё]+|[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.)\s*$")
MARKER = re.compile(r"\[[0-9]+(?:\.[0-9]+)*\]")


def _docx_files():
    from pathlib import Path
    return [p for p in sorted(Path(TEMPLATES).rglob("*.docx"))
            if not p.name.startswith("~$")]


def _paragraphs(path):
    return [re.sub(r"[ ]+", " ", p.text).strip() for p in Document(str(path)).paragraphs]


def test_no_real_case_numbers():
    hits = []
    for path in _docx_files():
        for i, text in enumerate(_paragraphs(path)):
            if CASE_NO.search(text):
                hits.append(f"{path.name} para {i}: {text[:60]!r}")
    assert not hits, "номер реального дела вместо маркера: " + "; ".join(hits)


def test_no_real_header_dates():
    hits = []
    for path in _docx_files():
        for i, text in enumerate(_paragraphs(path)):
            if HEADER_DATE.search(text):
                hits.append(f"{path.name} para {i}: {text[:60]!r}")
    assert not hits, "дата реального акта вместо маркера: " + "; ".join(hits)


def test_no_judge_names():
    hits = []
    for path in _docx_files():
        for i, text in enumerate(_paragraphs(path)):
            if JUDGE_IN_TEXT.search(text) or (JUDGE_SIGN.match(text) and "[415]" not in text):
                hits.append(f"{path.name} para {i}: {text[:60]!r}")
    assert not hits, "ФИО судьи вместо [415]: " + "; ".join(hits)


def test_every_template_has_markers():
    """Шаблон без маркеров — признак того, что в папку попал готовый акт."""
    poor = []
    for path in _docx_files():
        markers = set()
        for p in Document(str(path)).paragraphs:
            markers.update(MARKER.findall(p.text))
        if len(markers) <= 3:
            poor.append(f"{path.name}: маркеров {len(markers)}")
    assert not poor, "похоже на заполненные акты: " + "; ".join(poor)

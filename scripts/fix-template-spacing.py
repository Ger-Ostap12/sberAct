# -*- coding: utf-8 -*-
"""Вставляет пробел между маркером и «руб.»: «[14]руб.» → «[14] руб.».

В шаблонах местами пропущен пробел, и в готовом акте суммы слипаются с единицей
измерения: «2 438 262,70руб.». Подстановка тут ни при чём — так написано в самом
шаблоне.

Правка идёт по runs: Word рвёт текст произвольно, и «]» с «руб» нередко лежат в
соседних runs — тогда пробел дописываем в начало второго. Замена значения
целиком через `paragraph.text` недопустима: она сбрасывает форматирование.

Запуск:
    python scripts/fix-template-spacing.py [--dry-run] [--templates <путь>]
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

from docx import Document

GLUED = re.compile(r"\](руб|рублей)")


def fix_paragraph(paragraph) -> int:
    """Возвращает число исправленных мест в абзаце."""
    fixed = 0

    # 1. Слипание внутри одного run.
    for run in paragraph.runs:
        if GLUED.search(run.text):
            run.text, n = GLUED.subn(r"] \1", run.text)
            fixed += n

    # 2. Слипание на границе двух runs: «…]» + «руб…».
    runs = paragraph.runs
    for i in range(len(runs) - 1):
        left, right = runs[i], runs[i + 1]
        if left.text.endswith("]") and re.match(r"(руб|рублей)", right.text):
            right.text = " " + right.text
            fixed += 1
    return fixed


def process(path: Path, dry_run: bool, backup_root: Path, rel: Path) -> int:
    doc = Document(str(path))
    fixed = 0
    for i, paragraph in enumerate(doc.paragraphs):
        n = fix_paragraph(paragraph)
        if n:
            fixed += n
            print(f"    para {i}: {n} шт. -> {paragraph.text[:80]!r}")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    fixed += fix_paragraph(paragraph)
    if fixed and not dry_run:
        backup = backup_root / rel
        backup.parent.mkdir(parents=True, exist_ok=True)
        if not backup.exists():
            shutil.copy(path, backup)
        doc.save(str(path))
    return fixed


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--templates", default=str(Path(__file__).resolve().parents[1] / "Templates"))
    args = parser.parse_args()

    root = Path(args.templates)
    backup_root = Path(__file__).resolve().parents[1] / ".template-backups"
    total = 0
    for path in sorted(root.rglob("*.docx")):
        if path.name.startswith("~$"):
            continue
        rel = path.relative_to(root)
        doc = Document(str(path))
        if not any(GLUED.search(p.text) for p in doc.paragraphs):
            continue
        print(f"== {rel}")
        total += process(path, args.dry_run, backup_root, rel)
    print(f"\nисправлено мест: {total}{' (dry-run)' if args.dry_run else ''}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

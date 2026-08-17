# -*- coding: utf-8 -*-
"""Возвращает «в размере» в резолютивку актов о включении в РТК.

В шаблонах фраза написана так:

    Включить требование «[989]» (ИНН [991], …, адрес регистрации: [988]) [12] руб.,
    из которых: [13] руб. – просроченный основной долг, …

и в готовом акте читается как «…адрес регистрации: 117312, г. Москва, ул.
Вавилова, д. 19) 180 532,78 руб., из которых» — сумма повисает без управления.
В теле того же акта («образовалась просроченная задолженность в размере [12]
руб.») формулировка правильная — расходятся именно шаблоны.

Правим только этот оборот: маркер [12] сразу после закрывающей скобки реквизитов
кредитора. Обороты «составляет [12] рублей», «на сумму [12] руб.» грамматически
целые и не трогаются. Заодно чиним слипшееся «руб.из которых» в тех же абзацах.

Правка идёт по runs (Word рвёт текст произвольно), замена через paragraph.text
недопустима — она сбрасывает форматирование.

Запуск:
    python scripts/fix-template-vkl-v-razmere.py [--dry-run] [--templates <путь>]
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

from docx import Document

# «…) [12]» / «…[988] [12]»: маркер суммы сразу за реквизитами кредитора.
TARGET = re.compile(r"(?:\)|\[988\])\s*(?=\[12\])")
INSERT = "в размере "
# «руб.из которых» / «руб. из которых» без запятой — в тех же абзацах.
GLUED_TAIL = re.compile(r"руб\.\s*(?=из\s+которых)")


def _insert_at(paragraph, index: int, text: str) -> bool:
    """Вставляет text в позицию index абзаца, работая по runs."""
    offset = 0
    for run in paragraph.runs:
        length = len(run.text)
        if offset <= index <= offset + length:
            local = index - offset
            run.text = run.text[:local] + text + run.text[local:]
            return True
        offset += length
    return False


def fix_paragraph(paragraph) -> int:
    """Возвращает число исправленных мест в абзаце."""
    fixed = 0
    # Вставки сдвигают текст — каждый раз перечитываем абзац заново.
    while True:
        match = TARGET.search(paragraph.text)
        if not match:
            break
        if not _insert_at(paragraph, match.end(), INSERT):
            break
        fixed += 1
    while True:
        text = paragraph.text
        match = GLUED_TAIL.search(text)
        if not match:
            break
        # Запятая ставится сразу после «руб.»; пробел дописываем, только если
        # его там не было («руб.из которых»), иначе получится двойной.
        at = match.start() + len("руб.")
        addition = "," if text[at:at + 1].isspace() else ", "
        if not _insert_at(paragraph, at, addition):
            break
        fixed += 1
    return fixed


def process(path: Path, dry_run: bool, backup_root: Path, rel: Path) -> int:
    doc = Document(str(path))
    fixed = 0
    for i, paragraph in enumerate(doc.paragraphs):
        if "[12]" not in paragraph.text:
            continue
        before = paragraph.text
        n = fix_paragraph(paragraph)
        if n:
            fixed += n
            print(f"    para {i}: {n} шт.")
            print(f"      было:  {before[:120]!r}")
            print(f"      стало: {paragraph.text[:120]!r}")
    if fixed and not dry_run:
        backup = backup_root / rel
        backup.parent.mkdir(parents=True, exist_ok=True)
        if not backup.exists():
            shutil.copy(path, backup)
        doc.save(str(path))
    return fixed


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--templates", default=str(Path(__file__).resolve().parents[1] / "Templates"))
    args = parser.parse_args()

    root = Path(args.templates)
    backup_root = Path(__file__).resolve().parents[1] / ".template-backups"
    total = 0
    for path in sorted(root.rglob("*.docx")):
        if path.name.startswith("~$"):
            continue
        rel = path.relative_to(root)
        doc = Document(str(path))
        if not any(TARGET.search(p.text) or GLUED_TAIL.search(p.text) for p in doc.paragraphs):
            continue
        print(f"== {rel}")
        total += process(path, args.dry_run, backup_root, rel)
    print(f"\nисправлено мест: {total}{' (dry-run)' if args.dry_run else ''}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

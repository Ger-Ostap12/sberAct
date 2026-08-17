# -*- coding: utf-8 -*-
"""Размечает «продление БД.docx»: реальные данные -> маркеры.

Файл лежал в Templates как готовый акт по делу А53-41727-2/2025 (судья Панова
К.О., должник Белозорева В.Н., кредитор ООО ПКО «Айди Коллект»). Из-за этого в
любом продлении Б/Д печатались чужие данные — отсюда жалоба «судья всегда
Панова К.О.».

Разметка снята с соседнего «Определение БД иное.docx» — тот же тип акта,
маркеры те же.

Запуск:
    python scripts/markup-prodlenie-bd.py [--dry-run]
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Emu, Twips

RELATIVE_PATH = Path("промежуточные_особые/продление БД.docx")

# (что искать, чем заменить). Порядок важен: длинные фрагменты идут раньше,
# чтобы короткие не съели их часть.
REPLACEMENTS: list[tuple[str, str]] = [
    # Шапка: дата акта и номер дела.
    ("«18» марта 2026 года", "[66]"),
    ("Дело № А53-41727-2/2025", "Дело № [1]"),
    # Судья.
    ("в составе судьи Пановой К.О.", "в составе судьи [415]"),
    # Кредитор-заявитель.
    ("общества с ограниченной ответственностью профессиональной коллекторской "
     "организации «Айди Коллект» (ИНН 7730233723, ОГРН 1177746355225, "
     "почтовый адрес 420044, г. Казань, пр. Победы, д. 220Б)",
     "«[989]» (ОГРН [990], ИНН [991], место нахождения: [988])"),
    ("ООО ПКО «Айди Коллект» обратилось", "[989] обратилось"),
    # Должник.
    ("Белозоревой Виктории Николаевны (05.06.1991 года рождения, место рождения: "
     "гор. Свободный-18 Амурской обл., ИНН 282369891741, СНИЛС 150-367-354 47, "
     "адрес регистрации: Ростовская область, г. Ростов-на-Дону, ул. Евдокимова, "
     "д. 37Б, кв. 67)",
     "[2.1] ([3] года рождения, место рождения: [3.1], ИНН [4], СНИЛС [5], "
     "адрес регистрации: [6])"),
    # Решение о банкротстве и финансовый управляющий.
    ("решением Арбитражного суда Ростовской области от 20.01.2026",
     "решением Арбитражного суда Ростовской области от [7]"),
    ("Финансовым управляющим утвержден Сергиенко Иван Геннадьевич",
     "Финансовым управляющим утвержден [8]"),
    # Публикация.
    ("«Коммерсантъ» №16(8190) от 31.01.2026", "«Коммерсантъ» №[67] от [68]"),
    # Дата обращения (штамп канцелярии / «Мой Арбитр»).
    ("04.02.2026 (согласно штампу канцелярии; 04.02.2026 посредством сервиса",
     "[23] (согласно штампу канцелярии; [24] посредством сервиса"),
    # Требование и его расшифровка.
    ("по кредитному договору № 309617 от 07.08.2019 в размере 337 854,38 руб.",
     "по кредитному договору № [110] от [100] в размере [12] руб."),
    ("238 141,27 руб. – основной долг", "[13] руб. – основной долг"),
    ("93 205,63 руб. – проценты", "[14] руб. – проценты"),
    ("3 560,28 руб. – пени, штрафы, неустойки", "[15] руб. – пени, штрафы, неустойки"),
    ("2 947,20 руб. – государственная пошлина", "[16] руб. – государственная пошлина"),
    # Предыдущее определение об оставлении без движения.
    ("Определением от 11.02.2026 заявление оставлено без движения до 17.03.2026",
     "Определением от [97] заявление оставлено без движения до [98]"),
    # Новый (продлённый) срок. Своего маркера для него в легенде нет, поэтому
    # используем тот же [98]: поле «срок устранения недостатков» в интерфейсе одно.
    ("в срок до 21.04.2026 устранить", "в срок до [98] устранить"),
]


def replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    """Замена по склейке абзаца: Word рвёт текст на runs произвольно, поэтому
    ищем в полном тексте, а пишем в первый затронутый run (остальные чистим)."""
    text = paragraph.text
    idx = text.find(old)
    if idx < 0:
        return False
    end = idx + len(old)
    pos = 0
    written = False
    for run in paragraph.runs:
        start, stop = pos, pos + len(run.text)
        pos = stop
        if stop <= idx or start >= end:
            continue
        keep_left = run.text[: max(0, idx - start)]
        keep_right = run.text[max(0, end - start):]
        if not written:
            run.text = keep_left + new + keep_right
            written = True
        else:
            run.text = keep_left + keep_right
    return written


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--templates", default=str(Path(__file__).resolve().parents[1] / "Templates"))
    args = parser.parse_args()

    path = Path(args.templates) / RELATIVE_PATH
    if not path.exists():
        print(f"!! нет файла: {path}")
        return 1

    doc = Document(str(path))
    applied, missing = 0, []
    for old, new in REPLACEMENTS:
        hit = False
        for paragraph in doc.paragraphs:
            if replace_in_paragraph(paragraph, old, new):
                hit = True
                applied += 1
                print(f"  + {old[:50]!r} -> {new[:50]!r}")
                break
        if not hit:
            missing.append(old)

    # Шапка «дата ⇥ Дело №»: табстоп был левый, из-за него части слипались.
    section = doc.sections[0]
    edge = int(Emu(section.page_width - section.left_margin - section.right_margin).twips)
    for paragraph in doc.paragraphs:
        if "[66]" in paragraph.text and "Дело" in paragraph.text:
            fmt = paragraph.paragraph_format
            fmt.tab_stops.clear_all()
            fmt.tab_stops.add_tab_stop(Twips(edge), WD_TAB_ALIGNMENT.RIGHT)
            print(f"  + табстоп по правому краю {edge} twips")

    if missing:
        print("\n!! не найдено в документе:")
        for m in missing:
            print(f"   {m[:80]!r}")

    if applied and not args.dry_run:
        backup = Path(__file__).resolve().parents[1] / ".template-backups" / RELATIVE_PATH
        backup.parent.mkdir(parents=True, exist_ok=True)
        if not backup.exists():
            shutil.copy(path, backup)
        doc.save(str(path))

    print(f"\nзамен применено: {applied}{' (dry-run)' if args.dry_run else ''}")
    remaining = [p.text for p in Document(str(path)).paragraphs
                 if re.search(r"Панова|Белозорев|Айди Коллект|А53-41727", p.text)]
    if remaining and not args.dry_run:
        print("!! остались реальные данные:")
        for r in remaining:
            print(f"   {r[:90]!r}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
